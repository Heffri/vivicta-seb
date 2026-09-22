"""Offline regression checks: python -m pipeline.test_runtime. No live models or corpus writes."""
from concurrent.futures import ThreadPoolExecutor
import json
import os
from pathlib import Path
import tempfile
import subprocess
import unittest
import io
from copy import deepcopy
from pptx import Presentation
from unittest.mock import patch

from fastapi.testclient import TestClient
import pymupdf

import app
from . import kb, llm


def _image_page(doc, text, width=300, height=100):
    """A page with no text layer: `text` rendered as a raster image, so a real OCR pass (not a
    mock) can recover it -- the same shape a scanned report page has."""
    page = doc.new_page(width=width, height=height)
    src = pymupdf.open()
    src_page = src.new_page(width=width, height=height)
    src_page.insert_text((10, height / 2), text, fontsize=20)
    page.insert_image(page.rect, pixmap=src_page.get_pixmap(dpi=200))
    src.close()
    return page


def _text_page(doc, text, width=300, height=100):
    page = doc.new_page(width=width, height=height)
    page.insert_text((10, height / 2), text, fontsize=11)
    return page


class RuntimeChecks(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.env = patch.dict(os.environ, {"KB_DIR": str(self.root / "kb"), "LLM_PROVIDER": "codex", "LLM_MODEL": "gpt-5.6-terra",
                                          "EMBED_MODEL": "test-model", "EMBED_BASE_URL": "http://embedding.test/v1", "FEWSHOT": "0"})
        self.env.start()
        self.uploads = patch.object(app, "UPLOADS", self.root / "uploads")
        self.uploads.start()
        app.UPLOADS.mkdir()
        for cache in (app.reports, app.extractions, app.texts_cache, app.library_paths, kb._vecs, kb._pages_cache, kb._status_cache):
            cache.clear()
        self.stem = "up-test"
        kb.save_report(self.stem, {"company": "Test AB", "fiscal_year": 2025, "pages": 1, "sha256": "pdf"}, ["Revenue 100 90"])
        self.client = TestClient(app.app)

    def tearDown(self):
        self.uploads.stop()
        self.env.stop()
        self.temp.cleanup()

    def test_index_identity_reuse_corruption_and_failed_build(self):
        with patch.object(kb, "embed", side_effect=lambda texts: [[1.0, 0.0] for _ in texts]) as embed:
            self.assertFalse(kb.index(self.stem)["cached"])
            self.assertTrue(kb.index(self.stem)["cached"])
            self.assertEqual(embed.call_count, 1)
            self.assertNotIn("vec", kb.inspect_chunks(self.stem)["items"][0])
            self.assertEqual(self.client.get(f"/api/knowledge/{self.stem}/chunks?offset=-1").status_code, 422)
            with patch.dict(os.environ, {"EMBED_MODEL": "new-model"}):
                self.assertEqual(kb.index_status(self.stem)["status"], "outdated")
                kb.index(self.stem)
            self.assertEqual(embed.call_count, 2)
            path = kb.kb_dir() / self.stem / "embeddings.jsonl"
            before = path.read_bytes()
            with patch.object(kb, "embed", side_effect=RuntimeError("unavailable")):
                with self.assertRaises(RuntimeError):
                    kb.index(self.stem)
            self.assertEqual(path.read_bytes(), before)
            path.write_text('broken', encoding="utf-8")
            self.assertEqual(kb.index_status(self.stem)["status"], "invalid")
            kb.index(self.stem)
            self.assertEqual(kb.index_status(self.stem)["status"], "ready")
            (path.parent / "index.json").unlink()
            self.assertEqual(kb.index_status(self.stem)["status"], "outdated")
            kb.index(self.stem)
            with patch.object(kb, "embed", return_value=[[1.0, 0.0, 0.0]]):
                with self.assertRaisesRegex(ValueError, "dimensions"):
                    kb.search([self.stem], "Revenue")

    def test_concurrent_index_builds_once(self):
        with patch.object(kb, "embed", side_effect=lambda texts: [[1.0, 0.0] for _ in texts]) as embed:
            with ThreadPoolExecutor(max_workers=2) as pool:
                results = list(pool.map(kb.index, [self.stem, self.stem]))
            self.assertEqual(embed.call_count, 1)
            self.assertEqual(sum(r["cached"] for r in results), 1)

    def test_listing_does_not_wait_for_a_long_index_build(self):
        with ThreadPoolExecutor(max_workers=1) as pool:
            with kb.report_lock(self.stem):
                status = pool.submit(kb.index_status, self.stem).result(timeout=1)
                self.assertEqual(status["status"], "building")

    def test_source_changes_failed_publication_and_empty_report(self):
        with patch.object(kb, "embed", side_effect=lambda texts: [[1.0, 0.0] for _ in texts]):
            kb.index(self.stem)
            folder = kb.kb_dir() / self.stem
            old = (folder / "embeddings.jsonl").read_bytes()
            kb.save_report(self.stem, {"company": "Test AB", "pages": 1, "sha256": "changed"}, ["Revenue 200 100"])
            self.assertEqual(kb.index_status(self.stem)["status"], "outdated")
            write = kb.atomic_write
            def fail_manifest(path, content):
                if path.name == "index.json":
                    raise OSError("disk full")
                return write(path, content)
            with patch.object(kb, "atomic_write", side_effect=fail_manifest):
                with self.assertRaises(OSError):
                    kb.index(self.stem)
            self.assertEqual((folder / "embeddings.jsonl").read_bytes(), old)
            self.assertEqual(kb.index_status(self.stem)["status"], "outdated")
            kb.index(self.stem)
            self.assertEqual(kb.index_status(self.stem)["status"], "ready")
            kb.save_report("empty", {"pages": 1, "sha256": "empty"}, [""])
            with self.assertRaisesRegex(ValueError, "requires OCR"):
                kb.index("empty")

    def test_extraction_cache_force_and_failure_preserve_saved_result(self):
        report = {"report_id": self.stem, "stem": self.stem, "company": "Test AB", "fiscal_year": 2025, "pages": 1}
        app.reports[self.stem] = report
        app.texts_cache[self.stem] = ["Revenue 100 90"]
        fixture = json.loads(app.FIXTURE.read_text(encoding="utf-8"))
        fixture.update(report_id=self.stem, warnings=[], timings={"model": 0.1, "validate": 0.01, "attempts": 1})
        url = f"/api/reports/{self.stem}/extract"
        with patch.object(app.locate, "candidate_pages", return_value=[1]), patch.object(app.extract_mod, "extract", side_effect=lambda *a, **kw: json.loads(json.dumps(fixture))) as extract:
            first = self.client.post(url, json={"section": "income_statement"})
            self.assertEqual(first.status_code, 200, first.text)
            self.assertFalse(first.json()["cached"])
            self.assertTrue(self.client.post(url, json={"section": "income_statement"}).json()["cached"])
            self.assertEqual(extract.call_count, 1)
            self.client.post(url, json={"section": "income_statement", "force": True})
            self.assertEqual(extract.call_count, 2)
            with patch.dict(os.environ, {"LLM_REASONING": "medium"}):
                self.client.post(url, json={"section": "income_statement"})
            self.assertEqual(extract.call_count, 3)
            path = kb.kb_dir() / self.stem / "extractions/income_statement.json"
            before = path.read_bytes()
            fixture["warnings"] = ["llm: RuntimeError: login expired"]
            for field in fixture["fields"]:
                field["value"] = None
            response = self.client.post(url, json={"section": "income_statement", "force": True})
            self.assertEqual(response.status_code, 502)
            self.assertEqual(path.read_bytes(), before)

    def test_saved_upload_reopens_and_exports(self):
        with pymupdf.open() as doc:
            doc.new_page().insert_text((40, 40), "Test AB annual report 2025 Revenue 100")
            doc.save(app.UPLOADS / f"{self.stem}.pdf")
        response = self.client.post(f"/api/knowledge/{self.stem}/open")
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(self.client.get(f"/api/reports/{self.stem}/pdf").status_code, 200)
        fixture = json.loads(app.FIXTURE.read_text(encoding="utf-8"))
        kb.save_extraction(self.stem, "income_statement", fixture)
        self.assertEqual(self.client.get(f"/api/kb/{self.stem}/income_statement").status_code, 200)
        self.assertEqual(self.client.get(f"/api/reports/{self.stem}/extraction.csv").status_code, 200)
        self.assertEqual(self.client.get(f"/api/reports/{self.stem}/extraction.pptx").status_code, 200)


    def test_exports_stay_with_requested_section_and_preserve_missing_values(self):
        app.reports[self.stem] = {"report_id": self.stem, "stem": self.stem}
        income = json.loads(app.FIXTURE.read_text(encoding="utf-8"))
        income["fields"][0]["value"] = 11.77
        kb.save_extraction(self.stem, "income_statement", income)
        debt = deepcopy(income)
        debt.update(section="debt_maturity", debt_scope="Consolidated borrowings", fields=[
            dict(key=k, label=k, value=v, unit="MSEK", period="2025", raw_label=k,
                 source=None, confidence=0, evidence=[]) for k, v in
            [("total_debt", 100), ("due_within_1_year", 10), ("due_1_to_5_years", None), ("due_after_5_years", None)]])
        kb.save_extraction(self.stem, "debt_maturity", debt)
        app.extractions[self.stem] = debt  # opening another section must not change an income export
        url = f"/api/reports/{self.stem}"
        csv = self.client.get(url + "/extraction.csv?section=income_statement")
        self.assertIn("11.77", csv.text)
        self.assertNotIn("debt_maturity", csv.text)
        pptx = self.client.get(url + "/extraction.pptx?section=income_statement")
        presentation = Presentation(io.BytesIO(pptx.content))
        cells = [cell.text for shape in presentation.slides[0].shapes if shape.has_table for row in shape.table.rows for cell in row.cells]
        self.assertIn("11.77", cells)
        pptx = self.client.get(url + "/extraction.pptx?section=debt_maturity")
        slide = Presentation(io.BytesIO(pptx.content)).slides[0]
        self.assertFalse(any(shape.has_chart for shape in slide.shapes))
        self.assertIn("Not available", [c.text for s in slide.shapes if s.has_table for r in s.table.rows for c in r.cells])
        self.assertIn("Consolidated borrowings", slide.notes_slide.notes_text_frame.text)
        debt["fields"][2]["value"], debt["fields"][3]["value"] = 50, 40
        debt["checks"] = [{"passed": True}]
        complete = Presentation(io.BytesIO(app.ppt.build_pptx(debt))).slides[0]
        chart = next(s.chart for s in complete.shapes if s.has_chart)
        self.assertTrue(all(0 <= int(e.get("val")) < 2**32 for e in chart._chartSpace.xpath(".//c:axId | .//c:crossAx")))
        saved = app.saved_extraction(self.stem, "debt_maturity")
        self.assertTrue(saved["stale"])
        self.assertEqual(saved["timings"]["attempts"], 0)
        (kb.kb_dir() / self.stem / "extractions/debt_maturity.json").write_text("[]")
        self.assertEqual(self.client.get(url + "/extraction.csv?section=debt_maturity").status_code, 409)

    def test_ocr_cache_settings_and_reopen_keep_provenance(self):
        from . import parse
        self.assertIsNotNone(kb.cached_texts(self.stem, "pdf"))
        with patch.dict(os.environ, {"OCR_LANGUAGE": "eng"}):
            self.assertIsNone(kb.cached_texts(self.stem, "pdf"))
        with pymupdf.open() as doc:
            doc.new_page()
            doc.save(app.UPLOADS / f"{self.stem}.pdf")
        def ocr(_path, metadata):
            metadata.update(ocr_pages=[1], ocr_settings=parse.ocr_settings())
            return ["Revenue 100 90"]
        with patch.object(parse, "page_texts", side_effect=ocr) as read:
            response = self.client.post(f"/api/knowledge/{self.stem}/open")
            self.assertEqual(response.status_code, 200, response.text)
            self.assertEqual(kb._meta(self.stem)["ocr_pages"], [1])
            app.reports.clear()
            self.client.post(f"/api/knowledge/{self.stem}/open")
            self.assertEqual(read.call_count, 1)

    def test_upload_over_budget_422_then_ocr_full_retry(self):
        """v191(a)(b): a scanned upload whose bounded candidate set alone exceeds OCR_PAGE_BUDGET is
        refused with a structured 422 naming how many pages full OCR needs; the same bytes with
        ocr=full run it unconditionally, no budget check. The ocr=full half exercises real OCR (not
        a mocked textpage) -- skipped, not red, on a clone that never ran scripts/setup_ocr.py."""
        from . import parse
        if not parse.ocr_ready():
            self.skipTest("OCR language files missing -- run python scripts/setup_ocr.py")
        doc = pymupdf.open()
        _text_page(doc, "Cover page")
        _text_page(doc, "Second page")
        _image_page(doc, "Scanned page three")
        _image_page(doc, "Scanned page four")
        _image_page(doc, "Scanned page five")
        data = doc.tobytes()
        doc.close()
        with patch.object(parse, "OCR_PAGE_BUDGET", 2):
            bounded = self.client.post("/api/reports", files={"file": ("scan.pdf", data, "application/pdf")})
            self.assertEqual(bounded.status_code, 422, bounded.text)
            body = bounded.json()
            self.assertEqual(body["ocr_pages_needed"], 3)
            self.assertRegex(body["detail"], r"~\d+ min for 3 pages")

            full = self.client.post("/api/reports?ocr=full", files={"file": ("scan.pdf", data, "application/pdf")})
            self.assertEqual(full.status_code, 200, full.text)
            self.assertEqual(full.json()["ocr_pages"], [3, 4, 5])

    def test_extract_fills_pending_candidate_page_on_demand(self):
        """v191(b): a candidate page the bounded registration pass left ocr_pending (it sits past
        the front matter, with no outline hit) is OCR'd individually, right before the model call
        that needs it -- extract_mod.extract never sees a blank page for a page it was told to read.
        The on-demand top-up runs real OCR -- skipped, not red, on a clone that never ran
        scripts/setup_ocr.py."""
        from . import parse
        if not parse.ocr_ready():
            self.skipTest("OCR language files missing -- run python scripts/setup_ocr.py")
        doc = pymupdf.open()
        _text_page(doc, "Cover page")
        for n in range(2, 10):
            _text_page(doc, f"Body text page {n}")
        _image_page(doc, "Borrowings note text")  # page 10: past the front matter, no toc hit
        doc.save(app.UPLOADS / f"{self.stem}.pdf")
        doc.close()

        app.reports[self.stem] = {"report_id": self.stem, "stem": self.stem, "company": "Test AB", "fiscal_year": 2025, "pages": 10}
        fixture = json.loads(app.FIXTURE.read_text(encoding="utf-8"))
        fixture.update(report_id=self.stem, warnings=[], timings={"model": 0.1, "validate": 0.01, "attempts": 1})
        captured = {}

        def fake_extract(texts, pages, schema, report, **kw):
            captured["texts"] = list(texts)
            return json.loads(json.dumps(fixture))

        with patch.object(app.locate, "candidate_pages", return_value=[10]), patch.object(app.extract_mod, "extract", side_effect=fake_extract):
            response = self.client.post(f"/api/reports/{self.stem}/extract", json={"section": "income_statement"})
        self.assertEqual(response.status_code, 200, response.text)
        self.assertIn("Borrowings note text", captured["texts"][9])  # the model read the filled page, not a blank one
        self.assertEqual(kb._meta(self.stem)["ocr_pending"], [])
        self.assertEqual(kb._meta(self.stem)["ocr_pages"], [10])

    def test_ask_uses_report_ids_and_withholds_bad_citations(self):
        from . import extract
        other = "test_2024"
        kb.save_report(other, {"company": "Test AB", "fiscal_year": 2024, "pages": 1, "sha256": "old"}, ["Revenue 80 70"])
        hits = [{"stem": s, "page": 1, "text": t, "score": 0.9} for s, t in [(self.stem, "Revenue 100 90"), (other, "Revenue 80 70")]]
        raw = {"answer": "2024 revenue was 80.", "citations": [{"company": "Test AB", "report_stem": other, "fiscal_year": 2024, "page": 1, "quote": "Revenue 80 70"}]}
        with patch.object(kb, "search", return_value=hits), patch.object(kb.llm, "chat", side_effect=lambda *a: json.dumps(raw)):
            answer = kb.ask([self.stem, other], "Compare revenue")
            self.assertEqual(answer["citations"][0]["fiscal_year"], 2024)
            raw["citations"][0]["report_stem"] = self.stem
            answer = kb.ask([self.stem, other], "Compare revenue")
            self.assertEqual(answer["citations"], [])
            self.assertIn("withheld", answer["answer"])
        self.assertEqual(self.client.post("/api/ask", json={"question": " ", "report_ids": [self.stem]}).status_code, 400)

    def test_facts_require_printed_evidence_and_force_reembeds(self):
        field = dict(key="revenue", label="Revenue", value=100, period="2025", unit="MSEK", confidence=1,
                     source={"page": 1, "quote": "Revenue 100 90"})
        x = dict(section="income_statement", fields=[field], checks=[])
        kb.save_extraction(self.stem, "income_statement", x)
        self.assertEqual(len([c for c in kb.chunks(self.stem) if c["start"] == -1]), 1)
        field["value"] = 999
        kb.save_extraction(self.stem, "income_statement", x)
        self.assertEqual(len([c for c in kb.chunks(self.stem) if c["start"] == -1]), 0)
        field["value"] = 100
        field["components"] = [{"source": field["source"]}, {"source": field["source"]}]
        kb.save_extraction(self.stem, "income_statement", x)
        self.assertEqual(len([c for c in kb.chunks(self.stem) if c["start"] == -1]), 0)
        with patch.object(kb, "embed", side_effect=lambda ts: [[1., 0.] for _ in ts]) as embed:
            kb.index(self.stem)
            kb.index(self.stem, force=True)
            self.assertEqual(embed.call_count, 2)


    def test_cli_schema_is_enabled_by_default(self):
        with patch.dict(os.environ):
            os.environ.pop("LLM_STRICT_SCHEMA", None)
            with patch.object(llm, "_codex_chat", return_value='{}') as call:
                llm.chat("system", "user", {"type": "object"})
                self.assertEqual(call.call_args.kwargs["schema"], {"type": "object"})

    def test_codex_inference_is_isolated(self):
        from unittest.mock import Mock
        def run(args, **kwargs):
            Path(args[args.index("-o") + 1]).write_text('{"answer":"ok"}', encoding="utf-8")
            return Mock(returncode=0)
        with patch.object(llm, "_codex_executable", return_value="codex"), patch.object(llm.subprocess, "run", side_effect=run) as call:
            self.assertIn("ok", llm._codex_chat("system", "untrusted input"))
            args = call.call_args.args[0]
            self.assertIn("--ignore-user-config", args)
            self.assertIn("features.shell_tool=false", args)
            self.assertIn('web_search="disabled"', args)
            self.assertNotIn("untrusted input", args)


if __name__ == "__main__":
    unittest.main()
