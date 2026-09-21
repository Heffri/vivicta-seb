"""Self-check for evidence scoring (docs/CONFIDENCE.md). Run: python -m pipeline.test_confidence"""
from . import extract as x


def demo():
    schema = {"keywords": ["income statement"], "checks": [{"name": "gp", "expr": "abs((revenue + cost_of_sales) - gross_profit) <= 2", "identity": True}],
              "fields": [{"key": "revenue", "synonyms": ["revenue", "net sales"], "unit_hint": "currency_millions"},
                         {"key": "eps_basic", "synonyms": ["earnings per share"], "unit_hint": "currency_per_share"}]}
    checks = [{"name": "gp", "passed": True, "detail": ""}]
    f = {"key": "revenue", "value": 168343, "unit": "MSEK", "period": "2025", "raw_label": "Revenues 3",
         "source": {"page": 106, "quote": "Revenues 3 168 343 176 771"}, "confidence": 0, "evidence": ["quote_on_page"]}
    x.score_field(f, schema["fields"][0], checks, schema, "MSEK", 2025, {106})
    assert f["confidence"] == 1.0 and len(f["evidence"]) == 7, f
    f = {"key": "eps_basic", "value": 5.43, "unit": "SEK", "period": "2025", "raw_label": "Basic earnings per share",
         "source": {"page": 106, "quote": "Basic earnings per share, SEK 5.43 5.66"}, "confidence": 0, "evidence": ["quote_on_page"]}
    x.score_field(f, schema["fields"][1], checks, schema, "MSEK", 2025, {106})
    assert "unit_ok" in f["evidence"] and "value_in_quote" in f["evidence"] and "label_known" not in f["evidence"], f  # "Basic earnings..." is not a prefix of a synonym
    f = {"key": "revenue", "value": 138494, "unit": "MSEK", "period": "2025", "raw_label": "Net sales",
         "source": {"page": 189, "quote": "made up"}, "confidence": 0, "evidence": []}
    x.score_field(f, schema["fields"][0], checks, schema, "MSEK", 2025, set())
    assert f["confidence"] == 0.25 and "quote_on_page" not in f["evidence"], f  # cap: no provenance
    f = {"key": "revenue", "value": 100, "unit": "MSEK", "period": "2025", "raw_label": "Net sales",
         "source": {"page": 1, "quote": "Net sales 100"}, "confidence": 0, "evidence": ["quote_on_page"]}
    x.score_field(f, schema["fields"][0], [{"name": "gp", "passed": False, "detail": ""}], schema, "MSEK", 2025, {1})
    assert f["confidence"] == 0.5 and "arith_ok" not in f["evidence"], f  # cap: failed check
    assert x._value_in_quote(-7246, "Income tax (7,246) (6,910)") and x._value_in_quote(4.56, "4,56 kr") and not x._value_in_quote(1683, "3 168 343")
    assert not x._value_in_quote(1171, "Earnings per share 11.71 10.02") and x._value_in_quote(11.71, "11.71") and x._value_in_quote(168343, "168.343")
    assert x.repair_value(1171, "Basic earnings per share, SEK 11.71") == 11.71 and x.repair_value(5424600, "Net sales 5,424.6 5,401.1") == 5424.6
    assert x.repair_value(168343, "Revenues 3 168 343 176 771") is None and x.repair_value(309, "EPS 3.09") == 3.09
    from .parse import quote_on_page
    page = "Revenue\n6,7\n176,658\n176,481\nCost of sales"
    assert quote_on_page("Revenue 176,658", page)  # note refs between label and number
    assert quote_on_page("Försäljningsintäkter 79 146", "Försäljningsintäkter\n4\n79 146\n63 751")
    assert quote_on_page("Kostnad för sålda varor -61 978", "Kostnad för sålda varor\n5\n–61 978\n–50 088")  # en dash
    assert not quote_on_page("Revenue 176", page) and not quote_on_page("Revenue 999", page)
    assert not quote_on_page("176,658", "Revenue 176,658") and not quote_on_page("Revenue", "Revenue 176,658")  # bare number / bare label
    telia = "Continuing operations \nRevenue\n C5, C6\n80,982\n80,965\nGoods and services purchased\nC7\n-28,277\nOperating income\nC5\n10,426\n10,834"
    assert quote_on_page("Continuing operations Operating income C5 10,426 10,834", telia) == "Operating income C5 10,426 10,834"  # header prepended
    assert quote_on_page("Earnings per share\n1\nGroup\n15.76\n15.73", "EARNINGS PER SHARE 1\nsek\n2025\n2024\nGroup, excluding items 17.09 16.74\nGroup\n15.76\n15.73") == "Group 15.76 15.73"
    # column repair: Sandvik prints 2024 before 2025
    assert x._year_column("Consolidated income statement\nMSEK\nNote\n2024\n2025\nRevenue\nG2, G3\n122,878\n120,680", 2025) == (1, 2)
    assert x._year_column("Income Statements\nYear ended December 31 ($ in millions)\n2025\n2024\nSales", 2025) == (0, 2)
    assert x._year_column("Consolidated income statement for the year ended december 31 msek note revenues 168 343", 2025) is None
    assert x._year_column("Industrial Operations Financial Services Eliminations Volvo Group\n2025 2024 2025 2024 2025 2024 2025 2024\nNet sales", 2025) is None  # segments side by side
    assert x._row_amounts("Revenue\nG2, G3\n122,878\n120,680") == [122878, 120680] and x._row_amounts("Revenue 6, 7 176,658 176,481") == [176658, 176481]
    assert x._row_amounts("Försäljningsintäkter 4 79 146 63 751") == [79146, 63751] and x._row_amounts("Cost of goods sold 6,12 -1,827.9 -1,791.4") == [-1827.9, -1791.4]
    assert x._row_amounts("Income tax (7,246) (6,910)") == [-7246, -6910] and x._row_amounts("Resultat per aktie (SEK) 11,77 9,02") == [11.77, 9.02]
    assert x._row_amounts("Discontinued operations held for sale, net of income tax 18 (19) (37)", 2) == [-19, -37]  # Arion: a two-digit amount in parentheses is not a note reference
    assert x._row_amounts("Earnings per share, Euro cent 23.0 38.1") == [23, 38.1] and x._row_amounts("Net sales 3,5,24 5,424.6 5,401.1") == [5424.6, 5401.1]
    assert x.repair_value(230, "Earnings per share, Euro cent 23.0 38.1") == 23 and x.repair_value(50, "Note 5 Revenue") is None
    assert x._value_in_quote(0.9, "total C20 0.90 1.80") and x._value_in_quote(11.7, "(SEK) 14 11,70 7,74") and not x._value_in_quote(11.7, "11.75")
    # Swedish space-grouped rows need the column count: "155 054 161 900" is two amounts, not one
    assert x._row_amounts("Sales 155 054 161 900", 2) == [155054, 161900] and x._row_amounts("Total sales 6, 10 155 113 161 921", 2) == [155113, 161921]
    assert x._row_amounts("Innehav utan bestämmande inflytande 42 39", 2) == [42, 39] and x._row_amounts("Production expenses 11, 12, 13 -121 972 -127 935", 2) == [-121972, -127935]
    assert x._row_amounts("Operating income 9 073 9 296", 2) == [9073, 9296]
    assert x._year_column("Consolidated income statement\nSEK million\nNote\n2025\nJan-Dec\n2024\nRevenue", 2025) == (0, 2)  # Telia
    # page rows: the printed row behind a truncated quote, and the number under a synonym label the model renamed
    page = "Consolidated income statement\nSEK M\nNote\n2025\n2024\nRevenue\n6, 7\n176,658\n176,481\nCost of sales\n8, 9, 10, 33, 38\n-161,906\n-162,001\nGross income\n14,753\n14,480\nProfit after financial items\n3,145\n2,282\nEarnings per share, SEK1\n17\n8.29\n6.01"
    rows = x._page_rows(page)
    assert "Gross income 14,753 14,480" in rows and "Revenue 6, 7 176,658 176,481" in rows and "Cost of sales 8, 9, 10, 33, 38 -161,906 -162,001" in rows, rows
    assert x._row_label("Gross income 14,753 14,480") == "Gross income" and x._row_label("Earnings per share, SEK1 17 8.29 6.01") == "Earnings per share, SEK1"
    assert x._row_label("Production expenses 11, 12, 13 -121 972 -127 935") == "Production expenses"
    row = next(r for r in rows if quote_on_page("Gross income 14,480", r))
    assert x._row_amounts(row, 2) == [14753, 14480]
    pbt = {"key": "profit_before_tax", "synonyms": ["profit before tax", "profit after financial items"]}
    hit = next(r for r in rows if x._label_known(x._row_label(r), pbt) and x._value_in_quote(3145, r))
    assert x._label_known("Profit/loss before tax", pbt) and x._label_known("Profit (loss) before tax 9", pbt)  # Yubico
    assert x._clean_label("Result before Tax") == "profit before tax" and x._clean_label("Operating result 2.1") == "operating profit"  # SSAB, Stora Enso
    assert x._ccy("KSEK") == x._ccy("TSEK") == x._ccy("SEK thousand") == "SEK" and x._ccy("TEUR") == "EUR"  # Paradox reports in KSEK
    assert hit == "Profit after financial items 3,145 2,282", hit
    assert quote_on_page("Cost of sales -161,906", page), "long note-reference gap"
    assert x._ccy("MSEK") == x._ccy("SEKm") == x._ccy("SEK million") == "SEK" and x._ccy("USD m") == "USD" and x._ccy("MEUR") == x._ccy("Euro cent") == "EUR" and x._ccy("kr") == "SEK"
    # a check whose operand is missing is n/a, not failed; another year's figure is capped at 0.5
    f = {"key": "revenue", "value": 100, "unit": "MSEK", "period": "2025", "raw_label": "Net sales",
         "source": {"page": 1, "quote": "Net sales 100"}, "confidence": 0, "evidence": ["quote_on_page"]}
    x.score_field(f, schema["fields"][0], [{"name": "gp", "passed": False, "detail": "missing: gross_profit"}], schema, "MSEK", 2025, {1})
    assert f["confidence"] == 1.0, f
    f = {"key": "revenue", "value": 122878, "unit": "MSEK", "period": "2024", "raw_label": "Revenue",
         "source": {"page": 1, "quote": "Revenue 122,878 120,680"}, "confidence": 0, "evidence": ["quote_on_page"]}
    x.score_field(f, schema["fields"][0], checks, schema, "MSEK", 2025, {1})
    assert f["confidence"] == 0.5 and "period_ok" not in f["evidence"], f
    f = {"key": "revenue", "value": 4578, "unit": "MSEK", "period": "2025", "raw_label": "EBITDA",
         "source": {"page": 1, "quote": "EBITDA 6,564 7,143"}, "confidence": 0, "evidence": ["quote_on_page"]}
    x.score_field(f, schema["fields"][0], checks, schema, "MSEK", 2025, {1})
    assert f["confidence"] == 0.5 and "value_in_quote" not in f["evidence"], f  # number not in its own quote
    # end to end on a by-nature statement: a revenue subtotal is unwound, a null the model gave is filled from the page
    sca = ("Consolidated income statement\nSEKm\nNote\n2025\n2024\nNet sales\nB1\n20,427\n20,232\nOther operating income\nB1, B2\n3,020\n3,395\n"
           "Revenue\n23,447\n23,627\nOperating profit\n4,432\n5,027\nProfit before tax\n3,999\n4,521\nIncome taxes\nB5\n–794\n–882")
    full = {"name": "is", "keywords": ["income statement"], "checks": [], "fields": [
        {"key": "revenue", "label": "Revenue", "synonyms": ["revenue", "net sales"], "unit_hint": "currency_millions", "excludes": ["other operating income"]},
        {"key": "profit_before_tax", "label": "PBT", "synonyms": ["profit before tax"], "unit_hint": "currency_millions"},
        {"key": "income_tax", "label": "Tax", "synonyms": ["income taxes"], "unit_hint": "currency_millions"}]}
    x.call_llm = lambda *a, **k: {"fields": [
        {"key": "revenue", "value": 23447, "unit": "SEKm", "period": "2025", "raw_label": "Revenue", "source": {"page": 1, "quote": "Revenue 23,447 23,627"}},
        {"key": "profit_before_tax", "value": None, "unit": None, "period": None, "raw_label": None, "source": None},
        {"key": "income_tax", "value": -794, "unit": "SEKm", "period": "2025", "raw_label": "Income taxes", "source": {"page": 2, "quote": "Income taxes -794"}}]}
    out = x.extract([sca, "Note B5 Taxes\nblah"], [1, 2], full, {"fiscal_year": 2025})
    got = {f["key"]: (f["value"], f["confidence"], f["source"]["page"]) for f in out["fields"]}
    assert got == {"revenue": (20427, 1.0, 1), "profit_before_tax": (3999, 1.0, 1), "income_tax": (-794, 1.0, 1)}, (got, out["warnings"])
    assert out["fields"][1]["unit"] == "SEKm" and out["fields"][1]["source"]["quote"] == "Profit before tax 3,999 4,521", out["fields"][1]
    # variant rows: an adjusted / diluted EPS label is not the field, the plain statement row replaces it
    import json, pathlib
    eps = next(f for f in json.loads((pathlib.Path(__file__).parents[1] / "schemas" / "income_statement.json").read_text("utf-8"))["fields"] if f["key"] == "eps_basic")
    assert x._label_known("Earnings per share before and after dilution (SEK)", eps) and x._label_known("Earnings per share (SEK), basic and diluted, total C20", eps)
    assert not x._label_known("Earnings per share before and after dilution and before items affecting comparability (SEK)1", eps)
    assert not x._label_known("Resultat per aktie efter utspädning (SEK)", eps) and x._label_known("Resultat per aktie före utspädning (SEK)", eps)
    sec = "Consolidated statement of income\nNote\n2025\n2024\nNet income for the year\n5 144\n5 172\nEarnings per share before and after dilution (SEK)\n3\n8.93\n9.01\nEarnings per share before and after dilution and before items affecting \ncomparability (SEK)1\n3\n11.55\n10.81"
    x.call_llm = lambda *a, **k: {"fields": [{"key": "eps_basic", "value": 11.55, "unit": "SEK", "period": "2025", "raw_label": "Earnings per share before and after dilution and before items affecting comparability (SEK)",
                                             "source": {"page": 1, "quote": "comparability (SEK)1 3 11.55 10.81"}}]}
    out = x.extract([sec], [1], {"name": "is", "keywords": [], "checks": [], "fields": [eps]}, {"fiscal_year": 2025})
    f = out["fields"][0]
    assert (f["value"], f["confidence"], f["source"]["quote"]) == (8.93, 1.0, "Earnings per share before and after dilution (SEK) 3 8.93 9.01"), (f, out["warnings"])
    # Röko: the text layer prints "Profit before tax 1,01 923" (a digit lost); the three rows above sum to 1,010 / 923
    roko = ("Consolidated income statement\nMSEK\nNote\n2025\n2024\nOperating profit\n8, 9, 10, 11, 12\n1,051\n969\nFinancial income\n7, 13\n49\n66\n"
            "Financial expenses\n7, 12, 13\n-90\n-112\nProfit before tax\n1,01\n923\nTax on net profit for the year\n14\n-254\n-221\nNet profit for the year1\n755\n702\n")
    sch = {"name": "is", "keywords": [], "fields": [{"key": k, "label": k, "synonyms": [syn], "unit_hint": "currency_millions"} for k, syn in
           (("profit_before_tax", "profit before tax"), ("income_tax", "tax on net profit for the year"), ("net_profit", "net profit for the year"))],
           "checks": [{"name": "net_profit_arith", "expr": "abs((profit_before_tax + income_tax) - net_profit) <= 2", "identity": True}]}
    x.call_llm = lambda *a, **k: {"fields": [
        {"key": "profit_before_tax", "value": 923, "unit": "MSEK", "period": "2025", "raw_label": "Profit before tax", "source": {"page": 1, "quote": "Profit before tax 1,01 923"}},
        {"key": "income_tax", "value": -254, "unit": "MSEK", "period": "2025", "raw_label": "Tax on net profit for the year", "source": {"page": 1, "quote": "Tax on net profit for the year 14 -254 -221"}},
        {"key": "net_profit", "value": 755, "unit": "MSEK", "period": "2025", "raw_label": "Net profit for the year1", "source": {"page": 1, "quote": "Net profit for the year1 755 702"}}]}
    out = x.extract([roko], [1], sch, {"fiscal_year": 2025})
    assert [(f["value"], f["confidence"]) for f in out["fields"]] == [(1010, 1.0), (-254, 1.0), (755, 1.0)], (out["fields"], out["warnings"])
    assert out["checks"][0]["passed"] and "value_derived" in out["fields"][0]["evidence"] and out["fields"][0]["source"]["quote"].startswith("Operating profit"), out
    # NOBA (a bank): no "profit before tax" row, so PBT's fallback synonym "operating profit" applies; "before credit losses" is excluded
    real = {f["key"]: f for f in json.load(open(pathlib.Path(__file__).parent.parent / "schemas" / "income_statement.json", encoding="utf-8"))["fields"]}
    bank = ("INCOME STATEMENT, CONSOLIDATED\nSEKm\nNOTE\nJAN-DEC \n2025\nJAN-DEC \n2024\nTotal operating income\n11,276\n9,884\nProfit before credit losses\n8,519\n7,161\n"
            "Net credit losses\n14\n-3,780\n-4,149\nOperating profit before amortisation of transaction surplus values\n4,739\n3,012\nAmortisation of surplus values from transactions\n12\n-128\n-134\n"
            "Operating profit \n6\n4,610\n2,878\nTax on profit for the year\n15\n-999\n-676\nNet profit for the year\n3,611\n2,202\n")
    sch = {"name": "is", "keywords": [], "fields": [real[k] for k in ("revenue", "profit_before_tax", "income_tax", "net_profit")],
           "checks": [{"name": "net_profit_arith", "expr": "abs((profit_before_tax + income_tax) - net_profit) <= 2", "identity": True}]}
    x.call_llm = lambda *a, **k: {"fields": [
        {"key": "revenue", "value": 11276, "unit": "SEKm", "period": "2025", "raw_label": "Total operating income", "source": {"page": 1, "quote": "Total operating income 11,276 9,884"}},
        {"key": "profit_before_tax", "value": 8519, "unit": "SEKm", "period": "2025", "raw_label": "Profit before credit losses", "source": {"page": 1, "quote": "Profit before credit losses 8,519 7,161"}},
        {"key": "income_tax", "value": -999, "unit": "SEKm", "period": "2025", "raw_label": "Tax on profit for the year", "source": {"page": 1, "quote": "Tax on profit for the year 15 -999 -676"}},
        {"key": "net_profit", "value": 3611, "unit": "SEKm", "period": "2025", "raw_label": "Net profit for the year", "source": {"page": 1, "quote": "Net profit for the year 3,611 2,202"}}]}
    out = x.extract([bank], [1], sch, {"fiscal_year": 2025})
    assert [(f["value"], f["confidence"]) for f in out["fields"]] == [(11276, 1.0), (4610, 1.0), (-999, 1.0), (3611, 1.0)], (out["fields"], out["warnings"])
    assert out["fields"][1]["raw_label"] == "Operating profit" and out["checks"][0]["passed"], out
    # an industrial keeps "Operating profit" and "Profit before tax" apart: the fallback never switches on when the primary row exists
    ind = "Income statement\nMSEK\n2025\n2024\nOperating profit\n1,051\n969\nProfit before tax\n1,010\n923\nTax\n-254\n-221\nNet profit for the year\n755\n702\n"
    x.call_llm = lambda *a, **k: {"fields": [
        {"key": "profit_before_tax", "value": 1051, "unit": "MSEK", "period": "2025", "raw_label": "Operating profit", "source": {"page": 1, "quote": "Operating profit 1,051 969"}},
        {"key": "income_tax", "value": -254, "unit": "MSEK", "period": "2025", "raw_label": "Tax", "source": {"page": 1, "quote": "Tax -254 -221"}},
        {"key": "net_profit", "value": 755, "unit": "MSEK", "period": "2025", "raw_label": "Net profit for the year", "source": {"page": 1, "quote": "Net profit for the year 755 702"}}]}
    out = x.extract([ind], [1], {**sch, "fields": sch["fields"][1:]}, {"fiscal_year": 2025})
    assert (out["fields"][0]["value"], out["fields"][0]["confidence"]) == (1051, 0.5) and "label_known" not in out["fields"][0]["evidence"], out["fields"][0]
    # NCC: EPS "before items affecting comparability" is the adjusted figure; "after items affecting comparability" is the IFRS one
    ncc = ("Consolidated income statement\nSEK M\nNote\n2025\n2024\nNet profit for the year\n142\n1,571\nEarnings per share before and after dilution\n26\n"
           "Earnings per share, SEK, before items affecting comparability\n13.89 \n16.08 \nEarnings per share, SEK, after items affecting comparability\n1.45 \n16.08 \n")
    x.call_llm = lambda *a, **k: {"fields": [{"key": "eps_basic", "value": 13.89, "unit": "SEK", "period": "2025", "raw_label": "Earnings per share, SEK, before items affecting comparability",
                                             "source": {"page": 1, "quote": "Earnings per share, SEK, before items affecting comparability 13.89 16.08"}}]}
    out = x.extract([ncc], [1], {"name": "is", "keywords": [], "checks": [], "fields": [real["eps_basic"]]}, {"fiscal_year": 2025})
    assert (out["fields"][0]["value"], out["fields"][0]["confidence"]) == (1.45, 1.0), (out["fields"][0], out["warnings"])
    # Stora Enso: "Materials and services" is not cost of sales, and a by-nature statement has no gross profit to check it against
    x.call_llm = lambda *a, **k: {"fields": [{"key": "cost_of_sales", "value": -7020, "unit": "EUR million", "period": "2025", "raw_label": "Materials and services",
                                             "source": {"page": 1, "quote": "Materials and services1 2.2 -7,020 -6,738"}}]}
    out = x.extract(["Consolidated income statement\nEUR million\n2025\n2024\nSales\n8,999\n9,050\nMaterials and services1\n2.2\n-7,020\n-6,738\n"], [1],
                    {"name": "is", "keywords": [], "checks": [], "fields": [real["cost_of_sales"], real["gross_profit"]]}, {"fiscal_year": 2025})
    assert out["fields"][0]["value"] is None and "dropped" in out["warnings"][-1], (out["fields"], out["warnings"])
    # Catena: no total tax row; "Current tax" + "Deferred tax" = the model's -423, and PBT + tax = net profit holds in both columns
    cat = ("Consolidated statement of comprehensive income\nSEK M\nNote 3 01/01/2025 -31/12/2025 01/01/2024 -31/12/2024\nProfit before tax\n2,067\n1,344\n"
           "Current tax\n-56\n-53\nDeferred tax\n10\n-367\n-211\nNet profit for the year\n19\n1,644\n1,080\n")
    assert x._year_column(cat, 2025) == (0, 2)
    x.call_llm = lambda *a, **k: {"fields": [
        {"key": "profit_before_tax", "value": 2067, "unit": "SEK M", "period": "2025", "raw_label": "Profit before tax", "source": {"page": 1, "quote": "Profit before tax 2,067 1,344"}},
        {"key": "income_tax", "value": -423, "unit": "SEK M", "period": "2025", "raw_label": "Current tax\nDeferred tax", "source": {"page": 1, "quote": "Deferred tax 10 -367 -211"}},
        {"key": "net_profit", "value": 1644, "unit": "SEK M", "period": "2025", "raw_label": "Net profit for the year", "source": {"page": 1, "quote": "Net profit for the year 19 1,644 1,080"}}]}
    sch = {"name": "is", "keywords": [], "fields": [real[k] for k in ("profit_before_tax", "income_tax", "net_profit")],
           "checks": [{"name": "net_profit_arith", "expr": "abs((profit_before_tax + income_tax) - net_profit) <= 2", "identity": True}]}
    out = x.extract([cat], [1], sch, {"fiscal_year": 2025})
    tax = out["fields"][1]
    assert (tax["value"], tax["confidence"], tax["raw_label"]) == (-423, 1.0, "Current tax + Deferred tax") and "value_derived" in tax["evidence"], (tax, out["warnings"])
    # ... but not when the comparative column disagrees (a coincidence in the fiscal year is not a proof): then -423 is printed nowhere and is dropped
    out = x.extract([cat.replace("-53", "-99")], [1], sch, {"fiscal_year": 2025})
    assert out["fields"][1]["value"] is None and any("printed on none of pages" in w for w in out["warnings"]), (out["fields"][1], out["warnings"])
    # ... and "Net operating surplus" is not operating profit; with no operating profit row on the page the value is dropped
    x.call_llm = lambda *a, **k: {"fields": [{"key": "operating_profit", "value": 2198, "unit": "SEK M", "period": "2025", "raw_label": "Net operating surplus",
                                             "source": {"page": 1, "quote": "Net operating surplus 2,198 1,789"}}]}
    out = x.extract([cat.replace("Profit before tax", "Net operating surplus\n2,198\n1,789\nProfit before tax")], [1],
                    {"name": "is", "keywords": [], "checks": [], "fields": [real["operating_profit"]]}, {"fiscal_year": 2025})
    assert out["fields"][0]["value"] is None and "dropped" in out["warnings"][-1], (out["fields"], out["warnings"])
    # SEB: the model timed out, rows are filled from the page; "Income tax expense 7,835" is printed unsigned; unit from the header
    seb = ("Income statement\nSEB GROUP\nSEK m\nNote\n2025\n2024\nTotal operating income\n76,939\n81,887\nOperating profit before items affecting comparability\n39,314\n46,043\n"
           "Operating profit\n38,898\n46,043\nIncome tax expense\n14\n7,835\n10,178\nNET PROFIT\n31,063\n35,865\nBasic earnings per share, SEK\n15\n15.60\n17.51\n"
           "Diluted earnings per share, SEK\n15\n15.43\n17.33\n")
    def boom(*a, **k):
        raise TimeoutError("stuck")
    x.call_llm = boom
    sch = {"name": "is", "keywords": [], "fields": [real[k] for k in ("revenue", "operating_profit", "profit_before_tax", "income_tax", "net_profit", "eps_basic")],
           "checks": [{"name": "net_profit_arith", "expr": "abs((profit_before_tax + income_tax) - net_profit) <= 2", "identity": True}]}
    out = x.extract([seb, ""], [1, 2], sch, {"fiscal_year": 2025})
    assert [(f["value"], f["confidence"]) for f in out["fields"]] == [(76939, 1.0), (38898, 1.0), (38898, 1.0), (-7835, 1.0), (31063, 1.0), (15.6, 1.0)], (out["fields"], out["warnings"])
    assert out["currency"] == "SEK m" and sum("llm:" in w for w in out["warnings"]) == 2, out  # the spread and the single page; not the four-page window
    assert x._page_unit("Consolidated Statement of Operations\nUSD Thousands\nNote 2025 2024") == "USD Thousands" and x._page_unit("SEK M\nNote 2025") == "SEK M"
    assert not x._label_known("Earnings per share fully diluted - USD1", real["eps_basic"]) and x._label_known("Earnings per share - USD1", real["eps_basic"])  # IPC
    # SEB, model answered: profit before tax offered as "Profit before credit losses..." (a variant); the statement row is the
    # exact synonym "Operating profit", not the prefix match "Operating profit before items affecting comparability"
    seb2 = seb.replace("Operating profit\n38,898", "Profit before credit losses and imposed levies\n44,342\n51,000\nOperating profit\n38,898")
    x.call_llm = lambda *a, **k: {"fields": [
        {"key": "profit_before_tax", "value": 44342, "unit": "SEK m", "period": "2025", "raw_label": "Profit before credit losses and imposed levies", "source": {"page": 1, "quote": "Profit before credit losses and imposed levies 44,342"}},
        {"key": "income_tax", "value": 7835, "unit": "SEK m", "period": "2025", "raw_label": "Income tax expense", "source": {"page": 1, "quote": "Income tax expense 14 7,835"}},
        {"key": "net_profit", "value": 31063, "unit": "SEK m", "period": "2025", "raw_label": "NET PROFIT", "source": {"page": 1, "quote": "NET PROFIT 31,063"}}]}
    out = x.extract([seb2], [1], {**sch, "fields": [real[k] for k in ("profit_before_tax", "income_tax", "net_profit")]}, {"fiscal_year": 2025})
    assert [(f["value"], f["confidence"]) for f in out["fields"]] == [(38898, 1.0), (-7835, 1.0), (31063, 1.0)], (out["fields"], out["warnings"])
    # IPC: the model quoted the first row under the "Cost of sales" heading; the four rows sum to revenue - gross profit in both columns
    ipc = ("Consolidated Statement of Operations\nUSD Thousands\nNote 2025 2024\nRevenue 3 685,888 797,783\nCost of sales\nProduction costs 4 (427,623) (448,218)\n"
           "Depletion and decommissioning costs 3,9 (122,749) (128,392)\nDepreciation of other tangible fixed assets 3,9 (5,597) (8,933)\n"
           "Exploration and business development costs 3 (1,799) (2,069)\nGross profit 3 128,120 210,171\nOther income / (expense) 751 1,137\n")
    x.call_llm = lambda *a, **k: {"fields": [
        {"key": "revenue", "value": 685888, "unit": "USD Thousands", "period": "2025", "raw_label": "Revenue", "source": {"page": 1, "quote": "Revenue 3 685,888"}},
        {"key": "cost_of_sales", "value": -427623, "unit": "USD Thousands", "period": "2025", "raw_label": "Cost of sales", "source": {"page": 1, "quote": "Production costs 4 (427,623)"}},
        {"key": "gross_profit", "value": 128120, "unit": "USD Thousands", "period": "2025", "raw_label": "Gross profit", "source": {"page": 1, "quote": "Gross profit 3 128,120"}}]}
    gp = {"name": "gross_profit_arith", "expr": "abs((revenue + cost_of_sales) - gross_profit) <= 2", "identity": True}
    out = x.extract([ipc], [1], {"name": "is", "keywords": [], "checks": [gp], "fields": [real[k] for k in ("revenue", "cost_of_sales", "gross_profit")]}, {"fiscal_year": 2025})
    cos = out["fields"][1]
    assert [(f["value"], f["confidence"]) for f in out["fields"]] == [(685888, 1.0), (-557768, 1.0), (128120, 1.0)], (out["fields"], out["warnings"])
    assert cos["raw_label"] == "Cost of sales" and cos["source"]["quote"].startswith("Production costs") and "value_derived" in cos["evidence"], cos
    # Lundbergs (by nature, no gross profit): "Rörelseresultat" offered as gross profit as well as operating profit -> the twin without the label goes,
    # then the cost row it required
    lb = ("Resultaträkning\nKONCERNEN, mnkr\nNot 2025 2024\nNettoomsättning m m 3 28 781 29 311\nRåvaror, förnödenheter samt kostnad sålda lageraktier -12 455 -13 494\n"
          "Personalkostnader 5 -4 038 -3 996\nRörelseresultat 11 16 077 10 334\n")
    x.call_llm = lambda *a, **k: {"fields": [
        {"key": "revenue", "value": 28781, "unit": "mnkr", "period": "2025", "raw_label": "Nettoomsättning m m", "source": {"page": 1, "quote": "Nettoomsättning m m 3 28 781"}},
        {"key": "cost_of_sales", "value": -12455, "unit": "mnkr", "period": "2025", "raw_label": "Råvaror, förnödenheter samt kostnad sålda lageraktier", "source": {"page": 1, "quote": "Råvaror, förnödenheter samt kostnad sålda lageraktier -12 455"}},
        {"key": "gross_profit", "value": 16077, "unit": "mnkr", "period": "2025", "raw_label": "Rörelseresultat", "source": {"page": 1, "quote": "Rörelseresultat 11 16 077"}},
        {"key": "operating_profit", "value": 16077, "unit": "mnkr", "period": "2025", "raw_label": "Rörelseresultat", "source": {"page": 1, "quote": "Rörelseresultat 11 16 077"}}]}
    out = x.extract([lb], [1], {"name": "is", "keywords": [], "checks": [gp], "fields": [real[k] for k in ("revenue", "cost_of_sales", "gross_profit", "operating_profit")]}, {"fiscal_year": 2025})
    assert [(f["value"], f["confidence"]) for f in out["fields"]] == [(28781, 1.0), (None, 0.0), (None, 0.0), (16077, 1.0)], (out["fields"], out["warnings"])
    assert x._page_unit(lb) == "mnkr", x._page_unit(lb)
    # Lundbergs, second run: the model paired every label with the row below it ("Rörelseresultat" 16077 quoted as "Resultat efter
    # finansiella poster 15 465"), and "Nettoomsättning" with the 30 615 subtotal on the "Övriga intäkter" row
    lb2 = ("Resultaträkning\nKONCERNEN, mnkr\nNot 2025 2024\nNettoomsättning m m 3 28 781 29 311\nÖvriga intäkter m m 4 1 834 2 385 30 615 31 696\n"
           "Rörelseresultat 11 16 077 10 334\nFinansiella intäkter 12 58 79\nResultat efter finansiella poster 15 465 9 808\nSkatt 13 -1 044 -1 425\nÅrets resultat 14 421 8 383\n")
    x.call_llm = lambda *a, **k: {"fields": [
        {"key": "revenue", "value": 30615, "unit": "mnkr", "period": "2025", "raw_label": "Nettoomsättning m m", "source": {"page": 1, "quote": "Övriga intäkter m m 4 1 834 2 385 30 615"}},
        {"key": "operating_profit", "value": 16077, "unit": "mnkr", "period": "2025", "raw_label": "Rörelseresultat", "source": {"page": 1, "quote": "Resultat efter finansiella poster 15 465"}},
        {"key": "profit_before_tax", "value": 15465, "unit": "mnkr", "period": "2025", "raw_label": "Resultat efter finansiella poster", "source": {"page": 1, "quote": "Skatt 13 -1 044"}},
        {"key": "income_tax", "value": -1044, "unit": "mnkr", "period": "2025", "raw_label": "Skatt", "source": {"page": 1, "quote": "Årets resultat 14 421"}},
        {"key": "net_profit", "value": 14421, "unit": "mnkr", "period": "2025", "raw_label": "Årets resultat", "source": {"page": 1, "quote": "Årets resultat 14 421"}}]}
    out = x.extract([lb2], [1], {**sch, "fields": [real[k] for k in ("revenue", "operating_profit", "profit_before_tax", "income_tax", "net_profit")]}, {"fiscal_year": 2025})
    assert [(f["value"], f["confidence"]) for f in out["fields"]] == [(28781, 1.0), (16077, 1.0), (15465, 1.0), (-1044, 1.0), (14421, 1.0)], (out["fields"], out["warnings"])
    assert out["fields"][1]["source"]["quote"] == "Rörelseresultat 11 16 077 10 334" and sum("row prints" in w for w in out["warnings"]) == 4, out["warnings"]
    # Catena (property company): rental income - property expenses = net operating surplus is the gross profit
    assert x._label_known("Net operating surplus", real["gross_profit"]) and x._label_known("Property expenses", real["cost_of_sales"]) and not x._label_known("Net operating surplus", real["operating_profit"])
    # Addnode: "Purchases of goods and services" is not a cost of sales label, but net sales + it = gross profit in both years
    an = "Consolidated income statement\nSEK m\nNote 2025 2024\nNet sales 2, 3, 38 5,793 7,757\nPurchases of goods and services 38 -1,350 -3,559\nGross profit 4,443 4,198\n"
    x.call_llm = lambda *a, **k: {"fields": [
        {"key": "revenue", "value": 5793, "unit": "SEK m", "period": "2025", "raw_label": "Net sales", "source": {"page": 1, "quote": "Net sales 2, 3, 38 5,793"}},
        {"key": "cost_of_sales", "value": -1350, "unit": "SEK m", "period": "2025", "raw_label": "Purchases of goods and services", "source": {"page": 1, "quote": "Purchases of goods and services 38 -1,350"}},
        {"key": "gross_profit", "value": 4443, "unit": "SEK m", "period": "2025", "raw_label": "Gross profit", "source": {"page": 1, "quote": "Gross profit 4,443"}}]}
    out = x.extract([an], [1], {"name": "is", "keywords": [], "checks": [gp], "fields": [real[k] for k in ("revenue", "cost_of_sales", "gross_profit")]}, {"fiscal_year": 2025})
    assert [f["confidence"] for f in out["fields"]] == [1.0, 1.0, 1.0] and "identity_all_columns" in out["fields"][1]["evidence"], (out["fields"], out["warnings"])
    out = x.extract([an.replace("-3,559", "-3,000")], [1], {"name": "is", "keywords": [], "checks": [gp], "fields": [real[k] for k in ("revenue", "cost_of_sales", "gross_profit")]}, {"fiscal_year": 2025})
    assert out["fields"][1]["confidence"] == 0.9, out["fields"][1]  # only one year adds up: the label stays unknown
    # Volvo: "2025 2024" per segment; the model took income taxes from Industrial Operations and the rest from Volvo Group
    vo = ("Consolidated income statement\nIndustrial Operations Financial Services Eliminations Volvo Group\nSEK M Note 2025 2024 2025 2024 2025 2024 2025 2024\n"
          "Net sales 6, 7 457,509 504,975 26,469 26,982 -4,795 -5,140 479,183 526,816\nIncome after financial items 43,522 63,168 3,869 4,042 – – 47,391 67,210\n"
          "Income taxes 10 -11,669 -15,542 -1,016 -1,092 – – -12,685 -16,634\nIncome for the period * 31,853 47,626 2,853 2,949 – – 34,707 50,576\n")
    assert x._row_amounts("Income taxes 10 -11,669 -15,542 -1,016 -1,092 – – -12,685 -16,634", 8) == [-11669, -15542, -1016, -1092, 0, 0, -12685, -16634]
    x.call_llm = lambda *a, **k: {"fields": [
        {"key": "revenue", "value": 479183, "unit": "SEK M", "period": "2025", "raw_label": "Net sales", "source": {"page": 1, "quote": "Net sales 6, 7 457,509 504,975 26,469 26,982 -4,795 -5,140 479,183 526,816"}},
        {"key": "profit_before_tax", "value": 47391, "unit": "SEK M", "period": "2025", "raw_label": "Income after financial items", "source": {"page": 1, "quote": "Income after financial items 43,522 63,168 3,869 4,042 – – 47,391 67,210"}},
        {"key": "income_tax", "value": -11669, "unit": "SEK M", "period": "2025", "raw_label": "Income taxes", "source": {"page": 1, "quote": "Income taxes 10 -11,669 -15,542 -1,016 -1,092 – – -12,685 -16,634"}},
        {"key": "net_profit", "value": 34707, "unit": "SEK M", "period": "2025", "raw_label": "Income for the period", "source": {"page": 1, "quote": "Income for the period * 31,853 47,626 2,853 2,949 – – 34,707 50,576"}}]}
    out = x.extract([vo], [1], {**sch, "fields": [real[k] for k in ("revenue", "profit_before_tax", "income_tax", "net_profit")]}, {"fiscal_year": 2025})
    assert [(f["value"], f["confidence"]) for f in out["fields"]] == [(479183, 1.0), (47391, 1.0), (-12685, 1.0), (34707, 1.0)], (out["fields"], out["warnings"])
    assert sum("segment" in w for w in out["warnings"]) == 1, out["warnings"]
    vo_fields = x.call_llm()["fields"]  # second run: profit before tax answered as the net income row with a quote the page does not print
    x.call_llm = lambda *a, **k: {"fields": [dict(f, value=34707, raw_label="Income for the period *", source={"page": 1, "quote": "Income for the period * 34,707 50,576"}) if f["key"] == "profit_before_tax" else f for f in vo_fields]}
    out = x.extract([vo], [1], {**sch, "fields": [real[k] for k in ("revenue", "profit_before_tax", "income_tax", "net_profit")]}, {"fiscal_year": 2025})
    assert [(f["value"], f["confidence"]) for f in out["fields"]] == [(479183, 1.0), (47391, 1.0), (-12685, 1.0), (34707, 1.0)], (out["fields"], out["warnings"])
    assert any("not printed on the statement" in w for w in out["warnings"]) and out["fields"][1]["raw_label"] == "Income after financial items", out["warnings"]
    assert x._row_amounts("Cost of sales 3 –297,0421) –320,821", 2) == [-297042, -320821] and x._value_in_quote(-297042, "Cost of sales 3 –297,0421) –320,821")  # Volvo Cars footnote marker
    # Volvo Cars: "Net income" is a prefix of "net income from discontinued operations", not a discontinued-operations row (the null fill took it, then "repaired" net profit)
    assert not x._label_known("Net income", real["profit_discontinued"]) and x._label_known("Income before tax", real["profit_before_tax"]) and x._label_known("Net income", real["net_profit"])
    # Addnode after two LLM timeouts: nothing answered; the statement rows fill the fields, tax = current + deferred between profit after financial items and profit for the year
    an2 = an + "Operating profit 2, 7 607 598\nProfit after financial items 514 536\nCurrent tax 11, 12 -157 -154\nDeferred tax 11, 12 27 20\nProfit for the year 384 402\nTax attributable to items that may be reclassified -11 14\n"
    x.call_llm = lambda *a, **k: {"fields": []}
    npa = {"name": "net_profit_arith", "expr": "abs((profit_before_tax + income_tax + profit_discontinued) - net_profit) <= 2", "identity": True}
    out = x.extract([an2], [1], {"name": "is", "keywords": [], "checks": [gp, npa], "fields": [real[k] for k in ("revenue", "cost_of_sales", "gross_profit", "profit_before_tax", "income_tax", "profit_discontinued", "net_profit")]}, {"fiscal_year": 2025})
    assert [(f["value"], f["confidence"]) for f in out["fields"]] == [(5793, 1.0), (-1350, 1.0), (4443, 1.0), (514, 1.0), (-130, 1.0), (None, 0.0), (384, 1.0)], (out["fields"], out["warnings"])
    assert out["fields"][4]["raw_label"] == "Current tax + Deferred tax" and "value_derived" in out["fields"][4]["evidence"] and "identity_all_columns" in out["fields"][1]["evidence"], out["fields"]
    assert not x._label_known("Tax attributable to items that may be reclassified", real["income_tax"])
    # Arion Bank: a column-major text layer (every figure, then every label); lines rebuilt from word coordinates
    import pymupdf
    from . import parse
    pg = pymupdf.open().new_page()
    for y, (lab, a, b) in enumerate([("Interest income", "129,777", "132,259"), ("Net earnings", "32,507", "26,112")]):
        pg.insert_text((40, 100 + 14 * y), lab + " " + "." * 40), pg.insert_text((400, 100 + 14 * y), a), pg.insert_text((470, 100 + 14 * y), b)
    assert parse._lines_from_words(pg).splitlines() == ["Interest income 129,777 132,259", "Net earnings 32,507 26,112"], parse._lines_from_words(pg)
    assert parse._numeric_run("Notes" + "".join(f"{chr(10)}{i}" for i in range(12)) + f"{chr(10)}Label") == 12 and parse._numeric_run(f"Gross income{chr(10)}14,753{chr(10)}14,480") == 2
    from . import locate
    assert locate.strip_boilerplate(["Financial statements Group and Parent company_ _____124\nNotes ......... 136\nConsolidated income statement\nNet sales 26 46,021 45,052"])[0] == "Consolidated income statement\nNet sales 26 46,021 45,052"  # AAK nav bar
    # NCC: "Result from sales of Group companies 20" offered as discontinued operations; the identity holds without it
    ncc2 = ("Consolidated income statement\nSEK M\nNote\n2025\n2024\nResult from sales of Group companies\n8\n20\n3\n"
            "Profit after financial items\n630\n1,863\nTax on profit for the year\n23, 39\n–489\n–292\nNet profit for the year\n142\n1,571\n")
    sch = {"name": "is", "keywords": [], "fields": [real[k] for k in ("profit_before_tax", "income_tax", "profit_discontinued", "net_profit")],
           "checks": [{"name": "net_profit_arith", "expr": "abs((profit_before_tax + income_tax + profit_discontinued) - net_profit) <= 2", "identity": True}]}
    x.call_llm = lambda *a, **k: {"fields": [
        {"key": "profit_before_tax", "value": 630, "unit": "SEK M", "period": "2025", "raw_label": "Profit after financial items", "source": {"page": 1, "quote": "Profit after financial items 630 1,863"}},
        {"key": "income_tax", "value": -489, "unit": "SEK M", "period": "2025", "raw_label": "Tax on profit for the year", "source": {"page": 1, "quote": "Tax on profit for the year 23, 39 –489 –292"}},
        {"key": "profit_discontinued", "value": 20, "unit": "SEK M", "period": "2025", "raw_label": "Result from sales of Group companies", "source": {"page": 1, "quote": "Result from sales of Group companies 8 20 3"}},
        {"key": "net_profit", "value": 142, "unit": "SEK M", "period": "2025", "raw_label": "Net profit for the year", "source": {"page": 1, "quote": "Net profit for the year 142 1,571"}}]}
    out = x.extract([ncc2], [1], sch, {"fiscal_year": 2025})
    assert out["fields"][2]["value"] is None and out["checks"][0]["passed"] and all(f["confidence"] == 1.0 for f in out["fields"] if f["value"] is not None), (out["fields"], out["warnings"])
    # Ericsson (debt maturity): the note has no >5y row, so due_after_5_years is null — a correct extraction, not a failure.
    # A check's null_as_zero lists the operands that count as 0 while they are null, as long as at least one of them is real;
    # when every listed operand is null the sum proves nothing and the check stays missing (never 0 == total).
    dmc = {"name": "maturity_sums_to_total", "expr": "abs((due_within_1_year + due_1_to_5_years + due_after_5_years) - total_debt) <= 2",
           "detail": "due_within_1_year + due_1_to_5_years + due_after_5_years == total_debt (±2 rounding)", "identity": True,
           "null_as_zero": ["due_within_1_year", "due_1_to_5_years", "due_after_5_years"]}
    c = x._check(dmc, {"due_within_1_year": 3538, "due_1_to_5_years": 29165, "total_debt": 32703})
    assert c["passed"] and "due_after_5_years null" in c["detail"], c  # the detail names the bucket counted as 0
    c = x._check(dmc, {"total_debt": 32703})  # no bucket printed at all
    assert not c["passed"] and c["detail"].startswith("missing:"), c
    c = x._check({k: v for k, v in dmc.items() if k != "null_as_zero"}, {"due_within_1_year": 3538, "due_1_to_5_years": 29165, "total_debt": 32703})
    assert not c["passed"] and c["detail"] == "missing: due_after_5_years", c  # without the flag: the old behaviour
    # ... end to end: the identity passes with the null bucket as 0, and total_debt keeps arith_ok at full confidence
    dm = json.loads((pathlib.Path(__file__).parents[1] / "schemas" / "debt_maturity.json").read_text("utf-8"))
    # 2026-09-18: the shipped schema sets require_explicit_values on the identity check (team ruling: a null
    # bucket stays unknown, never an implicit 0 -- see docs/API.md). Everything below exercises the null_as_zero
    # mechanism itself, so the self-check runs the schema with that flag off; the flag's own behaviour is
    # asserted right above (the "without the flag"/"missing:" pair) and in backend/test_workbench.py.
    for _chk in dm.get("checks", []):
        _chk.pop("require_explicit_values", None)
    ericsson = "Note 20 Borrowings\nMSEK\n2025\nTotal borrowings 32,703\nWithin 1 year 3,538\n1-5 years 29,165\n"
    x.call_llm = lambda *a, **k: {"fields": [
        {"key": "total_debt", "value": 32703, "unit": "MSEK", "period": "2025", "raw_label": "Total borrowings", "source": {"page": 1, "quote": "Total borrowings 32,703"}},
        {"key": "due_within_1_year", "value": 3538, "unit": "MSEK", "period": "2025", "raw_label": "Within 1 year", "source": {"page": 1, "quote": "Within 1 year 3,538"}},
        {"key": "due_1_to_5_years", "value": 29165, "unit": "MSEK", "period": "2025", "raw_label": "1-5 years", "source": {"page": 1, "quote": "1-5 years 29,165"}},
        {"key": "due_after_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None}]}
    out = x.extract([ericsson], [1], dm, {"fiscal_year": 2025})
    c = out["checks"][0]
    assert c["passed"] and "due_after_5_years null" in c["detail"], (c, out["warnings"])
    total = next(f for f in out["fields"] if f["key"] == "total_debt")
    assert total["confidence"] == 1.0 and "arith_ok" in total["evidence"], total
    # v011b: a printed bucket label is a known synonym — _clean_label glues digits ("Within 1 year" -> within1year), so
    # _label_known must run the synonym through the same cleaning before the prefix match (en dash, ">", space variants).
    dmf = {f["key"]: f for f in dm["fields"]}
    assert x._label_known("Within 1 year", dmf["due_within_1_year"]) and x._label_known("Inom 1 år", dmf["due_within_1_year"]) and x._label_known("< 1 år", dmf["due_within_1_year"])
    assert x._label_known("1-5 years", dmf["due_1_to_5_years"]) and x._label_known("1–5 years", dmf["due_1_to_5_years"]) and x._label_known("1–5 år", dmf["due_1_to_5_years"])
    assert x._label_known("After 5 years", dmf["due_after_5_years"]) and x._label_known("> 5 år", dmf["due_after_5_years"]) and x._label_known("Later than 5 years", dmf["due_after_5_years"])
    assert not x._label_known("1-5 years", dmf["due_after_5_years"]) and not x._label_known("Total borrowings", dmf["due_1_to_5_years"])  # prefix match, not substring
    # ... so the Ericsson fixture's buckets read perfectly at full confidence too
    out = x.extract([ericsson], [1], dm, {"fiscal_year": 2025})
    assert [(f["value"], f["confidence"]) for f in out["fields"]] == [(32703, 1.0), (3538, 1.0), (29165, 1.0), (None, 0.0)], (out["fields"], out["warnings"])
    # v012: _statement_row finds the printed bucket rows. Two defects, both in its direct helpers: the exact-match set
    # held the synonyms verbatim (v011b fixed the prefix match only), and _row_label cut the label at the first
    # space-digit, so "Within 1 year 3 538 4 100" matched as "Within" and "Inom 1 år …" as "Inom".
    assert x._row_label("Within 1 year 3 538 4 100") == "Within 1 year" and x._row_label("Inom 1 år 3 538 4 100") == "Inom 1 år" \
        and x._row_label("< 1 år 3 538 4 100") == "< 1 år"  # a number with label words after it is part of the label
    assert x._row_label("1–5 years 29 165 30 000") == "1–5 years" and x._row_label("Total sales 6, 10 155 113 161 921") == "Total sales"  # amounts still start where the label ends
    amb = ["Total borrowings 32 703 34 500", "Within 1 year 3 538 4 100", "1-5 years 29 165 30 000", "1-5 years, of which fixed rate 10 000 9 000"]
    assert x._statement_row(amb, dmf["due_1_to_5_years"], 2) == "1-5 years 29 165 30 000", amb  # the exact row, not the sub-row under it
    assert x._statement_row(amb, dmf["due_within_1_year"], 2) == "Within 1 year 3 538 4 100"
    assert x._statement_row(["Summa räntebärande skulder 32 703 34 500", "Inom 1 år 3 538 4 100", "Inom 1 år, varav konvertibler 500 400"],
                            dmf["due_within_1_year"], 2) == "Inom 1 år 3 538 4 100"
    row = x._statement_row(["Total borrowings 32 703 34 500", "1–5 years 29 165 30 000"], dmf["due_1_to_5_years"], 2)  # en dash
    assert x._row_amounts(row, 2) == [29165, 30000], row  # the "1" and "5" are the label's, not amounts
    # ... end to end: the model answers nothing, the note's own printed rows fill every bucket it prints, and the identity closes
    x.call_llm = lambda *a, **k: {"fields": []}
    out = x.extract(["Note 20 Borrowings\nMSEK\n2025 2024\nTotal borrowings 32 703 34 500\nWithin 1 year 3 538 4 100\n"
                     "1-5 years 29 165 30 000\n1-5 years, of which fixed rate 10 000 9 000\n"], [1], dm, {"fiscal_year": 2025})
    assert [(f["key"], f["value"], f["confidence"]) for f in out["fields"]] == [
        ("total_debt", 32703, 1.0), ("due_within_1_year", 3538, 1.0), ("due_1_to_5_years", 29165, 1.0),
        ("due_after_5_years", None, 0.0)], (out["fields"], out["warnings"])
    assert out["checks"][0]["passed"] and "due_after_5_years null" in out["checks"][0]["detail"], out["checks"]
    # v014: the sum-repair's own_syns (extract(), the _derived_value call) held the synonyms verbatim while the
    # rows' labels go through _clean_label -- a part row named by a digit-bearing synonym ("Within 1 year" is
    # within1year) could never join the field's own row, so the Essity-style join silently died for digit labels.
    # "cost total" (digit-free) names the quote's own row, so own_row itself is true on both sides of the change;
    # the digit synonym names the part row that must be allowed to join.
    gp14 = {"name": "gp14", "expr": "abs((revenue + special_cost) - gross_profit) <= 2", "identity": True}
    syn14 = {"name": "is", "keywords": [], "fields": [
        {"key": "revenue", "label": "Revenue", "synonyms": ["revenue"], "unit_hint": "currency_millions"},
        {"key": "special_cost", "label": "Special cost", "synonyms": ["cost total", "within 1 year"], "unit_hint": "currency_millions"},
        {"key": "gross_profit", "label": "Gross profit", "synonyms": ["gross profit"], "unit_hint": "currency_millions"}],
        "checks": [gp14]}
    p14 = ("Consolidated income statement\nSEKm\nNote\n2025\n2024\nCost of sales\nWithin 1 year items -20 -15\n"
           "Cost total -100 -90\nRevenue 500 450\nGross profit 380 345\n")
    x.call_llm = lambda *a, **k: {"fields": [
        {"key": "revenue", "value": 500, "unit": "SEKm", "period": "2025", "raw_label": "Revenue", "source": {"page": 1, "quote": "Revenue 500 450"}},
        {"key": "special_cost", "value": -100, "unit": "SEKm", "period": "2025", "raw_label": "Cost total", "source": {"page": 1, "quote": "Cost total -100 -90"}},
        {"key": "gross_profit", "value": 380, "unit": "SEKm", "period": "2025", "raw_label": "Gross profit", "source": {"page": 1, "quote": "Gross profit 380 345"}}]}
    out = x.extract([p14], [1], syn14, {"fiscal_year": 2025})
    cost = out["fields"][1]
    assert (cost["value"], cost["confidence"]) == (-120, 1.0) and "value_derived" in cost["evidence"] \
        and cost["raw_label"] == "Cost of sales" \
        and cost["source"]["quote"] == "Within 1 year items -20 -15 Cost total -100 -90", (cost, out["warnings"])
    assert out["checks"][0]["passed"], out["checks"]
    # v018: a currency/unit token glued to a footnote marker ("SEK1)") defeated both the currency-strip regex (no
    # word boundary between the letter and the digit) and the trailing note-ref regex (the closing ")" sits past
    # the digit, short of the end-of-string anchor) -- so the token survived there while the same currency word
    # without a footnote was fully stripped. Superscript digits ("SEK¹") hit the same gap once NFKC turns them
    # into a plain digit; stripped pre-NFKC so a lowercase unit ("kr") is covered too, not only the uppercase
    # 3-letter codes the note-ref regex happens to rescue by accident.
    assert x._clean_label("EPS, SEK1)") == x._clean_label("EPS, SEK") == "eps"
    assert x._clean_label("Resultat per aktie, kr¹") == x._clean_label("Resultat per aktie, kr") == "resultat per aktie"
    assert x._clean_label("Earnings per share, MSEK2)") == x._clean_label("Earnings per share, MSEK")  # symmetric before v023 too (both kept ", msek"); after v023 both fully strip it instead, see below
    # not a footnote: digits that are part of the label's own meaning are untouched, exactly as before
    assert x._clean_label("1-5 years") == x._clean_label("1–5 years") == "1-5years" and x._clean_label("Within 1 year") == "within1year"
    # v023: a magnitude prefix (k/K/m/M/b/B/t/T -- thousand/million/billion, "t/T" also covering the Swedish
    # "tusen" convention) glued directly ahead of a currency code was never stripped: \bSEK\b has no word boundary
    # between "M" and "S" in "MSEK", so "Revenue, MSEK" kept ", msek" while plain "Revenue" had nothing to strip
    # (the gap v018's Findings flagged and left open). The bare-code strip ("EPS, SEK" -> "eps") and the
    # digit-preserving cases above are unaffected -- this only extends the same currency alternation, still
    # anchored on both sides by \b, to allow one optional magnitude letter immediately before the code.
    assert x._clean_label("Revenue, MSEK") == x._clean_label("Revenue") == "revenue"
    assert x._clean_label("Revenue, TSEK") == x._clean_label("Revenue, KSEK") == x._clean_label("Revenue, kSEK") == "revenue"
    assert x._clean_label("Net sales, MEUR") == x._clean_label("Net sales") == "net sales"
    assert x._clean_label("Net sales (MSEK)") == x._clean_label("Net sales")
    # v036: the dominant failure mode v035 found across a seed of 10 -- a maturity note that prints one row
    # (often "Total") with the buckets as *columns* instead of one row per bucket. _year_column finds no year in
    # a bucket header, so it and everything built on it (_column_values, _derived_value, _between_rows) bail
    # before trying. Cloetta's own Note 21 (real text, trimmed to the header + the one row that matters):
    cloetta = ("Note 21 Borrowings\n31 Dec 2025\nSEKm\nRemaining term\n< 1 year\nRemaining term\nRemaining term\n"
               "1–2 years\nRemaining term\n2–5 years\n> 5 years Total\nLoans from credit institutions - - 1,353 - 1,353\n"
               "Capitalised transaction costs -3 -5 -4 - -12\nCommercial papers 149 - - - 149\nAccrued interest 0 - - - 0\n"
               "Lease liabilities 51 27 28 9 115\nTotal 197 22 1,377 9 1,605\n")
    x.call_llm = lambda *a, **k: {"fields": [
        {"key": "total_debt", "value": 1605, "unit": "SEKm", "period": "2025", "raw_label": "Total", "source": {"page": 1, "quote": "Total 197 22 1,377 9 1,605"}},
        {"key": "due_within_1_year", "value": 197, "unit": "SEKm", "period": "2025", "raw_label": "Total", "source": {"page": 1, "quote": "Total 197 22 1,377 9 1,605"}},
        {"key": "due_1_to_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None},  # the model tried 22+1,377=1,399, correctly dropped: no literal quote for a computed sum
        {"key": "due_after_5_years", "value": 9, "unit": "SEKm", "period": "2025", "raw_label": "Total", "source": {"page": 1, "quote": "Total 197 22 1,377 9 1,605"}}]}
    out = x.extract([cloetta], [1], dm, {"fiscal_year": 2025})
    got = {f["key"]: (f["value"], f["confidence"]) for f in out["fields"]}
    assert got == {"total_debt": (1605, 1.0), "due_within_1_year": (197, 1.0), "due_1_to_5_years": (1399, 1.0), "due_after_5_years": (9, 1.0)}, (got, out["warnings"])  # the bare "Total" row closes: label_known for all four (identity_all_columns), no longer 0.9
    assert all("identity_all_columns" in f["evidence"] and "label_known" in f["evidence"] for f in out["fields"]), [f["evidence"] for f in out["fields"]]
    assert out["checks"][0]["passed"], out["checks"]  # null-filling the one bucket the model couldn't quote closes the identity, lifting the other three off their failed-check 0.5 cap too
    d15 = next(f for f in out["fields"] if f["key"] == "due_1_to_5_years")
    assert d15["source"]["quote"] == "Total 197 22 1,377 9 1,605" and "value_derived" in d15["evidence"] and "value_in_quote" not in d15["evidence"]  # a sum of two columns has no literal quote of its own
    # disagree: the model attributes the "1-2 years" column (22) to due_within_1_year instead of its own "< 1 year" column -- the page's column order wins, same as every other repair in this file
    x.call_llm = lambda *a, **k: {"fields": [
        {"key": "total_debt", "value": 1605, "unit": "SEKm", "period": "2025", "raw_label": "Total", "source": {"page": 1, "quote": "Total 197 22 1,377 9 1,605"}},
        {"key": "due_within_1_year", "value": 22, "unit": "SEKm", "period": "2025", "raw_label": "Total", "source": {"page": 1, "quote": "Total 197 22 1,377 9 1,605"}},
        {"key": "due_1_to_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None},
        {"key": "due_after_5_years", "value": 9, "unit": "SEKm", "period": "2025", "raw_label": "Total", "source": {"page": 1, "quote": "Total 197 22 1,377 9 1,605"}}]}
    out = x.extract([cloetta], [1], dm, {"fiscal_year": 2025})
    within = next(f for f in out["fields"] if f["key"] == "due_within_1_year")
    assert within["value"] == 197 and any("22 disagrees with the maturity table" in w for w in out["warnings"]), (within, out["warnings"])
    # a maturity table whose columns are calendar years rather than named buckets ("2026 2027 2028 Later"), per HANDOFF's "expect these": folded to <1y / 1-5y / >5y off the report's own fiscal year, not the model's
    yr = "Note 20 Borrowings\nMSEK\n2026 2027 2028 Later Total\nTotal borrowings 412 1,286 933 47 2,678\n"
    x.call_llm = lambda *a, **k: {"fields": [
        {"key": "total_debt", "value": 2678, "unit": "MSEK", "period": "2025", "raw_label": "Total borrowings", "source": {"page": 1, "quote": "Total borrowings 412 1,286 933 47 2,678"}},
        {"key": "due_within_1_year", "value": None, "unit": None, "period": None, "raw_label": None, "source": None},
        {"key": "due_1_to_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None},
        {"key": "due_after_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None}]}
    out = x.extract([yr], [1], dm, {"fiscal_year": 2025})
    got = {f["key"]: f["value"] for f in out["fields"]}
    assert got == {"total_debt": 2678, "due_within_1_year": 412, "due_1_to_5_years": 2219, "due_after_5_years": 47} and out["checks"][0]["passed"], (got, out["warnings"])  # 2027 + 2028 = 1,286 + 933
    # not every bucket-as-columns table is readable this way, and guessing wrong is worse than not fixing it -- three
    # real shapes from the same seed that stay exactly as they were (v035's own "not fixed" verdicts, unchanged):
    # XANO: the header wraps across two of _page_rows' lines ("Summa" / "inom 1 år"), scrambling column order out of
    # print order; the safety valve is the row's own column count (8) never matching however many bucket keys are found
    xano = ("FINANSIELLA SKULDER Förfallotid\nPER 2025-12-31 –30 dgr 31–90 dgr 91–360 dgr Summa Mellan 1 Mellan 3 Efter 5 år Totalt\n"
            "inom 1 år och 3 år och 5 år\nSumma räntebärande skulder 4 090 8 680 38 305 51 075 768 885 33 784 50 778 904 522\n")
    x.call_llm = lambda *a, **k: {"fields": [
        {"key": "total_debt", "value": 904522, "unit": "TSEK", "period": "2025", "raw_label": "Summa räntebärande skulder", "source": {"page": 1, "quote": "Summa räntebärande skulder 4 090 8 680 38 305 51 075 768 885 33 784 50 778 904 522"}},
        {"key": "due_within_1_year", "value": 51075, "unit": "TSEK", "period": "2025", "raw_label": "inom 1 år och 3 år och 5 år: Summa räntebärande skulder", "source": {"page": 1, "quote": "Summa räntebärande skulder 4 090 8 680 38 305 51 075 768 885 33 784 50 778 904 522"}},
        {"key": "due_1_to_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None},
        {"key": "due_after_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None}]}
    out = x.extract([xano], [1], dm, {"fiscal_year": 2025})
    got = {f["key"]: (f["value"], f["confidence"]) for f in out["fields"]}
    # v044: "Efter 5 år" is printed whole on the header line (only "Summa"/"Mellan 1"/"Mellan 3" wrap away), so
    # due_after_5_years is a bucket the table demonstrably has, just one the safety valve correctly declined to
    # read a number for -- the check reads missing: due_after_5_years, not a false failure against a fabricated
    # due_after_5_years=0, and total_debt/due_within_1_year (independently quote- and label-verified) are no
    # longer capped for a sibling bucket's gap. due_1_to_5_years has no matching synonym anywhere on the page
    # (its own header words wrapped into "inom 1 år och 3 år och 5 år", which reads as digit-glued "1"/"3"/"5"
    # tokens, none of which forms "mellan...och..." or a "1-5"/"1-2"/"2-5" run) and stays a confirmed 0.
    assert got == {"total_debt": (904522, 1.0), "due_within_1_year": (51075, 1.0), "due_1_to_5_years": (None, 0.0), "due_after_5_years": (None, 0.0)} \
        and out["checks"][0]["detail"] == "missing: due_after_5_years" and not out["warnings"], (got, out["checks"], out["warnings"])
    # Ework: the table's own finer split (<1 month / 1-3 / 3-12 months) never says "year", and an unexplained
    # 6th numeric column means the header's bucket-key hits never reach the row's own column count (8) either
    ework = ("kSEK Due < 1 month 1-3 months 3-12 months 1-5 years > 5 years Total undis- Carrying\ncounted value amount\n2025\n"
             "Short-term interest-bearing\nliabilities* – 153,761 971 1,677 – – 156,410 156,410\n")
    x.call_llm = lambda *a, **k: {"fields": [
        {"key": "total_debt", "value": 156410, "unit": "kSEK", "period": "2025", "raw_label": "Interest-bearing liabilities", "source": {"page": 1, "quote": "Short-term interest-bearing liabilities* – 153,761 971 1,677 – – 156,410 156,410"}},
        {"key": "due_within_1_year", "value": None, "unit": None, "period": None, "raw_label": None, "source": None},
        {"key": "due_1_to_5_years", "value": 1677, "unit": "kSEK", "period": "2025", "raw_label": "Short-term interest-bearing liabilities*", "source": {"page": 1, "quote": "Short-term interest-bearing liabilities* – 153,761 971 1,677 – – 156,410 156,410"}},
        {"key": "due_after_5_years", "value": 0, "unit": "kSEK", "period": "2025", "raw_label": "Short-term interest-bearing liabilities*", "source": {"page": 1, "quote": "Short-term interest-bearing liabilities* – 153,761 971 1,677 – – 156,410 156,410"}}]}
    out = x.extract([ework], [1], dm, {"fiscal_year": 2025})
    got = {f["key"]: f["value"] for f in out["fields"]}
    assert got == {"total_debt": 156410, "due_within_1_year": None, "due_1_to_5_years": 1677, "due_after_5_years": None}, (got, out["warnings"])  # unchanged from the model's own (repair-dropped) answer
    # Bergman & Beving: the only row that could be "the total" is "Total financial liabilities", which is not a
    # total_debt synonym (no "interest-bearing"/"borrowings") and not bare Total/Summa/Totalt either -- it also
    # includes non-interest-bearing accounts payable, so treating it as one would have been wrong, not just unproven
    assert x._bucket_total_row(x._page_rows("Total financial liabilities 3,111 3,380 778 454 2,116 32\n"), dmf["total_debt"]) == []
    # v040: Acast (seed 2, real Note text) -- two stacked calendar-year maturity tables (this year, then a
    # prior-year comparative), each ending in a bare "Total" row. The current year's own Total row is glued to
    # trailing sidebar prose by the PDF's column layout ("...19,250 approximation of fair value." then "2024"
    # on its own line merges onto the same row), so it never reads as >=2 amounts and drops out as a candidate;
    # only the prior year's row is left. Its own column count (6) still happens to match the *current* year's
    # own header (found first, scanning top-down over the whole page) closely enough to pass the safety valve,
    # producing a due_within_1_year bigger than its own total_debt. A bucket can never exceed the total it is
    # part of, in any report: reject the whole row's column reading when one does, rather than "fix" one field
    # out of a header that was never really aligned to that row to begin with.
    acast = ("SEK thousand Carrying amount 2026 2027 2028 2029 After 2029 The carrying amounts of the Group's finance lease\n"
             "Total 690,785 592,113 34,980 33,467 21,826 19,250 approximation of fair value.\n2024\n"
             "SEK thousand Carrying amount 2025 2026 2027 2028 After 2028\n"
             "Total 555,370 446,990 32,692 32,404 32,404 45,414\n")
    leases = "SEK thousand 31.12.2025 31.12.2024\nLease liabilities\nCurrent 32,052 23,443\nNon-current 103,330 117,709\nTotal lease Liability 135,382 141,152\n"
    x.call_llm = lambda *a, **k: {"fields": [
        {"key": "total_debt", "value": 135382, "unit": "SEK thousand", "period": "2025", "raw_label": "Total lease Liability",
         "source": {"page": 1, "quote": "Total lease Liability 135,382 141,152"}},
        {"key": "due_within_1_year", "value": None, "unit": None, "period": None, "raw_label": None, "source": None},
        {"key": "due_1_to_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None},
        {"key": "due_after_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None}]}
    out = x.extract([leases, acast], [1, 2], dm, {"fiscal_year": 2025})
    got = {f["key"]: (f["value"], f["confidence"]) for f in out["fields"]}
    assert got == {"total_debt": (135382, 0.9), "due_within_1_year": (None, 0.0), "due_1_to_5_years": (None, 0.0), "due_after_5_years": (None, 0.0)}, (got, out["warnings"])
    assert any("exceed its own total 45414" in w for w in out["warnings"]), out["warnings"]  # the column reading is rejected outright, not one field silently "fixed"
    # v040: Nelly Group (seed 2, real Note text) -- the model correctly reads due_within_1_year from its own
    # "Kortfristiga" (current) row, but the schema has no bare "kortfristiga" synonym (only the compound
    # "kortfristiga räntebärande skulder"), so the row was never recognised as due_within_1_year's own -- and
    # the neighbour-sum repair, left unrestricted, folded in "Långfristiga" (non-current) too, turning a
    # correct current-portion read into the current+non-current total. A row the model already read and
    # labelled faithfully is its own row whether or not the schema happens to recognise that label as a
    # synonym *yet* -- see own_row in extract().
    nelly = ("Leasingskulder (Miljoner kronor) 2025 2024\nKortfristiga 36,2 35,7\nLångfristiga 250,9 251,5\n"
             "Leasingskulder som ingår i\n\xadrapporten över finansiell ställning 287,1 287,2\n")
    x.call_llm = lambda *a, **k: {"fields": [
        {"key": "total_debt", "value": 287.1, "unit": "Miljoner kronor", "period": "2025",
         "raw_label": "Leasingskulder som ingår i rapporten över finansiell ställning",
         "source": {"page": 1, "quote": "\xadrapporten över finansiell ställning 287,1 287,2"}},
        {"key": "due_within_1_year", "value": 36.2, "unit": "Miljoner kronor", "period": "2025", "raw_label": "Kortfristiga",
         "source": {"page": 1, "quote": "Kortfristiga 36,2 35,7"}},
        {"key": "due_1_to_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None},
        {"key": "due_after_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None}]}
    out = x.extract([nelly], [1], dm, {"fiscal_year": 2025})
    got = {f["key"]: (f["value"], f["confidence"]) for f in out["fields"]}
    assert got == {"total_debt": (287.1, 0.5), "due_within_1_year": (36.2, 0.5), "due_1_to_5_years": (None, 0.0), "due_after_5_years": (None, 0.0)}, (got, out["warnings"])
    assert not out["warnings"], out["warnings"]  # no repair fires at all: the model's own, correctly-labelled row stands unmolested
    # v040: a maturity table's own Total row is not immune to the report's usual "-" nil convention either --
    # a dash inside the bucket columns (not just in the instrument rows above the total, which Cloetta's own
    # test above already covers) must not desync col_keys from amounts, trailing or mid-row. v078 refines the
    # reading: a dash in the selected row's OWN column for a bucket is the report's explicit 0 for that window
    # (evidence printed_nil) once the row's own arithmetic closes -- still never a fabricated number, and a
    # bucket the table prints no column for at all stays null.
    header = "Note 21 Borrowings\n31 Dec 2025\nSEKm\nRemaining term\n< 1 year\nRemaining term\nRemaining term\n1–2 years\nRemaining term\n2–5 years\n> 5 years Total\n"
    x.call_llm = lambda *a, **k: {"fields": [{"key": k, "value": None, "unit": None, "period": None, "raw_label": None, "source": None} for k in dmf]}
    out = x.extract([header + "Total 197 22 1,377 - 1,596\n"], [1], dm, {"fiscal_year": 2025})  # trailing dash: the >5y column prints nil
    got = {f["key"]: (f["value"], f["confidence"]) for f in out["fields"]}
    a5 = next(f for f in out["fields"] if f["key"] == "due_after_5_years")
    # v078: the dash column is the report's explicit 0 (printed_nil @0.5), not a null the check must guess around
    assert got == {"total_debt": (1596, 1.0), "due_within_1_year": (197, 1.0), "due_1_to_5_years": (1399, 1.0), "due_after_5_years": (0, 0.5)}, (got, out["warnings"])  # the closing row identifies its bare "Total" label (identity_all_columns)
    assert a5["evidence"] == ["quote_on_page", "printed_nil", "identity_all_columns", "arith_ok", "label_known", "period_ok", "page_is_statement", "unit_ok"] \
        and a5["source"]["quote"] == "Total 197 22 1,377 - 1,596", a5
    assert out["checks"][0]["passed"], out["checks"]
    out = x.extract([header + "Total 197 - 1,377 9 1,583\n"], [1], dm, {"fiscal_year": 2025})  # mid-row dash: no 1-2y bucket, 2-5y and >5y still line up
    got = {f["key"]: (f["value"], f["confidence"]) for f in out["fields"]}
    assert got == {"total_debt": (1583, 1.0), "due_within_1_year": (197, 1.0), "due_1_to_5_years": (1377, 1.0), "due_after_5_years": (9, 1.0)}, (got, out["warnings"])
    assert out["checks"][0]["passed"], out["checks"]
    # v134: a model-answered 0 whose citation is a row the report prints as dashes is the report's own printed
    # nil (v078's convention, one reading further: the model, not the bucket-column reader, did the reading).
    # Linc p.100 "Räntebärande skulder – –" (v090's standing gap) and Rejlers' "1-2 years - -" block (v097
    # finding 4) both died -- "computed, not read" / 0@0.25 "quote not found" -- because quote_on_page can
    # never verify a quote without a number token. Two conditions on the rows the quote matches, no arithmetic
    # gate (an all-dash row prints no total to close against, and nothing on it can be summed into a 0): a
    # known label, and a dash where the fiscal year reads. Refusals: an unknown label, a figure in the
    # fiscal-year column, no dash row behind the quote, a dashed component row beside a real debt row.
    linc = ("Substansvärde 2025-12-31 2024-12-31\nTotala tillgångar 4 442 632 4 839 415\nLångfristiga skulder – –\n"
            "Kortfristiga skulder -11 620 -8 757\nSubstansvärde 4 431 012 4 830 658\nNettoskuld 2025-12-31 2024-12-31\n"
            "Räntebärande skulder – –\nLikvida medel -283 982 -549 657\nNettoskuld -283 982 -549 657\n")
    x.call_llm = lambda *a, **k: {"fields": [
        {"key": "total_debt", "value": 0, "unit": "Tkr", "period": "2025", "raw_label": "Räntebärande skulder",
         "source": {"page": 1, "quote": "Räntebärande skulder – –"}}] + [
        {"key": k, "value": None, "unit": None, "period": None, "raw_label": None, "source": None}
        for k in ("due_within_1_year", "due_1_to_5_years", "due_after_5_years")]}
    out = x.extract([linc], [1], dm, {"fiscal_year": 2025})
    total = next(f for f in out["fields"] if f["key"] == "total_debt")
    assert (total["value"], total["confidence"]) == (0, 0.5) \
        and total["evidence"] == ["quote_on_page", "printed_nil", "arith_ok", "label_known", "period_ok", "page_is_statement", "unit_ok"] \
        and total["source"] == {"page": 1, "quote": "Räntebärande skulder – –"}, (total, out["warnings"])  # label_known: the row synonym scores; the nil keeps the cap
    assert any("total_debt: 0 kept -- 'Räntebärande skulder – –' prints a dash in the fiscal-year column under a known label (printed nil)" in w
               for w in out["warnings"]), out["warnings"]
    # ... the bare row-synonym label ("räntebärande skulder", no summa/totalt prefix) is what the total field is
    # recognised by here, the same vocabulary _bucket_total_row recognises the debt row by
    assert x._model_zero_on_dash_row({"value": 0, "source": {"page": 1, "quote": "Räntebärande skulder – –"}},
                                     dmf["total_debt"], [linc], [1], 2025) == (1, "Räntebärande skulder – –")
    assert x._model_zero_on_dash_row({"value": 0, "source": {"page": 1, "quote": "Avstämningspost – –"}},
                                     dmf["total_debt"], [linc.replace("Räntebärande", "Avstämningspost")], [1], 2025) is None  # an unknown label: any dash row is not the field
    assert x._model_zero_on_dash_row({"value": 0, "source": {"page": 1, "quote": "Räntebärande skulder - 100"}},
                                     dmf["total_debt"], ["Nettoskuld 2025-12-31 2024-12-31\nRäntebärande skulder - 100\nLikvida medel -50 -60\n"],
                                     [1], 2025) == (1, "Räntebärande skulder - 100")  # the fiscal-year column is the dash; the 100 prints in the prior-year one
    # Rejlers (v097 finding 4): the dash bucket rows quoted as one block; the page also prints "0.8 per cent",
    # which is exactly what kept the answer alive unproven at 0.25 -- the printed-nil rule outranks that fallback.
    rejlers = ("NOTE 24\nLiabilities to credit institutions\nThe Group has an overdraft facility with a limit of SEK 250\n"
               "million (250), of which SEK 27.8 million (12.1) is utilised. The interest margin is fixed at 0.8 per\n"
               "cent with an addition for STIBOR 3 months.\n"
               "Maturity analysis, liabilities to credit institutions, SEK\nmillion 2025 2024\n"
               "Within one year 400.3 349.5\n1-2 years - -\n2–3 years - -\n3-4 years - -\n4-5 years - -\n")
    x.call_llm = lambda *a, **k: {"fields": [
        {"key": "total_debt", "value": None, "unit": None, "period": None, "raw_label": None, "source": None},
        {"key": "due_within_1_year", "value": 400.3, "unit": "MSEK", "period": "2025", "raw_label": "Within one year",
         "source": {"page": 1, "quote": "Within one year 400.3 349.5"}},
        {"key": "due_1_to_5_years", "value": 0, "unit": "MSEK", "period": "2025", "raw_label": "1-2 years; 2–3 years; 3-4 years; 4-5 years",
         "source": {"page": 1, "quote": "1-2 years - -\n2–3 years - -\n3-4 years - -\n4-5 years - -"}},
        {"key": "due_after_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None}]}
    out = x.extract([rejlers], [1], dm, {"fiscal_year": 2025})
    b15 = next(f for f in out["fields"] if f["key"] == "due_1_to_5_years")
    assert (b15["value"], b15["confidence"]) == (0, 0.5) \
        and b15["evidence"] == ["quote_on_page", "printed_nil", "arith_ok", "label_known", "period_ok", "page_is_statement", "unit_ok"] \
        and b15["source"]["quote"] == "1-2 years - -", (b15, out["warnings"])  # the first matched dash row is the provenance
    assert any("due_1_to_5_years: 0 kept -- '1-2 years - -' prints a dash in the fiscal-year column under a known label (printed nil)" in w
               for w in out["warnings"]), out["warnings"]
    # refusals, end to end: all three stay null, dropped as computed, not read
    for page, key, quote in [("Nettoskuld 2025-12-31 2024-12-31\nAvstämningspost – –\n", "total_debt", "Avstämningspost – –"),
                             ("NOTE 24\nLiabilities to credit institutions\nMaturity analysis, SEK million\nBetween 1 and 5 years 250 200\n",
                              "due_1_to_5_years", "Between 1 and 5 years 250 200"),
                             ("Note 20 Borrowings\nMSEK\n2025 2024\nWithin 1 year 3,538\n1-5 years 29,165\n", "total_debt", "Total borrowings – –"),
                             ("Borrowings note\n2025 2024\nLease liabilities - -\nBank loans 500 400\n", "total_debt", "Lease liabilities - -")]:
        x.call_llm = lambda *a, _p=page, _k=key, _q=quote, **k: {"fields": [
            {"key": _k, "value": 0, "unit": "MSEK", "period": "2025", "raw_label": _q,
             "source": {"page": 1, "quote": _q}}]}
        out = x.extract([page], [1], dm, {"fiscal_year": 2025})
        bad = next(f for f in out["fields"] if f["key"] == key)
        assert bad["value"] is None and bad["confidence"] == 0.0, (bad, out["warnings"])
        assert any("dropped as computed, not read" in w for w in out["warnings"]), (page, out["warnings"])
        assert not any("printed nil" in w for w in out["warnings"]), (page, out["warnings"])
    # v044: null_as_zero, refined -- v041 finding 3 (MedCap, Ependion, same seed): a bucket the model failed to
    # extract reads identically to a bucket the report never prints -- both are null -- but only the second one
    # is really a 0. Before defaulting a null bucket to 0, _check now looks for its own synonym label (the same
    # normalize_ws digit-gluing v012/v014 already use for row labels) on the pages the check's other, real
    # operands were sourced from, or the candidate pages; found -> leave it out of the substitution, so the check
    # reads "missing: <field>" instead of a false failure; not found -> 0, exactly as before. v058: the search
    # runs over the operands' OWN table -- the printed rows their quotes anchor to, from the nearest year-header
    # row above the topmost down to the bottommost -- whenever at least one anchor exists, and over those whole
    # pages otherwise; a bucket word printed by another table on a candidate page no longer blocks the 0.
    # texts/pages/fields/schema are new, optional _check args: every call above this one omits them, so the
    # pre-v044 unconditional zero-fill is unchanged for all of them (the full suite above this block is
    # unmodified except the one XANO case, which hits this exact mechanism too -- see its own updated comment).
    c = x._check(dmc, {"total_debt": 718.7, "due_1_to_5_years": 616.5}, texts=["Totalt 718,7\n1 – 5 år 616,5\n",
                 "MSEK <1 år 1 – 2 år 2 – 4 år 4 – 5 år >5år\n"], pages=[1, 2],
                 fields=[{"key": "total_debt", "source": {"page": 1}}, {"key": "due_1_to_5_years", "source": {"page": 1}}], schema=dm)
    assert not c["passed"] and c["detail"] == "missing: due_within_1_year", c  # "<1 år" sits on the candidate page: not a confirmed 0
    c = x._check(dmc, {"due_within_1_year": 3538, "due_1_to_5_years": 29165, "total_debt": 32703},
                 texts=["Total borrowings 32,703\nWithin 1 year 3,538\n1-5 years 29,165\n"], pages=[1],
                 fields=[{"key": "due_within_1_year", "source": {"page": 1}}, {"key": "due_1_to_5_years", "source": {"page": 1}}], schema=dm)
    assert c["passed"] and "due_after_5_years null" in c["detail"], c  # no "after 5 years"/"efter 5 år" anywhere: still a real 0
    # end to end, MedCap's own page text (seed3-kb, pp.101-102, relabelled 1-2 here): the model reads total_debt
    # and due_1_to_5_years verbatim off the carrying-amount table (p.101); due_within_1_year is null (its true
    # value is the sum 54.3+48.0=102.3, printed only as two separate rows) and due_after_5_years is null too
    # (this table really has no >5y row). v044 found due_within_1_year's "<1 år" on the contractual table two
    # pages later and kept the check at missing; v052's derivation then filled 102.3 from the rows above the
    # "1 – 5 år" row. v058 scopes the present-label search to the operands' own table -- the contractual
    # table's "<1 år" decides nothing any more -- and the ruling adds the month rows to the schema itself
    # ("6 månader eller mindre", "6 – 12 månader"): the guard keeps due_within_1_year out on its own table's
    # wording, the check reads missing, the derivation fills 102.3, and with due_after_5_years a confirmed 0
    # the identity closes -- passed, where v044/v052 stalled at "missing: due_after_5_years".
    medcap101 = ("Förfallotidpunkt för upplåning Koncernen Moderbolaget\nMSEK 2025-12-31 2024-12-31 2025-12-31 2024-12-31\n"
                 "6 månader eller mindre 54,3 41,8 – –\n6 – 12 månader 48,0 19,0 – –\n1 – 5 år 616,5 316,7 – –\nTotalt 718,7 377,6 – –\n")
    medcap102 = ("Koncernen 2025-12-31\nMSEK <1 år 1 – 2 år 2 – 4 år 4 – 5 år >5år\nSkulder till kreditinstitut 59,3 97,5 59,1 – –\n"
                 "Leasingskulder 50,6 100,8 66,4 30,6 142,9\nÖvriga långfristiga skulder 80,9 187,1 – – –\n"
                 "Leverantörs- och övriga skulder 241,0 – – – –\nTotalt 431,8 385,4 125,5 30,6 142,9\n")
    x.call_llm = lambda *a, **k: {"fields": [
        {"key": "total_debt", "value": 718.7, "unit": "MSEK", "period": "2025", "raw_label": "Totalt", "source": {"page": 1, "quote": "Totalt 718,7 377,6 – –"}},
        {"key": "due_within_1_year", "value": None, "unit": None, "period": None, "raw_label": None, "source": None},
        {"key": "due_1_to_5_years", "value": 616.5, "unit": "MSEK", "period": "2025", "raw_label": "1 – 5 år", "source": {"page": 1, "quote": "1 – 5 år 616,5 316,7 – –"}},
        {"key": "due_after_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None}]}
    out = x.extract([medcap101, medcap102], [1, 2], dm, {"fiscal_year": 2025})
    assert out["checks"][0]["passed"] and "abs((102.3 + 616.5 + 0 (due_after_5_years null)) - 718.7) <= 2" in out["checks"][0]["detail"], out["checks"]
    within = next(f for f in out["fields"] if f["key"] == "due_within_1_year")
    assert within["value"] == 102.3 and "value_derived" in within["evidence"], (within, out["warnings"])
    total = next(f for f in out["fields"] if f["key"] == "total_debt")
    bucket15 = next(f for f in out["fields"] if f["key"] == "due_1_to_5_years")
    assert total["confidence"] == 0.9 and "arith_ok" in total["evidence"], total
    assert bucket15["confidence"] == 1.0 and "arith_ok" in bucket15["evidence"], bucket15
    # Ependion (same seed, v041 finding 3's other example): stays failed, correctly -- its bucket table (p.155)
    # never prints a digit-form boundary at all ("Between 1 and 2 years" / "Between 2 and 3 years", spelled out,
    # scrambled by a sidebar TOC column glued onto the same lines by parse.py), and the schema has no spelled-out
    # synonym for either bucket (docs/acrylic/evidence/v041.md's own finding: "the spelled-out ... wording
    # doesn't exist in the schema at all yet regardless"). A schema-vocabulary gap, out of this lane's territory
    # (schema untouched) -- not the null-vs-missing conflation this lane fixes. Recorded so a future schema
    # change adding that wording is the one that should flip this case, not a regression in this mechanism.
    ependion155 = ("Contracted terms\nSEK 000 31 Dec. 2025 31 Dec. 2024\nBetween Between\nWithin 12 1 and 2 and\n"
                   "months 2 years 3 years Total\nBorrowing 167,546 35,636 380,348 583,531\n")
    x.call_llm = lambda *a, **k: {"fields": [
        {"key": "total_debt", "value": 583531, "unit": "SEK 000", "period": "2025", "raw_label": "Borrowing", "source": {"page": 1, "quote": "Borrowing 167,546 35,636 380,348 583,531"}},
        {"key": "due_within_1_year", "value": 167546, "unit": "SEK 000", "period": "2025", "raw_label": "Borrowing", "source": {"page": 1, "quote": "Borrowing 167,546 35,636 380,348 583,531"}},
        {"key": "due_1_to_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None},
        {"key": "due_after_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None}]}
    out = x.extract([ependion155], [1], dm, {"fiscal_year": 2025})
    assert not out["checks"][0]["passed"] and not out["checks"][0]["detail"].startswith("missing:"), out["checks"]
    within = next(f for f in out["fields"] if f["key"] == "due_within_1_year")
    assert within["confidence"] == 0.5, within  # still capped: neither sibling bucket's label is anywhere on the page
    # v060: Ependion's *current* page text (post-v049 parse.py fix; not the sidebar-scrambled snapshot the case
    # above was recorded against, and a still-open gap this one doesn't touch). "Within 12 months" and the day/
    # month-range header wording the schema's synonyms deliberately exclude from row-label matching (v046) now
    # live in each bucket field's own header_synonyms, read only by _bucket_synonym_hits/_bucket_header, never
    # by _label_known -- label_regression stays 0 because of exactly this split. The bucket row is labelled
    # "Borrowing" (a total_debt fallback synonym, not "total"/"summa"), so _bucket_total_row's debt-row fallback
    # picks it -- and not the *other* "- Borrowing 583,531 557,173" row two lines above (a leading "- " that
    # breaks _label_known's prefix match) or "Accounts payable-trade" (not a debt synonym at all).
    ependion_real = ("Financial liabilities\n- Borrowing 583,531 557,173\n- Accounts payable-trade 164,155 154,411\n"
                      "Contracted terms\nSEK 000 Within 12 months\nBetween 1 and 2 years\nBetween 2 and 3 years Total\n"
                      "Borrowing 167,546 35,636 380,348 583,531\n"
                      "Accounts payable-trade 164,155 164,155 331,701 35,636 380,348 747,686\n")
    x.call_llm = lambda *a, **k: {"fields": [
        {"key": "total_debt", "value": None, "unit": None, "period": None, "raw_label": None, "source": None},
        {"key": "due_within_1_year", "value": None, "unit": None, "period": None, "raw_label": None, "source": None},
        {"key": "due_1_to_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None},
        {"key": "due_after_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None}]}
    out = x.extract([ependion_real], [1], dm, {"fiscal_year": 2025})
    got = {f["key"]: (f["value"], f["confidence"]) for f in out["fields"]}
    assert got == {"total_debt": (583531, 0.95), "due_within_1_year": (167546, 0.95), "due_1_to_5_years": (415984, 0.95),
                    "due_after_5_years": (None, 0.0)}, (got, out["warnings"])  # "Borrowing" closes on its own total: label_known
    assert out["checks"][0]["passed"], out["checks"]  # due_after_5_years null_as_zero: this table prints no such bucket
    total = next(f for f in out["fields"] if f["key"] == "total_debt")
    assert total["raw_label"] == "Borrowing" and total["source"]["quote"] == "Borrowing 167,546 35,636 380,348 583,531", total  # not the "- Borrowing" row, not Accounts payable-trade
    # v060: Boozt's bucket row is labelled "Lease liabilities" (also a debt-row fallback synonym: the report's own
    # total_debt *is* its lease liabilities, no other interest-bearing debt) with a header split across two of
    # _page_rows' own rows ("Maturity within" / "3 months", a bare word-wrap, not the header-transposition XANO/
    # Ework hit below) that header_synonyms now reads correctly -- both this row and the page's own bare "Total"
    # rows (whole-table sums, wrong scope regardless) read a 6-column header matching their own column count.
    # v064: due_within_1_year (=26+78, not printed) used to stay null here -- _row_amounts read the row's own
    # adjacent 2- and 3-digit columns ("78 273") as the single Swedish-grouped number 78273 (the same ambiguity
    # as "6, 10 155" elsewhere in this file), one column short of the header's six, so the safety valve declined
    # rather than guess. Fixed in _row_amounts itself (ncols known, exactly one column short, exactly one
    # "NN NNN" token that splits to land exactly on ncols -- see its docstring). The page's own bare "Total 2,358
    # 1,928 92 273 63 0" row has the identical shape ("92 273" is just as splittable) but must NOT also start
    # reading as a candidate: it sums every instrument on the page (accounts payables and other liabilities
    # included, not just the lease), so a naive fix here would have _fill_bucket_columns's own hits-before-
    # candidates row order (out of this lane's six-function territory, untouched) prefer it over the correct
    # row and corrupt total_debt to 2358. _row_amounts's extra guard -- only split when the ordinary reading
    # already has a nil in it -- is what keeps this row declined: Boozt's own liability-type rows print "-" for
    # a bucket they have nothing in (sparse, one instrument each), but a whole-table Total is a computed sum
    # across every instrument and is essentially never nil anywhere, so it stays a column short and gets
    # skipped, exactly as before.
    boozt = ("Total borrowing\nMaturity within\n3 months\nMaturity within three to twelve months\n"
             "Maturity within one to five years\nMaturity within five to nie years\nMaturity after nine years\n"
             "Maturity structure of borrowing Dec 31,2024\n"
             "Liabilities to credit institutions 380 - - 380 - -\nLease liabilities 499 29 85 251 135 -\n"
             "Accounts payables 1,235 1,225 10 - - -\nOther liabilities 531 527 4 - - -\n"
             "Total 2,645 1,781 99 631 135 0\n"
             "Maturity structure of borrowing Dec 31, 2025\n"
             "Liabilities to credit institutions 0 - - 0 - -\nLease liabilities 441 26 78 273 63 -\n"
             "Accounts payables 1,384 1,374 10 - - -\nOther liabilities 533 528 4 - - -\n"
             "Total 2,358 1,928 92 273 63 0\n")
    bucket_sfs = {k: dmf[k] for k in ("due_within_1_year", "due_1_to_5_years", "due_after_5_years")}
    boozt_rows = x._page_rows(boozt)
    assert x._bucket_header(boozt_rows, 15, bucket_sfs, 2025) == \
        ["total", "due_within_1_year", "due_within_1_year", "due_1_to_5_years", "due_after_5_years", "due_after_5_years"]  # the header itself now reads right
    assert x._row_amounts(boozt_rows[15], 6, nil=None) == [441, 26, 78, 273, 63, None]  # v064: ncols pins the split of "78 273" into 78 + 273
    assert boozt_rows[18] == "Total 2,358 1,928 92 273 63 0" and len(x._row_amounts(boozt_rows[18], 6, nil=None)) != 6  # v064: same "NN NNN" shape ("92 273"), stays declined -- no nil in its own ordinary reading
    x.call_llm = lambda *a, **k: {"fields": [
        {"key": "total_debt", "value": 441, "unit": "SEK million", "period": "2025", "raw_label": "Lease liabilities", "source": {"page": 1, "quote": "Lease liabilities 441 26 78 273 63 -"}},
        {"key": "due_within_1_year", "value": None, "unit": None, "period": None, "raw_label": None, "source": None},
        {"key": "due_1_to_5_years", "value": 273, "unit": "SEK million", "period": "2025", "raw_label": "Lease liabilities", "source": {"page": 1, "quote": "Lease liabilities 441 26 78 273 63 -"}},
        {"key": "due_after_5_years", "value": 63, "unit": "SEK million", "period": "2025", "raw_label": "Lease liabilities", "source": {"page": 1, "quote": "Lease liabilities 441 26 78 273 63 -"}}]}
    out = x.extract([boozt], [1], dm, {"fiscal_year": 2025})
    got = {f["key"]: (f["value"], f["confidence"]) for f in out["fields"]}
    # v064: due_within_1_year now fills to 26+78=104 (the bucket-columns cross-column sum, _row_amounts's own
    # limit fixed) -- and, closing the loop, maturity_sums_to_total now passes (104+273+63-441 = -1, within the
    # check's own +-2), which lifts total_debt/due_1_to_5_years/due_after_5_years off the 0.50 "contradicts its
    # neighbours" cap they were pinned to while due_within_1_year was null (score_field, untouched by this lane).
    # The model's own three answers quote that same closing row: they earn label_known by its arithmetic too
    # (identity_all_columns), not only the column-read one -- "Lease liabilities" is no synonym of total debt.
    assert got == {"total_debt": (441, 1.0), "due_within_1_year": (104, 1.0), "due_1_to_5_years": (273, 1.0), "due_after_5_years": (63, 1.0)}, (got, out["warnings"])
    assert all("identity_all_columns" in f["evidence"] for f in out["fields"]), [f["evidence"] for f in out["fields"]]
    assert out["warnings"] == ["due_within_1_year: model returned null; 104 read from 'Lease liabilities 441 26 78 273 63 -' by its column order"], out["warnings"]
    # v083: the debt-row fallback word list moved from a private extract.py constant (_DEBT_ROW_SYNONYMS) to
    # schema's total_debt.row_synonyms -- dmf["total_debt"] below is read straight from backend/schemas/
    # debt_maturity.json, so this proves the schema key is what _bucket_total_row now reads, not a stale copy.
    # Tången Industrikapital (seed8-kb/v082 Finding 1, real p.62 text): the maturity table's own row is labelled
    # "Lån Kreditinstitut" ("Loan, Credit institutions") -- not a total_debt synonym, not bare Total/Summa, and
    # not in the pre-v083 fallback list either, so the row was invisible to the fallback and never reached
    # _bucket_header at all. "lån kreditinstitut" is added to row_synonyms as the exact two-word phrase actually
    # printed (no "till"/"to"): deliberately NOT bare "lån" ("loan"), which v060's own Finding 2 already proved
    # unsafe (Norion Bank's asset-side "Loans to credit institutions" via bare "loan"/"loans" -- see the fallback
    # list's own comment history). The exact phrase "lån kreditinstitut" does not prefix-match a bank's own "lån
    # till kreditinstitut" (asset-side lending, "till" intervenes), so it stays direction-safe the same way
    # "bank loans" is. Two header words the same finding names are also missing from seed1-7's vocabulary:
    # "mindre än 12 månader" (due_within_1_year) and "mellan 1 och 2 år" (due_1_to_5_years) -- "mellan 3 och 5 år"
    # and "senare än 5 år" were already synonyms.
    tangen62_2025 = ("31 december 2025 (KSEK)\nMindre än 12 månader\nMellan 1 och 2 år\nMellan 3 och 5 år\n"
                     "Senare än 5 år Summa\nLeverantörsskulder och övriga skulder (exklusive icke finansiella skulder) 244 513 - - - 244 513\n"
                     "Lån Kreditinstitut 141 734 137 898 251 081 57 486 588 199\n"
                     "Leasingskulder 46 562 44 431 61 727 2 134 154 854\n")
    x.call_llm = lambda *a, **k: {"fields": [
        {"key": "total_debt", "value": None, "unit": None, "period": None, "raw_label": None, "source": None},
        {"key": "due_within_1_year", "value": None, "unit": None, "period": None, "raw_label": None, "source": None},
        {"key": "due_1_to_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None},
        {"key": "due_after_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None}]}
    out = x.extract([tangen62_2025], [1], dm, {"fiscal_year": 2025})
    got = {f["key"]: f["value"] for f in out["fields"]}
    assert got == {"total_debt": 588199, "due_within_1_year": 141734, "due_1_to_5_years": 388979, "due_after_5_years": 57486}, (got, out["warnings"])
    assert out["checks"][0]["passed"], out["checks"]  # 141734 + (137898+251081) + 57486 == 588199, exactly
    total = next(f for f in out["fields"] if f["key"] == "total_debt")
    assert total["raw_label"] == "Lån Kreditinstitut", total  # not "Leverantörsskulder..." (not a debt synonym), not "Leasingskulder" (not in row_synonyms, not this report's total)
    # v084: Finding 1's own real page prints BOTH fiscal years' tables, same row label each ("Lån
    # Kreditinstitut" under "31 december 2025", and again under "31 december 2024") -- v083 left this
    # ambiguous and unfilled (docs/acrylic/evidence/v083.md), since a bare index count can't tell which
    # candidate is the fiscal year's own row. _bucket_total_row now reuses _bucket_row_prior_year (v066's
    # own "this row names last year's header" test, already trusted at the write gate in
    # _fill_bucket_columns) to drop the 2024 row from contention first: it is the sole survivor, so it
    # wins outright, before known_total narrowing is even tried (which -- unrelated, still-open gap named
    # in v083 -- can't help here anyway: the un-hinted _row_amounts merges this table's Swedish thousands
    # into unrecognisable blobs). This is a real fix, not a special case for Tången: the rule is "one
    # candidate is provably last year's, drop it", which just happens to leave one row standing here.
    tangen62_full = tangen62_2025 + (
        "31 december 2024 (KSEK)\nMindre än 12 månader\nMellan 1 och 2 år Mellan 3 och 5 år Senare än 5 år Summa\n"
        "Leverantörsskulder och övriga skulder (exklusive icke finansiella skulder) 163 517 - - - 163 517\n"
        "Lån Kreditinstitut 90 970 68 704 188 828 4 520 353 022\n")
    tangen_rows = x._page_rows(tangen62_full)
    warnings: list = []
    assert x._bucket_total_row(tangen_rows, dmf["total_debt"], bucket_sfs, 2025, 588199, warnings) == [tangen_rows.index("Lån Kreditinstitut 141 734 137 898 251 081 57 486 588 199")]
    assert warnings == [], warnings  # resolved outright -- no "ambiguous" warning, nothing left for known_total to narrow
    # End to end (not just the row picker in isolation): the model returns all-null, same as v083's isolated
    # case above, but this time on the real two-table page -- extract() must still land on the fiscal year's
    # own four figures and close the identity, exactly like the single-year replay in docs/acrylic/evidence/v083.md.
    x.call_llm = lambda *a, **k: {"fields": [
        {"key": "total_debt", "value": None, "unit": None, "period": None, "raw_label": None, "source": None},
        {"key": "due_within_1_year", "value": None, "unit": None, "period": None, "raw_label": None, "source": None},
        {"key": "due_1_to_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None},
        {"key": "due_after_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None}]}
    out = x.extract([tangen62_full], [1], dm, {"fiscal_year": 2025})
    got = {f["key"]: f["value"] for f in out["fields"]}
    assert got == {"total_debt": 588199, "due_within_1_year": 141734, "due_1_to_5_years": 388979, "due_after_5_years": 57486}, (got, out["warnings"])
    assert out["checks"][0]["passed"], out["checks"]  # 141734 + (137898+251081) + 57486 == 588199, exactly -- the 2024 table (353022) never enters the sum
    # Counter-example (v084): two same-labelled candidate rows that are NOT resolvable by year -- both name
    # the SAME fiscal year (a report accidentally or genuinely repeats a table; Byggmax Group's real p.88 is
    # structurally the same shape, two same-labelled "Borrowing" rows, docs/acrylic/evidence/v084.md) --
    # _bucket_row_prior_year is false for both, so the new narrowing step drops neither, and the table must
    # stay exactly as ambiguous as pre-v084, not silently pick one.
    tangen62_sameyear = tangen62_2025 + (
        "31 december 2025 (KSEK)\nMindre än 12 månader\nMellan 1 och 2 år Mellan 3 och 5 år Senare än 5 år Summa\n"
        "Leverantörsskulder och övriga skulder (exklusive icke finansiella skulder) 163 517 - - - 163 517\n"
        "Lån Kreditinstitut 90 970 68 704 188 828 4 520 353 022\n")
    warnings = []
    assert x._bucket_total_row(x._page_rows(tangen62_sameyear), dmf["total_debt"], bucket_sfs, 2025, None, warnings) == []
    assert warnings == ["total_debt: 2 candidate debt rows for the bucket table ('Lån Kreditinstitut', 'Lån Kreditinstitut') -- ambiguous, none used"], warnings
    # v066 (a): Proact IT Group -- real Note 24 text (seed6-kb/proact_it_2025, p.103, trimmed to the two tables
    # that matter). The note's own per-instrument date breakdown lets the model correctly sum 216,360 (16 Jul
    # 2026) + 96,098 (2026) = 312,458 for due_within_1_year, closing the identity with the already-read
    # total_debt (478,611) and due_1_to_5_years (166,153) exactly -- but the page also prints "Group | Parent
    # company" (the note's other, unrelated sub-table), so _segment_column's own Group/Parent branch picks
    # column 1 of the page's 3-column year run; the single-instrument row that is due_within_1_year's own quote
    # reads as 0 in that column, and the pre-v066 guard threw the correct 312,458 away for it (docs/acrylic/
    # evidence/v063.md gap 2).
    proact103 = (
        "Notes\nNote 24 CONT.\nOther financial Group Parent company\nliabilities measured\nat accrued\n"
        "acquisition value 31 Dec 2025 31 Dec 2024 31 Dec 2025 31 Dec 2024\nNon-interest-bearing\n"
        "Liabilities on acquisi-\ntions 1) 139,365 - - -\nCurrency derivatives 1,146 408 29 -\n"
        "Other liabilities 1) 57,666 59,607 2,617 6,918\nAccounts payable 470,637 561,880 2,904 3,173\n"
        "Total\nnon-interest-bearing 668,814 621,895 5,550 10,091\nInterest-bearing\nBank loans, of which\n"
        "short-term portion 216,360 - 216,360 -\nBank loans, of which\nlong-term portion - 229,730 - 229,730\n"
        "Lease liabilities 2) 262,251 253,735 - -\nTotal interest-bearing 478,611 483,465 216,360 229,730\n"
        "Total other financial 1,147,425 1,105,360 221,910 239,821\nliabilities measured at\naccrued acquisition\n"
        "value\nProact Proact Proact Annual Annual Annual and and and Sustainability Sustainability Sustainability "
        "Report Report Report 2025 2025 2025\nInterest-bearing\nliabilities, Group, Reported\n"
        "31 Dec 2025 Interest Maturity value\nUtilised overdraft\nfacility, Nordea 1) 2) Base rate +2.0% 31 Dec 2025 -\n"
        "Bank loan, Nordea 2) STIBOR 3M +1.25% 16 Jul 2026 -\nBank loan, Nordea 2) EURIBOR 3M +1.25% 16 Jul 2026 -\n"
        "Bank loan, Svensk\nExportkredit 2) EURIBOR 3M + 1.8% 16 Jul 2026 216,360\n"
        "Lease liability 3) 3.86% - 5.59% 2026 96,098\nLease liability 3) 4.06% - 5.59% 2027-2030 166,153\n"
        "Total interest-bearing\nliabilities 478,611\n")
    x.call_llm = lambda *a, **k: {"fields": [
        {"key": "total_debt", "value": 478611, "unit": "SEK thousand", "period": "2025", "raw_label": "Total interest-bearing liabilities",
         "source": {"page": 1, "quote": "Total interest-bearing liabilities 478,611"}},
        {"key": "due_within_1_year", "value": 312458, "unit": "SEK thousand", "period": "2025",
         "raw_label": "Bank loan, Svensk Exportkredit 2); Lease liability 3)",
         "source": {"page": 1, "quote": "Lease liability 3) 3.86% - 5.59% 2026 96,098"}},
        {"key": "due_1_to_5_years", "value": 166153, "unit": "SEK thousand", "period": "2025", "raw_label": "Lease liability 3)",
         "source": {"page": 1, "quote": "Lease liability 3) 4.06% - 5.59% 2027-2030 166,153"}},
        {"key": "due_after_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None}]}
    out = x.extract([proact103], [1], dm, {"fiscal_year": 2025})
    got = {f["key"]: (f["value"], f["confidence"]) for f in out["fields"]}
    assert got == {"total_debt": (478611, 1.0), "due_within_1_year": (312458, 1.0), "due_1_to_5_years": (166153, 0.9),
                    "due_after_5_years": (None, 0.0)} and out["checks"][0]["passed"], (got, out["checks"], out["warnings"])
    within = next(f for f in out["fields"] if f["key"] == "due_within_1_year")
    assert "identity_kept" in within["evidence"] and "value_derived" in within["evidence"] \
        and any("kept: closes the identity exactly" in w for w in out["warnings"]), (within, out["warnings"])
    # the guard is an admission condition on that one correction, not a blanket exemption for the page: closing
    # the identity requires at least two OTHER real operands, so a lone total "closing" only against itself
    # (everything else null) must not count -- the supervisor's own named failure mode for this guard
    assert x._identity_closes("due_within_1_year", 312458,
        {"due_within_1_year": 312458, "due_1_to_5_years": 166153, "total_debt": 478611}, {}, dm)
    assert not x._identity_closes("due_within_1_year", 312458, {"due_within_1_year": 312458, "total_debt": 312458}, {}, dm)

    # v066 (b): Net Insight -- real Note 21 text (seed6-kb/net_insight_2025, p.101, trimmed). The model correctly
    # reads due_within_1_year=50,379 and due_after_5_years=0 off the note's own "<1 year"/">5 years" bucket rows,
    # but _fill_bucket_columns also finds the maturity table's own two-year "Total 81,489 57,647" row (2025's
    # total then 2024's, side by side) a plausible bucket-column candidate: its own header (the "<1 year"/"1-5
    # yeas"/">5 years" rows above it) has no recognised "total" column of its own (only two bucket keys, the
    # report's own "1-5 yeas" typo keeps the third out), so the pre-v066 over-valve has no total to compare the
    # row against and lets it overwrite both fields by column order -- 2025's own total (81,489) into
    # due_within_1_year, 2024's total (57,647) into due_after_5_years.
    net_insight101 = (
        "Note 21. Financial Assets and Liabilities\n"
        "Group's financial instruments by category\nAmounts in SEK thousands\nValue- tier\n"
        "Liabilities measured at amortized cost\nLiabilities measured at fair value trough profit and loss\n"
        "Value- tier\nLiabilities measured at amortized cost\nLiabilities measured at fair value trough profit and loss\n"
        "Liabilities in Balance Sheet\nDerivative instruments 2 - 2 3,790\n"
        "Accounts payable and other liabilities, excluding non-financial liabilities 42,074 44,354\n"
        "Lease liabilities 39,415 13,293\nTotal 81,489 - 57,647 3,790\n"
        "31 Dec 2025 31 Dec 2024\n31 Dec 2025 31 Dec 2024\n31 Dec 2025 31 Dec 2024\n"
        "<1 year 50,379 56,092\n1-5 yeas 31,110 1,555\n>5 years - -\nTotal 81,489 57,647\n")
    x.call_llm = lambda *a, **k: {"fields": [
        {"key": "total_debt", "value": 81489, "unit": "SEK thousands", "period": "2025", "raw_label": "Total",
         "source": {"page": 1, "quote": "Total 81,489 - 57,647 3,790"}},
        {"key": "due_within_1_year", "value": 50379, "unit": "SEK thousands", "period": "2025", "raw_label": "<1 year",
         "source": {"page": 1, "quote": "<1 year 50,379 56,092"}},
        {"key": "due_1_to_5_years", "value": 31110, "unit": "SEK thousands", "period": "2025", "raw_label": "1-5 yeas",
         "source": {"page": 1, "quote": "1-5 yeas 31,110 1,555"}},
        {"key": "due_after_5_years", "value": 0, "unit": "SEK thousands", "period": "2025", "raw_label": ">5 years",
         "source": {"page": 1, "quote": ">5 years - -"}}]}
    out = x.extract([net_insight101], [1], dm, {"fiscal_year": 2025})
    got = {f["key"]: f["value"] for f in out["fields"]}
    # due_after_5_years lands null, not a stored 0: the model's own 0 is not literally printed anywhere else on
    # the page either (a separate, pre-existing guard, out of this lane's territory) -- null_as_zero still closes
    # the identity with it, same as the report's own dash meaning no debt is due in that window
    assert got == {"total_debt": 81489, "due_within_1_year": 50379, "due_1_to_5_years": 31110, "due_after_5_years": None} \
        and out["checks"][0]["passed"], (got, out["checks"], out["warnings"])
    assert any("is a Total*/Summa* row with no total column" in w for w in out["warnings"]), out["warnings"]
    assert not any("read from 'Total 81,489 57,647' by its column order" in w for w in out["warnings"]), out["warnings"]

    # v066 (b), second valve: a bucket-column row whose own year is the fiscal year's predecessor must not lend
    # its figures to this year's buckets even when the row has its own recognised total column (so the first
    # valve above would not have caught it) -- synthetic, no seed6 company happened to hit this exact shape.
    assert x._bucket_row_prior_year(["Note 9 Borrowings 2024 2023", "Within 1 year 1-5 years Total", "Total 3,300 4,400 7,700"], 2, 2025)
    assert not x._bucket_row_prior_year(["Note 9 Borrowings 2025 2024", "Within 1 year 1-5 years Total", "Total 3,300 4,400 7,700"], 2, 2025)
    prior_year_bucket = "Note 9 Borrowings\n2024 2023\nWithin 1 year 1-5 years Total\nTotal 3,300 4,400 7,700\n"
    x.call_llm = lambda *a, **k: {"fields": [{"key": k, "value": None, "unit": None, "period": None, "raw_label": None, "source": None} for k in dmf]}
    out = x.extract([prior_year_bucket], [1], dm, {"fiscal_year": 2025})
    got = {f["key"]: f["value"] for f in out["fields"]}
    assert got == {"total_debt": None, "due_within_1_year": None, "due_1_to_5_years": None, "due_after_5_years": None}, (got, out["warnings"])
    assert any("names fiscal year 2024, not 2025" in w for w in out["warnings"]), out["warnings"]

    # v069: Ework/XANO's real header text, now correctly single-line per v068's parse.py fix (PARSER_VERSION 6,
    # docs/acrylic/evidence/v068.md) instead of matrix-transposed across two physical PDF lines -- the reason
    # both companies stayed unfixed from v036 through v066 (the synthetic scrambled-header cases above, kept
    # unchanged: still a real shape a header could be in, still correctly declined). With the header finally
    # readable in true column order, two pre-existing, unrelated-to-parsing gaps in this file's own hit
    # collection surfaced (v068's own "downstream" section, read-only, named for whoever owns extract.py/the
    # schema next -- this lane): XANO's plain synonym "inom 1 år" and the bare word "Summa" each also match a
    # piece of the one subtotal phrase "summa inom 1 år" already matched whole (10 hits, not 8); Ework's own
    # header_synonym "3 months" is a literal substring of its own "1-3 months" (1 spurious hit that happened to
    # roughly cancel out, in raw count only, the header's own separate, still-open "Due" gap below). Both fixed
    # the same way, _drop_nested_hits: a hit whose span sits entirely inside another, wider hit's own span is
    # dropped -- one printed phrase is one column, no matter how many schema synonyms happen to match pieces of
    # it, a general geometric rule, not a XANO- or Ework-specific patch.
    assert x._drop_nested_hits([(0, 5, "a"), (0, 16, "b:subtotal"), (6, 16, "a")]) == [(0, 16, "b:subtotal")]
    assert x._drop_nested_hits([(0, 5, "a"), (5, 10, "b")]) == [(0, 5, "a"), (5, 10, "b")]  # adjacent, not nested: both kept

    # "Due" (Ework's own column, v069's work order: add it to due_within_1_year's header_synonyms) turned out
    # unsafe as a plain header_synonym entry -- that list is matched case-insensitively, with no word boundary,
    # over the whole 25-row window _bucket_header searches above a candidate row, not just the header line
    # itself, and Ework's own p.70/71 carries a dozen+ ordinary-prose "due"s in that window ("...risk due to
    # assets...", "Past due accounts receivable...", "...not yet due..."). Confirmed end to end before choosing
    # _BUCKET_BOUNDARY instead (case-sensitive "Due", excludes "Due to ..."): a naive case-insensitive substitute
    # for the pattern below picks up "due" from that prose and pushes a wrong candidate row's own hit count to
    # coincidentally match its amount count, reaching _fill_bucket_columns's `over` valve instead of the safety
    # valve above it -- caught here, not guaranteed caught on every other company's own surrounding prose.
    assert not x._BUCKET_BOUNDARY["due_within_1_year"].search("The Company is exposed to translation risk due to assets")
    assert not x._BUCKET_BOUNDARY["due_within_1_year"].search("Past due accounts receivable 0-30 days")
    assert not x._BUCKET_BOUNDARY["due_within_1_year"].search("Accounts receivable not yet due")
    assert x._BUCKET_BOUNDARY["due_within_1_year"].search("kSEK Due < 1 month 1-3 months")

    xano_real = ("FINANSIELLA SKULDER Förfallotid\n"
                 "PER 2025-12-31 –30 dgr 31–90 dgr 91–360 dgr Summa inom 1 år Mellan 1 och 3 år Mellan 3 och 5 år Efter 5 år Totalt\n"
                 "Lån och leasingskulder 4 090 8 680 38 305 51 075 735 261 33 784 50 778 870 898\n"
                 "Summa räntebärande skulder 4 090 8 680 38 305 51 075 768 885 33 784 50 778 904 522\n")
    xano_rows = x._page_rows(xano_real)
    xano_idx = next(i for i, r in enumerate(xano_rows) if r.startswith("Summa räntebärande skulder"))
    # red, confirmed against unmodified origin/acrylic before writing _drop_nested_hits: 10 hits (2 spurious,
    # above) against the row's own 8 amounts -- declined, every field null
    assert x._bucket_header(xano_rows, xano_idx, bucket_sfs, 2025) == \
        ["_excluded", "_excluded", "_excluded", "due_within_1_year", "due_1_to_5_years", "due_1_to_5_years", "due_after_5_years", "total"]
    x.call_llm = lambda *a, **k: {"fields": [{"key": k, "value": None, "unit": None, "period": None, "raw_label": None, "source": None} for k in dmf]}
    out = x.extract([xano_real], [1], dm, {"fiscal_year": 2025})
    got = {f["key"]: f["value"] for f in out["fields"]}
    assert got == {"total_debt": 904522, "due_within_1_year": 51075, "due_1_to_5_years": 802669, "due_after_5_years": 50778} \
        and out["checks"][0]["passed"], (got, out["checks"], out["warnings"])  # matches eval/labels.csv (v057) exactly

    # Ework: "Due" (above -- a real, sometimes-nonzero column, the Accounts payable row's own Due figure is
    # 91,284, not a synonym for anything else in the schema) plus the same dedup now reads a clean, correct
    # 7-column hit list -- but the row has 8 real columns, and nothing in the schema recognises "Carrying
    # amount" (the report's own book-value total, v028's own total_debt basis) as distinct from "Total
    # undiscounted value" (the buckets' own subtotal, two columns to its left): v069 left that honestly
    # declined (its own doc, "not silently forced -- a second total slot risks summing both trailing columns
    # into total_debt on some other company's table where they are not equal"); v076 below resolves it with
    # exactly the missing half -- the two trailing columns named apart in the schema, counted, never summed.
    ework_real = ("kSEK Due < 1 month 1-3 months 3-12 months 1-5 years > 5 years Total undiscounted value Carrying amount\n"
                  "Short-term interest-bearing liabilities* – 153,761 971 1,677 – – 156,410 156,410\n")
    ework_rows = x._page_rows(ework_real)
    ework_idx = len(ework_rows) - 1
    assert x._bucket_header(ework_rows, ework_idx, bucket_sfs, 2025) == \
        ["due_within_1_year", "due_within_1_year", "due_within_1_year", "due_within_1_year", "due_1_to_5_years", "due_after_5_years", "total"]
    assert len(x._row_amounts(ework_rows[ework_idx], 7, nil=None)) == 8  # the row's own 8 amounts never shrink to fit a 7-column guess

    # v076: the two trailing total-shaped columns separate. Ework p.70's header carries BOTH "Total undiscounted
    # value" (the buckets' own undiscounted subtotal -- with future interest it can exceed the carrying total, so
    # it must never be summed into, or compared against, total_debt) and "Carrying amount" (v028's own stated
    # basis: the total that ties to the balance sheet). The schema now names both shapes -- carrying wording as
    # total_debt's own header_synonyms (-> the total slot), undiscounted/contractual wording as the schema's
    # explicit ignore_header_synonyms (-> an _ignore slot that is counted -- the column-count safety valve still
    # needs 8 hits for the row's 8 amounts -- but never assigned), and carrying wins the total slot when both
    # appear on the same header line (v028), while a header with only the undiscounted column keeps today's
    # behaviour: its own bare "Total" word stays the total slot, the identity closing against the report's basis.
    ework_8 = ["due_within_1_year", "due_within_1_year", "due_within_1_year", "due_within_1_year",
               "due_1_to_5_years", "due_after_5_years", "_ignore", "total"]
    assert x._bucket_header(ework_rows, ework_idx, bucket_sfs, 2025, total_sf=dmf["total_debt"],
                            ignore_syns=dm["ignore_header_synonyms"]) == ework_8
    # the prose guard, on Ework's own p.70 furniture: the note title ("…– undiscounted cash flows") and a
    # sentence ("…reflected in the carrying amount of…") sit inside the same 25-row header window -- v069's
    # bare-"due" lesson all over again -- and name no bucket on their own lines, so neither may add a slot.
    # (Acast/Nederman print "Carrying amount" over non-bucket year/instrument tables; same guard keeps those
    # headers reading exactly as before.)
    ework_page = ("At the balance sheet date, there is no significant concentration of credit exposure. "
                  "The maximum exposure to credit risk is reflected in the carrying amount of the trade receivables\n"
                  "Maturity structure financial liabilities – undiscounted cash flows\n"
                  "The Group kSEK Due < 1 month 1-3 months 3-12 months 1-5 years > 5 years Total undiscounted value Carrying amount 2025\n"
                  "Lease liabilities – – 5,230 5,332 22,169 794 33,525 33,403\n"
                  "Short-term interest-bearing liabilities* – 153,761 971 1,677 – – 156,410 156,410\n")
    ework70_rows = x._page_rows(ework_page)
    assert x._bucket_header(ework70_rows, len(ework70_rows) - 1, bucket_sfs, 2025, total_sf=dmf["total_debt"],
                            ignore_syns=dm["ignore_header_synonyms"]) == ework_8
    # undiscounted wording with no carrying column on the line: today's behaviour -- its own bare word is the total
    assert x._bucket_header(x._page_rows(
        "kSEK Due < 1 month 1-3 months 3-12 months 1-5 years > 5 years Total undiscounted value\n"
        "Short-term interest-bearing liabilities* – 153,761 971 1,677 – – 156,410\n"),
        1, bucket_sfs, 2025, total_sf=dmf["total_debt"], ignore_syns=dm["ignore_header_synonyms"]) == \
        ["due_within_1_year", "due_within_1_year", "due_within_1_year", "due_within_1_year", "due_1_to_5_years", "due_after_5_years", "total"]
    # end to end, the row read by column even with no model answer at all: the single debt-row candidate fills
    # total_debt (156,410, the Carrying amount column) and due_within_1_year (156,409 = 153,761 + 971 + 1,677,
    # the Due column nil). v078: the row's own 1-5 years and > 5 years cells print dashes -- the report saying
    # "no debt is due in those windows" -- so those buckets are the report's explicit 0 (printed_nil @0.5), and
    # the identity closes for real instead of stalling on the v044/v058 printed-bucket-label guard's "missing".
    x.call_llm = lambda *a, **k: {"fields": [{"key": k, "value": None, "unit": None, "period": None, "raw_label": None, "source": None} for k in dmf]}
    out = x.extract([ework_real], [1], dm, {"fiscal_year": 2025})
    got = {f["key"]: f["value"] for f in out["fields"]}
    assert got == {"total_debt": 156410, "due_within_1_year": 156409, "due_1_to_5_years": 0, "due_after_5_years": 0} \
        and out["checks"][0]["passed"], (got, out["checks"], out["warnings"])
    for k in ("due_1_to_5_years", "due_after_5_years"):
        f = next(f for f in out["fields"] if f["key"] == k)
        assert (f["confidence"], f["raw_label"], f["source"]["quote"]) == \
            (0.5, "Short-term interest-bearing liabilities* –",
             "Short-term interest-bearing liabilities* – 153,761 971 1,677 – – 156,410 156,410"), f
        assert f["evidence"] == ["quote_on_page", "printed_nil", "identity_all_columns", "arith_ok", "label_known", "period_ok", "page_is_statement", "unit_ok"], f  # the closing row identifies the label; the printed nil keeps the 0.5 cap
    assert any("a dash is printed in the row's own column for it" in w for w in out["warnings"]), out["warnings"]
    # end to end with the model's own answer restored (docs/acrylic/evidence/v051/ework_2025.partB.debt_maturity.json,
    # the real pass's stored shape): total_debt already reads 156,410 (the Carrying amount column's own figure, so
    # the column read agrees and the model's own evidence stands), and the column read fills due_within_1_year --
    # 0 (Due) + 153,761 + 971 + 1,677 = 156,409; the two dash columns are the report's explicit 0s (v078), and
    # 156,409 + 0 + 0 sits within +-2 of 156,410: maturity_sums_to_total passes.
    ework_answer = [
        {"key": "total_debt", "value": 156410, "unit": "kSEK", "period": "2025", "raw_label": "Short-term interest-bearing liabilities*",
         "source": {"page": 1, "quote": "Short-term interest-bearing liabilities* – 153,761 971 1,677 – – 156,410 156,410"}},
        {"key": "due_within_1_year", "value": None, "unit": None, "period": None, "raw_label": None, "source": None},
        {"key": "due_1_to_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None},
        {"key": "due_after_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None}]
    x.call_llm = lambda *a, **k: {"fields": json.loads(json.dumps(ework_answer))}
    out = x.extract([ework_real], [1], dm, {"fiscal_year": 2025})
    got = {f["key"]: f["value"] for f in out["fields"]}
    assert got == {"total_debt": 156410, "due_within_1_year": 156409, "due_1_to_5_years": 0, "due_after_5_years": 0} \
        and out["checks"][0]["passed"], (got, out["checks"], out["warnings"])
    w1y = next(f for f in out["fields"] if f["key"] == "due_within_1_year")
    assert "value_derived" in w1y["evidence"] and w1y["source"]["quote"] == \
        "Short-term interest-bearing liabilities* – 153,761 971 1,677 – – 156,410 156,410", (w1y, out["warnings"])
    assert any("156409 read from" in w for w in out["warnings"]), out["warnings"]
    d15 = next(f for f in out["fields"] if f["key"] == "due_1_to_5_years")
    assert "printed_nil" in d15["evidence"] and d15["confidence"] == 0.5, d15  # the report's own dash, not a fabricated number

    # v085: Instalco (seed8-kb/v082 Finding 2, real p.128 text): the same carrying-beside-undiscounted shape as
    # Ework above, but split across three lines instead of one -- "Total contractual cash flows" (the
    # undiscounted group's own header) and "31/12/2025 Carrying amount receivables/ payables" (the carrying
    # column's own header) each sit on their own line, directly above "Within 6 months / 6-12 months / 1-5
    # years / Later than 5 years", nothing between them. v076's own has_carry never sees either: both hits are
    # stripped before that check runs (their own row names no bucket), the same gate v082/v083 both traced to
    # and left blocked. Admitted by the contiguous-run fallback below: the walk backward from the header's own
    # last bucket-naming line crosses both, in order, and stops at the first line that carries neither a bucket
    # word nor a carrying/ignore word of its own ("Current Non-current").
    instalco128 = ("The Group\nCurrent Non-current\nTotal contractual cash flows\n"
                   "31/12/2025 Carrying amount receivables/ payables\n"
                   "Within 6 months 6–12 months 1–5 years Later than 5 years\n"
                   "Liabilities to credit institutions – – 3,209 – 3,209 3,122\n")
    instalco_rows = x._page_rows(instalco128)
    instalco_idx = instalco_rows.index("Liabilities to credit institutions – – 3,209 – 3,209 3,122")
    assert x._bucket_header(instalco_rows, instalco_idx, bucket_sfs, 2025, total_sf=dmf["total_debt"],
                            ignore_syns=dm["ignore_header_synonyms"]) == \
        ["due_within_1_year", "due_within_1_year", "due_1_to_5_years", "due_after_5_years", "_ignore", "total"]
    # end to end: the header now reads 6 columns matching the row's own 6 amounts (declined before this lane,
    # 5 against 6), but the fill still does not land -- due_1_to_5_years' own undiscounted figure (3,209)
    # exceeds total_debt's own discounted "Carrying amount" (3,122) by design (future interest over the 1-5-
    # year horizon, the same undiscounted-vs-carrying gap Ework's own table has above), which
    # _fill_bucket_columns's `over` valve (out of this lane's territory -- word-list/admission-condition only)
    # reads as a misaligned column and declines whole. Reported, not forced: the four fields stay exactly where
    # v082/v083 left them, now with an honest reason on record instead of a silent decline with no warning.
    x.call_llm = lambda *a, **k: {"fields": [
        {"key": "total_debt", "value": 3122, "unit": "SEK m", "period": "2025", "raw_label": "Liabilities to credit institutions",
         "source": {"page": 1, "quote": "Liabilities to credit institutions – – 3,209 – 3,209 3,122"}},
        {"key": "due_within_1_year", "value": None, "unit": None, "period": None, "raw_label": None, "source": None},
        {"key": "due_1_to_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None},
        {"key": "due_after_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None}]}
    out = x.extract([instalco128], [1], dm, {"fiscal_year": 2025})
    got = {f["key"]: f["value"] for f in out["fields"]}
    assert got == {"total_debt": 3122, "due_within_1_year": None, "due_1_to_5_years": None, "due_after_5_years": None}, (got, out["warnings"])
    assert any("due_1_to_5_years 3209 exceeds its own total 3122" in w for w in out["warnings"]), out["warnings"]

    # v085 guard, real page shapes that must NOT gain a fallback slot (the contiguous run stops before reaching
    # them): Ework's own real p.70 (not the synthetic ework_real above) wraps "Total undiscounted value" /
    # "Carrying amount" across a line-break so ragged pymupdf glues it into one row with the bucket words
    # themselves ("Total undis- Carrying kSEK Due < 1 month ... counted value amount 2025") -- neither phrase
    # is a contiguous substring any more, so this bucket-naming line has no carrying hit of its own, and its
    # own immediate predecessor ("The Group") has neither a bucket word nor a carrying/ignore word either: the
    # run stops on its very first step, never reaching the real prose danger six-plus lines further back
    # ("...reflected in the carrying amount...", a note title with "...undiscounted cash flows"). Confirmed
    # against the mechanism this lane could have broken, not assumed: without the stop, both leak in and this
    # exact page previously misread the *other* row, Lease liabilities, as the table's own total (33,525, its
    # own undiscounted-value column) -- caught by seed5-kb's own replay, not by inspection.
    ework_real70 = ("At the balance sheet date, there is no significant concentration of credit exposure. "
                    "The maximum exposure to credit risk is reflected in the carrying amount in the statement "
                    "of financial position for each financial asset.\n"
                    "Maturity structure financial liabilities – undiscounted cash flows\nThe Group\n"
                    "Total undis- Carrying kSEK Due < 1 month 1-3 months 3-12 months 1-5 years > 5 years counted value amount 2025\n"
                    "Lease liabilities – – 5,230 5,332 22,169 794 33,525 33,403\n"
                    "Short-term interest-bearing liabilities* – 153,761 971 1,677 – – 156,410 156,410\n")
    x.call_llm = lambda *a, **k: {"fields": [
        {"key": "total_debt", "value": 156410, "unit": "kSEK", "period": "2025", "raw_label": "Short-term interest-bearing liabilities*",
         "source": {"page": 1, "quote": "Short-term interest-bearing liabilities* – 153,761 971 1,677 – – 156,410 156,410"}},
        {"key": "due_within_1_year", "value": None, "unit": None, "period": None, "raw_label": None, "source": None},
        {"key": "due_1_to_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None},
        {"key": "due_after_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None}]}
    out = x.extract([ework_real70], [1], dm, {"fiscal_year": 2025})
    got = {f["key"]: f["value"] for f in out["fields"]}
    assert got == {"total_debt": 156410, "due_within_1_year": None, "due_1_to_5_years": None, "due_after_5_years": None} \
        and not out["warnings"], (got, out["warnings"])  # untouched: no candidate ever reaches the write step
    # Ependion's own real p.155 (seed5-kb): two unrelated sentences each end "...correspond[s] to carrying
    # amount[,] because..." between the page's own instrument rows and its "Contracted terms" bucket table --
    # bucket-less, and each one line further back than the last, so the run (starting at "Between 1 and 2
    # years", the header's own second-to-last bucket-naming line, which has no carrying/ignore hit of its own)
    # stops immediately, before it can reach either. Without the stop, both leak in, wrongly demote the header's
    # own real bare "Total" to _ignore, and turn this already-correct fill (v060's own ependion_real case above,
    # same table, clean text) into a false decline -- caught the same way, by seed5-kb's own replay.
    ependion155 = ("Financial liabilities measured at amortized cost\n- Borrowing 583,531 557,173\n"
                   "- Accounts payable-trade 164,155 154,411\n"
                   "The fair value of borrowing corresponds to carrying amount because interest on this "
                   "borrowing is on a par with current market interest rates or due to borrowing being short term.\n"
                   "Accounts payable are unsecured and normally paid within 30 days. The fair value of accounts "
                   "payable are considered to correspond to carrying amount, because they are inherently short term.\n"
                   "The following table states the contracted maturities of financial liabilities.\n"
                   "Contracted terms\nSEK 000 Within 12 months\nBetween 1 and 2 years\nBetween 2 and 3 years Total\n"
                   "Borrowing 167,546 35,636 380,348 583,531\n"
                   "Accounts payable-trade 164,155 164,155 331,701 35,636 380,348 747,686\n")
    x.call_llm = lambda *a, **k: {"fields": [
        {"key": "total_debt", "value": None, "unit": None, "period": None, "raw_label": None, "source": None},
        {"key": "due_within_1_year", "value": None, "unit": None, "period": None, "raw_label": None, "source": None},
        {"key": "due_1_to_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None},
        {"key": "due_after_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None}]}
    out = x.extract([ependion155], [1], dm, {"fiscal_year": 2025})
    got = {f["key"]: f["value"] for f in out["fields"]}
    assert got == {"total_debt": 583531, "due_within_1_year": 167546, "due_1_to_5_years": 415984, "due_after_5_years": None}, (got, out["warnings"])
    assert out["checks"][0]["passed"], out["checks"]

    # v089: the maturity basis is a user choice, not a hardcode -- DEBT_BASIS=carrying (default;
    # everything above, byte for byte, and the extraction now says so in its own top-level "maturity_basis")
    # or "undiscounted": the same bucket columns read as always, but total_debt comes from the
    # contractual undiscounted total column -- the schema's ignore_header_synonyms wording (v076)
    # becomes the total slot and the carrying wording becomes the ignored column -- so the identity
    # closes against the undiscounted total, and Instalco's 3,209-vs-3,122 gap, which the carrying
    # basis' over valve correctly rejects above, is on this basis the report's own read.
    assert x.debt_basis() == "carrying"  # default: no env set, no behaviour change anywhere
    x.call_llm = lambda *a, **k: {"fields": [
        {"key": "total_debt", "value": 3122, "unit": "SEK m", "period": "2025", "raw_label": "Liabilities to credit institutions",
         "source": {"page": 1, "quote": "Liabilities to credit institutions – – 3,209 – 3,209 3,122"}},
        {"key": "due_within_1_year", "value": None, "unit": None, "period": None, "raw_label": None, "source": None},
        {"key": "due_1_to_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None},
        {"key": "due_after_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None}]}
    out = x.extract([instalco128], [1], dm, {"fiscal_year": 2025})
    assert out["maturity_basis"] == "carrying" and out["checks"][0]["passed"] is False, (out["maturity_basis"], out["checks"])
    # the prompt names the basis (the schema carries both wordings; the carrying one is the one
    # every extraction has carried since v028, unchanged)
    sp = x.system_prompt(dm)
    assert "Read the carrying-amount table" in sp, sp[:200]
    basis_saved = x.debt_basis
    x.debt_basis = lambda: "undiscounted"
    assert x.system_prompt(dm) != sp and "Read the carrying-amount table" not in x.system_prompt(dm)
    assert "Read the contractual undiscounted cash flows table" in x.system_prompt(dm), x.system_prompt(dm)[:200]
    # header level: the two total-shaped columns swap slots
    assert x._bucket_header(instalco_rows, instalco_idx, bucket_sfs, 2025, total_sf=dmf["total_debt"],
                            ignore_syns=dm["ignore_header_synonyms"], basis="undiscounted") == \
        ["due_within_1_year", "due_within_1_year", "due_1_to_5_years", "due_after_5_years", "total", "_ignore"]
    assert x._bucket_header(ework_rows, ework_idx, bucket_sfs, 2025, total_sf=dmf["total_debt"],
                            ignore_syns=dm["ignore_header_synonyms"], basis="undiscounted") == \
        ["due_within_1_year", "due_within_1_year", "due_within_1_year", "due_within_1_year",
         "due_1_to_5_years", "due_after_5_years", "total", "_ignore"]
    # end to end, Instalco under undiscounted: total_debt 3,122 (the model's own carrying read) is
    # overwritten with the row's own contractual total 3,209, the dash buckets are the report's 0,
    # and 0 + 3,209 + 0 closes against 3,209 exactly -- passed, for the first time on this table.
    out = x.extract([instalco128], [1], dm, {"fiscal_year": 2025})
    got = {f["key"]: f["value"] for f in out["fields"]}
    assert got == {"total_debt": 3209, "due_within_1_year": 0, "due_1_to_5_years": 3209, "due_after_5_years": 0} \
        and out["checks"][0]["passed"] and out["maturity_basis"] == "undiscounted", (got, out["checks"], out["warnings"])
    assert any("total_debt: 3122 disagrees with the maturity table" in w for w in out["warnings"]), out["warnings"]
    nilw = next(f for f in out["fields"] if f["key"] == "due_after_5_years")
    assert nilw["confidence"] == 0.5 and "printed_nil" in nilw["evidence"], nilw  # the dash, still the report's own 0
    # Ework under undiscounted: the row prints the same figure in both total columns (156,410), so the
    # four fields land as under carrying -- now proven against the *undiscounted* column. The ragged
    # real p.70 header (neither phrase contiguous, v085's own guard shape) still declines on this
    # basis too: no basis column anywhere in it, no fill, no warning, both bases alike.
    x.call_llm = lambda *a, **k: {"fields": json.loads(json.dumps(ework_answer))}
    out = x.extract([ework_real], [1], dm, {"fiscal_year": 2025})
    got = {f["key"]: f["value"] for f in out["fields"]}
    assert got == {"total_debt": 156410, "due_within_1_year": 156409, "due_1_to_5_years": 0, "due_after_5_years": 0} \
        and out["checks"][0]["passed"] and out["maturity_basis"] == "undiscounted", (got, out["checks"], out["warnings"])
    x.call_llm = lambda *a, **k: {"fields": json.loads(json.dumps(ework_answer))}
    out = x.extract([ework_real70], [1], dm, {"fiscal_year": 2025})
    got = {f["key"]: f["value"] for f in out["fields"]}
    assert got == {"total_debt": 156410, "due_within_1_year": None, "due_1_to_5_years": None, "due_after_5_years": None} \
        and not out["warnings"], (got, out["warnings"])  # untouched: no candidate ever reaches the write step
    x.debt_basis = basis_saved  # restore: every later case runs on the default carrying basis

    # v095 (Karnell, seed9-kb, v088 Finding 1, real p.106 text): the maturity table's own header prints
    # "<1 year 1-3 years >3 years Total" -- "<1 year" is a bucket word, "1-3 years" a finer split of
    # due_1_to_5_years, but ">3 years" ([36 months, infinity)) is off the schema's 1/5-year grid: it crosses
    # the 5-year boundary without naming it, so 3 header keys stood against the row's 4 printed amounts and
    # the length valve declined everything (v088: four nulls on both passes of a real, well-formed table).
    # The off-grid column is now counted (4 keys, 4 amounts, alignment restored) and never assigned: on the
    # interest-bearing row its cell prints "-", which resolves the windows it covers outright -- the whole of
    # due_after_5_years sits inside the nil ">3 years" column, so it is 0, and 43.5 + 353.7 + 0 == 397.2
    # closes exactly.
    karnell106 = ("NOTE 47 - FINANCIAL LIABILITIES\n"
                  "Maturity analysis liabilities, 2025 <1 year 1-3 years >3 years Total\n"
                  "Liabilities to credit institutions 43.5 353.7 - 397.2\n"
                  "Contingent earn-outs 27.4 0.7 - 28.1\nPut/call options - - 107.0 107.0\n"
                  "Accounts payable 1.2 - - 1.2\nTotal 72.2 354.4 107.0 533.5\n"
                  "Maturity analysis liabilities, 2024 <1 year 1-3 years >3 years Total\n"
                  "Liabilities to credit institutions 71.4 351.7 - 423.2\n"
                  "Contingent earn-outs 3.9 31.0 - 34.9\nPut/call options 38.9 - 92.5 131.4\n"
                  "Accounts payable 0.8 - - 0.8\nTotal 115.1 382.7 92.5 590.3\n")
    # the scan claims only what no bucket word did: on-grid phrases ("> 5 years", "mer än 5 år", "1-5 years")
    # and finer splits the schema already owns ("1-3 years") stay untouched; bounds land in months, open-ended hi empty
    hline = "Maturity analysis liabilities, 2025 <1 year 1-3 years >3 years Total"
    syn = x._bucket_synonym_hits(hline, bucket_sfs, dmf["total_debt"], dm["ignore_header_synonyms"])
    assert [k for _, _, k in x._offgrid_hits(hline, [(s, e) for s, e, _ in syn])] == ["offgrid:36-|>3 years"], syn
    assert x._offgrid_hits("mer än 5 år > 5 years 1-5 years", []) == []
    assert [k for _, _, k in x._offgrid_hits("over 2 years 1-4 years", [])] == \
        ["offgrid:24-|over 2 years", "offgrid:12-48|1-4 years"]
    kr = x._page_rows(karnell106)
    debt_i = kr.index("Liabilities to credit institutions 43.5 353.7 - 397.2")
    assert x._bucket_header(kr, debt_i, bucket_sfs, 2025, total_sf=dmf["total_debt"]) == \
        ["due_within_1_year", "due_1_to_5_years", "offgrid:36-|>3 years", "total"]
    # the off-grid column never becomes one of the >=2 distinct bucket keys the header valve requires: prose
    # mentioning one bucket word plus ">3 years" is still not a bucket-column table
    assert x._bucket_header(x._page_rows("The loans mature over 3 years.\nWithin 1 year 5 7\n"), 1, bucket_sfs, 2025) is None
    # end to end, the model all-null exactly as v088's two passes answered: the debt row (a row_synonyms hit)
    # fills all four fields, the identity passes on explicit values, and the derived 0 carries value_derived --
    # the dash is in a column of the table's own naming (">3 years"), not the bucket's own column, so it is a
    # derived 0, not v078's printed_nil
    x.call_llm = lambda *a, **k: {"fields": [
        {"key": "total_debt", "value": None, "unit": None, "period": None, "raw_label": None, "source": None},
        {"key": "due_within_1_year", "value": None, "unit": None, "period": None, "raw_label": None, "source": None},
        {"key": "due_1_to_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None},
        {"key": "due_after_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None}]}
    out = x.extract([karnell106], [1], dm, {"fiscal_year": 2025})
    got = {f["key"]: f["value"] for f in out["fields"]}
    assert got == {"total_debt": 397.2, "due_within_1_year": 43.5, "due_1_to_5_years": 353.7, "due_after_5_years": 0} \
        and out["checks"][0]["passed"], (got, out["checks"], out["warnings"])
    total = next(f for f in out["fields"] if f["key"] == "total_debt")
    after5 = next(f for f in out["fields"] if f["key"] == "due_after_5_years")
    assert total["raw_label"] == "Liabilities to credit institutions" and after5["raw_label"] == total["raw_label"]
    assert "value_derived" in after5["evidence"] and "printed_nil" not in after5["evidence"], after5
    assert any("the row prints a dash in its '>3 years' column" in w for w in out["warnings"]), out["warnings"]
    # counterfactual (v095 CE1): the same table with a valued ">3 years" cell on the debt row stays honest --
    # four nulls and a warning naming the boundary straddled; the split (1-5-years money vs after-5-years
    # money) is a guess neither the check nor this file can arbitrate
    ce1 = karnell106.replace("Liabilities to credit institutions 43.5 353.7 - 397.2",
                             "Liabilities to credit institutions 43.5 353.7 173.0 570.2")
    out = x.extract([ce1], [1], dm, {"fiscal_year": 2025})
    assert all(f["value"] is None for f in out["fields"]), (out["fields"], out["warnings"])
    assert any("'>3 years' 173 straddles the 5-year boundary" in w for w in out["warnings"]), out["warnings"]
    # v095 rule 2, the row picker: a debt-scoped row (the total field's own wording or a row_synonyms hit)
    # outranks a bare table Total whose scope is only "whatever this table sums" (Karnell's 533.5 includes
    # earn-outs, put/call options and accounts payable). Only-bare-Totals pages keep today's order exactly.
    warnings: list = []
    order = x._bucket_total_row(kr, dmf["total_debt"], bucket_sfs, 2025, None, warnings)
    assert order == [debt_i, kr.index("Total 72.2 354.4 107.0 533.5"), kr.index("Total 115.1 382.7 92.5 590.3")], order
    assert warnings == [], warnings  # one surviving debt row (the 2024 one never aligns: both headers in its window)
    ce2 = karnell106.replace("Liabilities to credit institutions 43.5 353.7 - 397.2",
                             "Other financial liabilities 43.5 353.7 - 397.2")
    warnings = []
    assert x._bucket_total_row(x._page_rows(ce2), dmf["total_debt"], bucket_sfs, 2025, None, warnings) == \
        [x._page_rows(ce2).index("Total 72.2 354.4 107.0 533.5"), x._page_rows(ce2).index("Total 115.1 382.7 92.5 590.3")]
    assert warnings == [], warnings
    out = x.extract([ce2], [1], dm, {"fiscal_year": 2025})
    assert all(f["value"] is None for f in out["fields"]), (out["fields"], out["warnings"])  # no debt candidate: bare Totals only, today's order, declined on their valued ">3 years" -- baseline values
    # rule 2's own proof is v088's rejected experiment flipped harmless: claim ">3 years" for
    # due_after_5_years (schema copy, test-only) and the header opens either way -- before the reorder that
    # filled the bare Total 533.5 with the identity passing (the wrong-scope value endorsed by its own
    # check); now the debt row outranks it and the same claim fills 397.2, the dash now the bucket's own
    # column (v078's printed_nil, not value_derived)
    dm95 = json.loads(json.dumps(dm))
    next(f for f in dm95["fields"] if f["key"] == "due_after_5_years")["synonyms"].append(">3 years")
    out = x.extract([karnell106], [1], dm95, {"fiscal_year": 2025})
    got = {f["key"]: f["value"] for f in out["fields"]}
    assert got == {"total_debt": 397.2, "due_within_1_year": 43.5, "due_1_to_5_years": 353.7, "due_after_5_years": 0} \
        and out["checks"][0]["passed"], (got, out["checks"], out["warnings"])
    after5 = next(f for f in out["fields"] if f["key"] == "due_after_5_years")
    assert "printed_nil" in after5["evidence"] and "value_derived" not in after5["evidence"], after5

    qcalls = []
    # v054: EXTRACT_QUOTE_RETRY (default off) -- one follow-up call for the fields whose answer cites a quote
    # that is not printed on the page it names: the cited page's own _page_rows go out numbered, and the model
    # copies the printed row verbatim (or answers null: no such row). A retried field is adopted only when its
    # new quote verifies on the page it names -- null, missing or still-unverified replies keep the original
    # answer, so the retry can never leave a field worse than the first call alone. "Current portion of loans"
    # is deliberately not one of the schema's synonyms, so the label-repair inside the validation loop cannot
    # fix this quote either -- exactly the 0.25 shape the hardening rounds kept hitting (Storytel, MEKO).
    qret_page = ("Note 20 Borrowings\nMSEK\n2025 2024\nTotal borrowings 32 703 34 500\nCurrent portion of loans 3 538 4 100\n"
                 "1-5 years 29 165 30 000\n")

    def qret_answer():
        return {"fields": [
            {"key": "total_debt", "value": 32703, "unit": "MSEK", "period": "2025", "raw_label": "Total borrowings",
             "source": {"page": 1, "quote": "Total borrowings 32 703 34 500"}},
            {"key": "due_within_1_year", "value": 3538, "unit": "MSEK", "period": "2025", "raw_label": "Current portion of loans",
             "source": {"page": 1, "quote": "Current portion of loans, 3,538"}},  # reformatted: not the printed line
            {"key": "due_1_to_5_years", "value": 29165, "unit": "MSEK", "period": "2025", "raw_label": "1-5 years",
             "source": {"page": 1, "quote": "1-5 years 29 165 30 000"}},
            {"key": "due_after_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None}]}

    def qret_fix():
        return {"fields": [{"key": "due_within_1_year", "value": 3538, "source": {"page": 1, "quote": "Current portion of loans 3 538 4 100"}}]}

    qret_seen = {}

    def qret_llm(system, user, schema=x.RESPONSE_SCHEMA, name="extraction"):
        qcalls.append(name)
        qret_seen[name] = (system, user)
        return qret_answer() if name == "extraction" else qret_fix()

    import os
    os.environ["EXTRACT_QUOTE_RETRY"] = "1"
    try:
        x.call_llm = qret_llm
        qcalls.clear()
        out = x.extract([qret_page], [1], dm, {"fiscal_year": 2025})
        assert qcalls == ["extraction", "quote_retry"], qcalls  # one follow-up, only because a quote failed
        assert qret_seen["quote_retry"][0] == qret_seen["extraction"][0], "retry reuses the extraction system prompt"
        assert "4: Current portion of loans 3 538 4 100" in qret_seen["quote_retry"][1], qret_seen["quote_retry"][1]  # the page's own rows, numbered
        assert "your value: 3538" in qret_seen["quote_retry"][1], qret_seen["quote_retry"][1]  # the field's own earlier answer
        w1y = next(f for f in out["fields"] if f["key"] == "due_within_1_year")
        assert (w1y["value"], w1y["confidence"], w1y["source"]["quote"]) == (3538, 0.9, "Current portion of loans 3 538 4 100"), (w1y, out["warnings"])
        assert sum(w.startswith("quote_retry:") for w in out["warnings"]) == 1, out["warnings"]
        assert next(f for f in out["fields"] if f["key"] == "total_debt")["confidence"] == 1.0, out["fields"]  # verified fields are not retried

        # nothing to fix (every quote already verbatim): no follow-up call at all
        def qret_good(system, user, schema=x.RESPONSE_SCHEMA, name="extraction"):
            qcalls.append(name)
            return {"fields": [{**f, "source": {"page": 1, "quote": "Current portion of loans 3 538 4 100"}} if f["key"] == "due_within_1_year" else f
                               for f in qret_answer()["fields"]]}
        x.call_llm = qret_good
        qcalls.clear()
        out = x.extract([qret_page], [1], dm, {"fiscal_year": 2025})
        assert qcalls == ["extraction"], qcalls

        # the retry answers null (no printed row states the figure): the field keeps the first answer's judgment
        def qret_null(system, user, schema=x.RESPONSE_SCHEMA, name="extraction"):
            qcalls.append(name)
            return qret_answer() if name == "extraction" else {"fields": [{"key": "due_within_1_year", "value": None, "source": None}]}
        x.call_llm = qret_null
        qcalls.clear()
        out = x.extract([qret_page], [1], dm, {"fiscal_year": 2025})
        assert qcalls == ["extraction", "quote_retry"], qcalls
        w1y = next(f for f in out["fields"] if f["key"] == "due_within_1_year")
        assert (w1y["value"], w1y["confidence"]) == (3538, 0.25) and not any(w.startswith("quote_retry:") for w in out["warnings"]), (w1y, out["warnings"])

        # a retry quote that still does not verify keeps the original answer too
        def qret_still_bad(system, user, schema=x.RESPONSE_SCHEMA, name="extraction"):
            qcalls.append(name)
            return qret_answer() if name == "extraction" else \
                {"fields": [{"key": "due_within_1_year", "value": 3538, "source": {"page": 1, "quote": "Loans due soon 3538"}}]}
        x.call_llm = qret_still_bad
        qcalls.clear()
        out = x.extract([qret_page], [1], dm, {"fiscal_year": 2025})
        assert qcalls == ["extraction", "quote_retry"], qcalls
        w1y = next(f for f in out["fields"] if f["key"] == "due_within_1_year")
        assert (w1y["value"], w1y["confidence"]) == (3538, 0.25) and not any(w.startswith("quote_retry:") for w in out["warnings"]), (w1y, out["warnings"])

        # a stated zero (v050's prose-negation path) is already proven by the report's own words: not retried
        prose = ("Not 18 Klassificering av finansiella instrument\n"
                 "Investmentföretaget har varken räntebärande skulder eller kundfordringar.\n")

        def qret_zero(system, user, schema=x.RESPONSE_SCHEMA, name="extraction"):
            qcalls.append(name)
            return {"fields": [
                {"key": "total_debt", "value": 0, "unit": "SEK mn", "period": "2025",
                 "raw_label": "Investmentföretaget har varken räntebärande skulder eller kundfordringar.",
                 "source": {"page": 1, "quote": "Investmentföretaget har varken räntebärande skulder eller kundfordringar."}},
                {"key": "due_within_1_year", "value": None, "unit": None, "period": None, "raw_label": None, "source": None},
                {"key": "due_1_to_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None},
                {"key": "due_after_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None}]}
        x.call_llm = qret_zero
        qcalls.clear()
        out = x.extract([prose], [1], dm, {"fiscal_year": 2025})
        assert qcalls == ["extraction"], qcalls  # quote_on_page fails on the digit-free sentence; _stated_zero already proves it
        td = next(f for f in out["fields"] if f["key"] == "total_debt")
        assert td["value"] == 0 and td["confidence"] == 0.5 and "stated_zero" in td["evidence"], (td, out["warnings"])

        # a value the model returned without any source: the retry may supply the printed row (page 1 = candidates[0])
        def qret_nosrc(system, user, schema=x.RESPONSE_SCHEMA, name="extraction"):
            qcalls.append(name)
            if name == "extraction":
                return {"fields": [{**qret_answer()["fields"][0], "source": None}, *qret_answer()["fields"][1:]]}
            return {"fields": [{"key": "total_debt", "value": 32703, "source": {"page": 1, "quote": "Total borrowings 32 703 34 500"}}]}
        x.call_llm = qret_nosrc
        qcalls.clear()
        out = x.extract([qret_page], [1], dm, {"fiscal_year": 2025})
        assert qcalls == ["extraction", "quote_retry"], qcalls
        td = next(f for f in out["fields"] if f["key"] == "total_debt")
        assert (td["value"], td["confidence"]) == (32703, 1.0), (td, out["warnings"])  # was "value without source dropped"
    finally:
        del os.environ["EXTRACT_QUOTE_RETRY"]

    # switch off (the default): no follow-up call, the first answer's judgment stands
    x.call_llm = qret_llm
    qcalls.clear()
    out = x.extract([qret_page], [1], dm, {"fiscal_year": 2025})
    assert qcalls == ["extraction"], qcalls
    w1y = next(f for f in out["fields"] if f["key"] == "due_within_1_year")
    assert (w1y["value"], w1y["confidence"]) == (3538, 0.25) and not any(w.startswith("quote_retry:") for w in out["warnings"]), (w1y, out["warnings"])
    # v052: a note page can stack a second table above the one the field's row sits in (MedCap p.101: a
    # receivables-ageing table over the maturity table), and the maturity header itself repeats the year across
    # Koncernen | Moderbolaget pairs. Row derivations now anchor the year header to the quoted row -- the nearest
    # year run above it (_row_year_column) -- and, when the fiscal year appears twice in that run, take the Group
    # side when Group is named before Parent on the run's row or the one above it, the last pair when reversed.
    # The page-level _year_column default is untouched: still the page's first run, still None on repeats.
    assert x._year_column(medcap101, 2025) is None  # the page-level default still declines a repeated year, as before
    age101 = ("MSEK 2025-12-31 2024-12-31\nEj förfallet 241,6 180,5\nMindre än 3 månader 25,0 29,1\n"
              "Äldre än 3 månader 1,8 1,8\nAvsättningar -1,1 -0,8\nTotalt 267,3 210,6\n")  # the receivables-ageing table printed above the maturity one on the real p.101
    full101 = age101 + medcap101
    assert x._year_column(full101, 2025) == (0, 2)  # unchanged page-level default: the page's FIRST header, the ageing table's
    rows101 = x._page_rows(full101)
    assert x._row_year_column(rows101, rows101.index("Totalt 718,7 377,6 – –"), 2025) == (0, 4)  # the maturity table's own header, Group side first
    assert x._row_year_column(rows101, rows101.index("Totalt 267,3 210,6"), 2025) == (0, 2)  # the ageing row anchors to its own table
    norows101 = x._page_rows(full101.replace(" Koncernen Moderbolaget", ""))
    assert x._row_year_column(norows101, norows101.index("Totalt 718,7 377,6 – –"), 2025) is None  # repeated year, no Group/Parent words: unknowable, declines as before
    catrows = x._page_rows(cat)
    assert x._row_year_column(catrows, catrows.index("Profit before tax 2,067 1,344"), 2025) == (0, 2)  # single-table page: the anchored header is the page's
    # ... end to end: the model computes due_within_1_year = 54,3 + 48,0 = 102,3 correctly but fabricates its quote,
    # so the value is dropped as computed, not read -- yet both rows are printed and are now the schema's own
    # synonyms, so the table-scoped guard keeps the null operand out (missing), v052's above-window derivation
    # fills 102.3 with the verified two-row quote, and once due_after_5_years is a confirmed 0 (its ">5år" lives
    # on the contractual table, which no longer counts) the identity closes -- passed, where v044/v052 stalled
    # at "missing: due_after_5_years".
    x.call_llm = lambda *a, **k: {"fields": [
        {"key": "total_debt", "value": 718.7, "unit": "MSEK", "period": "2025", "raw_label": "Totalt", "source": {"page": 1, "quote": "Totalt 718,7 377,6 – –"}},
        {"key": "due_1_to_5_years", "value": 616.5, "unit": "MSEK", "period": "2025", "raw_label": "1 – 5 år", "source": {"page": 1, "quote": "1 – 5 år 616,5 316,7 – –"}},
        {"key": "due_within_1_year", "value": 102.3, "unit": "MSEK", "period": "2025", "raw_label": "6 månader eller mindre", "source": {"page": 1, "quote": "6 månader eller mindre 54,3 + 6 – 12 månader 48,0 = 102,3"}},
        {"key": "due_after_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None}]}
    out = x.extract([full101, medcap102], [1, 2], dm, {"fiscal_year": 2025})
    within = next(f for f in out["fields"] if f["key"] == "due_within_1_year")
    assert within["value"] == 102.3 and "value_derived" in within["evidence"] and within["source"]["quote"] \
        == "6 månader eller mindre 54,3 41,8 – – 6 – 12 månader 48,0 19,0 – –", (within, out["warnings"])
    assert within["confidence"] == 1.0 and within["raw_label"] == "6 månader eller mindre + 6 – 12 månader", within
    assert out["checks"][0]["passed"] and "abs((102.3 + 616.5 + 0 (due_after_5_years null)) - 718.7) <= 2" in out["checks"][0]["detail"], out["checks"]
    total = next(f for f in out["fields"] if f["key"] == "total_debt")
    bucket15 = next(f for f in out["fields"] if f["key"] == "due_1_to_5_years")
    assert (total["value"], total["confidence"]) == (718.7, 0.9) and (bucket15["value"], bucket15["confidence"]) == (616.5, 1.0), (total, bucket15)
    # v102: a date-form year header ("SEK million 31 Dec 2025 31 Dec 2024", Ambea p.119 Note G18) names its
    # year columns: each year may carry its balance-date wording (day + month name, either order, "31 Dec" /
    # "31 dec." / "31 december" / "Dec 31,"; the ISO and slash spellings the _year_run subs already reduce),
    # the day/month tokens are neither amount columns nor run-breakers, and the years keep print order -- the
    # first column is the first year printed. Before, the day digit broke the gap rule, the header formed no
    # run at all, and the page scan fell through to a bare-year pair in prose ("... issued in 2024 or 2025",
    # the note's own last sentence) -- which named 2025 as the SECOND column, so the year correction rewrote
    # the model's correct first-column figures into the prior-year comparatives (12,643 -> 10,757,
    # 2,308 -> 1,879; docs/acrylic/evidence/v099.md Finding 1, v088's own false "year-column corrected").
    # The reduction lives in _year_column/_row_year_column only: _year_run itself (and _segment_column,
    # _bucket_row_prior_year, _operand_table_rows through it) keep seeing the header as printed.
    assert x._year_run("SEK million 31 Dec 2025 31 Dec 2024") == []  # unchanged, by design: the selectors reduce, not the scan
    assert x._year_column("SEK million 31 Dec 2025 31 Dec 2024", 2025) == (0, 2)
    assert x._row_year_column(x._page_rows("SEK million 31 Dec 2025 31 Dec 2024\nTotal interest-bearing liabilities 12,643 10,757\n"), 1, 2025) == (0, 2)
    assert x._year_column("MSEK 31 december 2024 31 december 2025", 2025) == (1, 2)  # prior year printed first: order is print order
    assert x._year_column("USD million Dec 31, 2025 Dec 31, 2024", 2025) == (0, 2)  # month-first spelling
    assert x._year_column("MSEK per 31 Dec. 2025 per 31 Dec. 2024", 2025) == (0, 2)  # per/at prefixes ride the one-word gap rule, digits were the only breaker
    assert x._year_column("SEK 000 31 Dec. 2025 31 Dec. 2024", 2025) == (0, 2)  # Ependion p.155's dotted abbreviation
    # counterexamples, unchanged: a bare-year header, a lone year (no run -- "Note 20 2025" is one column),
    # Cloetta's single-date bucket note (one date, one column, still the bucket-column mechanism's own page),
    # and a non-December balance date, which stays a non-run because its calendar year is NOT the fiscal year:
    # Rusta's FY2025 ends 30 Apr 2026, so "30 Apr 2025" names the PRIOR year and the report's own fiscal-year
    # header ("2025/26 2024/25", the _SPLIT_YEAR rule) remains the authority (docs/acrylic/evidence/v096.md).
    assert x._year_column("Note 20 Borrowings\nMSEK\n2025 2024", 2025) == (0, 2)
    assert x._year_column("Note 20 Borrowings\nMSEK\n2025", 2025) is None
    assert x._year_column("Note 21 Borrowings 31 Dec 2025", 2025) is None
    assert x._year_column("Maturity analysis 30 Apr 2026 30 Apr 2025", 2025) is None
    # repeated date-form pairs (Proact p.103 prints Group | Parent side by side) flow into the v052 machinery
    # unchanged: the page-level default declines a repeated fiscal year, and the row-anchored header takes the
    # Group side when Group is named before Parent on the run's row or the one above it
    proact_hdr = x._page_rows("Other financial liabilities Group Parent company\n"
                              "acquisition value 31 Dec 2025 31 Dec 2024 31 Dec 2025 31 Dec 2024\n"
                              "Total interest-bearing 478,611 483,465 216,360 229,730\n")
    proact_page = "\n".join(proact_hdr)
    assert x._year_column(proact_page, 2025) is None
    assert x._row_year_column(proact_hdr, 2, 2025) == (0, 4)
    # ... end to end on Ambea's real p.119 (trimmed to Note G18; the model answer is v088 pass 1's correct
    # read): the figures stay in the first column -- no "is not the 2025 column" correction, period_ok and
    # value_in_quote on both fields -- and the current/non-current shape leaves the buckets honestly missing.
    # Runs on the SHIPPED schema (require_explicit_values on: a null bucket stays "missing", the stored
    # extraction's own convention), not the flag-off copy `dm` above.
    ambea119 = ("NOTE G18 Interest-bearing liabilities\nSEK million 31 Dec 2025 31 Dec 2024\n"
                "Non-current liabilities\nLiabilities to credit institutions 2,115 1,087\n"
                "Non-current lease liabilities 8,220 7,791\nTotal non-current interest-bearing liabilities 10,335 8,878\n"
                "Current liabilities\nCommercial paper 1,232 1,039\nCurrent lease liabilities 1,076 840\n"
                "Total current interest-bearing liabilities 2,308 1,879\nTotal interest-bearing liabilities 12,643 10,757\n"
                "rate fluctuations as well as payback periods are presented in Note G26\n"
                "company's participations in subsidiaries was issued in 2024 or 2025.\n")
    x.call_llm = lambda *a, **k: {"fields": [
        {"key": "total_debt", "value": 12643, "unit": "SEK million", "period": "2025", "raw_label": "Total interest-bearing liabilities",
         "source": {"page": 1, "quote": "Total interest-bearing liabilities 12,643 10,757"}},
        {"key": "due_within_1_year", "value": 2308, "unit": "SEK million", "period": "2025", "raw_label": "Total current interest-bearing liabilities",
         "source": {"page": 1, "quote": "Total current interest-bearing liabilities 2,308 1,879"}},
        {"key": "due_1_to_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None},
        {"key": "due_after_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None}]}
    dm_ship = json.loads((pathlib.Path(__file__).parents[1] / "schemas" / "debt_maturity.json").read_text("utf-8"))
    out = x.extract([ambea119], [1], dm_ship, {"fiscal_year": 2025})
    got = {f["key"]: (f["value"], f["confidence"]) for f in out["fields"]}
    assert got == {"total_debt": (12643, 1.0), "due_within_1_year": (2308, 0.9),
                   "due_1_to_5_years": (None, 0.0), "due_after_5_years": (None, 0.0)}, (got, out["warnings"], out["checks"])
    assert not any("is not the 2025 column" in w for w in out["warnings"]), out["warnings"]
    td = next(f for f in out["fields"] if f["key"] == "total_debt")
    assert "value_in_quote" in td["evidence"] and "period_ok" in td["evidence"], td
    assert out["checks"][0]["detail"] == "missing: due_1_to_5_years", out["checks"]
    # v051 (seed 5's live Codex rerun of this same page): a related but distinct gap from the fabricated-quote
    # case just above. There, the model's quote never verifies at all (no "+"/"=" arithmetic notation is
    # printed), so due_within_1_year is null by the time the "missing:"-driven repair loop runs, and the
    # PRE-EXISTING _between_rows call above already rescues it. Here the model instead quotes the two real rows
    # concatenated as if printed together -- a real quote_on_page match (they really are contiguous on the
    # page) -- so the field survives the first per-field gate with a value (102.3, truncated by the "recover the
    # full row" step down to just the second row's own quote). Once due_within_1_year is itself numeric, it
    # counts toward the check's own missing-operand detection, which reported "missing: due_after_5_years" (the
    # one bucket that really has no row) -- so the "missing:"-only repair loop above never even looks at
    # due_within_1_year, and it would otherwise reach the Sectra-pattern drop below with nothing to rescue it.
    # v058: the month rows are now the schema's own synonyms and the present-label search is table-scoped, so
    # due_after_5_years is a confirmed 0 and the check this case reports PASSES on the very value the
    # concatenation rescue just secured: 102.3+616.5+0 = 718.8 ≈ 718.7.
    x.call_llm = lambda *a, **k: {"fields": [
        {"key": "total_debt", "value": 718.7, "unit": "MSEK", "period": "2025", "raw_label": "Totalt", "source": {"page": 1, "quote": "Totalt 718,7 377,6 – –"}},
        {"key": "due_1_to_5_years", "value": 616.5, "unit": "MSEK", "period": "2025", "raw_label": "1 – 5 år", "source": {"page": 1, "quote": "1 – 5 år 616,5 316,7 – –"}},
        {"key": "due_within_1_year", "value": 102.3, "unit": "MSEK", "period": "2025", "raw_label": "Förfallotidpunkt för upplåning",
         "source": {"page": 1, "quote": "6 månader eller mindre 54,3 41,8 – – 6 – 12 månader 48,0 19,0 – –"}},
        {"key": "due_after_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None}]}
    out = x.extract([full101, medcap102], [1, 2], dm, {"fiscal_year": 2025})
    within = next(f for f in out["fields"] if f["key"] == "due_within_1_year")
    assert within["value"] == 102.3 and "value_derived" in within["evidence"] and within["confidence"] == 1.0 \
        and within["source"]["quote"] == "6 månader eller mindre 54,3 41,8 – – 6 – 12 månader 48,0 19,0 – –", (within, out["warnings"])
    assert out["checks"][0]["passed"] and "abs((102.3 + 616.5 + 0 (due_after_5_years null)) - 718.7) <= 2" in out["checks"][0]["detail"], out["checks"]
    # ... and a repeated year with no Group/Parent named above the rows must not fire: Volvo-style segment
    # repetition stays unknowable, the field stays null exactly as before the anchoring existed
    out = x.extract([full101.replace(" Koncernen Moderbolaget", ""), medcap102], [1, 2], dm, {"fiscal_year": 2025})
    within = next(f for f in out["fields"] if f["key"] == "due_within_1_year")
    assert within["value"] is None and within["confidence"] == 0.0 and not within["evidence"], (within, out["warnings"])
    assert not out["checks"][0]["passed"], out["checks"]
    # v058: v052's above-window derivation also fires when the within-1-year wording sits in the table's own
    # year-header row ("inom 1 år", glued to "inom1år" like v012/v014) rather than on the bucket rows -- the
    # guard keeps the null operand out either way, and the real table here is only replayed with one header
    # word changed, nothing else.
    medcap101_inom = medcap101.replace(
        "MSEK 2025-12-31 2024-12-31 2025-12-31 2024-12-31", "MSEK inom 1 år 2025-12-31 2024-12-31 2025-12-31 2024-12-31")
    x.call_llm = lambda *a, **k: {"fields": [
        {"key": "total_debt", "value": 718.7, "unit": "MSEK", "period": "2025", "raw_label": "Totalt", "source": {"page": 1, "quote": "Totalt 718,7 377,6 – –"}},
        {"key": "due_1_to_5_years", "value": 616.5, "unit": "MSEK", "period": "2025", "raw_label": "1 – 5 år", "source": {"page": 1, "quote": "1 – 5 år 616,5 316,7 – –"}},
        {"key": "due_within_1_year", "value": 102.3, "unit": "MSEK", "period": "2025", "raw_label": "6 månader eller mindre", "source": {"page": 1, "quote": "6 månader eller mindre 54,3 + 6 – 12 månader 48,0 = 102,3"}},
        {"key": "due_after_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None}]}
    out = x.extract([age101 + medcap101_inom, medcap102], [1, 2], dm, {"fiscal_year": 2025})
    within = next(f for f in out["fields"] if f["key"] == "due_within_1_year")
    assert within["value"] == 102.3 and "value_derived" in within["evidence"] and within["source"]["quote"] \
        == "6 månader eller mindre 54,3 41,8 – – 6 – 12 månader 48,0 19,0 – –", (within, out["warnings"])
    assert within["confidence"] == 1.0 and within["raw_label"] == "6 månader eller mindre + 6 – 12 månader", within
    assert out["checks"][0]["passed"] and "abs((102.3 + 616.5 + 0 (due_after_5_years null)) - 718.7) <= 2" in out["checks"][0]["detail"], out["checks"]
    # v058's synthetic discriminator, MedCap's carrying table (the two month rows collapsed into their single
    # "< 1 år" sum row and the Parent pair dropped, so the joined-quote/_segment_column detours stay out of the
    # way) plus the contractual cash-flow table relocated onto ONE page BELOW it: the bucket word sits past the
    # Totalt row the operands anchor, so both the page-level rule (found -> stay open) and the bottom-of-table
    # boundary are exercised at once. Scoped to the operands' own table the null bucket is a confirmed 0:
    # 102.3+616.5+0 = 718.8 ≈ 718.7.
    onepage = ("Förfallotidpunkt för upplåning\nMSEK 2025 2024\n"
               "< 1 år 102,3 60,8\n1 – 5 år 616,5 316,7\nTotalt 718,7 377,6\n"
               "Koncernen 2025-12-31\nMSEK <1 år 1 – 2 år 2 – 4 år 4 – 5 år >5år\nSkulder till kreditinstitut 59,3 97,5 59,1 – –\n"
               "Leasingskulder 50,6 100,8 66,4 30,6 142,9\nTotalt 431,8 385,4 125,5 30,6 142,9\n")
    x.call_llm = lambda *a, **k: {"fields": [
        {"key": "total_debt", "value": 718.7, "unit": "MSEK", "period": "2025", "raw_label": "Totalt", "source": {"page": 1, "quote": "Totalt 718,7 377,6"}},
        {"key": "due_within_1_year", "value": 102.3, "unit": "MSEK", "period": "2025", "raw_label": "< 1 år", "source": {"page": 1, "quote": "< 1 år 102,3 60,8"}},
        {"key": "due_1_to_5_years", "value": 616.5, "unit": "MSEK", "period": "2025", "raw_label": "1 – 5 år", "source": {"page": 1, "quote": "1 – 5 år 616,5 316,7"}},
        {"key": "due_after_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None}]}
    out = x.extract([onepage], [1], dm, {"fiscal_year": 2025})
    assert out["checks"][0]["passed"] and "0 (due_after_5_years null)" in out["checks"][0]["detail"], out["checks"]
    got = {f["key"]: (f["value"], f["confidence"]) for f in out["fields"]}
    assert got == {"total_debt": (718.7, 1.0), "due_within_1_year": (102.3, 1.0), "due_1_to_5_years": (616.5, 1.0), "due_after_5_years": (None, 0.0)}, (got, out["warnings"])  # total_debt: the identity holds with the bare-Totalt row in both columns -> identity_all_columns
    # v043: EXTRACT_TWO_PASS (default off) -- pass 1 is a small "which candidate page" question over a
    # head-of-page snippet of *every* candidate (not just the top-2 the single-pass window shows), pass 2
    # is the existing extraction prompt fed only the page(s) pass 1 picks. Candidate 3 (never reached by
    # the old pages[:2] window) is the real note; 1-2 are decoys.
    import os
    decoy1 = "Five-year summary\nMSEK\n2025 2024 2023 2022 2021\nNet sales 100 90 80 70 60\n"
    decoy2 = "Segment overview\nNothing about borrowings here.\n"
    real = ("Note 20 Borrowings\nMSEK\n2025 2024\nTotal borrowings 32 703 34 500\nWithin 1 year 3 538 4 100\n"
            "1-5 years 29 165 30 000\n1-5 years, of which fixed rate 10 000 9 000\n")
    decoy4 = "Appendix\nUnrelated content, never a pass-1 candidate.\n"  # page 4: not in `pages`, only reachable as page 3's default +1 companion
    full_answer = {"fields": [
        {"key": "total_debt", "value": 32703, "unit": "MSEK", "period": "2025", "raw_label": "Total borrowings", "source": {"page": 3, "quote": "Total borrowings 32,703"}},
        {"key": "due_within_1_year", "value": 3538, "unit": "MSEK", "period": "2025", "raw_label": "Within 1 year", "source": {"page": 3, "quote": "Within 1 year 3,538"}},
        {"key": "due_1_to_5_years", "value": 29165, "unit": "MSEK", "period": "2025", "raw_label": "1-5 years", "source": {"page": 3, "quote": "1-5 years 29,165"}},
        {"key": "due_after_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None}]}
    calls = []

    def fake_llm(system, user, schema=x.RESPONSE_SCHEMA, name="extraction"):
        calls.append((name, "=== PAGE 1 ===" in user, "=== PAGE 3 ===" in user))
        return {"pages": [3]} if name == "page_select" else full_answer
    x.call_llm = fake_llm

    calls.clear()
    out = x.extract([decoy1, decoy2, real, decoy4], [1, 2, 3], dm, {"fiscal_year": 2025})
    assert calls == [("extraction", True, False)], calls  # off: exactly the old single call, over pages[:2]; page 3 never shown
    assert not any(w.startswith("two_pass") for w in out["warnings"]), out["warnings"]
    # v157: ONE hole in the identity (this answer's due_after_5_years -- the ordinary report with no
    # >5-year row) keeps that single call. TWO holes is the shape of a wrong table read confidently
    # (BTS p.93's non-current-only note: a total and one bucket, nothing closed) and buys one widened
    # call over pages[:4], which is where the real maturity note sat
    two_holes = {"fields": [f if f["key"] != "due_within_1_year" else
                            {"key": "due_within_1_year", "value": None, "unit": None, "period": None,
                             "raw_label": None, "source": None} for f in full_answer["fields"]]}
    x.call_llm = lambda system, user, schema=x.RESPONSE_SCHEMA, name="extraction": (
        calls.append((name, "=== PAGE 1 ===" in user, "=== PAGE 3 ===" in user)), two_holes)[1]
    calls.clear()
    x.extract([decoy1, decoy2, real, decoy4], [1, 2, 3], dm, {"fiscal_year": 2025})
    assert calls == [("extraction", True, False), ("extraction", True, True)], calls
    x.call_llm = fake_llm

    os.environ["EXTRACT_TWO_PASS"] = "1"
    try:
        calls.clear()
        out = x.extract([decoy1, decoy2, real, decoy4], [1, 2, 3], dm, {"fiscal_year": 2025})
        assert calls == [("page_select", True, True), ("extraction", False, True)], calls  # pass 1 sees every candidate; pass 2 only the pages it picked
        # v045: a single-page reply ({"pages": [3]}) is no longer taken at face value -- it is repaired with
        # the default companion (3+1=4, page 4 was never itself a candidate), so the companion page single-pass
        # always got for free (locate.candidate_pages forces pages[0]+1 into position 2) is not lost here either.
        assert "two_pass: page [3, 4] selected from candidates [1, 2, 3]" in out["warnings"], out["warnings"]
        got = {f["key"]: (f["value"], f["confidence"]) for f in out["fields"]}
        assert got["total_debt"] == (32703, 1.0) and got["due_within_1_year"] == (3538, 1.0) and got["due_1_to_5_years"] == (29165, 1.0), (got, out["warnings"])

        def illegal_llm(system, user, schema=x.RESPONSE_SCHEMA, name="extraction"):
            calls.append((name, "=== PAGE 1 ===" in user, "=== PAGE 3 ===" in user))
            return {"pages": [99]} if name == "page_select" else full_answer  # 99 is not a candidate: illegal
        x.call_llm = illegal_llm
        calls.clear()
        out = x.extract([decoy1, decoy2, real, decoy4], [1, 2, 3], dm, {"fiscal_year": 2025})
        assert calls == [("page_select", True, True), ("extraction", True, False)], calls  # illegal pick falls back to pages[:2]
        assert any("page selection failed or illegal" in w for w in out["warnings"]), out["warnings"]
    finally:
        del os.environ["EXTRACT_TWO_PASS"]

    # v045: _select_pages itself -- the companion repair/validation rules, direct (no EXTRACT_TWO_PASS/extract()
    # wiring needed here, unlike the end-to-end checks above).
    four = ["p1 text", "p2 text", "p3 text", "p4 text"]
    x.call_llm = lambda *a, **k: {"pages": [2]}
    assert x._select_pages(dm, [1, 2, 3], four) == [2, 3]  # single-page reply repaired with the default +1
    x.call_llm = lambda *a, **k: {"pages": [4]}
    assert x._select_pages(dm, [1, 2, 4], four) == [3, 4]  # primary is the last page: +1 doesn't exist, falls back to -1
    x.call_llm = lambda *a, **k: {"pages": [2, 1]}
    assert x._select_pages(dm, [1, 2, 3], four) == [1, 2]  # the model's own -1 override, honoured as given
    x.call_llm = lambda *a, **k: {"pages": [2, 4]}
    assert x._select_pages(dm, [1, 2, 3, 4], four) is None  # companion not adjacent to the primary: illegal
    x.call_llm = lambda *a, **k: {"pages": [2, 3]}
    assert x._select_pages(dm, [1, 2], four) == [2, 3]  # companion (3) need not itself be a candidate -- only the primary must be
    x.call_llm = lambda *a, **k: {"pages": [99]}
    assert x._select_pages(dm, [1, 2, 3], four) is None  # primary outside the candidate list: illegal, same as v043
    x.call_llm = lambda *a, **k: {"pages": [2, 3, 4]}
    assert x._select_pages(dm, [1, 2, 3, 4], four) is None  # more than 2 pages: illegal, unchanged from v043

    # v045: pass-1's own snippet -- running headers (locate.strip_boilerplate) and bare page-number lines
    # dropped, ~1200-char head, plus (Apotea's own failure mode, docs/acrylic/evidence/v043.md) any schema
    # keyword line found beyond that head, numbered, so a heading pushed past the cutoff by filler still surfaces.
    captured = {}
    page_select_hints_env = os.environ.pop("PAGE_SELECT_HINTS", None)

    def capture_llm(system, user, schema=x.PAGE_SELECT_SCHEMA, name="page_select"):
        captured["system"] = system
        captured["user"] = user
        return {"pages": [1, 2]}
    x.call_llm = capture_llm
    running_header = "ACME GROUP ANNUAL REPORT 2025"
    filler = "\n".join(f"Filler line number {i} of running prose unrelated to the statement." for i in range(30))  # > PAGE_SELECT_SNIPPET chars
    p1 = f"{running_header}\n7\n{filler}\nNote 20 Borrowings\nTotal borrowings 32,703 34,500\n"
    p2 = f"{running_header}\n8\nSegment overview, nothing about borrowings here.\n"
    other_pages = [f"{running_header}\n{n}\nUnrelated page." for n in range(20, 29)]  # > 10% of the doc repeats running_header
    x._select_pages(dm, [1, 2], [p1, p2] + other_pages)
    user = captured["user"]
    assert running_header not in user, user  # running header stripped (locate.strip_boilerplate)
    assert "7" not in user.splitlines(), user  # bare page-number line stripped (not just "differs per page", see locate's own comment)
    assert "\n...\n" in user  # the head/keyword-hits separator only appears once something was found beyond the head
    assert any(l.endswith(": Note 20 Borrowings") for l in user.splitlines()), user  # the heading, past the head cutoff, surfaces numbered
    # v146 red: this is the page-select prompt assembly itself, before its zero-model page hints exist.
    # The compact pages below preserve the discriminating title/table lines of v124's real Avarda pages:
    # p.27 is the group balance sheet and p.81 says its liquidity table is contractual/undiscounted.
    avarda = [""] * 81
    avarda[26] = ("BALANCE SHEET - GROUP\nSEK thousand Note 31 Dec 2025 31 Dec 2024\n"
                  "TOTAL ASSETS 29,118,468 25,069,605\nLIABILITIES AND EQUITY\n"
                  "TOTAL LIABILITIES AND EQUITY 29,118,468 25,069,605\n")
    avarda[79] = "Remaining interest term to maturity\n31 Dec 2025 SEK thousand Up to 3 months\n"
    avarda[80] = ("Liquidity risk\nThe amounts shown in the table represent contractual, undiscounted liquidity flows.\n"
                  "Remaining maturity\n31 Dec 2025 SEK thousand Payable on demand Up to 3 months\n")
    prompt_dm = {**dm, "title": "Prompt test", "description": "Frozen prompt description."}
    captured.clear()
    x._select_pages(prompt_dm, [80, 81, 27], avarda)
    expected_default_page_select_prompt = """You are given the start of 3 candidate pages from a corporate annual report (Swedish or English), each labelled with its page number. Which page holds the Prompt test statement itself -- the printed table of figures -- not a table of contents, a note reference, or an unrelated table?

Frozen prompt description.

Always name TWO pages: the primary page (the one with the table itself, must be one of the candidates above) and a companion page next to it, since a table's header or rows often continue onto the neighbouring page. Default the companion to primary+1; use primary-1 instead only if the table's own heading or first rows actually sit on the page before the primary one -- the companion does not itself have to be one of the candidates above.

Return ONE JSON object {"pages": [primary, companion]}, primary first. Never invent a primary page number that is not listed above."""
    # v146-b: default off is the old rendered prompt byte-for-byte and never labels candidate pages.
    assert captured["system"].encode("utf-8") == expected_default_page_select_prompt.encode("utf-8"), captured["system"]
    assert "[balance sheet]" not in captured["user"], captured["user"]
    assert x.PAGE_SELECT_DEBT_MATURITY_HINT not in captured["system"], captured["system"]

    os.environ["PAGE_SELECT_HINTS"] = "1"
    captured.clear()
    x._select_pages(prompt_dm, [80, 81, 27], avarda)
    assert "=== PAGE 27 [balance sheet] ===" in captured["user"], captured["user"]
    assert x.PAGE_SELECT_DEBT_MATURITY_HINT in captured["system"], captured["system"]
    # v162: the retry route can override that process-wide opt-in for one of its two runs.
    captured.clear()
    x._select_pages(prompt_dm, [80, 81, 27], avarda, page_select_hints=False)
    assert "[balance sheet]" not in captured["user"], captured["user"]
    assert x.PAGE_SELECT_DEBT_MATURITY_HINT not in captured["system"], captured["system"]
    captured.clear()
    x._select_pages(prompt_dm, [80, 81, 27], avarda, page_select_hints=True)
    assert "=== PAGE 27 [balance sheet] ===" in captured["user"], captured["user"]
    assert x.PAGE_SELECT_DEBT_MATURITY_HINT in captured["system"], captured["system"]
    if page_select_hints_env is None:
        del os.environ["PAGE_SELECT_HINTS"]
    else:
        os.environ["PAGE_SELECT_HINTS"] = page_select_hints_env

    # v146 green: all fixtures are compact, frozen title/table lines from v124's cited real pages --
    # never a read of mutable data/kb (LESSONS 43). The range of expected tags proves title-zone scope:
    # Morrow p.74 and Synsam p.122/p.124 are relevant prose/adjacent pages but carry no qualifying title.
    assert x._page_select_tags(dm, [80, 81, 27], avarda) == {
        81: ["liquidity-risk / undiscounted"], 27: ["balance sheet"]}

    lime = [""] * 91
    lime[71] = ("Liquidity risk - Group\nAs of 31 December 2025 Less than 3 months Between 3 months and 1 year\n"
                "Borrowing (incl. overdraft) 15,000 51,396 60,000 25,000\n")
    lime[72] = "4. OPERATIONAL RISKS\nThe Group is exposed to various risks through its operations.\n"
    lime[90] = ("21. BANK LOANS, LEASE LIABILITIES AND OTHER LIABILITIES ADDITIONAL PURCHASE PRICE\n"
                "The Group The Parent\nNon-current liability 2025 2024 2025 2024\nBank loan 85,000 145,000 - 12,500\n"
                "Bank loan\nBank loans have been taken out by the parent company and run until 2026 and 2027.\n")
    assert x._page_select_tags(dm, [72, 73, 91], lime) == {
        72: ["liquidity-risk / undiscounted"], 91: ["borrowings note: carrying"]}

    morrow = [""] * 75
    morrow[42] = ("Balance sheet\nAmounts in NOK million Note\n31 December 2025 2024\n"
                  "Total assets 20,910.4 18,616.9\nTotal liabilities and equity 20,910.4 18,616.9\n")
    morrow[73] = ("Note 17 Liquidity and interest rate risk\nLiquidity risk\n"
                  "The liquidity risk of the Bank arises from or results from the maturity profile of the Bank's assets and liabilities. "
                  "Below follows an overview of different time intervals as to when the Bank's assets and liabilities mature.\n"
                  "DISTRIBUTION OF TERMS AS AT 31 DECEMBER 2024\nAmounts in NOK million\n"
                  "Loans and deposits with credit institutions 1,024.1 - - - - 1,024.1\n")
    morrow[74] = "Interest rate risk\nAmounts in NOK million 0 month < 3 months No interest Total\n"
    assert x._page_select_tags(dm, [74, 75, 43], morrow) == {43: ["balance sheet"]}

    synsam = [""] * 124
    synsam[121] = ("Loans from financial institutions amounted to SEK 2,718 million on 31 December 2025, of which current "
                   "liabilities amounted to SEK 0 million.\nThe loans are payable in full upon maturity and run for five years.\n")
    synsam[122] = ("Maturity structures for financial liabilities – undiscounted cash flows\nGroup\n"
                   "31 Dec 2025 Interest rate Less than 1 year 1–5 years More than 5 years Total\n"
                   "Loans from financial institutions 3.40–3.57% 94 3,060 – 3,154\n")
    synsam[123] = ("Loans from financial institutions by currency\nGroup\n31 Dec 2025 31 Dec 2024\n"
                   "Loans from financial institutions 2,733 2,619\n")
    assert x._page_select_tags(dm, [123, 124, 122], synsam) == {123: ["liquidity-risk / undiscounted"]}
    assert "Pages tagged liquidity-risk / undiscounted list contractual cash flows" in x.PAGE_SELECT_DEBT_MATURITY_HINT
    # v050: a prose no-debt statement has no digit for _value_in_quote, so both provenance gates dropped the
    # model's 0 as "computed, not read" (v048's recorded warning for Creades, seed 4, p.61 real text). The
    # total_debt schema field now opts in with "zero_if_stated" (the field-level analogue of a check's
    # null_as_zero): a digit-free quote that sits verbatim on the cited page (whitespace/NBSP-insensitive;
    # quote_on_page itself requires a number token, which prose never has), names the field's subject (a
    # schema keyword or one of the field's synonyms) and carries a negation word from the schema's own list,
    # proves a 0 the report stated in words. The buckets get no opt-in (their descriptions forbid a 0 for an
    # unprinted bucket), so they stay null -- and _check lifts its all-null guard only when the identity's
    # remaining operand is itself such a stated zero, making 0+0+0 == 0 a real pass.
    creades61 = ("Noter 59\nCreades årsredovisning 2025\n"
                 "Not 18 Klassificering av finansiella instrument värderade till verkligt värde\n"
                 "Enligt IFRS 9 ska ett företag klassificera sina finansiella tillgångar och skulder. Creades klassificering av sina finansiella tillgångar och skulder\n"
                 "framgår av följande matris.\n"
                 "Likvida medel, kundfordringar och leverantörsskulder har kort löptid och bedöms ha ett upplupet anskaffningsvärde som inte avviker\n"
                 "väsentligt från verkligt värde. Investmentföretaget har varken räntebärande skulder eller kundfordringar.\n"
                 "Finansiella tillgångar 251231, per värderingskategori enligt IFRS 9\nSEK mn\nTillgångar\n"
                 "Andelar i portföljbolag 10 879 – 10 879 10 879\nLångfristig fordran 72 – 72 72\nLikvida medel – 287 287 287\n"
                 "Summa tillgångar 10 951 287 11 237 11 237\n"
                 "Indelning i hierarkiska nivåer\nTillgångar och skulder värderade till verkligt värde via resultatet\n"
                 "Creades har inga finansiella tillgångar eller skulder hänförliga\ntill Nivå 2.\n")
    creades62 = "60 Noter\nCreades årsredovisning 2025\nVärderingen görs vanligen genom principen ”Multipelvärdering”,\n"
    creades_sentence = "Investmentföretaget har varken räntebärande skulder eller kundfordringar."

    def zero_answer():
        return {"fields": [
            {"key": "total_debt", "value": 0, "unit": "SEK mn", "period": "2025",
             "raw_label": creades_sentence, "source": {"page": 1, "quote": creades_sentence}},
            {"key": "due_within_1_year", "value": None, "unit": None, "period": None, "raw_label": None, "source": None},
            {"key": "due_1_to_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None},
            {"key": "due_after_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None}]}

    x.call_llm = lambda *a, **k: zero_answer()
    out = x.extract([creades61, creades62], [1, 2], dm, {"fiscal_year": 2025})
    td = next(f for f in out["fields"] if f["key"] == "total_debt")
    assert td["value"] == 0 and td["confidence"] == 0.5 and "stated_zero" in td["evidence"] \
        and td["source"]["quote"] == creades_sentence, (td, out["warnings"])  # 0.50: the sentence is on the page, the figure never is
    assert out["checks"][0]["passed"] and "null" in out["checks"][0]["detail"], out["checks"]  # 0+0+0 == 0, every operand named
    assert all(next(f for f in out["fields"] if f["key"] == k)["value"] is None for k in dmf if k != "total_debt"), out["fields"]  # buckets stay null
    # off-page quote: the negation sentence not printed on the cited page -> still "computed, not read"
    off = zero_answer()
    off["fields"][0]["source"]["quote"] = "Bolaget har absolut inga räntebärande skulder."
    x.call_llm = lambda *a, **k: off
    out = x.extract([creades61, creades62], [1, 2], dm, {"fiscal_year": 2025})
    td = next(f for f in out["fields"] if f["key"] == "total_debt")
    assert td["value"] is None and any("dropped as computed, not read" in w for w in out["warnings"]), (td, out["warnings"])
    # on-page quote about something else: "Creades klassificering ..." says skulder but holds no total_debt vocabulary word
    off = zero_answer()
    off["fields"][0]["source"]["quote"] = "Creades klassificering av sina finansiella tillgångar och skulder"
    x.call_llm = lambda *a, **k: off
    out = x.extract([creades61, creades62], [1, 2], dm, {"fiscal_year": 2025})
    td = next(f for f in out["fields"] if f["key"] == "total_debt")
    assert td["value"] is None and any("dropped as computed, not read" in w for w in out["warnings"]), (td, out["warnings"])
    # the buckets have no opt-in: 0 + the very same sentence still drops, the buckets' own "null if the table
    # has no such row" rule stands, and the identity reads missing, never 0 == nothing
    bucket = zero_answer()
    bucket["fields"][0] = {"key": "total_debt", "value": None, "unit": None, "period": None, "raw_label": None, "source": None}
    bucket["fields"][1] = {"key": "due_within_1_year", "value": 0, "unit": "SEK mn", "period": "2025",
                           "raw_label": creades_sentence, "source": {"page": 1, "quote": creades_sentence}}
    x.call_llm = lambda *a, **k: bucket
    out = x.extract([creades61, creades62], [1, 2], dm, {"fiscal_year": 2025})
    w1y = next(f for f in out["fields"] if f["key"] == "due_within_1_year")
    assert w1y["value"] is None and any(w.startswith("due_within_1_year: 0 is printed on none of pages") for w in out["warnings"]), (w1y, out["warnings"])
    assert not out["checks"][0]["passed"] and out["checks"][0]["detail"].startswith("missing:"), out["checks"]
    # the negation list lives in the schema, and it is what separates a stated zero from a bare row label:
    # "Summa räntebärande skulder" quoted alone off a column-major table (XANO's 8-column row) has the
    # vocabulary and sits verbatim on the page, but no negation word -- it must stay a drop, never become a 0
    assert not x._stated_zero({"value": 0, "source": {"page": 1, "quote": "Summa räntebärande skulder"}},
                              dmf["total_debt"], dm, ["Summa räntebärande skulder 4 090 8 680 38 305 51 075 768 885 33 784 50 778 904 522"])
    # a digit-bearing quote is the original gates' business, unchanged (Svolder, seed 4, real p.18 sentence +
    # p.19's printed axis 0): the quote verifies through quote_on_page, the 0 lands at 0.5 with no stated_zero
    svol18 = ("Riskhantering\nMSEK\ntillgångarnas värde och har en kreditfacilitet på upp \n"
              "till 500 MSEK hos en nordisk affärsbank, med aktier \n"
              "som säkerhet. Vid balansdagen den 31 augusti 2025 \nvar krediten oanvänd.\n")
    svol19 = "Not 19 Rörelsen\nMSEK\n−200 −100 0 100 200 300 400 500\n"
    svol = {"fields": [
        {"key": "total_debt", "value": 0, "unit": "MSEK", "period": "2025", "raw_label": "krediten oanvänd",
         "source": {"page": 1, "quote": "Vid balansdagen den 31 augusti 2025 var krediten oanvänd."}},
        {"key": "due_within_1_year", "value": None, "unit": None, "period": None, "raw_label": None, "source": None},
        {"key": "due_1_to_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None},
        {"key": "due_after_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None}]}
    x.call_llm = lambda *a, **k: json.loads(json.dumps(svol))
    out = x.extract([svol18, svol19], [1, 2], dm, {"fiscal_year": 2025})
    td = next(f for f in out["fields"] if f["key"] == "total_debt")
    assert td["value"] == 0 and td["confidence"] == 0.5 and "stated_zero" not in td["evidence"] \
        and "quote_on_page" in td["evidence"], (td, out["warnings"])  # Svolder's own stored read, byte-identical
    assert not out["checks"][0]["passed"] and out["checks"][0]["detail"].startswith("missing:"), out["checks"]
    # Flat Capital (v035, seed 1, real p.18): the no-debt sentence sits on p.33, never a locator candidate
    # ([18, 19] were) -- a p.33 sentence cited against a candidate page is not printed there, so the new rule
    # must not fire; p.18's own printed "interest-bearing liabilities ... 0 TSEK (0)" keeps the value on the
    # old printed-zero path, at the no-provenance 0.25 cap, whichever round of code ran
    fc18 = ("of 34,522 KSEK (0). See also Note 3 regarding changes in \n"
            "to 144,075 KSEK (158,832), of which interest-bearing liabili-\nties amounted to 0 TSEK (0).\n")
    fc = {"fields": [
        {"key": "total_debt", "value": 0, "unit": "KSEK", "period": "2025", "raw_label": "interest-bearing liabilities",
         "source": {"page": 1, "quote": "the investment company has neither interest-bearing liabilities nor accounts receivable."}},
        {"key": "due_within_1_year", "value": None, "unit": None, "period": None, "raw_label": None, "source": None},
        {"key": "due_1_to_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None},
        {"key": "due_after_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None}]}
    x.call_llm = lambda *a, **k: json.loads(json.dumps(fc))
    out = x.extract([fc18], [1], dm, {"fiscal_year": 2025})
    td = next(f for f in out["fields"] if f["key"] == "total_debt")
    assert td["value"] == 0 and td["confidence"] == 0.25 and "stated_zero" not in td["evidence"], (td, out["warnings"])
    assert any("quote not found on page 1" in w for w in out["warnings"]), out["warnings"]
    # v062: the subject gate (e) also reads zero_if_stated.subject_terms -- bare, generic debt words real
    # no-debt prose uses but the label vocabulary must not (v051: Vicore Pharma and BioGaia, seed 5, both
    # declined because their sentences carry only bare "loan(s)"; widening the field's own synonyms would
    # loosen _label_known's row matching). Real seed5 page text, the exact sentences v051 recorded,
    # constructed model answers -- the other four gates untouched, so this only opens sentences that are
    # verbatim on the cited page AND carry a negation word, at the same 0.5 tier as Creades.
    vicore46 = ("Refinancing risk refers to the risk that cash and cash equivalents are unavailable and that financing can\n"
                "only be obtained partially, not at all or at an elevated cost. Currently, the group is financed by shareholders’\n"
                "equity and is therefore not exposed to risks related to external loan financing. The main risks therefore\n"
                "entail the inability to obtain further equity investments from Vicore’s shareholders.\n")
    biogaia153 = "The group has no external loans. Excess liquidity is invested mainly in banks.\n"
    vicore_sentence = "Currently, the group is financed by shareholders’ equity and is therefore not exposed to risks related to external loan financing."
    biogaia_sentence = "The group has no external loans."

    def zero_loan_answer(sentence):
        return {"fields": [
            {"key": "total_debt", "value": 0, "unit": "SEK mn", "period": "2025",
             "raw_label": sentence, "source": {"page": 1, "quote": sentence}},
            {"key": "due_within_1_year", "value": None, "unit": None, "period": None, "raw_label": None, "source": None},
            {"key": "due_1_to_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None},
            {"key": "due_after_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None}]}

    x.call_llm = lambda *a, **k: zero_loan_answer(vicore_sentence)
    out = x.extract([vicore46], [1], dm, {"fiscal_year": 2025})
    td = next(f for f in out["fields"] if f["key"] == "total_debt")
    assert td["value"] == 0 and td["confidence"] == 0.5 and "stated_zero" in td["evidence"], (td, out["warnings"])
    assert out["checks"][0]["passed"] and "null" in out["checks"][0]["detail"], out["checks"]  # 0+0+0 == 0, buckets stay null
    x.call_llm = lambda *a, **k: zero_loan_answer(biogaia_sentence)
    out = x.extract([biogaia153], [1], dm, {"fiscal_year": 2025})
    td = next(f for f in out["fields"] if f["key"] == "total_debt")
    assert td["value"] == 0 and td["confidence"] == 0.5 and "stated_zero" in td["evidence"], (td, out["warnings"])
    # the subject terms buy no licence beyond the page: a bare-loan negation sentence that is NOT printed
    # on the cited page still drops (v050's off-page counter-example, now through the widened vocabulary)
    x.call_llm = lambda *a, **k: zero_loan_answer("The group has no bank loans whatsoever.")
    out = x.extract([vicore46], [1], dm, {"fiscal_year": 2025})
    td = next(f for f in out["fields"] if f["key"] == "total_debt")
    assert td["value"] is None and any("dropped as computed, not read" in w for w in out["warnings"]), (td, out["warnings"])
    # known boundary, pinned on purpose: the bare subject word cannot tell borrowing from lending --
    # loans to associates are an ASSET, so this sentence says nothing about borrowings, yet it passes
    # every gate (loan vocabulary, digit-free, verbatim on the page, negation). Reported honestly per
    # the work order, not forced away: word-level subject matching cannot see the "to associates"
    # complement; blocking it would need phrase-level context, a different mechanism.
    associates = "At year end the group had no loans to associates."
    x.call_llm = lambda *a, **k: zero_loan_answer(associates)
    out = x.extract([associates + "\nOther financial information follows here.\n"], [1], dm, {"fiscal_year": 2025})
    td = next(f for f in out["fields"] if f["key"] == "total_debt")
    assert td["value"] == 0 and td["confidence"] == 0.5 and "stated_zero" in td["evidence"], (td, out["warnings"])
    # v073: a third maturity shape besides bucket-rows and bucket-columns -- a note that gives each loan/lease
    # its own printed due date, year or year-range instead of a bucket label (Proact, docs/acrylic/evidence/v063.md
    # gap 2). _maturity_bucket buckets a single point directly; a range or bare year only counts when its start
    # and end land in the SAME bucket, else "straddle" -- and a straddle must abort the whole read, not just
    # that one row, since it proves nothing about which side of a boundary the debt sits on.
    from datetime import date as _date
    fye = _date(2025, 12, 31)
    assert x._maturity_bucket("Bank loan 16 Jul 2026 216,360", fye) == "due_within_1_year"
    assert x._maturity_bucket("Lease liability 2026 96,098", fye) == "due_within_1_year"  # bare year: its own Jan-1..Dec-31 span
    assert x._maturity_bucket("Lease liability 2027-2030 166,153", fye) == "due_1_to_5_years"  # inclusive at exactly 5 years
    assert x._maturity_bucket("Lease liability 2031-2032 100", fye) == "due_after_5_years"
    assert x._maturity_bucket("Loan 2026-2027 500", fye) == "straddle"  # start within 1y, end in 1-5y: no safe read
    # calendar-exact, not a fixed day count: a fixed 366-day cutoff would tip "2027" itself (exactly 366 days
    # after a 2025-12-31 fye) into due_within_1_year and manufacture a false straddle on "2027-2030" above
    assert x._bucket_for_date(_date(2027, 1, 1), fye) == "due_1_to_5_years"
    # a row naming its own maturity three times or more is a running header/watermark glued into one row by
    # _page_rows, never a printed instrument (Proact's own page carries exactly this artifact, v073)
    assert x._maturity_bucket("Proact Annual Report 2025 2025 2025", fye) is None
    assert x._maturity_bucket("liabilities measured 31 Dec 2025 31 Dec 2024 31 Dec 2025 31 Dec 2024", fye) is None
    # the balance sheet date defaults to 31 Dec of the fiscal year, or the day/month a caption states instead --
    # but only a caption with nothing _row_amounts would call an amount after its date, never an instrument's
    # own due-date row (a nil "-" included), which would otherwise be mistaken for the table's own caption
    assert x._balance_sheet_date(["Group maturity analysis as of 31 Mar 2026", "Loan Y 15 Jun 2027 900"], 2026) == _date(2026, 3, 31)
    assert x._balance_sheet_date(["Loan Z 31 Dec 2025 -"], 2025) == _date(2025, 12, 31)  # default: this is an instrument row, not a caption
    # end to end, Proact's own page 103 (real text, trimmed to the note's per-instrument list): the model
    # reads total_debt correctly (it is printed twice on the page; this is the carrying-note's own close) but
    # answers null on every bucket. No bucket-row or bucket-column mechanism reaches this shape -- there is no
    # bucket label anywhere and no single row prints all three buckets as columns -- so before this lane the
    # check stayed "missing: due_within_1_year" forever. The date-per-instrument read closes it exactly:
    # 216,360 (16 Jul 2026) + 96,098 (2026) = 312,458 due within 1 year; 166,153 (2027-2030) due 1-5 years;
    # 312,458 + 166,153 == 478,611 to the cent, and due_after_5_years is honestly left null (no row for it).
    proact103 = ("Notes\nNote 24 CONT.\nInterest-bearing\nliabilities, Group, Reported\n31 Dec 2025 Interest Maturity value\n"
                 "Utilised overdraft\nfacility, Nordea 1) 2) Base rate +2.0% 31 Dec 2025 -\n"
                 "Bank loan, Nordea 2) STIBOR 3M +1.25% 16 Jul 2026 -\n"
                 "Bank loan, Nordea 2) EURIBOR 3M +1.25% 16 Jul 2026 -\n"
                 "Bank loan, Svensk\nExportkredit 2) EURIBOR 3M + 1.8% 16 Jul 2026 216,360\n"
                 "Lease liability 3) 3.86% - 5.59% 2026 96,098\n"
                 "Lease liability 3) 4.06% - 5.59% 2027-2030 166,153\n"
                 "Total interest-bearing\nliabilities 478,611\n")
    x.call_llm = lambda *a, **k: {"fields": [
        {"key": "total_debt", "value": 478611, "unit": "TSEK", "period": "2025", "raw_label": "Total interest-bearing liabilities",
         "source": {"page": 1, "quote": "Total interest-bearing\nliabilities 478,611"}},
        {"key": "due_within_1_year", "value": None, "unit": None, "period": None, "raw_label": None, "source": None},
        {"key": "due_1_to_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None},
        {"key": "due_after_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None}]}
    out = x.extract([proact103], [1], dm, {"fiscal_year": 2025})
    got = {f["key"]: f for f in out["fields"]}
    assert got["due_within_1_year"]["value"] == 312458 and got["due_within_1_year"]["confidence"] == 1.0, (got, out["warnings"])
    assert got["due_1_to_5_years"]["value"] == 166153 and got["due_1_to_5_years"]["confidence"] == 0.9, (got, out["warnings"])
    assert got["due_after_5_years"]["value"] is None  # no row for it -- null_as_zero, not a fabricated 0
    assert out["checks"][0]["passed"], out["checks"]
    assert "value_derived" in got["due_within_1_year"]["evidence"] and "value_in_quote" not in got["due_within_1_year"]["evidence"]  # a sum of several rows has no literal quote
    assert "value_in_quote" in got["due_1_to_5_years"]["evidence"] and "value_derived" not in got["due_1_to_5_years"]["evidence"]  # one row's own printed figure -- never both (WEIGHTS)
    assert "16 Jul 2026" in got["due_within_1_year"]["source"]["quote"] and "2027-2030" in got["due_1_to_5_years"]["source"]["quote"]
    assert any("closing maturity_sums_to_total exactly" in w for w in out["warnings"]), out["warnings"]
    # must not misfire: three real seed7 shapes where the model also answers every bucket null, none of them
    # date-per-instrument (docs/acrylic/evidence/v072.md). HANZA's own row is bucket-COLUMN shaped (no date
    # anywhere near its total); Stillfront's rows are component/duration-labelled ("Repayment within 2-5 yr."),
    # not dated; Nederman's total is read off a sentence far ABOVE its own per-instrument due-date table, with
    # nothing date-shaped in the bounded window above it -- so the derivation correctly finds nothing in all
    # three, and the check stays exactly as honestly unresolved as it was before this lane.
    hanza132 = ("Förfallotidpunkt\nMellan\n2025-12-31 Redovisat Mindre än 6 månader Mellan 1 Senare än\n"
                "Typ av upplåning Valutor värde 6 månader och 1 år och 5 år\nBanklån SEK, EUR,\n"
                "CNY, CZK, PLN 1 332 15 9 1 306 2\nAvbetalnings-\nkontrakt 112 18 18 76 -\nSumma 1 444 33 27 1 382 2\n")
    stillfront109 = ("Notes\nNote 21\nInterest-bearing liabilities\nGroup\nMSEK 31 Dec 2025 31 Dec 2024\n"
                     "Contingent considerations for shares 1,294 2,032 Contingent considerations\n"
                     "Bond loans 2,835 2,829 Repayment within 2–5 yr. 620 1,170\n"
                     "Liabilities to credit institutions 984 1,376 Repayment after more than 5 yr. – –\n"
                     "Term loan 649 688 Non-current liability 620 1,170\n"
                     "Leasing liabilities 103 105 Current liability 675 862\n"
                     "Other interest-bearing liabilities 22 22 Total Contingent considerations 1,294 2,032\n"
                     "Total 5,887 7,053\n")
    nederman150 = ("150 CONSOLIDATED FINANCIAL STATEMENTS\nNote 21. Interest-bearing liabilities\n"
                   "As of 31 December 2025, the group had pension provisions of SEK 37.5m (42.3) Total interest-bearing liabilities 2,475.3 2,479.7\n"
                   "as recognised in note 24.\nLong-term liabilities, SEKm 2025 2024\nBank loans 1,915.8 1,859.8\nLease liabilities 442.2 483.7\nTotal 2,358.0 2,343.5\n"
                   "Short-term liabilities, SEKm 2025 2024\nCurrent part of bank loans 20.8 32.4\nCurrent part of lease liabilities 96.5 103.8\nTotal 117.3 136.2\n"
                   "Terms and repayment due dates\n2025, SEKm Currency Due date\nNominal interest rate Nominal amount in original currency Carrying amount\n"
                   "Bank loan* (revolving) SEK 24 Mar 2027 2.65%-3.86% 881.5 880.1\nBank loan* (revolving) EUR 24 Mar 2027 2.81%-2.95% 8.0 86.4\n"
                   "Bank loan* (revolving) USD 24 Mar 2027 5.05%-5.97% 49.0 450.1\nBank loan* (term loan) SEK 20 Dec 2027 3.05%-4.13% 500.0 499.2\n"
                   "Current bank loan CNY 5 Dec 2028 3.85%-3.95% 15.8 20.8\nLease liabilities 538.7\nTotal interest-bearing liabilities 2,475.3\n")
    for name, text, total_value, total_quote, unit in [
        ("HANZA", hanza132, 1444, "Summa 1 444 33 27 1 382 2", "MSEK"),
        ("Stillfront", stillfront109, 5887, "Total 5,887 7,053", "MSEK"),
        ("Nederman", nederman150, 2475.3,
         "As of 31 December 2025, the group had pension provisions of SEK 37.5m (42.3) Total interest-bearing liabilities 2,475.3 2,479.7", "SEKm"),
    ]:
        x.call_llm = lambda *a, tv=total_value, tq=total_quote, u=unit, **k: {"fields": [
            {"key": "total_debt", "value": tv, "unit": u, "period": "2025", "raw_label": "Total", "source": {"page": 1, "quote": tq}},
            {"key": "due_within_1_year", "value": None, "unit": None, "period": None, "raw_label": None, "source": None},
            {"key": "due_1_to_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None},
            {"key": "due_after_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None}]}
        out = x.extract([text], [1], dm, {"fiscal_year": 2025})
        assert all(f["value"] is None for f in out["fields"] if f["key"] != "total_debt"), (name, out["fields"])
        assert out["checks"][0]["detail"] == "missing: due_within_1_year", (name, out["checks"])
    # v096: the bucket-ROW finer split (Rusta Note 11, docs/acrylic/evidence/v088.md findings 2 and 3):
    # a maturity table whose rows are finer intervals than the three buckets -- "0–6 months 523 /
    # 7–12 months 516 / 1–2 years 969 / 2–5 years 2,184 / >5 years 2,431 / Total 6,624" -- needs the
    # sibling rows of one bucket SUMMED (within1y = 523+516 = 1,039; 1-5y = 969+2,184 = 3,153;
    # 1,039+3,153+2,431 = 6,623 ≈ 6,624), which no existing repair does: _statement_row declines two
    # rows matching one field (v058), _between_rows' window stops at "1–2 years" (a due_1_to_5_years
    # synonym, so its other-field guard fires -- exactly why Rusta stayed "missing" through v088), and
    # _fill_bucket_columns reads one row's columns, not rows. v088's schema decision also proved these
    # wordings must never join the synonym lists the column scanner reads ("0-6 months" in
    # due_within_1_year.synonyms leaks into _bucket_synonym_hits and changes _bucket_header's read),
    # so _bucket_span parses interval GEOMETRY, sharing nothing with that vocabulary. On the label:
    assert x._bucket_span("0–6 months") == (0, 6) and x._bucket_span("7–12 months") == (7, 12)
    assert x._bucket_span("6 months or less") == (0, 6) and x._bucket_span("6 månader eller mindre") == (0, 6)
    assert x._bucket_span("6 – 12 månader") == (6, 12) and x._bucket_span("1 – 5 år") == (12, 60)
    assert x._bucket_span("1–2 years") == (12, 24) and x._bucket_span("2–5 years") == (24, 60)
    assert x._bucket_span("between 1 and 2 years") == (12, 24) and x._bucket_span("mellan 1 och 5 år") == (12, 60)
    assert x._bucket_span("mellan 1 år och 2 år") == (12, 24) and x._bucket_span("mellan 2 år och 5 år") == (24, 60)  # v104: Svedbergs' long forms, a unit per operand
    assert x._bucket_span("mellan 3 månader och 1 år") == (3, 12) and x._bucket_span("mellan 18 månader och 2 år") == (18, 24)  # mixed units
    assert x._bucket_span("högst 2 år") == (0, 24) and x._bucket_span("mellan 5 år och 2 år") is None  # reversed bounds name no interval
    assert x._bucket_span("later than 1 year but within 3 years") == (12, 36)
    assert x._bucket_span(">5 years") == (60, float("inf")) and x._bucket_span("mer än 5 år") == (60, float("inf"))
    assert x._bucket_span("Within one year") == (0, 12) and x._bucket_span("mindre än 3 månader") == (0, 3)
    assert x._bucket_span("6-10 years") == (72, 120)  # the schema's own value_convention finer split of >5y
    assert x._bucket_span("3–7 years") == (36, 84)  # parses fine -- crossing a boundary is the caller's abort, not a parse failure
    assert x._bucket_span("Maturity analysis – lease liabilities") is None and x._bucket_span("Total") is None
    assert x._bucket_span("Lease liabilities") is None and x._bucket_span("Repayment within 2–5 yr.") is None  # duration wording, not a maturity interval
    assert x._bucket_span("-30 dgr") is None  # day wording is the subtotal-column mechanism's own territory (XANO, v078)
    assert x._span_bucket((7, 12)) == "due_within_1_year" and x._span_bucket((24, 60)) == "due_1_to_5_years"
    assert x._span_bucket((72, 120)) == "due_after_5_years" and x._span_bucket((36, 84)) is None  # crosses the 5-year boundary
    assert x._span_bucket((12, float("inf"))) is None  # ">1 year" spans 1-5y AND >5y: no safe split of the three buckets
    # ... end to end on Rusta's own page 115 (real text, trimmed to the leasing note; the model answer
    # is v088 pass 1 reconstructed from the stored warnings: total 6,669 year-corrected to 6,624, the
    # model's own bucket sums 975/3,062 computed-not-read, the repair locking the single "2–5 years"
    # row's 2,184 -- the stored red, byte-matched in docs/acrylic/evidence/v096.md). The page-level
    # year header is the P&L table's "2025/26 2024/25" below the maturity table (Rusta's own maturity
    # header "30 Apr 2026 30 Apr 2025" is not a _year_run), which is why the fixture keeps it.
    rusta115 = ("Lease liabilities\nGroup\nMaturity analysis – lease liabilities 30 Apr 2026 30 Apr 2025\n"
                "0–6 months 523 490\n7–12 months 516 485\n1–2 years 969 932\n2–5 years 2,184 2,130\n"
                ">5 years 2,431 2,633\nTotal 6,624 6,669\nGroup\n"
                "Amounts recognised in profit or loss 2025/26 2024/25\nDepreciation of right-of-use assets –798 –784\n"
                "Interest on lease liabilities –236 –242\nParent Company\n30 Apr 2026 30 Apr 2025\n"
                "Within one year 607 590\nBetween 1 and 5 years 1,924 1,949\nLater than 5 years 2,016 2,231\n"
                "Total 4,547 4,769\n")
    x.call_llm = lambda *a, **k: {"fields": [
        {"key": "total_debt", "value": 6669, "unit": "MSEK", "period": "2025", "raw_label": "Total",
         "source": {"page": 1, "quote": "Total 6,624 6,669"}},
        {"key": "due_within_1_year", "value": 975, "unit": "MSEK", "period": "2025", "raw_label": "Within 1 year",
         "source": {"page": 1, "quote": "Within 1 year 975"}},
        {"key": "due_1_to_5_years", "value": 3062, "unit": "MSEK", "period": "2025", "raw_label": "1–2 years; 2–5 years",
         "source": {"page": 1, "quote": "2–5 years 2,184 2,130"}},
        {"key": "due_after_5_years", "value": 2633, "unit": "MSEK", "period": "2025", "raw_label": ">5 years",
         "source": {"page": 1, "quote": ">5 years 2,431 2,633"}}]}
    out = x.extract([rusta115], [1], dm, {"fiscal_year": 2025})
    got = {f["key"]: f for f in out["fields"]}
    assert got["total_debt"]["value"] == 6624, (got, out["warnings"])
    w1 = got["due_within_1_year"]
    assert w1["value"] == 1039 and w1["confidence"] == 1.0 and "value_derived" in w1["evidence"], (w1, out["warnings"])
    assert w1["source"]["quote"] == "0–6 months 523 490 7–12 months 516 485" and w1["raw_label"] == "0–6 months + 7–12 months", w1
    w15 = got["due_1_to_5_years"]
    assert w15["value"] == 3153 and w15["source"]["quote"] == "1–2 years 969 932 2–5 years 2,184 2,130", (w15, out["warnings"])
    assert any("2184 reads one finer-split row" in w for w in out["warnings"]), out["warnings"]  # the replaced partial read is named
    assert got["due_after_5_years"]["value"] == 2431 and got["due_after_5_years"]["source"]["quote"] == ">5 years 2,431 2,633", got["due_after_5_years"]  # one row: untouched, status quo
    assert out["checks"][0]["passed"] and "abs((1039 + 3153 + 2431) - 6624) <= 2" in out["checks"][0]["detail"], out["checks"]
    # single-row buckets keep today's behaviour: the same page's table with ONE row per bucket ("Within
    # one year" / "Between 1 and 5 years" / "Later than 5 years", real text) -- every bucket's "sum" is
    # one row's own literal figure, so the finer-split derivation writes nothing and the fields are
    # exactly what the single-row paths already produce (statement-spread fills 607 and 2,016 off their
    # own printed rows, _between_rows closes 1,924 between them -- all pre-v096 machinery, every quote a
    # single row, no sibling join anywhere). v111: the block prints these rows under a "Parent Company"
    # section heading, which the wrong-table guard now refuses for the Group's section -- so the same
    # rows under a "Group" heading (the v096 mechanics unchanged, the entity switched) carry the
    # assertion, and the real parent block below asserts the refusal itself.
    rusta115_group = rusta115.replace("Parent Company\n", "Group\n")
    x.call_llm = lambda *a, **k: {"fields": [
        {"key": "total_debt", "value": 4547, "unit": "MSEK", "period": "2025", "raw_label": "Total",
         "source": {"page": 1, "quote": "Total 4,547 4,769"}},
        {"key": "due_within_1_year", "value": None, "unit": None, "period": None, "raw_label": None, "source": None},
        {"key": "due_1_to_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None},
        {"key": "due_after_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None}]}
    out = x.extract([rusta115_group], [1], dm, {"fiscal_year": 2025})
    got = {f["key"]: (f["value"], f["source"]["quote"]) for f in out["fields"]}
    assert got == {"total_debt": (4547, "Total 4,547 4,769"), "due_within_1_year": (607, "Within one year 607 590"),
                   "due_1_to_5_years": (1924, "Between 1 and 5 years 1,924 1,949"),
                   "due_after_5_years": (2016, "Later than 5 years 2,016 2,231")}, (got, out["warnings"])
    assert out["checks"][0]["passed"] and not any("finer-split" in w for w in out["warnings"]), (out["checks"], out["warnings"])  # all buckets single-row: the derivation stayed silent
    # v111: the real parent block (verbatim from the page) is the wrong entity for the Group's section --
    # every read of it, the model's own total citation included, is refused with the section named
    x.call_llm = lambda *a, **k: {"fields": [
        {"key": "total_debt", "value": 4547, "unit": "MSEK", "period": "2025", "raw_label": "Total",
         "source": {"page": 1, "quote": "Total 4,547 4,769"}},
        {"key": "due_within_1_year", "value": None, "unit": None, "period": None, "raw_label": None, "source": None},
        {"key": "due_1_to_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None},
        {"key": "due_after_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None}]}
    out = x.extract([rusta115], [1], dm, {"fiscal_year": 2025})
    got = {f["key"]: f["value"] for f in out["fields"]}
    assert all(v is None for v in got.values()), (got, out["warnings"])
    assert sum("Parent Company section" in w for w in out["warnings"]) >= 2, out["warnings"]
    # counter-example 1: sibling rows whose sums do NOT close the identity (100+50 vs Total 200) --
    # nothing adopted, fields exactly as before: the identity gate is the whole decision, no plausibility
    badsum = "Maturity analysis\nMSEK 2025 2024\n0–6 months 100 90\n7–12 months 50 40\nTotal 200 190\n"
    x.call_llm = lambda *a, **k: {"fields": [
        {"key": "total_debt", "value": 200, "unit": "MSEK", "period": "2025", "raw_label": "Total",
         "source": {"page": 1, "quote": "Total 200 190"}},
        {"key": "due_within_1_year", "value": None, "unit": None, "period": None, "raw_label": None, "source": None},
        {"key": "due_1_to_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None},
        {"key": "due_after_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None}]}
    out = x.extract([badsum], [1], dm, {"fiscal_year": 2025})
    assert all(f["value"] is None for f in out["fields"] if f["key"] != "total_debt"), (out["fields"], out["warnings"])
    assert not any("finer-split" in w for w in out["warnings"]) and out["checks"][0]["detail"].startswith("missing:"), out["checks"]
    # counter-example 2: an interval CROSSING a bucket boundary ("3–7 years" = [36,84] months) aborts
    # the whole derivation, even though 150+500 would numerically close -- the table's own granularity
    # cannot decide which side of 5 years that row's debt sits on, so no row here is safely bucketed
    straddle = "Maturity analysis\nMSEK 2025 2024\n0–6 months 100 90\n7–12 months 50 40\n3–7 years 500 400\nTotal 650 640\n"
    x.call_llm = lambda *a, **k: {"fields": [
        {"key": "total_debt", "value": 650, "unit": "MSEK", "period": "2025", "raw_label": "Total",
         "source": {"page": 1, "quote": "Total 650 640"}},
        {"key": "due_within_1_year", "value": None, "unit": None, "period": None, "raw_label": None, "source": None},
        {"key": "due_1_to_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None},
        {"key": "due_after_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None}]}
    out = x.extract([straddle], [1], dm, {"fiscal_year": 2025})
    assert all(f["value"] is None for f in out["fields"] if f["key"] != "total_debt"), (out["fields"], out["warnings"])
    assert not any("finer-split" in w for w in out["warnings"]), out["warnings"]
    # v104: the same derivation on a TORN side-by-side page -- Svedbergs p.132 prints the maturity table
    # and the financing-changes note in two columns, and pymupdf glues the note's lines onto the maturity
    # rows (the Summa row too), so _row_amounts' last-alpha read returns the OTHER table's figures and the
    # walk reads each row's own four columns label-anchored (em-dash parent-company nils included). The
    # labels are the Swedish long forms v097's seed-10 round watched the model sum correctly (1-5y
    # 423,946 + 11,896 = 435,842) and then drop as computed-not-read: the finer rows live in the Koncernen
    # 2025 column of "Koncernen Moderbolaget / 2025 2024 2025 2024", the parent pair dashes.
    sved132 = ("Not 33 Räntebärande skulder\nKoncernen Moderbolaget Kassaflödespåverkande förändringar:\n"
               "2025 2024 2025 2024 Förändringar övriga skulder –130 273 –188 385 –318 657 –121 155 — –121 155\n"
               "3 månader eller mindre 26 733 213 095 5 153 12 186 Förändringar leasingskuld — –29 364 –29 364 — — —\n"
               "Mellan 3 månader och 1 år 24 707 28 420 15 460 18 924 Ej kassaflödespåverkande förändringar:\n"
               "Mellan 1 år och 2 år 423 946 614 610 412 211 578 370 Förändringar övriga skulder –23 579 19 869 –3 709 — — —\n"
               "Mellan 2 år och 5 år 11 896 19 122 — — Förändringar leasingskuld 360 125 50 547 410 672 — — —\n"
               "Mer än 5 år 9 124 11 230 — — Valutakursdifferenser –46 145 –13 437 –59 583 –45 004 –2 374 –47 378\n"
               "Summa 496 405 886 478 432 824 609 480 Per 31 december 2025 865 828 101 399 967 226 412 211 20 613 432 824\n")
    x.call_llm = lambda *a, **k: {"fields": [
        {"key": "total_debt", "value": 496405, "unit": "Tkr", "period": "2025", "raw_label": "Summa",
         "source": {"page": 1, "quote": "Summa 496 405 886 478 432 824 609 480 Per 31 december 2025 865 828 101 399 967 226 412 211 20 613 432 824"}},
        {"key": "due_within_1_year", "value": 51440, "unit": "Tkr", "period": "2025",
         "raw_label": "3 månader eller mindre; Mellan 3 månader och 1 år",
         "source": {"page": 1, "quote": "Mellan 3 månader och 1 år 24 707 28 420 15 460 18 924 Ej kassaflödespåverkande förändringar:"}},
        {"key": "due_1_to_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None},
        {"key": "due_after_5_years", "value": 9124, "unit": "Tkr", "period": "2025", "raw_label": "Mer än 5 år",
         "source": {"page": 1, "quote": "Mer än 5 år 9 124 11 230 — — Valutakursdifferenser –46 145 –13 437 –59 583 –45 004 –2 374 –47 378"}}]}
    out = x.extract([sved132], [1], dm, {"fiscal_year": 2025})
    got = {f["key"]: f for f in out["fields"]}
    assert got["total_debt"]["value"] == 496405 and got["due_after_5_years"]["value"] == 9124, (got, out["warnings"])  # single-row bucket: the model's own read stands
    w15 = got["due_1_to_5_years"]
    assert w15["value"] == 435842 and "value_derived" in w15["evidence"], (w15, out["warnings"])
    assert w15["source"]["quote"] == ("Mellan 1 år och 2 år 423 946 614 610 412 211 578 370 Förändringar övriga skulder –23 579 19 869 –3 709 — — —"
                                      " Mellan 2 år och 5 år 11 896 19 122 — — Förändringar leasingskuld 360 125 50 547 410 672 — — —"), w15["source"]
    assert w15["raw_label"] == "Mellan 1 år och 2 år + Mellan 2 år och 5 år", w15
    w1 = got["due_within_1_year"]
    assert w1["value"] == 51440 and "value_derived" in w1["evidence"], (w1, out["warnings"])  # the model's own sum, now provable from the printed rows
    assert w1["source"]["quote"] == ("3 månader eller mindre 26 733 213 095 5 153 12 186 Förändringar leasingskuld — –29 364 –29 364 — — —"
                                     " Mellan 3 månader och 1 år 24 707 28 420 15 460 18 924 Ej kassaflödespåverkande förändringar:"), w1["source"]
    assert out["checks"][0]["passed"] and "abs((51440 + 435842 + 9124) - 496405) <= 2" in out["checks"][0]["detail"], out["checks"]
    # ... and the straddle counter-example in the new form: "Mellan 3 år och 7 år" = [36,84] crosses the
    # 5-year boundary, so the whole derivation is abandoned even though 150+500 would numerically close
    straddle_sv = "Maturity analysis\nMSEK 2025 2024\n3 månader eller mindre 100 90\nMellan 3 år och 7 år 500 400\nSumma 650 640\n"
    x.call_llm = lambda *a, **k: {"fields": [
        {"key": "total_debt", "value": 650, "unit": "MSEK", "period": "2025", "raw_label": "Summa",
         "source": {"page": 1, "quote": "Summa 650 640"}},
        {"key": "due_within_1_year", "value": None, "unit": None, "period": None, "raw_label": None, "source": None},
        {"key": "due_1_to_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None},
        {"key": "due_after_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None}]}
    out = x.extract([straddle_sv], [1], dm, {"fiscal_year": 2025})
    assert all(f["value"] is None for f in out["fields"] if f["key"] != "total_debt"), (out["fields"], out["warnings"])
    assert not any("finer-split" in w for w in out["warnings"]), out["warnings"]
    # v101: a debt total/bucket printed under a liabilities-negative sign convention -- net-debt and
    # capital-management presentations print cash positive and every debt row negative -- carries its
    # magnitude with a minus, and the fields (semantics: a carrying amount of borrowings) must record
    # the magnitude. Catella p.108 and Scandi Standard p.134 real table text (trimmed; the stored
    # baseline red -- -1,474@1.0 / -2,307 / -70 passing every gate -- is in docs/acrylic/evidence/v101.md).
    assert x._NET_DEBT_LABEL.search(x._clean_label("Net debt")) and x._NET_DEBT_LABEL.search(x._clean_label("Nettoskuld")) \
        and x._NET_DEBT_LABEL.search(x._clean_label("Net interest-bearing debt")) and x._NET_DEBT_LABEL.search(x._clean_label("Nettokassa / nettoskuld (-)"))
    assert not x._NET_DEBT_LABEL.search(x._clean_label("Gross debt December 31 2025 (Note 21)")) \
        and not x._NET_DEBT_LABEL.search(x._clean_label("Total interest-bearing liabilities")) \
        and not x._NET_DEBT_LABEL.search(x._clean_label("- repayable within one year"))
    cat108 = ("Information on the Group's net debt profile and a sensitivity analysis are presented below, with informa-\n"
              "tion on fixed interest periods.\n"
              "SEK M 31 Dec 31 Dec\n"
              "EUR liabilities −82 −135\n"
              "SEK liabilities −1,242 −1,342\n"
              "GBP liabilities −146 −150\n"
              "DKK liabilities −2 −1,105\n"
              "Liabilities in other currencies - -\n"
              "Total interest-bearing liabilities −1,474 −2,735\n"
              "Term (days) 91 90\n"
              "Total interest-bearing assets 2,164 1,472\n")
    x.call_llm = lambda *a, **k: {"fields": [
        {"key": "total_debt", "value": -1474, "unit": "SEK M", "period": "2025", "raw_label": "Total interest-bearing liabilities",
         "source": {"page": 1, "quote": "Total interest-bearing liabilities −1,474 −2,735"}},
        {"key": "due_within_1_year", "value": None, "unit": None, "period": None, "raw_label": None, "source": None},
        {"key": "due_1_to_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None},
        {"key": "due_after_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None}]}
    out = x.extract([cat108], [1], dm, {"fiscal_year": 2025})
    got = {f["key"]: f for f in out["fields"]}
    assert got["total_debt"]["value"] == 1474, (got["total_debt"], out["warnings"])
    assert "sign_normalized" in got["total_debt"]["evidence"] and got["total_debt"]["confidence"] == 1.0, got["total_debt"]
    assert any("total_debt: -1474 printed with a liabilities-negative sign convention; recorded as 1474" in w for w in out["warnings"]), out["warnings"]
    # Scandi Standard p.134: the financing reconciliation prints every debt row negative, buckets included
    ss134 = ("4) Reconciliation of Net interest-bearing debt\n"
             "Net interest-bearing debt1),\nMSEK 2025 2024\n"
             "Cash and cash equivalents 279 109\n"
             "Interest-bearing liabilities\n– repayable within one year –70 –64\n"
             "Interest-bearing liabilities\n– repayable after one year –2,238 –1,982\n"
             "Net interest-bearing debt –2,032 –1,935\n"
             "Gross debt – variable interest rates –2,307 –2,046\n"
             "Liabilities from financing activities\nChanges in gross debt, MSEK\n"
             "Gross debt December 31, 2024\n(Note 21) –1,733 –313 –2,046\n"
             "new loans –338 –53 –391\n"
             "repayments 97 69 165\n"
             "Gross debt December 31 2025\n(Note 21) –2,021 –286 –2,307\n")
    x.call_llm = lambda *a, **k: {"fields": [
        {"key": "total_debt", "value": -2307, "unit": "MSEK", "period": "2025", "raw_label": "Gross debt December 31 2025 (Note 21)",
         "source": {"page": 1, "quote": "(Note 21) –2,021 –286 –2,307"}},
        {"key": "due_within_1_year", "value": -70, "unit": "MSEK", "period": "2025", "raw_label": "– repayable within one year",
         "source": {"page": 1, "quote": "– repayable within one year –70 –64"}},
        {"key": "due_1_to_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None},
        {"key": "due_after_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None}]}
    out = x.extract([ss134], [1], dm, {"fiscal_year": 2025})
    got = {f["key"]: f for f in out["fields"]}
    assert got["total_debt"]["value"] == 2307 and got["due_within_1_year"]["value"] == 70, (got, out["warnings"])
    assert "sign_normalized" in got["total_debt"]["evidence"] and "sign_normalized" in got["due_within_1_year"]["evidence"], (got["total_debt"], got["due_within_1_year"])
    assert sum("printed with a liabilities-negative sign convention" in w for w in out["warnings"]) == 2, out["warnings"]
    # counter-example: a row labelled as net debt is a different shape -- its negative can be a genuine
    # net-cash position, and total_debt must not point there at all. Left exactly as printed, shape named.
    netdebt = ("Capital management\nMSEK\n2025 2024\n"
               "Borrowings -600 -550\n"
               "Cash and bank deposits 100 90\n"
               "Net debt (nettoskuld) -500 -460\n")
    x.call_llm = lambda *a, **k: {"fields": [
        {"key": "total_debt", "value": -500, "unit": "MSEK", "period": "2025", "raw_label": "Net debt (nettoskuld)",
         "source": {"page": 1, "quote": "Net debt (nettoskuld) -500 -460"}},
        {"key": "due_within_1_year", "value": None, "unit": None, "period": None, "raw_label": None, "source": None},
        {"key": "due_1_to_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None},
        {"key": "due_after_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None}]}
    out = x.extract([netdebt], [1], dm, {"fiscal_year": 2025})
    got = {f["key"]: f for f in out["fields"]}
    assert got["total_debt"]["value"] == -500 and "sign_normalized" not in got["total_debt"]["evidence"], (got["total_debt"], out["warnings"])
    assert any("net-debt label" in w for w in out["warnings"]), out["warnings"]

    # v103: the wrong-table guard. Two mis-scoped fills from v097 read their values out of the liquidity
    # note's tables -- RVRC's p.110 "Maturity analysis regarding non-discounted liabilities" (its Total row
    # sums trade payables 119, expected returns 34 and other current liabilities 54 into total_debt, and
    # the identity endorsed it: 210+10+0 == 220, check passed) and Lime's p.72 "Liquidity risk - Group"
    # table (its Borrowing row is contractual cash flow, not carrying). The schema now names the marker
    # words (table_scope_words) and _table_scope refuses such a table under the carrying basis, at every
    # fill mechanism's entry and on the model's own citations. Real page text throughout, trimmed.
    rvrc110 = ("Maturity analysis regarding non-discounted liabilities\n"
               "SEKm <6 months 6–12 months 1–3 years 3–5 years >5 years Total\n"
               "30 June 2025\n"
               "Leases1) 3 3 7 0 0 12\n"
               "Trade payables 119 — — — — 119\n"
               "Expected returns 34 — — — — 34\n"
               "Other liabilites 0 — — — — 0\n"
               "Other current liabilites 54 — — — — 54\n"
               "Total 210 3 7 0 0 220\n")
    rvrc_rows = x._page_rows(rvrc110)
    tsw = dm["table_scope_words"]
    # header level: the title row's wording (short standalone line, inside the 25-row window) classes the
    # table "undiscounted" under carrying; under the undiscounted basis the same table's bare-Total shape
    # is "all_liabilities" (trade payables among its rows -- never debt, on either basis), while its debt
    # row (Leases, the report's own interest-bearing scope) stays readable
    assert x._table_scope(rvrc_rows, None, 8, "carrying", tsw) == "undiscounted"
    assert x._table_scope(rvrc_rows, None, 8, "undiscounted", tsw) == "all_liabilities"
    assert x._table_scope(rvrc_rows, None, 3, "undiscounted", tsw, debt_scoped=True) == "unknown"
    assert x._table_scope(rvrc_rows, None, 3, "carrying", tsw, debt_scoped=True) == "undiscounted"  # the table is refused whole
    lime72 = ("Lime's financial policy states that Lime shall not use any\n"
              "excess liquidity for trading in financial assets and that cash and cash equivalents over\n"
              "time shall amount to at least 8% of annual sales.\n"
              "Liquidity risk - Group\n"
              "As of 31 December 2025 Less than 3 months Between 3 months and 1 year Between 1 and 2 years Between 2 and 5 years\n"
              "Borrowing (incl. overdraft) 15,000 51,396 60,000 25,000\n"
              "Liabilities related to leasing 5,112 11,298 8,000 4,664\n"
              "Acquisition-related liabilities - - - 35,275\n"
              "Accounts payable 12,806 - - -\n"
              "Amount 32,918 62,694 68,000 64,939\n")
    lime_rows = x._page_rows(lime72)
    li = lime_rows.index("Borrowing (incl. overdraft) 15,000 51,396 60,000 25,000")
    ai = lime_rows.index("Amount 32,918 62,694 68,000 64,939")
    assert x._table_scope(lime_rows, None, li, "carrying", tsw) == "undiscounted"
    assert x._table_scope(lime_rows, None, ai, "undiscounted", tsw) == "all_liabilities"  # accounts payable above the grand Amount row
    assert x._table_scope(lime_rows, None, li, "undiscounted", tsw, debt_scoped=True) == "unknown"  # v089's own read
    # the carrying-column exemption: a table whose title says undiscounted but which prints a carrying
    # column is readable under both bases (Ework/Instalco, v076/v085 -- the title names the OTHER column)
    assert x._table_scope(ework_rows, None, ework_idx, "carrying", tsw) == "carrying"
    assert x._table_scope(instalco_rows, None, instalco_idx, "carrying", tsw) == "carrying"
    assert x._table_scope(instalco_rows, None, instalco_idx, "undiscounted", tsw) == "carrying"
    # a "liquidity risk" section heading above prose above the carrying table never brands the table
    # (MedCap p.101: the walk stops at the prose that names amounts, before reaching LIKVIDITETSRISK);
    # Boozt p.121's all-liabilities maturity table has NO marker on its own title line -- the marker
    # ("contractual undiscounted amounts") sits in a 181-char prose sentence the title cap excludes, and
    # the debt row the label reads stays readable (the label pins 273/63 from exactly that row)
    medcap101 = ("LIKVIDITETSRISK\n"
                 "Likviditetsrisk är risken att koncernen inte kan finansiera lånebetalningar eller andra finansiella\n"
                 "åtaganden i den takt de förfaller till betalning och hanteras genom uppläggning av\n"
                 "av tillgängliga likvida medel, kortfristiga placeringar och outnyttjat utrymme på checkräkningskredit.\n"
                 "Per 2025-12-31 uppgick driftlikviditeten till 370,4 (370,1)\n"
                 "MSEK. 100 procent av koncernens krediter löper med så\n"
                 "kallad rörlig ränta med 3 månaders bindningstid.\n"
                 "Förfallotidpunkt för upplåning Koncernen Moderbolaget\n"
                 "MSEK 2025-12-31 2024-12-31 2025-12-31 2024-12-31\n"
                 "6 månader eller mindre 54,3 41,8 – –\n"
                 "6 – 12 månader 48,0 19,0 – –\n"
                 "1 – 5 år 616,5 316,7 – –\n"
                 "Totalt 718,7 377,6 – –\n")
    mc_rows = x._page_rows(medcap101)
    assert x._table_scope(mc_rows, None, len(mc_rows) - 1, "carrying", tsw) == "unknown"
    assert x._table_scope(mc_rows, None, len(mc_rows) - 1, "undiscounted", tsw) == "unknown"
    boozt121 = ("LIQUIDITY RISK\n"
                "The maturity structure for all of the Group's financial liabilities, including principal and interest, is shown in the table below. The table shows contractual undiscounted amounts.\n"
                "MATURITY STRUCTURE OF OUTSTANDNING ACCOUNT PAYABLES AND OTHER LIABILITIES\n"
                "Total borrowing\n"
                "Maturity within\n"
                "3 months\n"
                "Maturity within three to twelve months\n"
                "Maturity within one to five years\n"
                "Maturity within five to nie years\n"
                "Maturity after nine years\n"
                "Maturity structure of borrowing Dec 31,2024\n"
                "Liabilities to credit institutions 380 - - 380 - -\n"
                "Lease liabilities 499 29 85 251 135 -\n"
                "Accounts payables 1,235 1,225 10 - - -\n"
                "Other liabilities 531 527 4 - - -\n"
                "Total 2,645 1,781 99 631 135 0\n"
                "Maturity structure of borrowing Dec 31, 2025\n"
                "Liabilities to credit institutions 0 - - 0 - -\n"
                "Lease liabilities 441 26 78 273 63 -\n"
                "Accounts payables 1,384 1,374 10 - - -\n"
                "Other liabilities 533 528 4 - - -\n"
                "Total 2,358 1,928 92 273 63 0\n")
    bz_rows = x._page_rows(boozt121)
    bzi = bz_rows.index("Lease liabilities 441 26 78 273 63 -")
    assert x._table_scope(bz_rows, None, bzi, "carrying", tsw) == "unknown"  # the label's own read, kept
    assert x._table_scope(bz_rows, None, bzi, "undiscounted", tsw) == "unknown"
    # Smartcraft p.59: the table's own title names it ("Total undiscounted lease liabilities") -- refused
    # under carrying even though the rows are the label-declined lease buckets the stored answer read
    smartcraft59 = ("2025 ANNUAL REPORT\n"
                    "Total undiscounted lease liabilities\n"
                    "Expenses related to the right of use assets and lease liabilites recongized in the P&L\n"
                    "Amounts in NOK (thousands) Maturity analysis Total 2025\n"
                    "Less than 1 year 13 825 13 825\n"
                    "1-2-years 8 088 8 088\n"
                    "2-3 years 7 462 7 462\n"
                    "3-4 years 2 155 2 155\n"
                    "4-5 years - -\n"
                    "More than 5 years - -\n"
                    "Total undiscounted lease liability 31 529 31 529\n"
                    "Amounts in NOK (thousands) 2025 2024\n"
                    "Total lease expenses related to short-term or low value leases 629 774\n"
                    "Depreciation 13 253 12 457\n"
                    "Interest on lease liabilites 2 467 1 825\n"
                    "Total expenses from leases recognized in the P&L 16 348 15 056\n")
    sc_rows = x._page_rows(smartcraft59)
    assert x._table_scope(sc_rows, None, 4, "carrying", tsw) == "undiscounted"
    # end to end, RVRC -- the stored v097 pass shape: the model answered the Total row of the refused
    # table for all four fields (the repair had overwritten its own 12 with 220). Every one of them is
    # refused now, with a warning naming the table; the check honestly reports its missing operands
    x.call_llm = lambda *a, **k: {"fields": [
        {"key": "total_debt", "value": 220, "unit": "SEKm", "period": "2025", "raw_label": "Total",
         "source": {"page": 1, "quote": "Total 210 3 7 0 0 220"}},
        {"key": "due_within_1_year", "value": 210, "unit": "SEKm", "period": "2025", "raw_label": "Total",
         "source": {"page": 1, "quote": "Total 210 3 7 0 0 220"}},
        {"key": "due_1_to_5_years", "value": 10, "unit": "SEKm", "period": "2025", "raw_label": "Total",
         "source": {"page": 1, "quote": "Total 210 3 7 0 0 220"}},
        {"key": "due_after_5_years", "value": 0, "unit": "SEKm", "period": "2025", "raw_label": "Total",
         "source": {"page": 1, "quote": "Total 210 3 7 0 0 220"}}]}
    out = x.extract([rvrc110], [1], dm, {"fiscal_year": 2025})
    got = {f["key"]: f["value"] for f in out["fields"]}
    assert got == {"total_debt": None, "due_within_1_year": None, "due_1_to_5_years": None, "due_after_5_years": None}, (got, out["warnings"])
    assert sum("refused" in w for w in out["warnings"]) >= 4, out["warnings"]
    assert any("table 'Maturity analysis regarding non-discounted liabilities' is undiscounted" in w for w in out["warnings"]), out["warnings"]
    assert out["checks"][0]["detail"].startswith("missing:"), out["checks"]
    # and the model's own Note 21 read (p.109, one page earlier) survives untouched: a carrying row of
    # the borrowings disclosure is not in any refused table -- 12 kept, buckets honestly null
    rvrc109 = ("The table below shows conditions and maturity dates for respective interest-bearing liabilities:\n"
               "SEKm Reported value Currency Matures Interest\n"
               "Long-term liabilities to credit institutions as of 30 June 2025\n"
               "Bank loan Facility B 0 SEK — Variable\n"
               "Lease liabilities (see Note 18 Lease agreements) 12 SEK 1) Variable\n"
               "Reported value\n")
    x.call_llm = lambda *a, **k: {"fields": [
        {"key": "total_debt", "value": 12, "unit": "SEKm", "period": "2025", "raw_label": "Lease liabilities (see Note 18 Lease agreements)",
         "source": {"page": 2, "quote": "Lease liabilities (see Note 18 Lease agreements) 12 SEK 1) Variable"}},
        {"key": "due_within_1_year", "value": None, "unit": None, "period": None, "raw_label": None, "source": None},
        {"key": "due_1_to_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None},
        {"key": "due_after_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None}]}
    out = x.extract([rvrc110, rvrc109], [1, 2], dm, {"fiscal_year": 2025})
    got = {f["key"]: f["value"] for f in out["fields"]}
    assert got["total_debt"] == 12 and got["due_within_1_year"] is None, (got, out["warnings"])
    # end to end, Lime -- the stored v097 shape (buckets citing the liquidity row, total honestly null):
    # both citations refused, and the column repair's own attempt on the same table declines with the
    # same reason instead of filling 66,396/85,000
    x.call_llm = lambda *a, **k: {"fields": [
        {"key": "total_debt", "value": None, "unit": None, "period": None, "raw_label": None, "source": None},
        {"key": "due_within_1_year", "value": 66396, "unit": "Tkr", "period": "2025", "raw_label": "Borrowing (incl. overdraft)",
         "source": {"page": 1, "quote": "Borrowing (incl. overdraft) 15,000 51,396 60,000 25,000"}},
        {"key": "due_1_to_5_years", "value": 85000, "unit": "Tkr", "period": "2025", "raw_label": "Borrowing (incl. overdraft)",
         "source": {"page": 1, "quote": "Borrowing (incl. overdraft) 15,000 51,396 60,000 25,000"}},
        {"key": "due_after_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None}]}
    out = x.extract([lime72], [1], dm, {"fiscal_year": 2025})
    got = {f["key"]: f["value"] for f in out["fields"]}
    assert got == {"total_debt": None, "due_within_1_year": None, "due_1_to_5_years": None, "due_after_5_years": None}, (got, out["warnings"])
    assert any("table 'Liquidity risk - Group' is undiscounted" in w for w in out["warnings"]), out["warnings"]
    # and the fill with the model declined everywhere (the live v097 shape): the column repair declines
    # on the refused table -- no 66,396/85,000 fill, the refusal warning instead
    x.call_llm = lambda *a, **k: {"fields": [
        {"key": "total_debt", "value": None, "unit": None, "period": None, "raw_label": None, "source": None},
        {"key": "due_within_1_year", "value": None, "unit": None, "period": None, "raw_label": None, "source": None},
        {"key": "due_1_to_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None},
        {"key": "due_after_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None}]}
    out = x.extract([lime72], [1], dm, {"fiscal_year": 2025})
    got = {f["key"]: f["value"] for f in out["fields"]}
    assert all(v is None for v in got.values()), (got, out["warnings"])
    assert any("column reading of 'Borrowing (incl. overdraft) 15,000 51,396 60,000 25,000' declined" in w
               and "refused under the carrying basis" in w for w in out["warnings"]), out["warnings"]
    # Smartcraft end to end, the live v097 shape (model declined, the statement-row fill read p.59's
    # "Less than 1 year 13 825" rows): the guard refuses what that fill cited -- the table's own title
    # names it ("Total undiscounted lease liabilities"; the labels pin the carrying 13,439/14,809 from
    # p.58, the stored 13,825-family reads are the undiscounted ones)
    x.call_llm = lambda *a, **k: {"fields": [
        {"key": "total_debt", "value": None, "unit": None, "period": None, "raw_label": None, "source": None},
        {"key": "due_within_1_year", "value": None, "unit": None, "period": None, "raw_label": None, "source": None},
        {"key": "due_1_to_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None},
        {"key": "due_after_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None}]}
    out = x.extract([smartcraft59], [1], dm, {"fiscal_year": 2025})
    got = {f["key"]: f["value"] for f in out["fields"]}
    assert all(v is None for v in got.values()), (got, out["warnings"])
    assert sum("refused" in w for w in out["warnings"]) >= 2, out["warnings"]
    assert any("table 'Total undiscounted lease liabilities' is undiscounted" in w for w in out["warnings"]), out["warnings"]
    # the undiscounted basis reads the liquidity table it is told to read (v089's Instalco case, above,
    # still passes); RVRC's all-liabilities Total stays refused on that basis too
    basis_saved = x.debt_basis
    x.debt_basis = lambda: "undiscounted"
    x.call_llm = lambda *a, **k: {"fields": [
        {"key": "total_debt", "value": 220, "unit": "SEKm", "period": "2025", "raw_label": "Total",
         "source": {"page": 1, "quote": "Total 210 3 7 0 0 220"}},
        {"key": "due_within_1_year", "value": None, "unit": None, "period": None, "raw_label": None, "source": None},
        {"key": "due_1_to_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None},
        {"key": "due_after_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None}]}
    out = x.extract([rvrc110], [1], dm, {"fiscal_year": 2025})
    got = {f["key"]: f["value"] for f in out["fields"]}
    assert all(v is None for v in got.values()), (got, out["warnings"])
    assert any("is an all-liabilities table" in w for w in out["warnings"]), out["warnings"]
    x.debt_basis = basis_saved  # restore: the carrying default governs everything after this block
    # v091: the prior fiscal year's four figures ride along as top-level `prior_year` metadata -- never a
    # model answer, only a deterministic re-read of the same table the current year came from, written
    # only when its own identity closes on explicit values (the shipped require_explicit_values rule,
    # applied to FY-1: a bucket with no prior-year figure is absent from the sum, never zero-filled).
    # Two shapes, exactly the two that read today's values deterministically:
    # (a) buckets-as-COLUMNS (_fill_bucket_columns' selected row): the same-labelled row of the prior
    #     year's own block on the same page (Tången p.62's "31 december 2024" table, Ework p.70's "2024"
    #     block), read with the SAME col_keys through _bucket_assign -- a bucket whose prior-year cells
    #     all print dashes is the report's explicit 0 (v078, one year back);
    # (b) buckets-as-ROWS (the year-column tables): the prior-year COLUMN of the very rows that supplied
    #     the current values (MedCap p.101's Koncernen 2024 column), each field admitted only when the
    #     fiscal-year column of its own rows reproduces its current value -- _column_values' admission.
    # Date-per-instrument notes (Proact, v073) and model-only answers give no prior year: no year
    # header, no twin row, no admission -- and the >=2-buckets + closure gate has the final say anyway.
    # Tången, end to end on v084's own two-table fixture (the model answers nothing; the current-year
    # read is v084's): the prior year's own table prints 90,970 / (68,704 + 188,828) / 4,520 / 353,022,
    # and 90,970 + 257,532 + 4,520 == 353,022 exactly.
    x.call_llm = lambda *a, **k: {"fields": [
        {"key": k, "value": None, "unit": None, "period": None, "raw_label": None, "source": None} for k in dmf]}
    out = x.extract([tangen62_full], [1], dm, {"fiscal_year": 2025})
    py = out["prior_year"]
    assert py["fiscal_year"] == 2024 and py["check"]["passed"] and "353022" in py["check"]["detail"], py
    assert {k: v["value"] for k, v in py["fields"].items()} == {
        "total_debt": 353022, "due_within_1_year": 90970, "due_1_to_5_years": 257532, "due_after_5_years": 4520}, py
    assert all(v["source"] == {"page": 1, "quote": "Lån Kreditinstitut 90 970 68 704 188 828 4 520 353 022"}
               for v in py["fields"].values()), py  # the prior year's own printed row, verbatim
    out = x.extract([tangen62_full], [1], dm_ship, {"fiscal_year": 2025})  # the shipped schema's explicit-values rule decides the gate identically
    assert {k: v["value"] for k, v in out["prior_year"]["fields"].items()} == {
        "total_debt": 353022, "due_within_1_year": 90970, "due_1_to_5_years": 257532, "due_after_5_years": 4520}, out["prior_year"]
    # The single-year page alone (v083's fixture): no prior-year twin row exists, and the bucket-column
    # table has no year header for a column read either -- no prior_year key at all, never a null one.
    out = x.extract([tangen62_2025], [1], dm, {"fiscal_year": 2025})
    assert "prior_year" not in out, out.get("prior_year")
    # Ework, end to end on the real p.70 shape (seed5-kb: the 2025 block, its Accounts payable and
    # Total rows, then the 2024 block; the year labels are lone lines that _page_rows glues onto the
    # row above -- v076's own account). The model's own answer is v078's (total_debt already reads the
    # Carrying amount column); the current-year fields are v078's, and the prior year is the 2024
    # block's same-labelled row: Due nil + 30,971 + 152,377 + 11,319 = 194,667 in the carrying column,
    # both dash buckets the report's explicit 0s, closing 194,667 + 0 + 0 == 194,667 exactly.
    ework_two_years = ("Maturity structure financial liabilities – undiscounted cash flows\n"
                       "The Group\n"
                       "kSEK Due < 1 month 1-3 months 3-12 months 1-5 years > 5 years Total undiscounted value Carrying amount\n"
                       "2025\n"
                       "Lease liabilities – – 5,230 5,332 22,169 794 33,525 33,403\n"
                       "Short-term interest-bearing liabilities* – 153,761 971 1,677 – – 156,410 156,410\n"
                       "Accounts payable 91,284 1,472,001 992,513 120,851 – – 2,676,650 2,676,650\n"
                       "Total 91,284 1,625,763 998,715 127,860 22,169 794 2,866,584 2,866,462\n"
                       "2024\n"
                       "Lease liabilities – – 3,000 193 7,518 19,973 30,684 27,918\n"
                       "Short-term interest-bearing liabilities* – 30,971 152,377 11,319 – – 194,667 194,667\n"
                       "Accounts payable 103,815 1,838,069 1,088,419 47,791 – – 3,078,094 3,078,094\n"
                       "Total 103,815 1,869,039 1,243,796 59,303 7,518 19,973 3,303,445 3,300,679\n")
    x.call_llm = lambda *a, **k: {"fields": [
        {"key": "total_debt", "value": 156410, "unit": "kSEK", "period": "2025", "raw_label": "Short-term interest-bearing liabilities*",
         "source": {"page": 1, "quote": "Short-term interest-bearing liabilities* – 153,761 971 1,677 – – 156,410 156,410"}},
        {"key": "due_within_1_year", "value": None, "unit": None, "period": None, "raw_label": None, "source": None},
        {"key": "due_1_to_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None},
        {"key": "due_after_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None}]}
    out = x.extract([ework_two_years], [1], dm, {"fiscal_year": 2025})
    got = {f["key"]: f["value"] for f in out["fields"]}
    assert got == {"total_debt": 156410, "due_within_1_year": 156409, "due_1_to_5_years": 0, "due_after_5_years": 0} \
        and out["checks"][0]["passed"], (got, out["checks"], out["warnings"])  # v078's current-year read, unchanged on the bigger page
    py = out["prior_year"]
    assert py["fiscal_year"] == 2024 and py["check"]["passed"], py
    assert {k: v["value"] for k, v in py["fields"].items()} == {
        "total_debt": 194667, "due_within_1_year": 194667, "due_1_to_5_years": 0, "due_after_5_years": 0}, py
    assert all(v["source"]["quote"] == "Short-term interest-bearing liabilities* – 30,971 152,377 11,319 – – 194,667 194,667"
               for v in py["fields"].values()), py
    # MedCap, end to end on the v052/v058 flow (real p.101 shape; the model reads total_debt and
    # due_1_to_5_years off the carrying table, the derivation fills due_within_1_year from the two
    # month rows): the prior year is the Koncernen 2024 COLUMN of the same rows -- 41.8 + 19.0 = 60.8
    # within 1 year, 316.7 in 1-5 years, 377.6 total. The table prints no >5y row for EITHER year, so
    # due_after_5_years is absent from the prior-year fields and named in the check detail; the gate
    # closes on the explicit values alone (60.8 + 316.7 = 377.5 ~= 377.6, within the check's own +-2).
    x.call_llm = lambda *a, **k: {"fields": [
        {"key": "total_debt", "value": 718.7, "unit": "MSEK", "period": "2025", "raw_label": "Totalt", "source": {"page": 1, "quote": "Totalt 718,7 377,6 – –"}},
        {"key": "due_within_1_year", "value": None, "unit": None, "period": None, "raw_label": None, "source": None},
        {"key": "due_1_to_5_years", "value": 616.5, "unit": "MSEK", "period": "2025", "raw_label": "1 – 5 år", "source": {"page": 1, "quote": "1 – 5 år 616,5 316,7 – –"}},
        {"key": "due_after_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None}]}
    out = x.extract([medcap101, medcap102], [1, 2], dm, {"fiscal_year": 2025})
    py = out["prior_year"]
    assert py["fiscal_year"] == 2024 and py["check"]["passed"] and "due_after_5_years" in py["check"]["detail"], py
    assert {k: v["value"] for k, v in py["fields"].items()} == {
        "total_debt": 377.6, "due_within_1_year": 60.8, "due_1_to_5_years": 316.7}, py  # no due_after_5_years key: no >5y row exists
    assert py["fields"]["due_within_1_year"]["source"] == {
        "page": 1, "quote": "6 månader eller mindre 54,3 41,8 – – 6 – 12 månader 48,0 19,0 – –"}, py["fields"]["due_within_1_year"]
    assert py["fields"]["total_debt"]["source"] == {"page": 1, "quote": "Totalt 718,7 377,6 – –"}, py["fields"]["total_debt"]
    # Counter-example 1: a date-per-instrument note (Proact, v073's own fixture and answer) has no
    # deterministic prior-year read -- the current-year buckets are date sums, the year-header table on
    # the page is a different table, and the admission check keeps its rows out. No prior_year key.
    x.call_llm = lambda *a, **k: {"fields": [
        {"key": "total_debt", "value": 478611, "unit": "SEK thousand", "period": "2025", "raw_label": "Total interest-bearing liabilities",
         "source": {"page": 1, "quote": "Total interest-bearing liabilities 478,611"}},
        {"key": "due_within_1_year", "value": 312458, "unit": "SEK thousand", "period": "2025",
         "raw_label": "Bank loan, Svensk Exportkredit 2); Lease liability 3)",
         "source": {"page": 1, "quote": "Lease liability 3) 3.86% - 5.59% 2026 96,098"}},
        {"key": "due_1_to_5_years", "value": 166153, "unit": "SEK thousand", "period": "2025", "raw_label": "Lease liability 3)",
         "source": {"page": 1, "quote": "Lease liability 3) 4.06% - 5.59% 2027-2030 166,153"}},
        {"key": "due_after_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None}]}
    out = x.extract([proact103], [1], dm, {"fiscal_year": 2025})
    assert "prior_year" not in out, out.get("prior_year")
    # Counter-example 2: the prior year's own buckets must close against the prior year's own total --
    # a twin row whose figures do not sum (a misglued or wrong-scope row) is declined whole, exactly
    # like every other identity-gated write in this file.
    tangen_bad = tangen62_2025 + (
        "31 december 2024 (KSEK)\nMindre än 12 månader\nMellan 1 och 2 år Mellan 3 och 5 år Senare än 5 år Summa\n"
        "Leverantörsskulder och övriga skulder (exklusive icke finansiella skulder) 163 517 - - - 163 517\n"
        "Lån Kreditinstitut 90 970 68 704 188 828 4 520 999 999\n")
    x.call_llm = lambda *a, **k: {"fields": [
        {"key": k, "value": None, "unit": None, "period": None, "raw_label": None, "source": None} for k in dmf]}
    out = x.extract([tangen_bad], [1], dm, {"fiscal_year": 2025})
    assert "prior_year" not in out, out.get("prior_year")  # 90,970 + 257,532 + 4,520 != 999,999: no key, not a failing one
    # v110: a borrowings note split into a non-current and a current section, each closed by its own
    # Summa row, with NO third grand-total row in the block (NOTE's Not 19, real p.107 text trimmed
    # to the note; the model answer is v088 pass 2's shape -- within1y 595,420 quote-verified,
    # total_debt 824,281 dropped as computed-not-read, the round's named failure). The two Summa rows
    # sum to 824,281 (228,861 + 595,420), each section closing on its own rows in both columns; the
    # report's own total is printed only as p.109's prose "räntebärande skulder 824,2 (630,0) MSEK" --
    # the thousands-to-millions conversion _value_in_quote structurally cannot see (v099 finding 4).
    # total_debt is written value_derived (the sum is printed nowhere); the quote is the contiguous
    # span from Summa A through Summa B, carrying both Summa rows verbatim -- a two-row join of
    # non-adjacent rows is not a page substring under quote_on_page's gap discipline, and the span is
    # the _between_rows precedent. due_within_1_year already equals the current section's own Summa,
    # so the model's own evidence stands untouched; the two finer buckets stay honestly missing.
    note107 = ("Not 19. Räntebärande skulder\n"
               "Långfristiga skulder 2025-12-31 2024-12-31\n"
               "Banklån 7 241 9 112\n"
               "Leasingskulder hänförlig till materiella anläggningstillgångar 118 027 134 803\n"
               "Leasingskulder hänförlig till nyttjanderättstillgångar för hyrda fastigheter 103 593 106 108\n"
               "Summa 228 861 250 023\n"
               "Kortfristiga skulder\n"
               "Checkräkningskredit 135 399 -\n"
               "Factoring (fakturabelåning) 322 054 316 185\n"
               "Kortfristig del av banklån 375 375\n"
               "Kortfristig del av leasingskulder hänförliga till materiella anläggningstillgångar 41 070 38 064\n"
               "Kortfristig del av leasingskulder hänförlig till nyttjanderättstillgångar\n"
               "för hyrda fastigheter 23 992 25 484\n"
               "Tilläggsköpeskilling 72 530 -\n"
               "Summa 595 420 380 108\n"
               "Ställda säkerheter\n"
               "Säkerhet för banklån, leasingskulder och checkräkningskrediter är utställda med ett belopp av 15 833 (16 812) i företa­\n"
               "gets mark och byggnader (se även not 13) samt 198 176 (198 539) i rörelsen.\n")
    note109 = ("Likviditetsrisker\n"
               "Med likviditetsrisk avses risken att inte kunna uppfylla betalningsförpliktelser till följd av otillräcklig likviditet eller\n"
               "svårigheter att uppta externa lån. Verksamheten finansieras bland annat genom eget kapital 1 648,2 (1 638,5) MSEK\n"
               "och räntebärande skulder 824,2 (630,0) MSEK, utnyttjad checkräkningskredit ingår med 135,4 (0) MSEK.\n")
    x.call_llm = lambda *a, **k: {"fields": [
        {"key": "total_debt", "value": 824281, "unit": "TSEK", "period": "2025", "raw_label": "Räntebärande skulder",
         "source": {"page": 1, "quote": "Räntebärande skulder 824 281 630 023"}},
        {"key": "due_within_1_year", "value": 595420, "unit": "TSEK", "period": "2025", "raw_label": "Summa",
         "source": {"page": 1, "quote": "Summa 595 420 380 108"}},
        {"key": "due_1_to_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None},
        {"key": "due_after_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None}]}
    out = x.extract([note107, note109], [1, 2], dm_ship, {"fiscal_year": 2025})
    got = {f["key"]: f for f in out["fields"]}
    td = got["total_debt"]
    assert td["value"] == 824281 and td["confidence"] == 1.0 and "value_derived" in td["evidence"] \
        and "value_in_quote" not in td["evidence"], td
    assert td["raw_label"] == "Summa (Långfristiga skulder) + Summa (Kortfristiga skulder)", td
    assert td["source"] == {"page": 1, "quote": "Summa 228 861 250 023 Kortfristiga skulder Checkräkningskredit 135 399 - "
                                                "Factoring (fakturabelåning) 322 054 316 185 Kortfristig del av banklån 375 375 "
                                                "Kortfristig del av leasingskulder hänförliga till materiella anläggningstillgångar 41 070 38 064 "
                                                "Kortfristig del av leasingskulder hänförlig till nyttjanderättstillgångar för hyrda fastigheter 23 992 25 484 "
                                                "Tilläggsköpeskilling 72 530 - Summa 595 420 380 108"}, td["source"]
    w1 = got["due_within_1_year"]
    assert w1["value"] == 595420 and w1["confidence"] == 0.9 and "value_in_quote" in w1["evidence"] \
        and "value_derived" not in w1["evidence"] and w1["source"] == {"page": 1, "quote": "Summa 595 420 380 108"}, w1
    assert got["due_1_to_5_years"]["value"] is None and got["due_after_5_years"]["value"] is None
    assert out["checks"][0]["detail"] == "missing: due_1_to_5_years", out["checks"]
    assert any("sum to 824281 (228861 + 595420)" in w and "corroborated by page 2 prose" in w for w in out["warnings"]), out["warnings"]
    # counter-example 1 (Karnov p.135, real text trimmed to the table): current / non-current / total
    # are three COLUMNS of one row, not two sectioned subtotals -- no section headings anywhere, so
    # the walk never starts. v088 pass 1's own read stands, byte-identically.
    karnov135 = ("* at December 31\nType of borrowing Interest rate* Currency\nNominal value\nin currency\n"
                 "Carrying amount,\ncurrent\nCarrying amount,\nnon-current\nTotal, carrying\namount\n"
                 "Borrowings from credit institutions December 31, 2025\n"
                 "Facility Agreement, Facility A Loan 1,7% + EURIBOR EUR 12.5 135.2 - 135.2\n"
                 "Facility Agreement, Facility B Loan 1,7% + EURIBOR EUR 170.0 - 1,839.0 1,839.0\n"
                 "Other Loan N/A EUR 0.2 - 1.7 1.7\nAmortised loan costs - -10.9 -10.9\n"
                 "Total, carrying amount 135.2 1,829.8 1,965.0\n"
                 "Borrowings from credit institutions December 31, 2024\n"
                 "Facility Agreement, Facility A Loan 2,3% + EURIBOR EUR 22.5 114.9 143.5 258.4\n"
                 "Other Loan N/A EUR 0.2 2.4 2.4\nAmortised loan costs -14.3 -14.3\n"
                 "Total, carrying amount 114.9 2,570.9 2,685.8\n")
    x.call_llm = lambda *a, **k: {"fields": [
        {"key": "total_debt", "value": 1965, "unit": "MEUR", "period": "2025", "raw_label": "Total, carrying amount",
         "source": {"page": 1, "quote": "Total, carrying amount 135.2 1,829.8 1,965.0"}},
        {"key": "due_within_1_year", "value": 135.2, "unit": "MEUR", "period": "2025", "raw_label": "Total, carrying amount",
         "source": {"page": 1, "quote": "Total, carrying amount 135.2 1,829.8 1,965.0"}},
        {"key": "due_1_to_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None},
        {"key": "due_after_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None}]}
    out = x.extract([karnov135], [1], dm_ship, {"fiscal_year": 2025})
    got = {f["key"]: (f["value"], f["confidence"]) for f in out["fields"]}
    assert got == {"total_debt": (1965, 0.9), "due_within_1_year": (135.2, 0.9),
                   "due_1_to_5_years": (None, 0.0), "due_after_5_years": (None, 0.0)}, (got, out["warnings"])
    assert not any("section subtotal" in w for w in out["warnings"]), out["warnings"]
    # counter-example 2 (Ambea G18 again, the model answering nothing): both sections and both
    # subtotals are there, but the block prints its own grand total right after the current section
    # ("Total interest-bearing liabilities 12,643") -- that row is total_debt's own read, and the
    # pair shape declines. The pre-existing statement-spread fill reads the printed total instead.
    x.call_llm = lambda *a, **k: {"fields": [
        {"key": k, "value": None, "unit": None, "period": None, "raw_label": None, "source": None} for k in dmf]}
    out = x.extract([ambea119], [1], dm_ship, {"fiscal_year": 2025})
    got = {f["key"]: f["value"] for f in out["fields"]}
    assert got["total_debt"] == 12643 and got["due_within_1_year"] is None, (got, out["warnings"])
    assert not any("section subtotal" in w for w in out["warnings"]), out["warnings"]
    # counter-example 3 (synthetic): a third Summa -- another table's own -- printed between the two
    # section Summas, both sections closing arithmetically so only the structure can decline it:
    # the walk from Summa A to the next section heading meets printed figures, the "one note"
    # premise fails, nothing is adopted.
    interleaved = ("Not 19. Räntebärande skulder\nLångfristiga skulder 2025-12-31 2024-12-31\n"
                   "Banklån 7 241 9 112\nLeasingskulder 221 620 240 911\nSumma 228 861 250 023\n"
                   "Ansvarsförbindelser\nAnsvarsbelopp 100 200\nSumma 100 200\nKortfristiga skulder\n"
                   "Checkräkningskredit 135 399 -\nFactoring 460 021 380 108\nSumma 595 420 380 108\n")
    out = x.extract([interleaved], [1], dm_ship, {"fiscal_year": 2025})
    assert all(f["value"] is None for f in out["fields"]), out["fields"]
    assert not any("section subtotal" in w for w in out["warnings"]), out["warnings"]
    # the same shape on a smaller synthetic note: the current-section fill and replacement. The
    # model read one instrument row (135,399, quote-verified) for within1y -- the section's own
    # Summa (457,453) wins, warning naming the read it replaced; and a VERIFIED model total that
    # disagrees with the two subtotals (696,629, a balance-sheet row) declines silently: the note
    # and the model then name different scopes, the model's answer stands.
    pair_plain = ("Not 19. Räntebärande skulder\nLångfristiga skulder 2025-12-31 2024-12-31\n"
                  "Banklån 7 241 9 112\nLeasingskulder 118 027 134 803\nSumma 125 268 143 915\n"
                  "Kortfristiga skulder\nCheckräkningskredit 135 399 -\nFactoring 322 054 316 185\n"
                  "Summa 457 453 316 185\n")
    x.call_llm = lambda *a, **k: {"fields": [
        {"key": "total_debt", "value": None, "unit": None, "period": None, "raw_label": None, "source": None},
        {"key": "due_within_1_year", "value": 135399, "unit": "TSEK", "period": "2025", "raw_label": "Checkräkningskredit",
         "source": {"page": 1, "quote": "Checkräkningskredit 135 399 -"}},
        {"key": "due_1_to_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None},
        {"key": "due_after_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None}]}
    out = x.extract([pair_plain], [1], dm_ship, {"fiscal_year": 2025})
    got = {f["key"]: f for f in out["fields"]}
    assert got["total_debt"]["value"] == 582721 and "value_derived" in got["total_debt"]["evidence"], (got, out["warnings"])
    w1 = got["due_within_1_year"]
    assert w1["value"] == 457453 and w1["source"] == {"page": 1, "quote": "Summa 457 453 316 185"} and "value_in_quote" in w1["evidence"], w1
    assert any("135399 is not the current section's own subtotal" in w for w in out["warnings"]), out["warnings"]
    x.call_llm = lambda *a, **k: {"fields": [
        {"key": "total_debt", "value": 696629, "unit": "TSEK", "period": "2025", "raw_label": "Räntebärande skulder",
         "source": {"page": 1, "quote": "Balansräkningen Räntebärande skulder 696 629 571 428"}},
        {"key": "due_within_1_year", "value": None, "unit": None, "period": None, "raw_label": None, "source": None},
        {"key": "due_1_to_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None},
        {"key": "due_after_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None}]}
    out = x.extract([pair_plain + "Balansräkningen Räntebärande skulder 696 629 571 428\n"], [1], dm_ship, {"fiscal_year": 2025})
    got = {f["key"]: f["value"] for f in out["fields"]}
    assert got == {"total_debt": 696629, "due_within_1_year": None, "due_1_to_5_years": None, "due_after_5_years": None}, (got, out["warnings"])
    assert not any("section subtotal" in w for w in out["warnings"]), out["warnings"]
    # v131: Enea's Note 23 continuation (p.76, real text trimmed to the note) -- the second clean
    # instance of v108's subtotals ruling (v119), and the first whose note names NO debt word at
    # all: the table header row is "Financial liabilities 2025 2024 2025 2024" (the schema's lists
    # carry totals and instrument rows, never the subject line), and the section headings print the
    # non-breaking hyphen U+2011 ("Non‑current liabilities, interest‑bearing") -- untranslated, the
    # ASCII tail "current" filed the NON-current heading under B (A rules first only while the
    # hyphen normalizes). Two gates blocked the walk; both widen inside _subtotal_pair_fill: the
    # title gate gains the note-family subject wording, _vocab folds U+2011. The pair then fires:
    # 169,152 + 136,403 = 305,555 (printed nowhere -- p.35's prose "SEK 305.6 million" sits outside
    # the candidate window), each Summa closing on its own row in all four columns (Group and
    # Parent side by side, the Group's 2025 first); the model's own bucket reads -- 136,403 within
    # 1y, 169,152 1-5y, both exactly the note's two sections -- stand untouched.
    enea76 = ("Note 23 Financial assets and liabilities, cont.\n"
              "Group Parent Company\n"
              "Financial liabilities 2025 2024 2025 2024\n"
              "Non‑current liabilities, interest‑bearing\n"
              "Liabilities to credit institutions1) 169,152 226,747 169,152 226,747\n"
              "Total non‑current liabilities, interest‑bearing 169,152 226,747 169,152 226,747\n"
              "Current liabilities, interest-bearing\n"
              "Liabilities to credit institutions1) 136,403 51,314 136,403 51,314\n"
              "Total current liabilities, interest-bearing 136,403 51,314 136,403 51,314\n"
              "Current liabilities, non-interest bearing\n"
              "Accounts payable 10,123 21,296 2,651 7,402\n"
              "Other liabilities 10,675 10,538 5,419 947\n"
              "Accrued expenses, supplier-related 8,406 16,477 - -\n"
              "Total current liabilities, non-interest bearing 29,204 48,311 8,070 8,349\n"
              "Total financial liabilities 334,759 326,372 313,625 286,410\n"
              "1) Enea has a bank loan with a term of 3 years. The interest rate is EURIBOR 3M (with a floor) plus a market-based margin.\n"
              "Group Parent Company\n"
              "Maturity analysis 2025 2024 2025 2024\n"
              "Non-current and current interest-bearing liabilities\n"
              "Within 1 year after the balance sheet date 136,403 51,314 136,403 51,314\n"
              "Later than 1 year but within 3 years after the balance sheet date 169,152 226,747 169,152 226,747\n"
              "Later than 3 years but within 5 years after the balance sheet date - - - -\n")
    x.call_llm = lambda *a, **k: {"fields": [
        {"key": "total_debt", "value": None, "unit": None, "period": None, "raw_label": None, "source": None},
        {"key": "due_within_1_year", "value": 136403, "unit": "SEK thousands", "period": "2025",
         "raw_label": "Within 1 year after the balance sheet date",
         "source": {"page": 1, "quote": "Within 1 year after the balance sheet date 136,403 51,314 136,403 51,314"}},
        {"key": "due_1_to_5_years", "value": 169152, "unit": "SEK thousands", "period": "2025",
         "raw_label": "Later than 1 year but within 3 years after the balance sheet date",
         "source": {"page": 1, "quote": "Later than 1 year but within 3 years after the balance sheet date 169,152 226,747 169,152 226,747"}},
        {"key": "due_after_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None}]}
    out = x.extract([enea76], [1], dm_ship, {"fiscal_year": 2025})
    got = {f["key"]: f for f in out["fields"]}
    td = got["total_debt"]
    assert td["value"] == 305555 and td["confidence"] == 1.0 and "value_derived" in td["evidence"] \
        and "value_in_quote" not in td["evidence"], td
    assert td["raw_label"] == ("Total non‑current liabilities, interest‑bearing "
                               "(Non‑current liabilities, interest‑bearing) + Total current "
                               "liabilities, interest-bearing (Current liabilities, interest-bearing)"), td
    assert td["source"] == {"page": 1, "quote": "Total non‑current liabilities, interest‑bearing "
                                                "169,152 226,747 169,152 226,747 Current liabilities, interest-bearing "
                                                "Liabilities to credit institutions1) 136,403 51,314 136,403 51,314 "
                                                "Total current liabilities, interest-bearing 136,403 51,314 136,403 51,314"}, td["source"]
    w1 = got["due_within_1_year"]
    assert w1["value"] == 136403 and w1["confidence"] == 1.0 and "value_in_quote" in w1["evidence"] \
        and "value_derived" not in w1["evidence"] and w1["source"]["quote"] == \
        "Within 1 year after the balance sheet date 136,403 51,314 136,403 51,314", w1
    assert got["due_1_to_5_years"]["value"] == 169152 and got["due_after_5_years"]["value"] is None
    assert out["checks"][0]["detail"] == "missing: due_after_5_years", out["checks"]  # no >5y row printed: honestly missing
    assert any("sum to 305555 (169152 + 136403)" in w and "the two closures alone prove it" in w
               for w in out["warnings"]), out["warnings"]
    # the same note with the total pretending a citation: 305,555 quoted off a row that is not
    # printed. The fabricated read is dropped as computed-not-read, and the note's own sections
    # still close the total -- the adoption stands on the note's proof, not the model's.
    x.call_llm = lambda *a, **k: {"fields": [
        {"key": "total_debt", "value": 305555, "unit": "SEK thousands", "period": "2025",
         "raw_label": "Total interest-bearing liabilities",
         "source": {"page": 1, "quote": "Total interest-bearing liabilities 305,555"}},
        {"key": "due_within_1_year", "value": 136403, "unit": "SEK thousands", "period": "2025",
         "raw_label": "Within 1 year after the balance sheet date",
         "source": {"page": 1, "quote": "Within 1 year after the balance sheet date 136,403 51,314 136,403 51,314"}},
        {"key": "due_1_to_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None},
        {"key": "due_after_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None}]}
    out = x.extract([enea76], [1], dm_ship, {"fiscal_year": 2025})
    got = {f["key"]: f for f in out["fields"]}
    assert got["total_debt"]["value"] == 305555 and got["total_debt"]["confidence"] == 1.0 \
        and got["total_debt"]["source"]["quote"].startswith("Total non‑current liabilities"), got["total_debt"]
    assert any("dropped as computed, not read" in w for w in out["warnings"]), out["warnings"]
    # counter-example 1: the widened title word manufactures nothing on its own -- a "Financial
    # liabilities" header over sections that are not the current/non-current pair (here a flat
    # non-interest-bearing list closed by its own grand total) never starts the walk: no A heading.
    flat = ("Financial liabilities 2025 2024\n"
            "Accounts payable 10,123 21,296\nOther liabilities 10,675 10,538\n"
            "Accrued expenses 8,406 16,477\nTotal financial liabilities 29,204 48,311\n")
    x.call_llm = lambda *a, **k: {"fields": [
        {"key": k, "value": None, "unit": None, "period": None, "raw_label": None, "source": None} for k in dmf]}
    out = x.extract([flat], [1], dm_ship, {"fiscal_year": 2025})
    assert all(f["value"] is None for f in out["fields"]), out["fields"]
    assert not any("section subtotal" in w for w in out["warnings"]), out["warnings"]
    # counter-example 2: the U+2011 hyphen alone neither fires nor widens the subject -- a
    # receivables note printing the same non-breaking-hyphen section headings has no title word
    # ("Financial assets" is not one), so the walk never starts however well-formed its pair.
    assets = ("Financial assets 2025 2024\n"
              "Non‑current receivables\n"
              "Tenant loans 500 400\nSumma 500 400\n"
              "Current receivables\n"
              "Trade receivables 300 200\nSumma 300 200\n")
    out = x.extract([assets], [1], dm_ship, {"fiscal_year": 2025})
    assert all(f["value"] is None for f in out["fields"]), out["fields"]
    assert not any("section subtotal" in w for w in out["warnings"]), out["warnings"]
    # counter-example 3: with both gates open, the note's own arithmetic still decides -- section
    # A's detail row no longer summing to its Summa declines the whole block silently.
    broken = enea76.replace("Liabilities to credit institutions1) 169,152 226,747 169,152 226,747",
                            "Liabilities to credit institutions1) 160,000 226,747 160,000 226,747")
    x.call_llm = lambda *a, **k: {"fields": [
        {"key": k, "value": None, "unit": None, "period": None, "raw_label": None, "source": None} for k in dmf]}
    out = x.extract([broken], [1], dm_ship, {"fiscal_year": 2025})
    assert all(f["value"] is None for f in out["fields"]), out["fields"]
    assert not any("section subtotal" in w for w in out["warnings"]), out["warnings"]
    # v156: Sdiptech's Note 24 (real p.110 text trimmed to the note) -- v110's pair shape whose
    # section rows print footnote stars and small cells the no-ncols read cannot place:
    # "Contingent considerations * 597 910" reads fused as 597910 (the star breaks the all-digits
    # tail the ncols split needs, and its split-back is a nil-gated Boozt case), while "Other
    # liabilities** 2 4" and "Liabilities to credit institutions 10 10" read [] and passed for
    # wrapped label lines -- three reading defects, each enough to fail a section closure. The
    # section's own ncols read of the star-stripped row arbitrates; the closures still decide.
    sdiptech110 = ("Note 24 Interest-bearing liabilities\n"
                   "Group interest-bearing liabilities\n"
                   "Non-current liabilities 2025 2024\n"
                   "Liabilities to credit institutions 2,309 1,910\n"
                   "Bond liabilities 800 811\n"
                   "Leasing liabilities 342 393\n"
                   "Contingent considerations * 597 910\n"
                   "Other liabilities** 2 4\n"
                   "Total 4,049 4,027\n"
                   "Current liabilities 2025 2024\n"
                   "Liabilities to credit institutions 10 10\n"
                   "Leasing liabilities 119 120\n"
                   "Contingent considerations * 317 406\n"
                   "Other liabilities** 1 1\n"
                   "Total 446 537\n")
    x.call_llm = lambda *a, **k: {"fields": [
        {"key": k, "value": None, "unit": None, "period": None, "raw_label": None, "source": None} for k in dmf]}
    out = x.extract([sdiptech110], [1], dm_ship, {"fiscal_year": 2025})
    got = {f["key"]: f for f in out["fields"]}
    td = got["total_debt"]
    assert td["value"] == 4495 and "value_derived" in td["evidence"], td
    assert td["raw_label"] == "Total (Non-current liabilities) + Total (Current liabilities)", td
    w1 = got["due_within_1_year"]
    assert w1["value"] == 446 and "value_in_quote" in w1["evidence"] and "value_derived" not in w1["evidence"], w1
    assert any("sum to 4495 (4049 + 446), each closing on its own rows" in w for w in out["warnings"]), out["warnings"]
    # counter: the star row's own figures no longer closing its section -- the un-fusion alone
    # proves nothing, the per-column closures decide, nothing is adopted and nothing warns
    broken_pair = sdiptech110.replace("Contingent considerations * 597 910", "Contingent considerations * 500 910")
    x.call_llm = lambda *a, **k: {"fields": [
        {"key": k, "value": None, "unit": None, "period": None, "raw_label": None, "source": None} for k in dmf]}
    out = x.extract([broken_pair], [1], dm_ship, {"fiscal_year": 2025})
    assert all(f["value"] is None for f in out["fields"]), out["fields"]
    assert not any("section subtotal" in w for w in out["warnings"]), out["warnings"]
    # v111: the wrong-table guard's second batch, two more _table_scope refusals under BOTH bases.
    # (a) non-debt-subject tables -- Volati p.177's "Timing of revenue recognition, contract liabilities":
    #     its "Within 1 year 87 87 1" is contract-liability timing, and the statement-spread fill wrote it
    #     over the model's own correct current-total 192 (Note 17 prints no maturity table at all; Salix,
    #     the same report family, is the label-pinned instance -- the label names 192 and calls the stored
    #     87 "Note 19 contract liabilities ... not debt"); (b) Parent Company sections -- Momentum p.111's
    #     parent lease maturity under a "Parent Company" heading: 2/2/0 against the Group's 622 balance-
    #     sheet total. Every counter-example keeps its read: v052's paired Koncernen|Moderbolaget column
    #     headers, a Group heading above the table, debt words in the title or rows, the carrying column.
    dsw = x._debt_subject_words(dm)
    pc = tsw["parent_company"]
    # the section-heading test itself: the entity word alone, trailing unit/date furniture excepted
    assert x._section_heading("Parent Company", pc) == "parent"
    assert x._section_heading("Parent Company, MSEK 31 Dec 2025 31 Dec 2024", pc) == "parent"  # Momentum p.117
    assert x._section_heading("Moderbolaget 2025 2024", pc) == "parent"  # Svedbergs p.132's parent block
    assert x._section_heading("Moderföretaget", pc) == "parent"
    assert x._section_heading("Group, MSEK 31 Dec 2025 31 Dec 2024", pc) == "group"  # the return to the Group
    assert x._section_heading("Group", pc) == "group" and x._section_heading("Koncernen", pc) == "group"
    assert x._section_heading("Koncernen Moderbolaget", pc) is None  # v052's paired columns: one table, not a section
    assert x._section_heading("Förfallotidpunkt för upplåning Koncernen Moderbolaget", pc) is None  # MedCap p.101
    assert x._section_heading("Moderbolaget har inga räntebärande skulder", pc) is None  # prose, not a heading
    assert x._section_heading("Parent Company financial statements", pc) is None  # a TOC line names no section
    # (a) header level, Volati p.177 trimmed to the timing table (real text; _page_rows glues the wrapped
    # title into one row): refused under both bases, whichever bucket row anchors the walk
    volati_timing = ("NOTE 19 | Contract assets and liabilities\n"
                     "Contract assets Contract liabilities\n"
                     "Opening balance, 1 Jan 2025 54 131 9 140 2\n"
                     "Projects accrued during the year 314 –81 –1 –83\n"
                     "Closing balance, 31 Dec 2025 104 87 8 95 1\n"
                     "Timing of revenue recognition, contract\n"
                     "liabilities\n"
                     "Within 1 year 87 87 1\n"
                     "1–2 years 2 2 –\n"
                     "2–5 years 3 3 –\n"
                     "After 5 years 2 2 –\n"
                     "87 8 95 1\n")
    vt_rows = x._page_rows(volati_timing)
    assert "Timing of revenue recognition, contract liabilities" in vt_rows, vt_rows
    for basis in ("carrying", "undiscounted"):
        assert x._table_scope(vt_rows, None, vt_rows.index("Within 1 year 87 87 1"), basis, tsw, debt_words=dsw) == "non_debt"
        assert x._table_scope(vt_rows, None, vt_rows.index("1–2 years 2 2 –"), basis, tsw, debt_words=dsw) == "non_debt"
    # debt words win: the work order's made-up page -- same "Contract liabilities" title, but the table's
    # own rows are the borrowings note's ("Interest-bearing liabilities") -- is NOT refused
    made_debt = ("Contract liabilities, timing of revenue\n"
                 "MSEK 2025 2024\n"
                 "Interest-bearing liabilities 500 480\n"
                 "Within 1 year 120 100\n"
                 "Between 1 and 5 years 380 380\n"
                 "Total 500 480\n")
    md_rows = x._page_rows(made_debt)
    assert x._table_scope(md_rows, None, len(md_rows) - 1, "carrying", tsw, debt_words=dsw) == "unknown", md_rows
    made_pure = ("Contract liabilities, timing of revenue\n"
                 "MSEK 2025 2024\n"
                 "Within 1 year 120 100\n"
                 "Between 1 and 5 years 380 380\n"
                 "Total 500 480\n")
    mp_rows = x._page_rows(made_pure)
    assert x._table_scope(mp_rows, None, len(mp_rows) - 1, "carrying", tsw, debt_words=dsw) == "non_debt", mp_rows
    # (b) header level, Momentum p.111 trimmed (real text): the parent lease table under "Parent Company"
    # is the parent's own, whichever of its rows anchors; a Rusta-shaped page (bare "Group" heading above
    # the Group's maturity table, the parent's own table below it) reads the Group rows untouched
    momentum111 = ("11 Leases\n"
                   "Right-of-use assets\n"
                   "Group, MSEK Premises Vehicles Other Total\n"
                   "Closing balance, 31 Dec 2025 187 46 0 233\n"
                   "Lease liabilities\n"
                   "Amounts recognised in profit or loss\n"
                   "MSEK 2025 2024\n"
                   "Depreciation of right-of-use assets –95 –83\n"
                   "Disclosure concerning operating leases in the Parent Company\n"
                   "Parent Company\n"
                   "Non-cancellable lease payments amount to: 31 Dec 2025 31 Dec 2024\n"
                   "Leases in which the company is the lessee\n"
                   "Within 1 year 2 2\n"
                   "Between 1 and 5 years 2 1\n"
                   "Later than 5 years – –\n"
                   "Total 4 3\n")
    mo_rows = x._page_rows(momentum111)
    for basis in ("carrying", "undiscounted"):
        for r in ("Within 1 year 2 2", "Between 1 and 5 years 2 1", "Later than 5 years – –", "Total 4 3"):
            assert x._table_scope(mo_rows, None, mo_rows.index(r), basis, tsw, debt_words=dsw) == "parent", r
    rusta_shape = ("Lease liabilities\n"
                   "Group\n"
                   "Maturity analysis – lease liabilities 30 Apr 2026 30 Apr 2025\n"
                   "0–6 months 523 490\n"
                   ">5 years 2,431 2,633\n"
                   "Total 6,624 6,669\n"
                   "On the balance sheet date, the Parent Company had outstanding commitments\n"
                   "Parent Company\n"
                   "30 Apr 2026 30 Apr 2025\n"
                   "Within one year 607 590\n"
                   "Total 4,547 4,769\n")
    ru_rows = x._page_rows(rusta_shape)
    assert x._table_scope(ru_rows, None, ru_rows.index("Total 6,624 6,669"), "carrying", tsw, debt_words=dsw) == "unknown"
    assert x._table_scope(ru_rows, None, ru_rows.index("Within one year 607 590"), "carrying", tsw, debt_words=dsw) == "parent"
    assert x._table_scope(mc_rows, None, len(mc_rows) - 1, "carrying", tsw, debt_words=dsw) == "unknown"  # MedCap's paired columns, again with the v111 words live
    # end to end, Volati (real p.177 shape, trimmed): the model's own 192 (Note 17's current total, quote
    # not a printed row) survives -- the fill that used to overwrite it with the timing table's 87 refuses,
    # warning that names the table; the finer buckets stay honestly null, the total keeps its verified 3,246.
    # The model's own quote is not stored (only its "quote not found on page 177" warning is), so the stub's
    # quote is shaped to fail verification the way the real one did -- a value the page prints ("192 191")
    # under a quote no suffix rule can find; note "Current liabilities 192 191" would NOT do: its suffix
    # "liabilities ... 192 ... 191" matches the printed "Lease liabilities 180 185 / 192 191" rows through
    # quote_on_page's note-reference gap, and a verified quote means no fill ever fires
    volati177 = ("NOTE 17 | Interest-bearing liabilities\n"
                 "Non-current liabilities 2025 2024\n"
                 "Lease liabilities 375 402\n"
                 "Liabilities to credit institutions 2,678 2,350\n"
                 "3,053 2,753\n"
                 "Current liabilities 2025 2024\n"
                 "Liabilities to credit institutions 12 7\n"
                 "Lease liabilities 180 185\n"
                 "192 191\n"
                 "31 December 3,246 2,944\n"
                 "NOTE 19 | Contract assets and liabilities\n"
                 "Contract assets Contract liabilities\n"
                 "Closing balance, 31 Dec 2025 104 87 8 95 1\n"
                 "Timing of revenue recognition, contract\n"
                 "liabilities\n"
                 "Within 1 year 87 87 1\n"
                 "1–2 years 2 2 –\n"
                 "2–5 years 3 3 –\n"
                 "After 5 years 2 2 –\n")
    x.call_llm = lambda *a, **k: {"fields": [
        {"key": "total_debt", "value": 3246, "unit": "SEK m", "period": "2025", "raw_label": "31 December",
         "source": {"page": 1, "quote": "31 December 3,246 2,944"}},
        {"key": "due_within_1_year", "value": 192, "unit": "SEK m", "period": "2025", "raw_label": "Current liabilities",
         "source": {"page": 1, "quote": "Note 17, current portion 192"}},
        {"key": "due_1_to_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None},
        {"key": "due_after_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None}]}
    out = x.extract([volati177], [1], dm_ship, {"fiscal_year": 2025})
    got = {f["key"]: f["value"] for f in out["fields"]}
    assert got == {"total_debt": 3246, "due_within_1_year": 192, "due_1_to_5_years": None, "due_after_5_years": None}, (got, out["warnings"])
    assert any("quote not found on page 1" in w for w in out["warnings"]), out["warnings"]  # the real run's own first warning, kept
    assert any("fill from page 1 row 'Within 1 year 87 87 1' refused" in w
               and "table 'Timing of revenue recognition, contract liabilities' is a non-debt subject table" in w
               for w in out["warnings"]), out["warnings"]
    # end to end, the stored-seed11 shape (the model itself answered the timing row, quote-verified): the
    # citation gate refuses it -- 87 goes null, the honest outcome the replay showed on the real page
    x.call_llm = lambda *a, **k: {"fields": [
        {"key": "total_debt", "value": 3246, "unit": "SEK m", "period": "2025", "raw_label": "31 December",
         "source": {"page": 1, "quote": "31 December 3,246 2,944"}},
        {"key": "due_within_1_year", "value": 87, "unit": "SEK m", "period": "2025", "raw_label": "Within 1 year",
         "source": {"page": 1, "quote": "Within 1 year 87 87 1"}},
        {"key": "due_1_to_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None},
        {"key": "due_after_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None}]}
    out = x.extract([volati177], [1], dm_ship, {"fiscal_year": 2025})
    got = {f["key"]: f["value"] for f in out["fields"]}
    assert got == {"total_debt": 3246, "due_within_1_year": None, "due_1_to_5_years": None, "due_after_5_years": None}, (got, out["warnings"])
    assert any("87 from 'Within 1 year 87 87 1' refused" in w and "non-debt subject table" in w for w in out["warnings"]), out["warnings"]
    # end to end, Momentum (p.111 + p.117 trimmed): the buckets the seed11 run filled out of the parent's
    # lease table are refused at the fill (the model returned null), the Group's own 622 on p.117 -- under
    # a "Group, MSEK" heading, not the Parent block above it -- keeps its verified read
    momentum117 = ("23 Expected recovery periods for assets, provisions and liabilities\n"
                   "Parent Company, MSEK 31 Dec 2025 31 Dec 2024\n"
                   "Financial liabilities\n"
                   "Liabilities to credit institutions 395 268\n"
                   "Total financial liabilities 884 642\n"
                   "22 Financial assets and liabilities\n"
                   "Group, MSEK 31 Dec 2025 31 Dec 2024\n"
                   "Financial liabilities\n"
                   "Interest-bearing liabilities 622 486\n"
                   "Accounts payable 248 246\n"
                   "Total financial liabilities 1,006 862\n")
    x.call_llm = lambda *a, **k: {"fields": [
        {"key": "total_debt", "value": 622, "unit": "MSEK", "period": "2025", "raw_label": "Interest-bearing liabilities",
         "source": {"page": 2, "quote": "Interest-bearing liabilities 622 486"}},
        {"key": "due_within_1_year", "value": None, "unit": None, "period": None, "raw_label": None, "source": None},
        {"key": "due_1_to_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None},
        {"key": "due_after_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None}]}
    out = x.extract([momentum111, momentum117], [1, 2], dm_ship, {"fiscal_year": 2025})
    got = {f["key"]: (f["value"], f["confidence"]) for f in out["fields"]}
    assert got == {"total_debt": (622, 1.0), "due_within_1_year": (None, 0.0), "due_1_to_5_years": (None, 0.0),
                   "due_after_5_years": (None, 0.0)}, (got, out["warnings"])  # 1.0: "Interest-bearing liabilities" is a row synonym, known for scoring
    assert any("the Parent Company section 'Parent Company' holds this table" in w for w in out["warnings"]), out["warnings"]
    assert not any("filled from page 1 row" in w and "Within 1 year 2 2" in w for w in out["warnings"]), out["warnings"]
    # and the stored-seed11 shape (the model itself cited the parent rows): every bucket citation refused
    x.call_llm = lambda *a, **k: {"fields": [
        {"key": "total_debt", "value": 622, "unit": "MSEK", "period": "2025", "raw_label": "Interest-bearing liabilities",
         "source": {"page": 2, "quote": "Interest-bearing liabilities 622 486"}},
        {"key": "due_within_1_year", "value": 2, "unit": "MSEK", "period": "2025", "raw_label": "Within 1 year",
         "source": {"page": 1, "quote": "Within 1 year 2 2"}},
        {"key": "due_1_to_5_years", "value": 2, "unit": "MSEK", "period": "2025", "raw_label": "Between 1 and 5 years",
         "source": {"page": 1, "quote": "Between 1 and 5 years 2 1"}},
        {"key": "due_after_5_years", "value": 0, "unit": "MSEK", "period": "2025", "raw_label": "Later than 5 years – –",
         "source": {"page": 1, "quote": "Later than 5 years – –"}}]}
    out = x.extract([momentum111, momentum117], [1, 2], dm_ship, {"fiscal_year": 2025})
    got = {f["key"]: f["value"] for f in out["fields"]}
    assert got == {"total_debt": 622, "due_within_1_year": None, "due_1_to_5_years": None, "due_after_5_years": None}, (got, out["warnings"])
    assert sum("refused" in w and "Parent Company" in w for w in out["warnings"]) >= 3, out["warnings"]
    # v127: CellaVision p.84's "Maturity analysis of lease liabilities" joins the refuse vocabulary. v115's
    # live pass1 read it four-up at confidence 1.0 through a passing identity (23,300 = 12,213+8,109+2,978;
    # raw labels unknown to the vocabulary, endorsed by identity_all_columns) against labels 25,651/13,680/
    # 8,768/3,203 off p.92's Note C9 -- a wrong-scope table whose title carries a debt word (so the v111
    # non-debt rule's rescue keeps it readable by design) and no refuse-marker (v103's list only had RVRC's
    # "maturity analysis regarding"). Both new forms are grep-real (the "of" form: 24 stems/25 pages, the
    # "of the" form: 4 stems/4 pages); the work order's Swedish forms have zero corpus hits and are not added.
    assert "maturity analysis of lease liabilities" in tsw["undiscounted"]
    assert "maturity analysis of the lease liabilities" in tsw["undiscounted"]
    # header level, CellaVision p.84 trimmed to its real rows: refused under the carrying basis anchored at
    # any of the four rows the pass1 answer cited. Under the undiscounted basis the family's own semantics
    # keep the read open -- the marker alone never refuses there (RVRC refuses through its payables rows,
    # all_liabilities; a table without them falls to "unknown"), so the verdict is "unknown", unchanged.
    cellavision84 = ("Note B8. Leasing\n"
                     "SEK thousands 2025 2024\n"
                     "Amounts recognized in the income statement Group Group\n"
                     "Buildings and land 10,646 10,141\n"
                     "Equipment, tools, fixtures and fittings 1,602 1,827\n"
                     "Depreciation on right of use 12,248 11,968\n"
                     "Interest expenses for leasing liabilities 485 545\n"
                     "Costs attributable to short-term and leasing contracts of low value 3,684 6,035\n"
                     "As of December 31, 2025, the Group has obligations regarding short-term and leasing agreements of low value of SEK 3,316\n"
                     "thousand (5,431).\n"
                     "SEK thousands 2025 2024\n"
                     "Maturity analysis of lease liabilities Group Group\n"
                     "- Within one year 12,213 18,646\n"
                     "- Later than one but within five years 8,109 8,876\n"
                     "- Later than within five years 2,978 -\n"
                     "Total 23,300 27,522\n"
                     "SEK thousands 2025 2024\n"
                     "Cash flow Group Group\n"
                     "Amortization of leasing liabilities 13,852 12,463\n")
    cv_rows = x._page_rows(cellavision84)
    for a in ("Total 23,300 27,522", "- Within one year 12,213 18,646",
              "- Later than one but within five years 8,109 8,876", "- Later than within five years 2,978 -"):
        assert x._table_scope(cv_rows, None, cv_rows.index(a), "carrying", tsw, debt_words=dsw) == "undiscounted", a
        assert x._table_scope(cv_rows, None, cv_rows.index(a), "undiscounted", tsw, debt_words=dsw) == "unknown", a
    # end to end, the v115 pass1 shape: every field's citation is a printed row of the lease table -- the
    # citation gate nulls all four, each warning naming the table, the check honestly missing
    x.call_llm = lambda *a, **k: {"fields": [
        {"key": "total_debt", "value": 23300, "unit": "SEK thousands", "period": "2025", "raw_label": "Total",
         "source": {"page": 1, "quote": "Total 23,300 27,522"}},
        {"key": "due_within_1_year", "value": 12213, "unit": "SEK thousands", "period": "2025",
         "raw_label": "- Within one year", "source": {"page": 1, "quote": "- Within one year 12,213 18,646"}},
        {"key": "due_1_to_5_years", "value": 8109, "unit": "SEK thousands", "period": "2025",
         "raw_label": "- Later than one but within five years",
         "source": {"page": 1, "quote": "- Later than one but within five years 8,109 8,876"}},
        {"key": "due_after_5_years", "value": 2978, "unit": "SEK thousands", "period": "2025",
         "raw_label": "- Later than within five years",
         "source": {"page": 1, "quote": "- Later than within five years 2,978 -"}}]}
    out = x.extract([cellavision84], [1], dm_ship, {"fiscal_year": 2025})
    got = {f["key"]: f["value"] for f in out["fields"]}
    assert got == {"total_debt": None, "due_within_1_year": None, "due_1_to_5_years": None, "due_after_5_years": None}, (got, out["warnings"])
    assert out["checks"][0]["passed"] is False and "missing: due_within_1_year" in out["checks"][0]["detail"], out["checks"]
    assert sum("refused" in w and "table 'Maturity analysis of lease liabilities Group Group' is undiscounted" in w
               for w in out["warnings"]) == 4, out["warnings"]
    # the narrowing the corpus audit forced: Rusta's Group lease table (dash-form title, lease-only issuer,
    # the label's own source) stays readable -- the new word matches CellaVision's "of" form, not the dash
    # form; BioInvent p.66's labelled rows (its only debt is leases) stay readable too, and its page's own
    # pointer line ("For maturity analysis of lease liabilities, see Note 21 ...") never reaches a title
    # block -- the note reference in it is a digit, and the title walk's digit break stops above such lines
    # (the corpus's other pointer lines, "see Note 2"/"see note 22"/"Note 27", share that furniture), so the
    # new words only ever fire on real title rows like CellaVision's own digit-free one
    rusta_group = ("Lease liabilities\n"
                   "Group\n"
                   "Maturity analysis – lease liabilities 30 Apr 2026 30 Apr 2025\n"
                   "0–6 months 523 490\n"
                   ">5 years 2,431 2,633\n"
                   "Total 6,624 6,669\n")
    rg_rows = x._page_rows(rusta_group)
    assert x._table_scope(rg_rows, None, rg_rows.index("Total 6,624 6,669"), "carrying", tsw, debt_words=dsw) == "unknown"
    bioinvent66 = ("Lease liabilities\n"
                   "SEK thousand 2025 2024 Note 24 Information about the Parent Company\n"
                   "Long term 1,157 8,215\n"
                   "Short term 7,370 9,198\n"
                   "Lease liabilities included in statement of financial position for the Group 8,527 17,413\n"
                   "For maturity analysis of lease liabilities, see Note 21 Financial assets and liabilities.\n"
                   "Amounts reported in the statement of comprehensive income for the Group\n"
                   "SEK thousand 2025 2024\n"
                   "Depreciation of rights of use assets -8,179 -8,058\n"
                   "Interest costs, leases -389 -533\n"
                   "Costs of low value leases -188 -177\n"
                   "Total -8,756 -8,768\n")
    bi_rows = x._page_rows(bioinvent66)
    for a in ("Long term 1,157 8,215", "Short term 7,370 9,198"):
        assert x._table_scope(bi_rows, None, bi_rows.index(a), "carrying", tsw, debt_words=dsw) == "unknown", a
    assert x._table_scope(bi_rows, None, bi_rows.index("Total -8,756 -8,768"), "carrying", tsw, debt_words=dsw) == "unknown"
    # v109: the report's own calendar-year maturity columns ride along as top-level buckets_by_year
    # metadata -- never a model answer, only the year columns the bucket-column reader itself
    # aligned (the _bucket_year_hits header fallback behind the recorded pick), re-read on the
    # pick's own row, one entry per printed column (a dash = the report's explicit 0, v078 one
    # window finer), gated on the schema's own identity against the extraction's total_debt within
    # the check's own +-2. The three bucket fields and their summation of the year columns are
    # v036/v095's, byte-identical; row-shaped year tables and date/interval notes give no key.
    # Electrolux Professional, end to end on the real p.174 shape (data/kb pages.jsonl, Note 18
    # "Repayment schedule for long-term borrowings, December 31, 2025"): header "SEKm 2026 2027
    # 2028 2029 2030 2031- Total", the Total row printing two dashes (v078's explicit 0s). The
    # corpus holds exactly one company whose year table passes every existing valve and closes
    # (docs/acrylic/evidence/v109.md); the model answers nothing here, so the read is purely
    # deterministic.
    electrolux174 = ("Note 18 FINANCIAL INSTRUMENTS, CONT. Financial\n"
                     "Repayment schedule for long-term borrowings, December 31, 2025 information\n"
                     "SEKm 2026 2027 2028 2029 2030 2031– Total\n"
                     "Bond loans 400 650 – 250 – – 1,300\n"
                     "Bank and other loans 144 744 144 – – – 1,032\n"
                     "Total 544 1,394 144 250 – – 2,332\n")
    x.call_llm = lambda *a, **k: {"fields": [
        {"key": k, "value": None, "unit": None, "period": None, "raw_label": None, "source": None} for k in dmf]}
    out = x.extract([electrolux174], [1], dm, {"fiscal_year": 2025})
    got = {f["key"]: f["value"] for f in out["fields"]}
    assert got == {"total_debt": 2332, "due_within_1_year": 544, "due_1_to_5_years": 1788, "due_after_5_years": 0} \
        and out["checks"][0]["passed"], (got, out["checks"], out["warnings"])  # v036/v078's own read, unchanged
    by = out["buckets_by_year"]
    assert by["basis"] == "carrying", by
    assert [(y["label"], y["value"]) for y in by["years"]] == [
        ("2026", 544), ("2027", 1394), ("2028", 144), ("2029", 250), ("2030", 0), ("2031–", 0)], by
    assert all(y["source"] == {"page": 1, "quote": "Total 544 1,394 144 250 – – 2,332"} for y in by["years"]), by
    # the same page through the model's own answer (the production shape: total_debt quoted from the
    # row): the years still come off the page and close on the field, not on the row's own column
    x.call_llm = lambda *a, **k: {"fields": [
        {"key": "total_debt", "value": 2332, "unit": "SEKm", "period": "2025", "raw_label": "Total",
         "source": {"page": 1, "quote": "Total 544 1,394 144 250 – – 2,332"}},
        *[{"key": k, "value": None, "unit": None, "period": None, "raw_label": None, "source": None}
          for k in ("due_within_1_year", "due_1_to_5_years", "due_after_5_years")]]}
    out = x.extract([electrolux174], [1], dm, {"fiscal_year": 2025})
    assert [(y["label"], y["value"]) for y in out["buckets_by_year"]["years"]][-1] == ("2031–", 0), out["buckets_by_year"]
    # synthetic, for the paths the one real page cannot reach: a bare "Later" tail word as its own
    # label, dash columns inside the 1-5y span, and closure only within the check's own +-2 (300
    # printed across the years, 301 in the total column -- the same tolerance the identity uses)
    yr301 = "Note 20 Borrowings\nMSEK\n2026 2027 2028 2029 2030 Later Total\nTotal borrowings 120 80 – 40 – 60 301\n"
    x.call_llm = lambda *a, **k: {"fields": [
        {"key": k, "value": None, "unit": None, "period": None, "raw_label": None, "source": None} for k in dmf]}
    out = x.extract([yr301], [1], dm, {"fiscal_year": 2025})
    assert [(y["label"], y["value"]) for y in out["buckets_by_year"]["years"]] == [
        ("2026", 120), ("2027", 80), ("2028", 0), ("2029", 40), ("2030", 0), ("Later", 60)], out["buckets_by_year"]
    # counter-example 1 (real, Cloetta's own Note 21 shape): named bucket columns, no year header
    # -- nothing to read per year, no key, exactly as before
    out = x.extract([cloetta], [1], dm, {"fiscal_year": 2025})
    assert "buckets_by_year" not in out, out.get("buckets_by_year")
    # counter-example 2 (real, Acast p.65's own shape): the year columns ARE on the page, but the
    # carrying column printed before them misaligns the header (v040's over-valve rejection) and
    # the years sum to an undiscounted 701,636 against a carrying 690,785 anyway -- no key
    x.call_llm = lambda *a, **k: {"fields": [
        {"key": "total_debt", "value": 135382, "unit": "SEK thousand", "period": "2025", "raw_label": "Total lease Liability",
         "source": {"page": 1, "quote": "Total lease Liability 135,382 141,152"}},
        *[{"key": k, "value": None, "unit": None, "period": None, "raw_label": None, "source": None}
          for k in ("due_within_1_year", "due_1_to_5_years", "due_after_5_years")]]}
    out = x.extract([leases, acast], [1, 2], dm, {"fiscal_year": 2025})
    assert "buckets_by_year" not in out, out.get("buckets_by_year")
    # counter-example 3 (synthetic): the pick records and the buckets fill, but the years do not
    # close on the extraction's total_debt -- no key, not a failing one
    yr999 = "Note 20 Borrowings\nMSEK\n2026 2027 2028 2029 2030 Later Total\nTotal borrowings 120 80 – 40 – 60 999\n"
    x.call_llm = lambda *a, **k: {"fields": [
        {"key": k, "value": None, "unit": None, "period": None, "raw_label": None, "source": None} for k in dmf]}
    out = x.extract([yr999], [1], dm, {"fiscal_year": 2025})
    assert {f["key"]: f["value"] for f in out["fields"]}["total_debt"] == 999, out["warnings"]
    assert "buckets_by_year" not in out, out.get("buckets_by_year")  # 120+80+0+40+0+60 != 999
    # v109 continuation -- the ROW-shaped year table (seed 11 landed after the lane's first pass): the
    # report prints one figure per calendar year and pymupdf tears the rows into a single line of
    # "<year> <figure>" pairs. BTS p.92 (data/kb pages.jsonl, real text, trimmed to the two tables):
    # "Maturity analyses for liabilities to credit institutions / SEK thousands 12-31-25 2026 77,141
    # 2027 39 2028 300,039 2029 202,539 2030 39 / Total 579,797" -- the years sum exactly onto the
    # model's own verified total (the liabilities table's row above), and the key appears. The pairs
    # chain from fiscal_year+1 exactly; the table's own one-figure Total row must print below them.
    bts92 = ("Liabilities to credit institutions\n"
             "SEK thousands 12-31-25 12-31-24\n"
             "Non-current liabilities 502,656 202,500\n"
             "Current liabilities 77,141 218,453\n"
             "Total 579,797 420,953\n"
             "Maturity analyses for liabilities to credit institutions\n"
             "SEK thousands 12-31-25 2026 77,141 2027 39 2028 300,039 2029 202,539 2030 39\n"
             "Total 579,797\n")
    x.call_llm = lambda *a, **k: {"fields": [
        {"key": "total_debt", "value": 579797, "unit": "SEK thousands", "period": "2025", "raw_label": "Total",
         "source": {"page": 1, "quote": "Total 579,797 420,953"}},
        *[{"key": k, "value": None, "unit": None, "period": None, "raw_label": None, "source": None}
          for k in ("due_within_1_year", "due_1_to_5_years", "due_after_5_years")]]}
    out = x.extract([bts92], [1], dm, {"fiscal_year": 2025})
    assert {f["key"]: f["value"] for f in out["fields"]}["total_debt"] == 579797, out["warnings"]
    by = out["buckets_by_year"]
    assert by["basis"] == "carrying", by
    assert [(y["label"], y["value"]) for y in by["years"]] == [
        ("2026", 77141), ("2027", 39), ("2028", 300039), ("2029", 202539), ("2030", 39)], by
    assert all(y["source"] == {"page": 1, "quote": "SEK thousands 12-31-25 2026 77,141 2027 39 2028 300,039 2029 202,539 2030 39"}
               for y in by["years"]), by
    # v157: the same ladder as ANSWERS, not just metadata -- 2026 -> within 1 year, 2027-2030 summed
    # into 1-5 years. No tail row is printed, so due_after_5_years stays null; never a fabricated 0
    vals = {f["key"]: f["value"] for f in out["fields"]}
    assert (vals["due_within_1_year"], vals["due_1_to_5_years"], vals["due_after_5_years"]) == (77141, 502656, None), vals
    # the tail-year path the BTS page cannot reach (Ericsson p.97's real shape carries it, and its own
    # printed total -- but a LEASE table: below it declines on total_debt, the gate doing its work)
    ericsson97 = ("Total interest-bearing liabilities 32,703 38,041\n"
                  "Future payment obligations for leases\n"
                  "Operating leases 2026 530 2027 386 2028 308 2029 252 2030 193\n"
                  "2031 and later 169\n"
                  "Total 1,838\n")
    x.call_llm = lambda *a, **k: {"fields": [
        {"key": "total_debt", "value": 32703, "unit": "SEK m", "period": "2025", "raw_label": "Total interest-bearing liabilities",
         "source": {"page": 1, "quote": "Total interest-bearing liabilities 32,703 38,041"}},
        *[{"key": k, "value": None, "unit": None, "period": None, "raw_label": None, "source": None}
          for k in ("due_within_1_year", "due_1_to_5_years", "due_after_5_years")]]}
    out = x.extract([ericsson97], [1], dm, {"fiscal_year": 2025})
    vals = {f["key"]: f["value"] for f in out["fields"]}
    assert vals["total_debt"] == 32703, out["warnings"]
    assert "buckets_by_year" not in out, out.get("buckets_by_year")  # 530+386+308+252+193+169 = 1,838 != 32,703
    # v157: and the year ladder must not ANSWER the buckets off it either -- a lease subset is strictly
    # smaller than total_debt, which is the fill's first gate (1,838 < 32,703, and 94% away from it)
    assert all(vals[k] is None for k in ("due_within_1_year", "due_1_to_5_years", "due_after_5_years")), vals
    # the same machinery with years that DO close on total_debt: the open-end tail row rides along as
    # the last bucket, its label as printed, its quote its own row
    yr_tail = "Note 20 Borrowings\nMSEK\n2026 120 2027 80 2028 – 2029 40\nSenare 60\nTotal 300\n"
    x.call_llm = lambda *a, **k: {"fields": [
        {"key": "total_debt", "value": 300, "unit": "MSEK", "period": "2025", "raw_label": "Total",
         "source": {"page": 1, "quote": "Total 300"}},
        *[{"key": k, "value": None, "unit": None, "period": None, "raw_label": None, "source": None}
          for k in ("due_within_1_year", "due_1_to_5_years", "due_after_5_years")]]}
    out = x.extract([yr_tail], [1], dm, {"fiscal_year": 2025})
    by = out["buckets_by_year"]
    assert [(y["label"], y["value"]) for y in by["years"]] == [
        ("2026", 120), ("2027", 80), ("2028", 0), ("2029", 40), ("Senare", 60)], by  # the dash is the report's explicit 0 (v078)
    assert by["years"][-1]["source"]["quote"] == "Senare 60", by
    vals = {f["key"]: f["value"] for f in out["fields"]}  # v157: the tail row is the after-5-years bucket
    assert (vals["due_within_1_year"], vals["due_1_to_5_years"], vals["due_after_5_years"]) == (120, 120, 60), vals
    # ABB p.89 (real text, trimmed): a US-GAAP borrowings note whose maturity table is a bare year
    # ladder -- no bucket synonym anywhere on it, so every bucket came back null however well the page
    # parsed. The ladder restates the total on its own basis (principal 8,247) while the model cited
    # the instrument table's carrying 7,905 on the facing page: filling the buckets alone would turn
    # three honest nulls into three right numbers plus a FAILING identity, so the bare-Total citation
    # moves to the row the buckets actually sum to.
    abb89 = ("Note 13 Debt\n"
             "Long-term debt\n"
             "Bonds and notes 6,180 5,006\n"
             "Total $ 7,905 $ 6,644\n"
             "Maturities of long-term debt outstanding\n"
             "($ in millions) 2026 442 2027 1,143 2028 591 2029 837 2030 1,221\n"
             "Thereafter 4,013\n"
             "Total 8,247\n")
    x.call_llm = lambda *a, **k: {"fields": [
        {"key": "total_debt", "value": 7905, "unit": "USD millions", "period": "2025", "raw_label": "Total",
         "source": {"page": 1, "quote": "Total $ 7,905 $ 6,644"}},
        *[{"key": k, "value": None, "unit": None, "period": None, "raw_label": None, "source": None}
          for k in ("due_within_1_year", "due_1_to_5_years", "due_after_5_years")]]}
    out = x.extract([abb89], [1], dm, {"fiscal_year": 2025})
    vals = {f["key"]: f["value"] for f in out["fields"]}
    assert (vals["due_within_1_year"], vals["due_1_to_5_years"], vals["due_after_5_years"]) == (442, 3792, 4013), vals
    assert vals["total_debt"] == 8247, (vals, out["warnings"])
    assert out["checks"][0]["passed"], out["checks"]
    assert any("re-cited" in w for w in out["warnings"]), out["warnings"]
    # the same page with a total the model could NAME: a recognised debt row outranks the ladder's own
    # Total and is left alone, buckets or not (the identity then fails loudly, which is the honest state)
    x.call_llm = lambda *a, **k: {"fields": [
        {"key": "total_debt", "value": 7905, "unit": "USD millions", "period": "2025",
         "raw_label": "Total borrowings", "source": {"page": 1, "quote": "Total borrowings 7,905"}},
        *[{"key": k, "value": None, "unit": None, "period": None, "raw_label": None, "source": None}
          for k in ("due_within_1_year", "due_1_to_5_years", "due_after_5_years")]]}
    out = x.extract([abb89 + "Total borrowings 7,905\n"], [1], dm, {"fiscal_year": 2025})
    assert {f["key"]: f["value"] for f in out["fields"]}["total_debt"] == 7905, out["warnings"]
    # the real ABB re-run's shape: the model reads the ends of the ladder and leaves the middle null,
    # because four printed years are one bucket. The middle is completed; answers that AGREE row for row
    # are left as the model wrote them
    half = lambda within, after: (lambda *a, **k: {"fields": [
        {"key": "total_debt", "value": 7905, "unit": "USD millions", "period": "2025", "raw_label": "Total",
         "source": {"page": 1, "quote": "Total $ 7,905 $ 6,644"}},
        {"key": "due_within_1_year", "value": within, "unit": "USD millions", "period": "2025", "raw_label": "2026",
         "source": {"page": 1, "quote": "($ in millions) 2026 442 2027 1,143 2028 591 2029 837 2030 1,221"}},
        {"key": "due_1_to_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None},
        {"key": "due_after_5_years", "value": after, "unit": "USD millions", "period": "2025", "raw_label": "Thereafter",
         "source": {"page": 1, "quote": "Thereafter 4,013"}}]})
    x.call_llm = half(442, 4013)
    out = x.extract([abb89], [1], dm, {"fiscal_year": 2025})
    vals = {f["key"]: f["value"] for f in out["fields"]}
    assert (vals["due_1_to_5_years"], vals["total_debt"]) == (3792, 8247), (vals, out["warnings"])
    assert out["checks"][0]["passed"], out["checks"]
    # but a model answer that does NOT match the ladder's own rows means two different tables are in
    # play: nothing is filled and nothing is re-cited. Tested on the fill itself -- end to end, the
    # column repairs upstream of it quietly correct a wrong bucket that sits on the cited page
    flds = [{"key": "total_debt", "value": 7905, "source": {"page": 1, "quote": "Total $ 7,905 $ 6,644"}},
            {"key": "due_within_1_year", "value": 442, "source": {"page": 1, "quote": "2026 442"}},
            {"key": "due_1_to_5_years", "value": None, "source": None},
            {"key": "due_after_5_years", "value": 3500, "source": {"page": 1, "quote": "Thereafter 3,500"}}]
    vals, warns = {f["key"]: f["value"] for f in flds}, []
    x._year_ladder_fill(flds, dm["fields"], dm, [abb89], [1], 2025, {}, warns, vals, set())
    assert [f["value"] for f in flds] == [7905, 442, None, 3500] and not warns, (flds, warns)
    flds[3]["value"] = 4013  # the same call with the ladder agreed to, as the check above proves end to end
    vals["due_after_5_years"] = 4013
    x._year_ladder_fill(flds, dm["fields"], dm, [abb89], [1], 2025, {}, warns, vals, set())
    assert [f["value"] for f in flds] == [8247, 442, 3792, 4013], flds
    # a two-figure Total row below the years is a stacked table's own row (BTS p.92's liabilities
    # table prints exactly that shape): no one-column total, no proof, no key
    stacked = "Maturity analyses\nSEK thousands 12-31-25 2026 100 2027 50\nTotal 150 140\n"
    x.call_llm = lambda *a, **k: {"fields": [
        {"key": "total_debt", "value": 150, "unit": "SEK thousands", "period": "2025", "raw_label": "Total",
         "source": {"page": 1, "quote": "Total 150 140"}},
        *[{"key": k, "value": None, "unit": None, "period": None, "raw_label": None, "source": None}
          for k in ("due_within_1_year", "due_1_to_5_years", "due_after_5_years")]]}
    out = x.extract([stacked], [1], dm, {"fiscal_year": 2025})
    assert "buckets_by_year" not in out, out.get("buckets_by_year")
    # BTS p.93's NOTE 21, the untorn one-row-per-year form (real text): non-current only, first year
    # 2027 = fiscal_year+2 -- the inline reader never fires and nothing else reads the shape; the
    # report's own narrower table (total 548,320) is not total_debt's answer
    bts93 = ("NOTE 21 | Non-current liabilities\n"
             "Group\n"
             "SEK thousands 12-31-25\n"
             "2027 18,415\n"
             "2028 318,414\n"
             "2029 211,452\n"
             "2030 39\n"
             "Total 548,320\n")
    x.call_llm = lambda *a, **k: {"fields": [
        {"key": "total_debt", "value": 548320, "unit": "SEK thousands", "period": "2025", "raw_label": "Total",
         "source": {"page": 1, "quote": "Total 548,320"}},
        *[{"key": k, "value": None, "unit": None, "period": None, "raw_label": None, "source": None}
          for k in ("due_within_1_year", "due_1_to_5_years", "due_after_5_years")]]}
    out = x.extract([bts93], [1], dm, {"fiscal_year": 2025})
    assert {f["key"]: f["value"] for f in out["fields"]}["total_debt"] == 548320, out["warnings"]
    assert "buckets_by_year" not in out, out.get("buckets_by_year")
    # v126: the note that classifies each instrument's balance by REPAYMENT TIMING instead of printing
    # bucket rows -- Stillfront Note 21, p.109 (the fixture above is this note's top half; here it is in
    # full). The model summed the note's own rows correctly (710 = the current-classified rows 675 + 35;
    # 5,152 = the five "Repayment within 2–5 yr." component rows 620 + 2,835 + 649 + 984 + 64) but no row
    # prints either figure, so the computed-not-read guard dropped both (v118 finding 3: labels 710/5152,
    # stored null). The rescue is window geometry, not vocabulary: every row of the page whose label
    # parses inside ONE bucket's window (_row_bucket_span -> _span_bucket), at least two of them, summed
    # in the one column where the note's own total row ties to the verified total_debt, equal to the
    # model's value -> kept as value_derived. "Repayment within 2–5 yr." parses ONLY behind the
    # "repayment" prefix (bare yr/y stays a non-unit in _bucket_span, v096's exclusion above).
    # The glued two-column page: _row_bucket_span reads the wording at the row's figure edge.
    assert x._row_bucket_span("0–6 months 523 490") == ((0, 6), "0–6 months")  # clean interval rows: unchanged _bucket_span territory
    assert x._row_bucket_span("Bond loans 2,835 2,829 Repayment within 2–5 yr. 620 1,170") == ((24, 60), "Repayment within 2–5 yr.")
    assert x._row_bucket_span("Repayment within 2–5 yr. 2,835 2,829") == ((24, 60), "Repayment within 2–5 yr.")
    assert x._row_bucket_span("Repayment after more than 5 yr. 4 –") == ((60, float("inf")), "Repayment after more than 5 yr.")
    assert x._row_bucket_span("Leasing liabilities 103 105 Current liability 675 862") == ((0, 12), "Current liability")  # the current classification IS the 1-year window
    assert x._row_bucket_span("Contingent considerations interest 46 56 Current liability (overdraft facilities) – –") is None  # another row's label follows: not the figure edge
    assert x._row_bucket_span("Total 1,294 2,032 Repayment within 2–5 yr. 620 1,170") is None  # a total never joins the sum of its own parts
    assert x._row_bucket_span("Total within 1 year 710") is None and x._row_bucket_span("Term loan 649 688 Non-current liability 620 1,170") is None
    assert x._close(5152, 5150) and x._close(710, 712.5) and not x._close(5152, 5120)  # ±2 or ±0.5%
    stillfront109_full = stillfront109 + (
        "Movement in the year\n"
        "Opening balance 7,053 7,003 Repayment after more than 5 yr. – –\n"
        "Cashflows\n"
        "Proceeds from borrowings – 1,833 Current liability – –\n"
        "Repayment of loans –1 –1,500 Total bond loans 2,835 2,829\n"
        "Net change in revolving credit facility –272 –404 Term Loans\n"
        "Net change in overdraft facility – –27 Repayment within 2–5 yr. 649 688\n"
        "Contingent considerations paid out in cash –576 –432 Non-current liability 649 688\n"
        "Payment of lease liabilities –41 –39\n"
        "Non cash changes\n"
        "New/changed IFRS16 lease liabilities 40 69 Non-current liability 984 1,376\n"
        "Contingent considerations interest 46 56 Current liability (overdraft facilities) – –\n"
        "Contingent considerations settled –221 –163 Total liabilities to credit institutions 984 1,376\n"
        "Contingent considerations revaluation 266 368\n"
        "Equity swap – 3\n"
        "Translation differences –408 286 Repayment within 2–5 yr. 64 65\n"
        "Closing balance 5,887 7,053 Repayment after more than 5 yr. 4 –\n"
        "Maturity structure Group\n"
        "MSEK 31 Dec 2025 31 Dec 2024\n"
        "Bond loans\n"
        "Repayment within 2–5 yr. 2,835 2,829\n"
        "Non-current liability 2,835 2,829\n"
        "Liabilities to credit institutions\n"
        "Repayment within 2–5 yr. 984 1,376\n"
        "Repayment after more than 5 yr. – –\n"
        "Other non-current liabilities and non-current lease liabilities\n"
        "Non-current liability 67 65\n"
        "Current liability 35 40\n"
        "Total other non-current liabilities and non-current lease liabilities 103 105\n")
    sf_v126 = [
        {"key": "total_debt", "value": 5887, "unit": "MSEK", "period": "2025", "raw_label": "Total",
         "source": {"page": 1, "quote": "Total 5,887 7,053"}},
        {"key": "due_within_1_year", "value": 710, "unit": "MSEK", "period": "2025", "raw_label": "Current liability",
         "source": {"page": 1, "quote": "Current liability 675 862"}},  # the model's own read, first row of its sum
        {"key": "due_1_to_5_years", "value": 5152, "unit": "MSEK", "period": "2025", "raw_label": "Repayment within 2–5 yr.",
         "source": {"page": 1, "quote": "Repayment within 2–5 yr. 620 1,170"}},
        {"key": "due_after_5_years", "value": 4, "unit": "MSEK", "period": "2025", "raw_label": "Repayment after more than 5 yr.",
         "source": {"page": 1, "quote": "Closing balance 5,887 7,053 Repayment after more than 5 yr. 4 –"}}]
    x.call_llm = lambda *a, **k: {"fields": [dict(f) for f in sf_v126]}
    out = x.extract([stillfront109_full], [1], dm, {"fiscal_year": 2025})
    got = {f["key"]: f for f in out["fields"]}
    assert got["total_debt"]["value"] == 5887 and got["due_after_5_years"]["value"] == 4, (got, out["warnings"])
    w1 = got["due_within_1_year"]
    assert w1["value"] == 710 and "value_derived" in w1["evidence"] and "quote_on_page" in w1["evidence"], (w1, out["warnings"])
    assert w1["source"]["quote"] == "Leasing liabilities 103 105 Current liability 675 862", w1["source"]  # scattered rows: the run's first row is the quote
    assert w1["raw_label"] == "Current liability + Current liability + Current liability", w1  # 675 + the nil bond row + 35
    w15 = got["due_1_to_5_years"]
    assert w15["value"] == 5152 and "value_derived" in w15["evidence"], (w15, out["warnings"])
    assert w15["source"]["quote"] == "Bond loans 2,835 2,829 Repayment within 2–5 yr. 620 1,170", w15["source"]
    assert w15["raw_label"].count("Repayment within 2–5 yr.") == 5, w15["raw_label"]  # 620 + 2,835 + 649 + 984 + 64
    assert any("is the sum of 5 rows on page 1 inside its window ('Repayment within 2–5 yr.' … 'Repayment within 2–5 yr.'); kept as value_derived" in w
               for w in out["warnings"]), out["warnings"]
    assert any("is the sum of 3 rows on page 1 inside its window ('Current liability' … 'Current liability')" in w
               for w in out["warnings"]), out["warnings"]
    # the identity stays honestly unresolved: the note classes 710 + 5,152 + 4 = 5,866 of the 5,887 total
    # ("Other interest-bearing liabilities 22" has no timing split on the page) -- the labels' own values
    assert not out["checks"][0]["passed"] and "abs((710 + 5152 + 4) - 5887) <= 2" in out["checks"][0]["detail"], out["checks"]
    # the same answers with the model's multi-row-merged quotes: the suffix of a merged quote verifies
    # (v051's MedCap mechanism), so both forms reach the same guard -- and the same rescue
    merged = [dict(f, source={"page": 1, "quote": q}) for f, q in zip(sf_v126, [
        "Total 5,887 7,053",
        "Current liability 675 862 Current liability 35 40",
        "Repayment within 2–5 yr. 620 1,170 Repayment within 2–5 yr. 2,835 2,829 Repayment within 2–5 yr. 649 688 "
        "Repayment within 2–5 yr. 984 1,376 Repayment within 2–5 yr. 64 65",
        "Closing balance 5,887 7,053 Repayment after more than 5 yr. 4 –"])]
    x.call_llm = lambda *a, **k: {"fields": merged}
    out = x.extract([stillfront109_full], [1], dm, {"fiscal_year": 2025})
    got = {f["key"]: f["value"] for f in out["fields"]}
    assert got["due_within_1_year"] == 710 and got["due_1_to_5_years"] == 5152, (got, out["warnings"])
    # counter-examples: a model value that equals a sum the page does not prove stays null --
    # the Total row never completes the family, the prior-year column never substitutes, rows on
    # another page never join, and another window's rows never mix in (docs/acrylic/evidence/v126.md)
    counter_p1 = ("Note 20 Borrowings\nMSEK 2025 2024\n"
                  "Bond A Repayment within 2–5 yr. 100 90\n"
                  "Bond B Repayment within 2–5 yr. 50 40\n"
                  "Total Repayment within 2–5 yr. 200 190\n"  # a total row carrying the window wording: refused as a summand
                  "Total 150 999\n")
    counter_p2 = "Bond C Repayment within 2–5 yr. 60 55\n"  # another page's same-window row: never a summand either
    for name, value, quote in [
        ("total row completes the sum", 350, "Bond A Repayment within 2–5 yr. 100 90"),  # 100 + 50 + 200(total)
        ("prior-year column substitutes", 130, "Bond A Repayment within 2–5 yr. 100 90"),  # 90 + 40, the non-total column
        ("another page's row completes the sum", 210, "Bond A Repayment within 2–5 yr. 100 90"),  # 100 + 50 + 60(p.2)
        ("another window's rows mix in", 464, "Bond A Repayment within 2–5 yr. 100 90"),  # 100 + 50 + 60 + 200 + 4(no >5y row exists)
        ("within-1-year value off the current rows", 710, "Bond A Repayment within 2–5 yr. 100 90"),  # no current rows here at all
    ]:
        x.call_llm = lambda *a, v=value, q=quote, **k: {"fields": [
            {"key": "total_debt", "value": 150, "unit": "MSEK", "period": "2025", "raw_label": "Total",
             "source": {"page": 1, "quote": "Total 150 999"}},
            {"key": "due_within_1_year", "value": None, "unit": None, "period": None, "raw_label": None, "source": None},
            {"key": "due_1_to_5_years", "value": v, "unit": "MSEK", "period": "2025", "raw_label": "Repayment within 2–5 yr.",
             "source": {"page": 1, "quote": q}},
            {"key": "due_after_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None}]}
        out = x.extract([counter_p1, counter_p2], [1, 2], dm, {"fiscal_year": 2025})
        got = {f["key"]: f for f in out["fields"]}
        assert got["due_1_to_5_years"]["value"] is None, (name, got["due_1_to_5_years"], out["warnings"])
        assert any("dropped as computed, not read" in w for w in out["warnings"]), (name, out["warnings"])
        assert not any("kept as value_derived" in w for w in out["warnings"]), (name, out["warnings"])
    # v156: the window family derives on a NULL answer (v126's family, v096's adoption) -- the same
    # family and column gates, and the closure gate in place of the model's own value: family sum +
    # the other buckets' already-read values (null ones at 0) must reach the tied total through
    # maturity_sums_to_total. Two "Repayment within 1 yr." rows closing on the verified 150 prove
    # the whole 150 is current; the >5y row here is a single row, so no family, null stays null.
    win_null = ("Note 20 Borrowings\nMSEK 2025 2024\n"
                "Repayment within 1 yr. 100 90\n"
                "Repayment within 1 yr. 50 40\n"
                "Repayment after more than 5 yr. - -\n"
                "Total 150 999\n")

    def win_answer(total_value, total_quote):
        return [{"key": "total_debt", "value": total_value, "unit": "MSEK", "period": "2025",
                 "raw_label": "Total", "source": {"page": 1, "quote": total_quote}},
                *({"key": k, "value": None, "unit": None, "period": None, "raw_label": None, "source": None}
                  for k in ("due_within_1_year", "due_1_to_5_years", "due_after_5_years"))]

    x.call_llm = lambda *a, **k: {"fields": win_answer(150, "Total 150 999")}
    out = x.extract([win_null], [1], dm, {"fiscal_year": 2025})
    got = {f["key"]: f for f in out["fields"]}
    w1 = got["due_within_1_year"]
    assert w1["value"] == 150 and "value_derived" in w1["evidence"], (w1, out["warnings"])
    assert w1["raw_label"] == "Repayment within 1 yr. + Repayment within 1 yr.", w1
    assert w1["source"]["quote"] == "Repayment within 1 yr. 100 90 Repayment within 1 yr. 50 40", w1
    assert any("derived as the sum of 2 rows inside its window on page 1 ('Repayment within 1 yr.' … "
               "'Repayment within 1 yr.'); closes on 150" in w for w in out["warnings"]), out["warnings"]
    # counter (the printed-subtotal guard): the same counter page with EVERY bucket null. Its
    # window's own subtotal row is printed ("Total Repayment within 2–5 yr. 200 190"), so the
    # family's 150 -- which would close on the total -- is never derived over it: the printed row
    # is the bucket's read, and the family disagreeing with it is the report's inconsistency.
    x.call_llm = lambda *a, **k: {"fields": win_answer(150, "Total 150 999")}
    out = x.extract([counter_p1], [1], dm, {"fiscal_year": 2025})
    got = {f["key"]: f["value"] for f in out["fields"]}
    assert got["due_within_1_year"] is None and got["due_1_to_5_years"] is None \
        and got["due_after_5_years"] is None, (got, out["warnings"])
    assert not any("derived as the sum of" in w for w in out["warnings"]), out["warnings"]
    # counter (closure): the family reads 150 but 45 of the total sits in another window whose row
    # is alone there (no family of one) -- 150 + nothing is 45 short of 195 and nothing is derived
    win_short = ("Note 20 Borrowings\nMSEK 2025 2024\n"
                 "Repayment within 1 yr. 100 90\n"
                 "Repayment within 1 yr. 50 40\n"
                 "Repayment after more than 5 yr. 45 40\n"
                 "Total 195 170\n")
    x.call_llm = lambda *a, **k: {"fields": win_answer(195, "Total 195 170")}
    out = x.extract([win_short], [1], dm, {"fiscal_year": 2025})
    got = {f["key"]: f["value"] for f in out["fields"]}
    assert got["due_within_1_year"] is None and got["due_after_5_years"] is None, (got, out["warnings"])
    assert not any("derived as the sum of" in w for w in out["warnings"]), out["warnings"]
    # counter (prior-year column): the family closes only in the column the total row does NOT tie
    # to -- the verified 130 ties column 0, the family's 90 + 40 lives in column 1, and the 150 in
    # the tied column closes nothing, so nothing is derived
    x.call_llm = lambda *a, **k: {"fields": win_answer(130, "Total 130 150")}
    out = x.extract([win_null], [1], dm, {"fiscal_year": 2025})
    got = {f["key"]: f["value"] for f in out["fields"]}
    assert got["due_within_1_year"] is None, (got, out["warnings"])
    assert not any("derived as the sum of" in w for w in out["warnings"]), out["warnings"]
    # counter (cross-page): the second same-window row sits on another page -- one row per page is
    # the single-row paths' territory, never a family
    x.call_llm = lambda *a, **k: {"fields": win_answer(150, "Total 150 999")}
    out = x.extract(["Note 20 Borrowings\nMSEK 2025 2024\nRepayment within 1 yr. 100 90\n",
                     "Repayment within 1 yr. 50 40\nTotal 150 999\n"], [1, 2], dm, {"fiscal_year": 2025})
    got = {f["key"]: f["value"] for f in out["fields"]}
    assert got["due_within_1_year"] is None, (got, out["warnings"])
    assert not any("derived as the sum of" in w for w in out["warnings"]), out["warnings"]
    # unit pins: the window half of the wording grammar (split out of _row_bucket_span), and the
    # printed-subtotal guard's accept/refuse table
    assert x._wording_window("Repayment within 2–5 yr.") == (24, 60)
    assert x._wording_window("Repayment within 1 yr.") == (0, 12)
    assert x._wording_window("Repayment after more than 5 yr.") == (60, float("inf"))
    assert x._wording_window("Current liability") == (0, 12)
    assert x._window_subtotal_row("Total Repayment within 2–5 yr. 200 190", "due_1_to_5_years")
    assert not x._window_subtotal_row("Total Repayment within 2–5 yr. 200 190", "due_within_1_year")
    assert not x._window_subtotal_row("Total 150 999", "due_1_to_5_years")  # no window wording: no subtotal of any window
    assert not x._window_subtotal_row("Summa inom 1 år 100", "due_within_1_year")  # no v126 wording: nobody's guard
    # v125: the model's own printed sub-total survives a column-order read from another table. Viscaria
    # (real pages, trimmed; renumbered p.100 -> 1, p.101 -> 2, p.110 -> 3, the walk order the stored
    # seed-rerun-v118 run had): the model answers Note 22's own printed current sub-total 661.8 and
    # total 677.8 (the label's values); the bucket-column read on p.100 -- the liquidity note's
    # undiscounted per-instrument table, reached through the lease row's row_synonyms candidacy
    # (v088's Volati family) -- used to overwrite w1y with 2.2 ("661.8 disagrees with the maturity
    # table; 2.2 read from 'Lease liabilities 0.7 2.2 3.0 2.6 -' by its column order"). Three gates,
    # all required: 661.8 printed verbatim in the candidate window, the row printing it a sub-total/
    # total shape (Total/current), and the read from another table (another page, or a same-page table
    # boundary between the rows). The null bucket still fills from the p.100 row: the rule protects an
    # answered value, not a missing one.
    viscaria100 = ("3 FINANCIAL RISK FACTORS\n"
                   "Maturity analysis\n"
                   "LESS BETWEEN 3 MONTHS AND 1 YEAR BETWEEN 1 AND 2 YEARS BETWEEN 2 AND 5 YEARS MORE\n"
                   "2025-12-31, SEK MILLION THAN 3 MONTHS THAN 5 YEARS\n"
                   "Convertible debentures 0.3 7.2 6.4 11.4 -\n"
                   "Lease liabilities 0.7 2.2 3.0 2.6 -\n"
                   "Shareholder loans - 708.0 - - -\n"
                   "Accounts payable 85.1 - - - -\n"
                   "LESS BETWEEN 3 MONTHS AND 1 YEAR BETWEEN 1 AND 2 YEARS BETWEEN 2 AND 5 YEARS MORE\n"
                   "2024-12-31, SEK MILLION THAN 3 MONTHS THAN 5 YEARS\n"
                   "Lease liabilities 0.7 2.0 2.7 2.0 -\n")
    viscaria101 = ("5 REMUNERATION OF AUDITORS\n"
                   "GROUP PARENT COMPANY\n"
                   "SEK MILLION 2025 2024 2025 2024\n"
                   "Audit engagement 1.0 0.9 1.0 0.9\n"
                   "Other services - 0.2 - 0.2\n"
                   "Total remuneration of auditors 1.2 1.2 1.2 1.2\n"
                   "6 REMUNERATION OF EMPLOYEES, ETC.\n"
                   "Other employees 30 19 23 16\n"
                   "Total 40 25 32 20\n")
    viscaria110 = ("22 INTEREST-BEARING LIABILITIES\n"
                   "GROUP PARENT COMPANY\n"
                   "SEK MILLION 2025 2024 2025 2024\n"
                   "Interest-bearing liabilities Non-current liabilities\n"
                   "Lease liability 5.0 4.2 - -\n"
                   "Convertible debentures 11.0 15.5 11.0 15.5\n"
                   "Total non-current liabilities 16.0 19.7 11.0 15.5\n"
                   "Current liabilities\n"
                   "Lease liability 2.6 2.5 - -\n"
                   "Convertible debentures 4.9 - 4.9 -\n"
                   "Shareholder loans 654.3 - 654.3 -\n"
                   "Total current liabilities 661.8 2.5 659.2 -\n"
                   "Total interest-bearing liabilities 677.8 22.2 670.2 15.5\n")
    viscaria_answer = lambda w1y_val, w1y_page, w1y_quote, w1y_label: {"fields": [
        {"key": "total_debt", "value": 677.8, "unit": "SEK MILLION", "period": "2025", "raw_label": "Total interest-bearing liabilities",
         "source": {"page": 3, "quote": "Total interest-bearing liabilities 677.8 22.2 670.2 15.5"}},
        {"key": "due_within_1_year", "value": w1y_val, "unit": "SEK MILLION", "period": "2025", "raw_label": w1y_label,
         "source": {"page": w1y_page, "quote": w1y_quote}},
        {"key": "due_1_to_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None},
        {"key": "due_after_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None}]}
    x.call_llm = lambda *a, **k: viscaria_answer(
        661.8, 3, "Total current liabilities 661.8 2.5 659.2 -", "Total current liabilities")
    out = x.extract([viscaria100, viscaria101, viscaria110], [1, 2, 3], dm, {"fiscal_year": 2025})
    got = {f["key"]: f["value"] for f in out["fields"]}
    assert got == {"total_debt": 677.8, "due_within_1_year": 661.8, "due_1_to_5_years": 5.6, "due_after_5_years": None}, (got, out["warnings"])
    assert any("kept the model's own 661.8 -- printed on page 3 as 'Total current liabilities 661.8 2.5 659.2 -'; "
               "column-order read 2.2 from 'Lease liabilities 0.7 2.2 3.0 2.6 -' not applied" in w
               for w in out["warnings"]), out["warnings"]
    w1y = next(f for f in out["fields"] if f["key"] == "due_within_1_year")
    assert w1y["source"] == {"page": 3, "quote": "Total current liabilities 661.8 2.5 659.2 -"} \
        and "value_in_quote" in w1y["evidence"], w1y  # the model's own evidence, untouched
    assert out["checks"][0]["passed"] is False and "due_after_5_years" in out["checks"][0]["detail"], out["checks"]  # honest: 661.8 + 5.6 is not 677.8 and after 5y is not printed
    # (b) fails: the model's value printed on a row that is no sub-total of its own (shareholder loans is
    # neither a w1y synonym nor a total/current label) -- the column-order read keeps winning, as for any
    # model misread that no printed total vouches for
    x.call_llm = lambda *a, **k: viscaria_answer(
        708.0, 1, "Shareholder loans - 708.0 - - -", "Shareholder loans")
    out = x.extract([viscaria100, viscaria101, viscaria110], [1, 2, 3], dm, {"fiscal_year": 2025})
    got = {f["key"]: f["value"] for f in out["fields"]}
    assert got == {"total_debt": 677.8, "due_within_1_year": 2.2, "due_1_to_5_years": 5.6, "due_after_5_years": None}, (got, out["warnings"])
    assert not any("kept the model's own" in w for w in out["warnings"]), out["warnings"]
    # (c)'s same-page branch: the sub-total sits in a table stacked BELOW the read row on the same page
    # (its own heading line between them) -- still another table, still kept, warning names that page
    viscaria_stacked = viscaria100 + ("NOTE 22 INTEREST-BEARING LIABILITIES\n"
                                      "SEK MILLION 2025 2024\n"
                                      "Total current liabilities 661.8 2.5\n"
                                      "Total interest-bearing liabilities 677.8 22.2\n")
    x.call_llm = lambda *a, **k: viscaria_answer(
        661.8, 1, "Total current liabilities 661.8 2.5", "Total current liabilities")
    out = x.extract([viscaria_stacked, viscaria101, viscaria110], [1, 2, 3], dm, {"fiscal_year": 2025})
    got = {f["key"]: f["value"] for f in out["fields"]}
    assert got == {"total_debt": 677.8, "due_within_1_year": 661.8, "due_1_to_5_years": 5.6, "due_after_5_years": None}, (got, out["warnings"])
    assert any("kept the model's own 661.8 -- printed on page 1 as 'Total current liabilities 661.8 2.5'" in w
               for w in out["warnings"]), out["warnings"]
    # v144: MTG p.142 is a comparative gross-values table.  The 2024 lease row's final two
    # cells are separate values, but the four recognised headers make _row_amounts read them
    # as the Swedish-looking 85,151.  Its three bucket cells total only 225, so that row cannot
    # replace the model's own 3,487 carrying amount.  This frozen fixture is the relevant real
    # text; do not couple the test to the mutable data/kb corpus (v130).
    mtg142 = ("Group notes\n"
              "Note 16 cont.\n"
              "Terms and payback period, gross values\n"
              "2025 Maturity 2028 Carrying\n"
              "(SEK million) Total Maturity 2026 Maturity 2027 or later amount\n"
              "Liabilities to financial institutions 3,900 582 3,318 — 3,487\n"
              "Lease liabilities 292 69 66 156 253\n"
              "Put and call option liability 261 261 — — 250\n"
              "Contingent consideration 1,217 1,062 — 156 1,145\n"
              "Other non-interest bearing liabilities 785 785 — — 785\n"
              "Accounts payable 464 464 — — 464\n"
              "Total 6,919 3,222 3,384 312 6,384\n"
              "2024 Maturity 2027 Carrying\n"
              "(SEK million) Total Maturity 2025 Maturity 2026 or later amount\n"
              "Lease liabilities 155 37 33 85 151\n"
              "Put and call option liability 332 194 138 — 311\n"
              "Contingent consideration 1,847 1,063 718 66 1,674\n"
              "Other non-interest bearing liabilities 602 602 — — 601\n"
              "Accounts payable 386 386 — — 386\n"
              "Total 3,322 2,281 890 151 3,123\n")
    mtg_texts = [""] * 143
    mtg_texts[141] = mtg142  # candidate page 142; p.143 has no debt-maturity row
    mtg_row = "Lease liabilities 155 37 33 85 151"
    mtg_answer = lambda total: {"fields": [
        {"key": "total_debt", "value": total, "unit": "SEK million", "period": "2025",
         "raw_label": "Liabilities to financial institutions",
         "source": {"page": 142, "quote": "Liabilities to financial institutions 3,900 582 3,318 — 3,487"}} if total is not None else
        {"key": "total_debt", "value": None, "unit": None, "period": None, "raw_label": None, "source": None},
        *[{"key": key, "value": value, "unit": "SEK million", "period": "2025", "raw_label": "Lease liabilities",
           "source": {"page": 142, "quote": mtg_row}}
          for key, value in (("due_within_1_year", 155), ("due_1_to_5_years", 37), ("due_after_5_years", 33))]
    ]}
    x.call_llm = lambda *a, **k: mtg_answer(3487)
    out = x.extract(mtg_texts, [142, 143], dm, {"fiscal_year": 2025})
    got = {f["key"]: f["value"] for f in out["fields"]}
    assert got["total_debt"] == 3487, (got, out["warnings"])
    assert any("total_debt: column reading of 'Lease liabilities 155 37 33 85 151' declined -- gross-value maturity block is headed 2024"
               in w for w in out["warnings"]), out["warnings"]
    # A null model answer has no safe deterministic fill: the only aligned row belongs to the
    # prior gross-value block, not the fiscal-year carrying extraction.
    x.call_llm = lambda *a, **k: mtg_answer(None)
    out = x.extract(mtg_texts, [142, 143], dm, {"fiscal_year": 2025})
    assert next(f for f in out["fields"] if f["key"] == "total_debt")["value"] is None, out["warnings"]
    # Stored runs contain the old repair as if it were the model answer: its total points straight
    # back to the same unclosed column-order row.  There is no independent 3,487 left to recover
    # in a replay, so the poisoned stored value must become null rather than retain 85,151.
    mtg_stored = mtg_answer(85151)
    mtg_stored["fields"][0].update(raw_label="Lease liabilities", source={"page": 142, "quote": mtg_row})
    x.call_llm = lambda *a, **k: mtg_stored
    out = x.extract(mtg_texts, [142, 143], dm, {"fiscal_year": 2025})
    assert next(f for f in out["fields"] if f["key"] == "total_debt")["value"] is None, out["warnings"]
    assert any("gross-value maturity block is headed 2024, not 2025"
               in w for w in out["warnings"]), out["warnings"]
    # The space pair is split only where the header actually has five cells; with MTG's four
    # recognised headers it remains ambiguous and the closure gate above decides whether it may
    # overwrite.  A genuine Swedish thousands group is one number, not a pair to be split for a
    # one-column header.
    assert x._row_amounts(mtg_row, 5, nil=None) == [155, 37, 33, 85, 151]
    assert x._row_amounts(mtg_row, 4, nil=None) == [155, 37, 33, 85151]
    assert x._row_amounts("Lease liabilities 1 234 567", 1, nil=None) == [1234567]
    # Conversely, a row that does close still corrects an answered model value.
    closed = ("Maturity analysis\nOther debt 900\nWithin 1 year\n1-5 years\nAfter 5 years\nTotal\n"
              "Lease liabilities 200, 300, 500, 1,000\n")
    x.call_llm = lambda *a, **k: {"fields": [
        {"key": "total_debt", "value": 900, "unit": "SEK million", "period": "2025", "raw_label": "Other debt",
         "source": {"page": 1, "quote": "Other debt 900"}},
        *[{"key": key, "value": value, "unit": "SEK million", "period": "2025", "raw_label": "Lease liabilities",
           "source": {"page": 1, "quote": "Lease liabilities 200, 300, 500, 1,000"}}
          for key, value in (("due_within_1_year", 200), ("due_1_to_5_years", 300), ("due_after_5_years", 500))]
    ]}
    out = x.extract([closed], [1], dm, {"fiscal_year": 2025})
    assert next(f for f in out["fields"] if f["key"] == "total_debt")["value"] == 1000, out["warnings"]
    # and the read inside the very table the model cited is untouched (Cloetta's disagree case above:
    # the cited row IS the read row) -- v052's paired columns and v068's prior-year column along with it
    print("confidence self-check ok")


def test_confidence_never_exceeds_one():
    """v097 found 1.2: a derived value whose stitched quote also contained the number earned value_in_quote AND
    value_derived. score_field keeps the derived marker only and clamps the sum at 1.0."""
    from . import extract as x
    f = {"key": "due_within_1_year", "value": 1039, "unit": "MSEK", "period": "2025", "raw_label": "0-6 months + 7-12 months",
         "source": {"page": 1, "quote": "0-6 months 523 7-12 months 516 1039"}, "evidence": ["quote_on_page", "value_derived"], "confidence": 0.0}
    sf = {"key": "due_within_1_year", "synonyms": ["0-6 months + 7-12 months"], "unit_hint": "MSEK"}
    x.score_field(f, sf, [], {"checks": []}, "MSEK", 2025, {1})
    assert "value_in_quote" not in f["evidence"] and "value_derived" in f["evidence"], f["evidence"]
    assert 0 < f["confidence"] <= 1.0, f["confidence"]
    print("confidence clamp ok")


def test_torn_bucket_headers():
    """v151: torn CTT bucket labels and BICO's orphan year rows are rejoined conservatively."""
    import json
    import pathlib

    dm = json.loads((pathlib.Path(__file__).parents[1] / "schemas" / "debt_maturity.json").read_text("utf-8"))
    dmf = {f["key"]: f for f in dm["fields"]}
    buckets = {key: dmf[key] for key in ("due_within_1_year", "due_1_to_5_years", "due_after_5_years")}
    bico = ("Note 20 Interest-bearing liabilities\nACCOUNTING PRINCIPLES\n"
            "Loans are classified as interest-bearing non-current or current liabilities in the balance sheet.\n"
            "Group Parent Company\nDec 31,\nDec 31, 2025\nDec 31, 2024\nDec 31, 2025 2024\n"
            "Non-current liabilities\nLiabilities to credit institutions1 - 2.9 - -\n"
            "Convertible bonds - 1,332.3 - 1,332.3\nLease liabilities 218.1 332.5 - -\n"
            "Other interest-bearing liabilities 0.2 1.9 - -\nTotal 218.3 1,669.6 - 1,332.3\n"
            "Current liabilities\nLiabilities to credit institutions1 2.7 5.7 - -\n"
            "Convertible bonds 1,001.5 - 1,001.5\nLease liabilities 63.8 87.1 - -\n"
            "Other interest-bearing liabilities 0.5 1.0 - -\nTotal 1,068.5 93.9 1,001.5 -\n")
    bico_rows = x._page_rows(bico)
    bico_sum = bico_rows.index("Total 218.3 1,669.6 - 1,332.3")
    assert x._row_year_column(bico_rows, bico_sum, 2025) == (0, 4)
    assert x._torn_orphan_year_header(bico_rows, bico_sum)
    x.call_llm = lambda *a, **k: {"fields": [
        {"key": f["key"], "value": None, "unit": None, "period": None, "raw_label": None, "source": None}
        for f in dm["fields"]]}
    out = x.extract([bico], [1], dm, {"fiscal_year": 2025})
    got = {f["key"]: f for f in out["fields"]}
    assert got["total_debt"]["value"] == 1286.8 and "value_derived" in got["total_debt"]["evidence"], (got, out["warnings"])
    assert got["due_within_1_year"]["value"] == 1068.5, (got, out["warnings"])
    assert any("sum to 1286.8 (218.3 + 1068.5)" in w for w in out["warnings"]), out["warnings"]
    # Orphan years must be contiguous and exactly double the visible run; a spacer or an uneven
    # fragment count leaves the old two-column reading alone.
    for broken in (bico.replace("Dec 31, 2024\nDec 31, 2025 2024", "Dec 31, 2024\nHeader furniture\nDec 31, 2025 2024"),
                   bico.replace("Dec 31, 2024\nDec 31, 2025 2024", "Dec 31, 2025 2024")):
        rows = x._page_rows(broken)
        subtotal = rows.index("Total 218.3 1,669.6 - 1,332.3")
        assert not x._torn_orphan_year_header(rows, subtotal), rows
        assert x._row_year_column(rows, subtotal, 2025) == (0, 2), rows
    ctt = ("NOTE 26 MATURITY OF FINANCIAL LIABILITIES\n"
           "Financial liabilities as of 31/12/2025 Within 1 Later than\n"
           "Due for payment as follows: year 1–2 years 2–3 years 3–4 years 4–5 years 5 years Total\n"
           "Interest-bearing financial liabilities including interest rates\nInterest-bearing financial liabilities\n"
           "Interest-bearing liabilities to credit institutions 35.5 - - - - - 35.5\n"
           "Total 35.5 - - - - - 35.5\n")
    ctt_rows = x._page_rows(ctt)
    debt_row = ctt_rows.index("Interest-bearing liabilities to credit institutions 35.5 - - - - - 35.5")
    assert x._bucket_header(ctt_rows, debt_row, buckets, 2025, total_sf=dmf["total_debt"]) == \
        ["due_within_1_year", "due_1_to_5_years", "due_1_to_5_years", "due_1_to_5_years", "due_1_to_5_years", "due_after_5_years", "total"]
    # A row boundary, an extra top fragment, or a non-bucket continuation must not be rejoined.
    for broken in (ctt.replace("Later than\nDue", "Later than\nTable\nDue"),
                   ctt.replace("Within 1 Later than", "Within 1 Later than After"),
                   ctt.replace("5 years Total", "final instalments Total")):
        rows = x._page_rows(broken)
        assert x._rejoin_torn_bucket_headers(rows) == rows, rows
    x.call_llm = lambda *a, **k: {"fields": [
        {"key": f["key"], "value": 35.5, "unit": "SEK m", "period": "2025",
         "raw_label": "Interest-bearing liabilities to credit institutions",
         "source": {"page": 1, "quote": "Interest-bearing liabilities to credit institutions 35.5 - - - - - 35.5"}}
        if f["key"] == "total_debt" else
        {"key": f["key"], "value": None, "unit": None, "period": None, "raw_label": None, "source": None}
        for f in dm["fields"]]}
    out = x.extract([ctt], [1], dm, {"fiscal_year": 2025})
    got = {f["key"]: f["value"] for f in out["fields"]}
    assert got == {"total_debt": 35.5, "due_within_1_year": 35.5, "due_1_to_5_years": 0, "due_after_5_years": 0}, (got, out["warnings"])
    print("torn header self-check ok")


def test_financial_liabilities_rollforward_total():
    """v155: AcadeMedia's current roll-forward beats a sibling generic currency total."""
    import copy
    import json
    import pathlib

    dm = json.loads((pathlib.Path(__file__).parents[1] / "schemas" / "debt_maturity.json").read_text("utf-8"))
    academedia87 = ("Belopp i MSEK 2025-06-30 2024-06-30\n"
                    "SEK 7 660 8 450\nNOK 1 648 1 906\nEUR 2 806 2 741\nSUMMA 12 114 13 097\n"
                    "K28: Skulder\nFÖRÄNDRING FINANSIELLA SKULDER 2024/25\n"
                    "1 juli 2024 Kassaflöde Förvärv av dotterbolag Orealiserade valutakursdifferenser Andra förändringar 30 juni 2025\n"
                    "Skulder till kreditinstitut exkl. fastighetslån 1 419 -518 0 -25 -2 874\n"
                    "Fastighetslån 693 -39 14 -38 0 630\n"
                    "Leasingskulder 10 982 -1 897 207 93 1 220 10 605\n"
                    "Övriga räntebärande skulder 0 0 0 0 0 0\n"
                    "Aktiverade lånekostnader -3 -5 0 0 4 -5\n"
                    "SUMMA 13 090 -2 460 220 30 1 222 12 103\n"
                    "FÖRÄNDRING FINANSIELLA SKULDER 2023/24\n"
                    "1 juli 2023 Kassaflöde Förvärv av dotterbolag Orealiserade valutakursdifferenser Andra förändringar 30 juni 2024\n"
                    "Skulder till kreditinstitut exkl. fastighetslån 842 582 0 -14 9 1 419\n"
                    "Fastighetslån 727 -25 0 -10 0 693\nLeasingskulder 9 511 -1 705 1 282 0 1 893 10 982\n"
                    "Övriga räntebärande skulder 27 -20 0 0 -7 0\nAktiverade lånekostnader -4 0 0 0 0 -3\n"
                    "SUMMA 11 104 -1 168 1 282 -23 1 895 13 090\n")
    rows = x._page_rows(academedia87)
    assert x._financial_liabilities_rollforward_total(rows, 2025) == (
        12103, "SUMMA 13 090 -2 460 220 30 1 222 12 103", "SUMMA"), rows
    assert x._financial_liabilities_rollforward_total(rows, 2026) is None, rows
    answer = [{"key": sf["key"], "value": 12114 if sf["key"] == "total_debt" else None,
               "unit": "MSEK" if sf["key"] == "total_debt" else None,
               "period": "2025" if sf["key"] == "total_debt" else None,
               "raw_label": "SUMMA" if sf["key"] == "total_debt" else None,
               "source": {"page": 1, "quote": "SUMMA 12 114 13 097"} if sf["key"] == "total_debt" else None}
              for sf in dm["fields"]]
    x.call_llm = lambda *a, **k: {"fields": copy.deepcopy(answer)}
    out = x.extract([academedia87], [1], dm, {"fiscal_year": 2025})
    total = next(f for f in out["fields"] if f["key"] == "total_debt")
    assert total["value"] == 12103 and total["source"] == {
        "page": 1, "quote": "SUMMA 13 090 -2 460 220 30 1 222 12 103"}, (total, out["warnings"])
    assert any("financial-liabilities roll-forward closing total 12103" in warning for warning in out["warnings"]), out["warnings"]
    print("financial-liabilities roll-forward self-check ok")


def test_cash_flow_citation_refusal():
    """v155: a cash-flow movement closing balance is not carrying debt."""
    import copy
    import json
    import pathlib

    dm = json.loads((pathlib.Path(__file__).parents[1] / "schemas" / "debt_maturity.json").read_text("utf-8"))
    dynavox135 = ("Note 27. Supplementary disclosures to the statement of cash flows\n"
                   "Liabilities related to financing activities\n"
                   "Consolidated statement of cash flows Interest-bearing borrowings 691.5 202.4 893.8\n"
                   "Consolidated statement of changes in equity Total 790.8 169.3 141.4 -7.1 9.0 -2.1 1,101.3\n")
    rows = x._page_rows(dynavox135)
    cited = "Consolidated statement of cash flows Interest-bearing borrowings 691.5 202.4 893.8"
    assert x._table_scope(rows, None, rows.index(cited), "carrying", dm["table_scope_words"],
                          debt_words=x._debt_subject_words(dm)) == "cash_flow", rows
    # A navigation/header mention cannot brand an ordinary balance-sheet row.
    sidebar = x._page_rows("Consolidated statement of cash flows\nInterest-bearing loans 896.2 691.5\n")
    assert x._table_scope(sidebar, None, 1, "carrying", dm["table_scope_words"],
                          debt_words=x._debt_subject_words(dm)) == "unknown", sidebar
    dustin117 = ("Changes in financing activities\nNot affecting cash flow\n"
                  "Aug 31, 2025 Opening balance Cash flow Reclassifications Change in leases Closing balance\n"
                  "Liabilities to credit institutions 3,619 -1,016 - - 2,538\n"
                  "Lease liabilities 569 -202 - 158 517\n"
                  "Total 4,188 -1,217 - 158 3,055\n")
    dustin_rows, dustin_total = x._page_rows(dustin117), "Total 4,188 -1,217 - 158 3,055"
    assert x._table_scope(dustin_rows, None, dustin_rows.index(dustin_total), "carrying", dm["table_scope_words"],
                          debt_words=x._debt_subject_words(dm)) == "cash_flow", dustin_rows
    answer = [{"key": sf["key"], "value": 893.8 if sf["key"] == "total_debt" else None,
               "unit": "MSEK" if sf["key"] == "total_debt" else None,
               "period": "2025" if sf["key"] == "total_debt" else None,
               "raw_label": "Interest-bearing borrowings" if sf["key"] == "total_debt" else None,
               "source": {"page": 1, "quote": cited} if sf["key"] == "total_debt" else None}
              for sf in dm["fields"]]
    x.call_llm = lambda *a, **k: {"fields": copy.deepcopy(answer)}
    out = x.extract([dynavox135], [1], dm, {"fiscal_year": 2025})
    total = next(f for f in out["fields"] if f["key"] == "total_debt")
    assert total["value"] is None and total["source"] is None, (total, out["warnings"])
    assert any("cash-flow statement" in warning and "893.8" in warning for warning in out["warnings"]), out["warnings"]
    dustin_answer = [{"key": sf["key"], "value": 3055 if sf["key"] == "total_debt" else None,
                      "unit": "MSEK" if sf["key"] == "total_debt" else None,
                      "period": "2025" if sf["key"] == "total_debt" else None,
                      "raw_label": "Total" if sf["key"] == "total_debt" else None,
                      "source": {"page": 1, "quote": dustin_total} if sf["key"] == "total_debt" else None}
                     for sf in dm["fields"]]
    x.call_llm = lambda *a, **k: {"fields": copy.deepcopy(dustin_answer)}
    out = x.extract([dustin117], [1], dm, {"fiscal_year": 2025})
    total = next(f for f in out["fields"] if f["key"] == "total_debt")
    assert total["value"] is None and any("cash-flow statement" in warning for warning in out["warnings"]), (total, out["warnings"])
    print("cash-flow citation refusal self-check ok")


def test_gross_value_prior_year_citation_refusal():
    """v155: an MTG-style prior gross-value block cannot stand in for FY2025 debt."""
    import copy
    import json
    import pathlib

    dm = json.loads((pathlib.Path(__file__).parents[1] / "schemas" / "debt_maturity.json").read_text("utf-8"))
    mtg142 = ("Terms and payback period, gross values\n"
              "2025 Maturity 2028 Carrying\n"
              "(SEK million) Total Maturity 2026 Maturity 2027 or later amount\n"
              "Liabilities to financial institutions 3,900 582 3,318 - 3,487\n"
              "Lease liabilities 292 69 66 156 253\n"
              "2024 Maturity 2027 Carrying\n"
              "(SEK million) Total Maturity 2025 Maturity 2026 or later amount\n"
              "Lease liabilities 155 37 33 85 151\n")
    rows, old_lease = x._page_rows(mtg142), "Lease liabilities 155 37 33 85 151"
    assert x._gross_value_prior_year_block(rows, rows.index(old_lease), 2025) == 2024, rows
    current_lease = "Lease liabilities 292 69 66 156 253"
    assert x._gross_value_prior_year_block(rows, rows.index(current_lease), 2025) is None, rows
    answer = [{"key": sf["key"], "value": 155 if sf["key"] == "due_within_1_year" else None,
               "unit": "MSEK" if sf["key"] == "due_within_1_year" else None,
               "period": "2025" if sf["key"] == "due_within_1_year" else None,
               "raw_label": "Lease liabilities" if sf["key"] == "due_within_1_year" else None,
               "source": {"page": 1, "quote": old_lease} if sf["key"] == "due_within_1_year" else None}
              for sf in dm["fields"]]
    x.call_llm = lambda *a, **k: {"fields": copy.deepcopy(answer)}
    out = x.extract([mtg142], [1], dm, {"fiscal_year": 2025})
    current = next(f for f in out["fields"] if f["key"] == "due_within_1_year")
    assert current["value"] is None and current["source"] is None, (current, out["warnings"])
    assert any("gross-value maturity block is headed 2024, not 2025" in warning for warning in out["warnings"]), out["warnings"]
    print("gross-value prior-year citation refusal self-check ok")


def test_note_citation_preference():
    """v157: a literal note row can improve provenance without changing the figure."""
    schema = {
        "name": "debt_maturity",
        "keywords": ["borrowings"],
        "fields": [{"key": "total_debt", "synonyms": ["total borrowings"]},
                   {"key": "due_within_1_year", "synonyms": ["within 1 year"]}],
    }
    # NCAB's stored total is literally printed both in its capital-management prose and Note 27.
    # These frozen excerpts keep this assertion independent of the mutable stored KB corpus.
    ncab_prose = ("Note 3, cont.\n3.2 MANAGEMENT OF CAPITAL\n"
                  + "The Group reviews its capital structure and funding policy each quarter. " * 3
                  + "\nTotal borrowings (Note 27) 1,090,075 998,103")
    ncab_note = ("NOTE 27 BORROWINGS\nCarrying amount\n"
                 "Total borrowings 1,090,075 998,103")
    fields = [{"key": "total_debt", "value": 1090075,
               "source": {"page": 1, "quote": "Total borrowings (Note 27) 1,090,075 998,103"},
               "evidence": ["quote_on_page"]}]
    warnings = []
    x._prefer_note_citations(fields, schema["fields"][:1], schema, [ncab_prose, ncab_note], [1, 2], warnings)
    assert fields[0]["value"] == 1090075 and fields[0]["source"] == {
        "page": 2, "quote": "Total borrowings 1,090,075 998,103"}, (fields, warnings)
    assert warnings == ["total_debt: cited page 2 'Total borrowings 1,090,075 998,103' (note row) instead of page 1 (balance sheet)"], warnings

    # This Boozt-shaped 104 is derived from the liquidity table, but a note-titled candidate prints it. The
    # source moves to that literal occurrence and stops carrying a derived-value marker.
    boozt_liquidity = ("LIQUIDITY RISK\nMaturity structure of borrowing\n"
                       "Lease liabilities 441 26 78 273 63 -")
    boozt_note = "NOTE 8 BORROWINGS\nWithin 1 year 104 97"
    fields = [{"key": "due_within_1_year", "value": 104,
               "source": {"page": 1, "quote": "Lease liabilities 441 26 78 273 63 -"},
               "evidence": ["quote_on_page", "value_derived"]}]
    warnings = []
    x._prefer_note_citations(fields, schema["fields"][1:], schema, [boozt_liquidity, boozt_note], [1, 2], warnings)
    assert fields[0]["value"] == 104 and fields[0]["source"] == {
        "page": 2, "quote": "Within 1 year 104 97"}, (fields, warnings)
    assert fields[0]["evidence"] == ["quote_on_page"], fields
    assert warnings == ["due_within_1_year: cited page 2 'Within 1 year 104 97' (note row) instead of page 1 (derived)"], warnings

    # Two equally specific note rows are intentionally stable; the preference is not reversible.
    fields = [{"key": "total_debt", "value": 100,
               "source": {"page": 1, "quote": "Total borrowings 100 90"},
               "evidence": ["quote_on_page"]}]
    warnings = []
    x._prefer_note_citations(fields, schema["fields"][:1], schema,
                              ["NOTE 10 BORROWINGS\nTotal borrowings 100 90",
                               "NOTE 11 BORROWINGS\nTotal borrowings 100 90"], [1, 2], warnings)
    assert fields[0]["source"]["page"] == 1 and not warnings, (fields, warnings)

    # A familiar row under an unrelated title is not self-proving as a borrowing-note citation.
    # This avoids moving otherwise page-correct records merely because the number is duplicated.
    fields = [{"key": "due_within_1_year", "value": 104,
               "source": {"page": 1, "quote": "Other liabilities 104 97"},
               "evidence": ["quote_on_page"]}]
    warnings = []
    x._prefer_note_citations(fields, schema["fields"][1:], schema,
                              ["STATEMENT\nOther liabilities 104 97", "Within 1 year 104 97"], [1, 2], warnings)
    assert fields[0]["source"]["page"] == 1 and not warnings, (fields, warnings)

    # An unknown-label row already in a borrowing note is likewise stable against a second note.
    fields = [{"key": "due_within_1_year", "value": 104,
               "source": {"page": 1, "quote": "Other liabilities 104 97"},
               "evidence": ["quote_on_page"]}]
    warnings = []
    x._prefer_note_citations(fields, schema["fields"][1:], schema,
                              ["NOTE 7 BORROWINGS\nOther liabilities 104 97",
                               "NOTE 8 BORROWINGS\nWithin 1 year 104 97"], [1, 2], warnings)
    assert fields[0]["source"]["page"] == 1 and not warnings, (fields, warnings)
    print("note citation preference ok")


def test_bucket_row_label_known():
    """A buckets-as-columns row earns label_known only by closing on its own total (identity_all_columns): a bare
    "Total" is no synonym of any field, and the same row that does not add up proves nothing -- never unconditional."""
    import json
    import pathlib

    dm = json.loads((pathlib.Path(__file__).parents[1] / "schemas" / "debt_maturity.json").read_text("utf-8"))
    header = "Note 21 Borrowings\n31 Dec 2025\nSEKm\n< 1 year\n1–5 years\n> 5 years Total\n"
    nulls = {"fields": [{"key": sf["key"], "value": None, "unit": None, "period": None, "raw_label": None, "source": None} for sf in dm["fields"]]}
    x.call_llm = lambda *a, **k: nulls
    out = x.extract([header + "Total 197 1,399 9 1,605\n"], [1], dm, {"fiscal_year": 2025})
    for f in out["fields"]:
        assert f["raw_label"] == "Total" and not x._label_known(f["raw_label"], next(sf for sf in dm["fields"] if sf["key"] == f["key"]))
        assert "identity_all_columns" in f["evidence"] and "label_known" in f["evidence"] and f["confidence"] == 1.0, f
    assert out["checks"][0]["passed"], out["checks"]
    out = x.extract([header + "Total 197 1,399 9 1,700\n"], [1], dm, {"fiscal_year": 2025})  # 1,605 != 1,700: the row proves nothing
    assert all("identity_all_columns" not in f["evidence"] and "label_known" not in f["evidence"] for f in out["fields"] if f["value"] is not None), [f["evidence"] for f in out["fields"]]
    assert not out["checks"][0]["passed"], out["checks"]
    # the model's own answer quoting that closing row earns the same proof; quoting another row does not
    own = lambda quote: {"fields": [{"key": "total_debt", "value": 1605, "unit": "SEKm", "period": "2025", "raw_label": "Total", "source": {"page": 1, "quote": quote}}] + nulls["fields"][1:]}
    x.call_llm = lambda *a, **k: own("Total 197 1,399 9 1,605")
    out = x.extract([header + "Total 197 1,399 9 1,605\n"], [1], dm, {"fiscal_year": 2025})
    assert "label_known" in out["fields"][0]["evidence"] and out["fields"][0]["confidence"] == 1.0, out["fields"][0]
    x.call_llm = lambda *a, **k: own("Interest-bearing liabilities 1,605")
    out = x.extract(["Interest-bearing liabilities 1,605\n" + header + "Total 197 1,399 9 1,605\n"], [1], dm, {"fiscal_year": 2025})
    assert "identity_all_columns" not in out["fields"][0]["evidence"], out["fields"][0]  # its own quote is another row; label_known, if any, comes from the synonym list alone
    # a markerless all-liabilities table (Karnell's payables and earn-outs) closes on its own arithmetic too --
    # no borrowing word in title or rows, so closing identifies nothing: no marker, no label_known
    x.call_llm = lambda *a, **k: nulls
    karnell = "Maturity of financial liabilities\nSEKm\n< 1 year\n1–5 years\n> 5 years Total\nAccounts payable 20.0 - - 20.0\nEarn-outs 9.4 - - 9.4\nTotal 72.2 454.4 - 526.6\n"
    out = x.extract([karnell], [1], dm, {"fiscal_year": 2025})
    assert all("identity_all_columns" not in f["evidence"] and "label_known" not in f["evidence"] for f in out["fields"] if f["value"] is not None), [f["evidence"] for f in out["fields"]]
    print("bucket-row label_known self-check ok")


def test_balance_sheet_tie():
    """v166: balance-sheet rows are independent zero-weight evidence for a debt total.

    Frozen excerpts preserve the real KABE, Net Insight, Cavotec, Ericsson and Dynavox shapes without
    coupling this test to mutable KB text.  KABE demonstrates the two narrowly added Swedish row labels;
    Cavotec proves a carrying total is not tied merely because a balance sheet mentions some debt rows.
    """
    import json
    import pathlib

    dm = json.loads((pathlib.Path(__file__).parents[1] / "schemas" / "debt_maturity.json").read_text("utf-8"))
    kabe = ("Koncernens rapport över finansiell ställning\nMkr Not 2025 2024\nSumma tillgångar 2 397 2 502\n"
            "Skulder\nLångfristiga skulder\nSkulder till kreditinstitut 23 12 17\n"
            "Långfristiga leaseskulder 23 8 14\nKortfristiga skulder\n"
            "Skulder till kreditinstitut 23 22 19\nKortfristiga leaseskulder 23 8 12\n"
            "Summa eget kapital och skulder 2 397 2 502")
    rows, with_leases = x._balance_sheet_tie([kabe], 2025, 50, dm)
    assert with_leases and [r["value"] for r in rows] == [12, 8, 22, 8], rows

    net_insight = ("Consolidated Balance Sheet\nAmounts in SEK thousands NOTE 31 Dec 2025 31 Dec 2024\n"
                   "Total assets 799,594 834,779\nNon-current liabilities\nLease liabilities 10 31,110 1,555\n"
                   "Current liabilities\nLease liabilities 10 8,305 11,738\n"
                   "TOTAL EQUITY AND LIABILITIES 799,594 834,779")
    rows, with_leases = x._balance_sheet_tie([net_insight], 2025, 39415, dm)
    assert with_leases and [r["value"] for r in rows] == [31110, 8305], rows

    cavotec = ("Consolidated balance sheet\nEuro thousands Notes 2025 2024\nTOTAL ASSETS 147,717 144,700\n"
               "Long-term liabilities\nLoans from credit institutions 27 9,811 13,601\n"
               "Lease liabilities 15 10,378 10,160\nCurrent liabilities\nLease liabilities 15 3,324 2,566\n"
               "TOTAL EQUITY AND LIABILITIES 147,717 144,700")
    assert x._balance_sheet_tie([cavotec], 2025, 23702, dm) is None

    ericsson = ("Consolidated balance sheet\nSEK million Dec 31 2025 Dec 31 2024\nTotal assets 279,223 292,374\n"
                "Non-current liabilities\nBorrowings, non-current F4 29,165 31,904\n"
                "Current liabilities\nBorrowings, current F4 3,538 6,137\n"
                "Total equity and liabilities 279,223 292,374")
    rows, with_leases = x._balance_sheet_tie([ericsson], 2025, 32703, dm)
    assert not with_leases and [r["value"] for r in rows] == [29165, 3538], rows
    total_sf = next(f for f in dm["fields"] if f["key"] == "total_debt")
    model_total = {"key": "total_debt", "value": 32703, "unit": "MSEK", "period": "2025",
                   "raw_label": "Total interest-bearing liabilities",
                   "source": {"page": 1, "quote": "Borrowings, non-current F4 29,165 31,904"},
                   "evidence": ["quote_on_page"]}
    x.score_field(model_total, total_sf, [], dm, "MSEK", 2025, {1}, [ericsson])
    assert "bs_tie" in model_total["evidence"] and "ties to BS: Borrowings, non-current F4 + Borrowings, current F4" in model_total["raw_label"], model_total

    # Ratos has exact current/non-current labels, but the note-reference tail "16, 25, 26 912"
    # is ambiguous in the shared row parser (26,912 versus note 26 + 912). It must not manufacture
    # a 33,914 balance-sheet total from that uncertainty.
    ratos = ("Consolidated statement of financial position\nSEKm Note 4 31 Dec 2025 31 Dec 2024\n"
             "Total assets 27,885 34,536\nLiabilities\n"
             "Non-current interest-bearing liabilities 16, 25, 26 7,002 7,613\n"
             "Current interest-bearing liabilities 16, 25, 26 912 1,390\n"
             "Total equity and liabilities 27,885 34,536")
    assert x._balance_sheet_tie([ratos], 2025, None, dm) is None

    dynavox = ("Consolidated statement of financial position\nSEK m Note Dec. 31, 2025 Dec. 31, 2024\n"
               "Total assets 2,474.5 1,997.8\nNon-current liabilities\nBorrowings, non-current 24 828.8 647.5\n"
               "Current liabilities\nBorrowings, current 24 67.5 44.0\n"
               "Total equity and liabilities 2,474.5 1,997.8")
    x.call_llm = lambda *a, **k: {"fields": [
        {"key": f["key"], "value": None, "unit": None, "period": None, "raw_label": None, "source": None}
        for f in dm["fields"]]}
    out = x.extract([dynavox], [1], dm, {"fiscal_year": 2025})
    total = next(f for f in out["fields"] if f["key"] == "total_debt")
    assert total["value"] == 896.3 and "value_derived" in total["evidence"] and "bs_tie" in total["evidence"], total
    assert "ties to BS: Borrowings, non-current + Borrowings, current" in total["raw_label"], total
    print("balance-sheet tie ok")

def test_absent_in_table_bucket():
    """v165: a bucket the maturity table prints no column for is the report's own explicit absence, not the
    model's silence. Ependion's contracted-terms table (real p.155 shape): columns Within 12 months / Between
    1 and 2 years / Between 2 and 3 years / Total -- no after-five-years column exists, so due_after_5_years
    stays null but carries evidence ["absent_in_table"] with the header row as its source, and the identity
    counts it as 0 (167,546 + 415,984 + 0 == 583,531 within the check's own rounding). A bucket the table
    DOES print a column for (here as a dash) keeps v078's honest rules: no dash-to-0 without the row's own
    arithmetic, and no absent_in_table either -- its null is a different, printed kind."""
    import json
    import pathlib

    dm = json.loads((pathlib.Path(__file__).parents[1] / "schemas" / "debt_maturity.json").read_text("utf-8"))
    ependion = ("Contracted terms\n"
                "SEK 000 Within 12 months\n"
                "Between 1 and 2 years\n"
                "Between 2 and 3 years Total\n"
                "Borrowing 167,546 35,636 380,348 583,531\n")
    x.call_llm = lambda *a, **k: {"fields": [
        {"key": f["key"], "value": None, "unit": None, "period": None, "raw_label": None, "source": None}
        for f in dm["fields"]]}
    out = x.extract([ependion], [1], dm, {"fiscal_year": 2025})
    got = {f["key"]: f for f in out["fields"]}
    assert got["due_within_1_year"]["value"] == 167546 and got["due_1_to_5_years"]["value"] == 415984 \
        and got["total_debt"]["value"] == 583531, (got, out["warnings"])  # the column read itself is unchanged
    absent = got["due_after_5_years"]
    assert absent["value"] is None, absent
    assert absent["evidence"] == ["absent_in_table"], (absent, out["warnings"])
    assert absent["source"] == {"page": 1, "quote": "Between 2 and 3 years Total"}, (absent, out["warnings"])
    assert absent["missing_reason"] == {
        "code": "absent_in_table",
        "detail": "The maturity table does not print a column for this window.",
        "page": 1,
        "quote": "Between 2 and 3 years Total",
    }, absent
    assert any(w.startswith("due_after_5_years: the maturity table on page 1 prints no column for this window")
               and "'Between 2 and 3 years Total'" in w for w in out["warnings"]), out["warnings"]
    check = out["checks"][0]
    assert check["passed"] is True, (check, out["warnings"])
    assert "missing" not in check["detail"] and "due_after_5_years" in check["detail"], check

    # A column the table DOES print but whose cell is a dash is not "absent": the header names it, so the
    # bucket has its own printed column and keeps the v078 rules -- no 0 without the row's own arithmetic,
    # no absent_in_table. And with another operand still an unanswered null, the identity's verdict stays
    # "missing: <that field>" instead of silently passing on the one proven absence.
    dashed = ("Contracted terms\n"
              "SEK 000 Within 12 months\n"
              "Between 1 and 2 years Total\n"
              "Borrowing 167,546 - 583,531\n")
    out = x.extract([dashed], [1], dm, {"fiscal_year": 2025})
    got = {f["key"]: f for f in out["fields"]}
    assert got["due_within_1_year"]["value"] == 167546, (got, out["warnings"])
    assert got["due_1_to_5_years"]["value"] is None and got["due_1_to_5_years"]["evidence"] == [], (got, out["warnings"])
    assert got["due_after_5_years"]["value"] is None and got["due_after_5_years"]["evidence"] == ["absent_in_table"], (got, out["warnings"])
    check = out["checks"][0]
    assert check["passed"] is False and check["detail"].startswith("missing: due_1_to_5_years"), (check, out["warnings"])

    # The marking is bucket-column only: a row-per-bucket note has no column header to prove an absence
    # with, so a null bucket there stays exactly what it was -- plain null, empty evidence, missing check.
    rows = ("Maturity of borrowings\n"
            "Within 1 year 100\n"
            "1-5 years 200\n"
            "Total 300\n")
    x.call_llm = lambda *a, **k: {"fields": [
        {"key": f["key"], "value": None, "unit": None, "period": None, "raw_label": None, "source": None}
        for f in dm["fields"]]}
    out = x.extract([rows], [1], dm, {"fiscal_year": 2025})
    got = {f["key"]: f for f in out["fields"]}
    assert got["due_after_5_years"]["value"] is None and got["due_after_5_years"]["evidence"] == [], (got, out["warnings"])
    assert out["checks"][0]["detail"].startswith("missing:"), out["checks"]
    print("absent-in-table bucket ok")


def test_fixed_pages_skip_selection():
    """v177: analyst-supplied pages are the complete model window, even when two-pass is enabled."""
    import json
    import os
    import pathlib

    schema = json.loads((pathlib.Path(__file__).parents[1] / "schemas" / "debt_maturity.json").read_text("utf-8"))
    selected = ["not selected", "Analyst evidence page two", "Analyst evidence page three"]
    old_select, old_call = x._select_pages, x.call_llm
    old_two_pass = os.environ.get("EXTRACT_TWO_PASS")
    calls = []
    try:
        os.environ["EXTRACT_TWO_PASS"] = "1"
        x._select_pages = lambda *args, **kwargs: (_ for _ in ()).throw(AssertionError("fixed pages must not select pages"))
        x.call_llm = lambda _system, user: calls.append(user) or {"fields": []}
        out = x.extract(selected, [2, 3], schema, {"fiscal_year": 2025}, fixed_pages=True)
        assert len(calls) == 1, calls
        assert "=== PAGE 2 ===\nAnalyst evidence page two" in calls[0], calls[0]
        assert "=== PAGE 3 ===\nAnalyst evidence page three" in calls[0], calls[0]
        assert "not selected" not in calls[0] and not any(w.startswith("two_pass:") for w in out["warnings"]), out
    finally:
        x._select_pages, x.call_llm = old_select, old_call
        if old_two_pass is None:
            os.environ.pop("EXTRACT_TWO_PASS", None)
        else:
            os.environ["EXTRACT_TWO_PASS"] = old_two_pass


def test_missing_reasons_for_honest_debt_nulls():
    """v181: known mapping gaps explain a null without changing any debt value."""
    import json
    import pathlib

    dm = json.loads((pathlib.Path(__file__).parents[1] / "schemas" / "debt_maturity.json").read_text("utf-8"))
    empty = lambda: [{"key": f["key"], "value": None, "unit": None, "period": None,
                      "raw_label": None, "source": None} for f in dm["fields"]]

    # A valued off-grid column crosses the five-year boundary. It remains unassigned, but the
    # original span and amount must be visible alongside the still-null standard buckets.
    offgrid = ("NOTE 47 - FINANCIAL LIABILITIES\n"
               "Maturity analysis liabilities, 2025 <1 year 1-3 years >3 years Total\n"
               "Liabilities to credit institutions 43.5 353.7 173.0 570.2\n")
    x.call_llm = lambda *a, **k: {"fields": empty()}
    out = x.extract([offgrid], [1], dm, {"fiscal_year": 2025})
    got = {f["key"]: f for f in out["fields"]}
    assert all(f["value"] is None for f in got.values()), (got, out["warnings"])
    assert out["checks"][0]["detail"].startswith("missing:"), out["checks"]
    reason = got["due_1_to_5_years"].get("missing_reason")
    assert reason and reason["code"] == "offgrid_span", (got, out["warnings"])
    assert reason["disclosed"] == [{"span": ">3 years", "amount": 173.0, "unit": None,
                                      "page": 1, "quote": "Liabilities to credit institutions 43.5 353.7 173.0 570.2"}], reason
    assert got["due_after_5_years"].get("missing_reason", {}).get("code") == "offgrid_span", got

    # Row-shaped intervals use the same honest boundary rule. The standard buckets remain null and
    # carry the report's original 3–7 years row rather than allocating its 500 to either side of 5 years.
    straddle = ("Maturity analysis\nMSEK 2025 2024\n0–6 months 100 90\n"
                "7–12 months 50 40\n3–7 years 500 400\nTotal 650 640\n")
    x.call_llm = lambda *a, **k: {"fields": [
        {"key": "total_debt", "value": 650, "unit": "MSEK", "period": "2025", "raw_label": "Total", "source": {"page": 1, "quote": "Total 650 640"}},
        *[{"key": key, "value": None, "unit": None, "period": None, "raw_label": None, "source": None}
          for key in ("due_within_1_year", "due_1_to_5_years", "due_after_5_years")],
    ]}
    out = x.extract([straddle], [1], dm, {"fiscal_year": 2025})
    got = {f["key"]: f for f in out["fields"]}
    assert got["due_1_to_5_years"]["value"] is None and got["due_after_5_years"]["value"] is None, got
    reason = got["due_1_to_5_years"].get("missing_reason")
    assert reason and reason["code"] == "straddle" and reason["disclosed"][0]["span"] == "3–7 years", (got, out["warnings"])
    import io
    from pptx import Presentation
    from pipeline import ppt
    slide = Presentation(io.BytesIO(ppt.build_pptx(out))).slides[0]
    slide_text = '\n'.join(
        [shape.text for shape in slide.shapes if hasattr(shape, "text")] +
        [cell.text for shape in slide.shapes if shape.has_table for row in shape.table.rows for cell in row.cells]
    )
    assert f"Why empty: {reason['detail']}" in slide_text, slide_text

    # No candidate reading or deterministic repair found a figure: leave every value null and say
    # only that the supplied candidate page was searched, rather than inventing a number or zero.
    x.call_llm = lambda *a, **k: {"fields": empty()}
    out = x.extract(["Management report\nNo borrowing schedule appears here."], [1], dm, {"fiscal_year": 2025})
    got = {f["key"]: f for f in out["fields"]}
    assert all(f["value"] is None for f in got.values()), (got, out["warnings"])
    assert got["due_within_1_year"].get("missing_reason", {}).get("code") == "not_found", got
    assert "page 1" in got["due_within_1_year"]["missing_reason"]["detail"], got

    # Scope gates retain their existing refusal. Parent-company and lease-only tables respectively
    # explain the empty standard buckets; neither table is repurposed as a Group borrowing schedule.
    parent = ("Parent Company\nInterest-bearing liabilities 622 486\n"
              "Within 1 year 2 2\n1-5 years 620 484\nTotal 622 486\n")
    x.call_llm = lambda *a, **k: {"fields": [
        {"key": "total_debt", "value": 622, "unit": "MSEK", "period": "2025", "raw_label": "Interest-bearing liabilities", "source": {"page": 1, "quote": "Interest-bearing liabilities 622 486"}},
        {"key": "due_within_1_year", "value": 2, "unit": "MSEK", "period": "2025", "raw_label": "Within 1 year", "source": {"page": 1, "quote": "Within 1 year 2 2"}},
        {"key": "due_1_to_5_years", "value": 620, "unit": "MSEK", "period": "2025", "raw_label": "1-5 years", "source": {"page": 1, "quote": "1-5 years 620 484"}},
        {"key": "due_after_5_years", "value": None, "unit": None, "period": None, "raw_label": None, "source": None},
    ]}
    out = x.extract([parent], [1], dm, {"fiscal_year": 2025})
    got = {f["key"]: f for f in out["fields"]}
    assert all(f["value"] is None for f in got.values()), (got, out["warnings"])
    assert any(f.get("missing_reason", {}).get("code") == "parent_only" for f in got.values()), (got, out["warnings"])

    lease = ("Maturity analysis of lease liabilities\nGroup\n"
             "Within one year 12 213\nLater than one but within five years 8 109\n"
             "Later than within five years 2 978\nTotal 23 300\n")
    x.call_llm = lambda *a, **k: {"fields": [
        {"key": "total_debt", "value": 23300, "unit": "MSEK", "period": "2025", "raw_label": "Total", "source": {"page": 1, "quote": "Total 23 300"}},
        {"key": "due_within_1_year", "value": 12213, "unit": "MSEK", "period": "2025", "raw_label": "Within one year", "source": {"page": 1, "quote": "Within one year 12 213"}},
        {"key": "due_1_to_5_years", "value": 8109, "unit": "MSEK", "period": "2025", "raw_label": "Later than one but within five years", "source": {"page": 1, "quote": "Later than one but within five years 8 109"}},
        {"key": "due_after_5_years", "value": 2978, "unit": "MSEK", "period": "2025", "raw_label": "Later than within five years", "source": {"page": 1, "quote": "Later than within five years 2 978"}},
    ]}
    out = x.extract([lease], [1], dm, {"fiscal_year": 2025})
    got = {f["key"]: f for f in out["fields"]}
    assert all(f["value"] is None for f in got.values()), (got, out["warnings"])
    assert any(f.get("missing_reason", {}).get("code") == "lease_table_only" for f in got.values()), (got, out["warnings"])
    print("missing debt reasons red/green cases ok")


if __name__ == "__main__":
    test_bucket_row_label_known()
    test_confidence_never_exceeds_one()
    test_torn_bucket_headers()
    test_financial_liabilities_rollforward_total()
    test_cash_flow_citation_refusal()
    test_gross_value_prior_year_citation_refusal()
    test_note_citation_preference()
    test_absent_in_table_bucket()
    test_fixed_pages_skip_selection()
    test_missing_reasons_for_honest_debt_nulls()
    test_balance_sheet_tie()
    demo()
