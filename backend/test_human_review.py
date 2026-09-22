"""Run: python test_human_review.py. Uses an isolated temporary KB, no model calls."""
import copy
import io
import json
import os
import tempfile
from pathlib import Path
from fastapi import HTTPException
from fastapi.testclient import TestClient
from pptx import Presentation
import app
from pipeline import kb

with tempfile.TemporaryDirectory() as tmp:
    os.environ['KB_DIR'] = tmp
    client = TestClient(app.app)
    f = dict(key='revenue', label='Revenue', value=100, unit='MSEK', period='2025', raw_label='Revenue', source={'page': 1, 'quote': 'Revenue 100'}, confidence=.8, evidence=['arith_ok', 'value_in_quote'])
    x = dict(report_id='lib-review_test', company='Test', fiscal_year=2025, currency='MSEK', section='income_statement', fields=[f], checks=[dict(name='net_profit_arith', passed=True, detail='old calculation')], warnings=[])
    app.reports[x['report_id']] = dict(report_id=x['report_id'], stem='review_test', pages=2)
    app.texts_cache[x['report_id']] = ['Revenue 100', 'Revenue 120\nCurrent debt 40\nNon-current debt 80']
    kb.save_extraction('review_test', 'income_statement', x)
    def save(field, decision, **kwargs):
        return app.review_field(x['report_id'], app.ReviewBody(section='income_statement', key='revenue', expected=field, decision=decision, reviewer='Tester', note='Checked page 1', **kwargs))
    confirmed = save(copy.deepcopy(f), 'confirmed')
    assert confirmed['fields'][0]['human_review']['decision'] == 'confirmed'
    assert 'source_verified' not in confirmed['fields'][0]['human_review']
    # One confirmed figure, no basis saved, its check unavailable: nothing left for a human, the basis form is only suggested
    assert confirmed['ready'] and confirmed['issues'] == [] and confirmed['checks'][0]['status'] == 'unavailable'
    assert confirmed['basis_issues'] and confirmed['basis_suggested']['entity'] == 'Test' and 'basis' not in confirmed
    try:
        save(f, 'unresolved')
        raise AssertionError('stale write accepted')
    except HTTPException as e:
        assert e.status_code == 409
    corrected = save(copy.deepcopy(confirmed['fields'][0]), 'corrected', value=120, unit='MSEK', period='2025', source_page=2, source_quote='Revenue 120')
    current = corrected['fields'][0]
    assert current['value'] == 120 and current['evidence'] == ['human_reviewed']
    assert current['source'] == {'page': 2, 'quote': 'Revenue 120'}
    assert current['human_review']['source_verified'] is True
    assert current['review_history'][1]['previous']['value'] == 100
    assert current['review_history'][1]['previous']['source'] == {'page': 1, 'quote': 'Revenue 100'}
    assert corrected['checks'][0]['status'] == 'unavailable'
    assert corrected['check_history'][0]['previous'][0]['detail'] == 'old calculation'
    try:
        save(copy.deepcopy(current), 'corrected', value=121, unit='MSEK', period='2025', source_page=2, source_quote='Revenue 121')
        raise AssertionError('citation not on page accepted')
    except HTTPException as e:
        assert e.status_code == 400 and 'page 2' in str(e.detail)
    component_sum = save(copy.deepcopy(current), 'corrected', unit='MSEK', period='2025', components=[
        {'label': 'Current debt', 'value': 40, 'page': 2, 'quote': 'Current debt 40'},
        {'label': 'Non-current debt', 'value': 80, 'page': 2, 'quote': 'Non-current debt 80'},
    ])
    current = component_sum['fields'][0]
    assert current['value'] == 120 and current['evidence'] == ['human_reviewed', 'components_sum']
    assert current['source'] == {'page': 2, 'quote': 'Current debt 40'}
    assert current['components'] == [
        {'label': 'Current debt', 'value': 40.0, 'page': 2, 'quote': 'Current debt 40'},
        {'label': 'Non-current debt', 'value': 80.0, 'page': 2, 'quote': 'Non-current debt 80'},
    ]
    assert current['review_history'][2]['previous']['source'] == {'page': 2, 'quote': 'Revenue 120'}
    csv_text = app.extraction_csv(x['report_id']).body.decode()
    assert 'human_source_page,human_source_quote,components' in csv_text
    assert 'Current debt 40' in csv_text and 'components_sum' not in csv_text
    slide_text = '\n'.join(shape.text for shape in Presentation(io.BytesIO(app.extraction_pptx(x['report_id']).body)).slides[0].shapes if hasattr(shape, 'text'))
    assert "Human review: p.2 'Current debt 40' — components 2" in slide_text
    persisted = json.loads((Path(tmp)/'review_test/extractions/income_statement.json').read_text())
    assert persisted['fields'][0] == current
    app.extractions.clear()
    app.texts_cache.pop(x['report_id'])
    reopened = save(copy.deepcopy(persisted['fields'][0]), 'unresolved')
    assert reopened['fields'][0]['human_review']['decision'] == 'unresolved'
    assert reopened['fields'][0]['source'] == {'page': 2, 'quote': 'Current debt 40'}
    assert len(reopened['fields'][0]['review_history']) == 4
    assert not reopened['ready'] and [i['kind'] for i in reopened['issues']] == ['field']
    try:
        app.run_extract(x['report_id'], app.ExtractBody(section='income_statement'))
        raise AssertionError('review overwritten by extraction')
    except HTTPException as e:
        assert e.status_code == 409
    try:
        save(copy.deepcopy(reopened['fields'][0]), 'corrected', value='NaN', unit='MSEK', period='2025')
        raise AssertionError('invalid numeric value accepted')
    except HTTPException as e:
        assert e.status_code == 422

    # An analyst-directed retry is a candidate only. It can use no page other than the
    # explicitly supplied window, cannot change the stored extraction, and keeps the normal
    # reviewed-section barrier ahead of even fixture-mode handling.
    def fill_request(report_id, field, pages):
        return client.post(f'/api/reports/{report_id}/fill?section=income_statement', json={'field': field, 'pages': pages})

    reviewed_fill = fill_request(x['report_id'], 'revenue', [2])
    assert reviewed_fill.status_code == 409, reviewed_fill.text

    fill_field = dict(key='revenue', label='Revenue', value=None, unit='MSEK', period='2025', raw_label='Revenue', source=None, confidence=0, evidence=[])
    untouched = dict(key='cost_of_sales', label='Cost of sales', value=-80, unit='MSEK', period='2025', raw_label='Cost of sales', source={'page': 1, 'quote': 'Cost of sales -80'}, confidence=.8, evidence=['quote_on_page'])
    fill = dict(report_id='lib-fill_test', company='Fill Test', fiscal_year=2025, currency='MSEK', section='income_statement', fields=[fill_field, untouched], checks=[], warnings=[])
    app.reports[fill['report_id']] = dict(report_id=fill['report_id'], stem='fill_test', pages=2)
    app.texts_cache[fill['report_id']] = ['Revenue 100\nCost of sales -80', 'Revenue 120']
    kb.save_extraction('fill_test', 'income_statement', fill)
    stored_before = (Path(tmp) / 'fill_test/extractions/income_statement.json').read_text(encoding='utf-8')
    configured, extract = app._llm_configured, app.extract_mod.extract
    try:
        app._llm_configured = lambda: False
        fixture_fill = fill_request(fill['report_id'], 'revenue', [2])
        assert fixture_fill.status_code == 422 and 'Demo mode does not support targeted field fill' in fixture_fill.text

        seen = {}
        def fixed_page_extract(texts, pages, schema, report, fixed_pages=False):
            seen.update(pages=pages, fixed_pages=fixed_pages)
            return {'fields': [
                {**fill_field, 'value': 120, 'source': {'page': 2, 'quote': 'Revenue 120'}, 'confidence': .8, 'evidence': ['quote_on_page']},
                {**untouched, 'value': -999},
            ], 'warnings': ['stubbed response']}
        app._llm_configured = lambda: True
        app.extract_mod.extract = fixed_page_extract
        response = fill_request(fill['report_id'], 'revenue', [2])
        assert response.status_code == 200, response.text
        candidate = response.json()
        assert candidate == {'candidate': {**fill_field, 'value': 120, 'source': {'page': 2, 'quote': 'Revenue 120'}, 'confidence': .8, 'evidence': ['quote_on_page']}, 'warnings': ['stubbed response']}
        assert seen == {'pages': [2], 'fixed_pages': True}
        assert (Path(tmp) / 'fill_test/extractions/income_statement.json').read_text(encoding='utf-8') == stored_before

        def outside_page_extract(texts, pages, schema, report, fixed_pages=False):
            return {'fields': [{**fill_field, 'value': 100, 'source': {'page': 1, 'quote': 'Revenue 100'}, 'confidence': .8, 'evidence': ['quote_on_page']}], 'warnings': []}
        app.extract_mod.extract = outside_page_extract
        rejected_response = fill_request(fill['report_id'], 'revenue', [2])
        assert rejected_response.status_code == 200, rejected_response.text
        rejected = rejected_response.json()
        assert rejected['candidate'] is None and any('outside the supplied pages' in warning for warning in rejected['warnings'])
    finally:
        app._llm_configured, app.extract_mod.extract = configured, extract
print('Human review persistence, conflict, correction, export and re-extraction checks passed')
