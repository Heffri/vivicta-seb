"""Run with python test_workbench.py. Isolated files, no network; the locate block builds the one
synthetic PDF (pymupdf, into a temp directory) and exercises it through a real upload."""
import copy
import csv
import io
import json
import os
import tempfile
from pathlib import Path
from unittest.mock import patch
from fastapi import HTTPException
from fastapi.testclient import TestClient
from pptx import Presentation
import pymupdf
from pptx.enum.shapes import MSO_SHAPE_TYPE
import app
from pipeline import kb, workbench, ppt

def statement(section, year=2025):
    schema = app.load_schema(section)
    values = dict(revenue=100, cost_of_sales=-60, gross_profit=40, operating_profit=30, profit_before_tax=25, income_tax=-5, profit_discontinued=0, net_profit=20, total_debt=100, due_within_1_year=20, due_1_to_5_years=50, due_after_5_years=30)
    fields = [dict(key=f['key'], label=f['label'], value=values.get(f['key'], 10), unit='MSEK', period=str(year), raw_label=f['label'], source={'page': 1, 'quote': f['label']}, confidence=0, evidence=[], human_review={'decision': 'confirmed', 'reviewer': 'Analyst', 'at': 'now', 'note': 'Source checked'}) for f in schema['fields']]
    basis = dict(entity='Atlas Copco AB', consolidation='Group', period=str(year), currency='SEK', scale='Millions', source='Page 1, group statement', restatement='As reported', debt_basis='Carrying amounts', leases='Included', bucket_mapping='Under 1, 1 to 5, over 5')
    return dict(report_id=f'lib-test_{year}', stem=f'test_{year}', company='Atlas Copco AB', fiscal_year=year, section=section, currency='MSEK', fields=fields, checks=[], warnings=[], basis={'values': {k: basis[k] for k in workbench.required(section)}, 'reviewer': 'Analyst', 'note': 'Confirmed against source', 'at': 'now'})

for section in ('income_statement', 'debt_maturity'):
    x = statement(section)
    schema = app.load_schema(section)
    assert workbench.decorate(x, schema)['ready'] and x['issues'] == [] and x['basis_issues'] == [] and x['not_reported'] == []
    x['fields'][0]['value'] = None
    workbench.decorate(x, schema)
    assert not x['ready'] and any(c['status'] == 'unavailable' for c in x['checks'])
    assert [i['kind'] for i in x['issues']] == ['field']  # an unavailable check is a missing operand, already listed; only a failed check is a task
    # An unconfirmed basis is a form to fill, not a blocker: it lives in basis_issues with a prefill, and never counts against ready.
    x = statement(section)
    x.pop('basis')
    x.update(maturity_basis='carrying')
    workbench.decorate(x, schema)
    assert x['ready'] and x['issues'] == [] and [i['key'] for i in x['basis_issues']] == workbench.required(section)
    expected = dict(entity='Atlas Copco AB', consolidation='Group', period='2025', currency='SEK', scale='Millions', source='Annual report', restatement='As reported')
    assert x['basis_suggested'] == expected | (dict(debt_basis='Carrying amounts', leases='', bucket_mapping='') if section == 'debt_maturity' else {}), x['basis_suggested']
    assert x['basis_suggested'].get('debt_basis', '') in workbench.CHOICES['debt_basis'] | {''}
    # Automated evidence resolves a field only with value_in_quote or exactly one stand-in for it, plus every other code and a source.
    x = statement(section)
    del x['fields'][0]['human_review']
    for evidence, ok in ((['quote_on_page', 'value_in_quote', 'label_known', 'period_ok', 'page_is_statement', 'unit_ok'], True),
                         (['quote_on_page', 'value_derived', 'label_known', 'period_ok', 'page_is_statement', 'unit_ok'], True),
                         (['quote_on_page', 'stated_zero', 'label_known', 'period_ok', 'page_is_statement', 'unit_ok'], True),
                         (['quote_on_page', 'printed_nil', 'label_known', 'period_ok', 'page_is_statement', 'unit_ok'], True),
                         (['quote_on_page', 'value_derived', 'stated_zero', 'label_known', 'period_ok', 'page_is_statement', 'unit_ok'], False),
                         (['quote_on_page', 'label_known', 'period_ok', 'page_is_statement', 'unit_ok'], False)):
        x['fields'][0]['evidence'] = evidence
        assert workbench.decorate(x, schema)['ready'] == ok, evidence
    # label_known is re-derived from the field's own label/synonyms/row synonyms (saved fields predate that vocabulary); an unknown label stays a task
    x['fields'][0].update(evidence=['quote_on_page', 'value_in_quote', 'period_ok', 'page_is_statement', 'unit_ok'], raw_label='Total')
    assert not workbench.decorate(x, schema)['ready']
    x['fields'][0]['raw_label'] = schema['fields'][0]['label']
    assert workbench.decorate(x, schema)['ready']
    x['fields'][0].update(evidence=['quote_on_page', 'value_in_quote', 'label_known', 'period_ok', 'page_is_statement', 'unit_ok'], source=None)
    assert not workbench.decorate(x, schema)['ready']
    if section == 'debt_maturity':
        # One printed row answers all four fields (Karnell p.106: "Liabilities to credit institutions
        # 43.5 353.7 - 397.2", one label over four columns). The report labels it once, for the total's
        # vocabulary, so every bucket inherits a label its own field cannot name and gets flagged though
        # the very same quote is already trusted. The buckets summing to the total is what proves the row
        # is really shared, so the grant is gated on that identity -- break it and the tasks come back.
        x = statement(section)
        row = {'page': 1, 'quote': 'Liabilities to credit institutions 20 50 30 100'}
        for f in x['fields']:
            f.pop('human_review')
            f.update(raw_label='Liabilities to credit institutions', source=dict(row),
                     evidence=['quote_on_page', 'value_in_quote', 'period_ok', 'page_is_statement', 'unit_ok'])
        assert workbench.decorate(x, schema)['ready'], x['issues']
        x['fields'][0]['value'] = 999  # the buckets no longer sum to the total: nothing proves the shared row
        assert [i['key'] for i in workbench.decorate(x, schema)['issues'] if i['kind'] == 'field'] == \
            ['due_within_1_year', 'due_1_to_5_years', 'due_after_5_years'], x['issues']
        x = statement(section)  # a bucket citing its OWN row is not covered by the total's label
        for f in x['fields']:
            f.pop('human_review')
            f.update(raw_label='Liabilities to credit institutions', source=dict(row),
                     evidence=['quote_on_page', 'value_in_quote', 'period_ok', 'page_is_statement', 'unit_ok'])
        x['fields'][1].update(raw_label='Förfaller', source={'page': 1, 'quote': 'Förfaller 20'})
        assert [i['key'] for i in workbench.decorate(x, schema)['issues']] == ['due_within_1_year'], x['issues']
    # A null on a schema-optional line (income statement cost of sales, gross profit, discontinued operations) is "not reported",
    # never an issue and never a zero; every other null stays a task. A reviewer marking it unresolved makes it one again.
    x = statement(section)
    optional = [f['key'] for f in schema['fields'] if f.get('optional')]
    assert optional == (['cost_of_sales', 'gross_profit', 'profit_discontinued'] if section == 'income_statement' else [])
    for f in x['fields']:
        if f['key'] in optional:
            f['value'] = None
            del f['human_review']
    workbench.decorate(x, schema)
    assert x['ready'] and x['not_reported'] == optional and not any(i['kind'] == 'field' for i in x['issues']), x['issues']
    if optional:
        assert any(c['status'] == 'unavailable' for c in x['checks']) and not any(i['kind'] == 'check' for i in x['issues'])
        x['fields'][1]['human_review'] = {'decision': 'unresolved', 'reviewer': 'Analyst', 'at': 'now', 'note': 'Which row?'}
        workbench.decorate(x, schema)
        assert not x['ready'] and x['issues'][0]['key'] == 'cost_of_sales' and x['not_reported'] == ['gross_profit', 'profit_discontinued']
    x = statement(section)
    x['fields'][0]['unit'] = 'EUR'
    assert any(c['status'] == 'unavailable' for c in workbench.checks(x, schema))
    x = statement(section)
    x['fields'][0]['period'] = '2024'
    assert any(c['status'] == 'unavailable' for c in workbench.checks(x, schema))
    # A null bucket the maturity table prints no column for (evidence "absent_in_table") joins the
    # reconciliation as 0 instead of holding it unavailable, and is not a queue issue; a reviewer marking
    # it unresolved re-opens it. A null income-statement field stays unavailable -- the participation is
    # require_explicit_values checks only, never a general null-to-zero.
    x = statement(section)
    null_key = 'due_after_5_years' if section == 'debt_maturity' else x['fields'][0]['key']
    next(f for f in x['fields'] if f['key'] == null_key).update(value=None, unit=None, period=None, source=None, evidence=['absent_in_table'])
    if section == 'debt_maturity':
        next(f for f in x['fields'] if f['key'] == 'total_debt')['value'] = 70  # 20 + 50 + the absent bucket's 0
        assert all(c['passed'] and c['status'] != 'unavailable' for c in workbench.checks(x, schema)), workbench.checks(x, schema)
        assert not any(i['kind'] == 'field' and i['key'] == null_key for i in workbench.decorate(x, schema)['issues'])
        next(f for f in x['fields'] if f['key'] == null_key)['human_review'] = {'decision': 'unresolved', 'reviewer': 'Analyst', 'at': 'now', 'note': 'Rechecked'}
        assert any(i['kind'] == 'field' and i['key'] == null_key for i in workbench.decorate(x, schema)['issues'])
    else:
        assert any(c['status'] == 'unavailable' for c in workbench.checks(x, schema))
    current, previous = statement(section), statement(section, 2024)
    current['fields'][0]['value'], previous['fields'][0]['value'] = 20, -10
    row = workbench.compare(current, previous)['rows'][0]
    assert row['delta'] == 30 and row['percent'] == 300 and row['sign_change']
    previous['fields'][0]['value'] = 0
    assert workbench.compare(current, previous)['rows'][0]['percent'] is None
    previous['basis']['values']['scale'] = 'Thousands'
    assert workbench.compare(current, previous)['rows'][0]['delta'] is None
    previous.pop('basis')
    assert workbench.compare(current, previous)['rows'][0]['delta'] is None

# Basis suggestions are a source-backed starting point, never a hidden confirmation.
# The standard maturity-header quote is deliberately present on every cited debt field here so
# the positive case proves the strict bucket-mapping gate rather than a label-name coincidence.
x = statement('debt_maturity')
x['maturity_basis'] = 'carrying'
header = 'Maturity profile: due within 1 year, due 1 to 5 years and due after 5 years (carrying amount)'
for field in x['fields']:
    field['source'] = {'page': 7, 'quote': header}
workbench.decorate(x, app.load_schema('debt_maturity'))
suggestions = {item['key']: item for item in x['basis_suggestions']}
assert set(suggestions) == {'entity', 'period', 'currency', 'scale', 'source', 'debt_basis', 'bucket_mapping'}
assert suggestions['entity'] == {'key': 'entity', 'value': 'Atlas Copco AB', 'source': 'report metadata'}
assert suggestions['period'] == {'key': 'period', 'value': '2025', 'source': 'report metadata'}
assert suggestions['currency']['value'] == 'SEK' and suggestions['scale']['value'] == 'Millions'
assert suggestions['source']['value'] == 'Cited pages p. 7'
assert suggestions['debt_basis']['value'] == 'Carrying amounts'
assert suggestions['bucket_mapping']['value'] == 'Under 1, 1 to 5, over 5'
assert all(item['source'] == 'report metadata' or item['source'] == {'page': 7, 'quote': header} for item in suggestions.values())
assert not {'consolidation', 'leases', 'restatement'} & set(suggestions)

# A mixed printed scale is a material ambiguity, not a reason to pick whichever field came first.
mixed = statement('debt_maturity')
for field in mixed['fields']:
    field['source'] = {'page': 7, 'quote': header}
next(field for field in mixed['fields'] if field['key'] == 'due_within_1_year')['unit'] = 'TSEK'
workbench.decorate(mixed, app.load_schema('debt_maturity'))
assert not {'currency', 'scale'} & {item['key'] for item in mixed['basis_suggestions']}

# No cited page means no field-derived value or native-bucket claim. Metadata remains explicitly
# labelled metadata; it does not turn the missing field evidence into a confirmation.
uncited = statement('debt_maturity')
for field in uncited['fields']:
    field['source'] = None
workbench.decorate(uncited, app.load_schema('debt_maturity'))
assert not {'currency', 'scale', 'source', 'bucket_mapping'} & {item['key'] for item in uncited['basis_suggestions']}

# Marking a missing value unresolved is never equivalent to zero or an arithmetic pass.
missing = statement('debt_maturity')
field = next(field for field in missing['fields'] if field['key'] == 'due_after_5_years')
field.update(value=None, unit=None, period=None, source=None, evidence=[], human_review={'decision': 'unresolved', 'reviewer': 'Analyst', 'at': 'now', 'note': 'Not disclosed'})
workbench.decorate(missing, app.load_schema('debt_maturity'))
assert not missing['ready'] and not missing['checks'][0]['passed']
assert any(issue['kind'] == 'field' and issue['key'] == 'due_after_5_years' for issue in missing['issues'])

# maturity_wall -- deterministic upcoming-maturities list.
# Pure function, no disk. statement('debt_maturity') is fully confirmed by default: total_debt=100,
# due_within_1_year=20 (both MSEK) -> share 0.2 is the baseline every case below tweaks one thing in.
x = statement('debt_maturity')
wall = workbench.maturity_wall([x])
row = wall['rows'][0]
assert row['comparable'] and row['share'] == 0.2 and row['reason'] == ''
assert wall['coverage'] == {'total': 1, 'comparable': 1, 'missing_total': 0, 'missing_w1y': 0, 'basis_unconfirmed': 0}

# Basis not confirmed: the ratio is still arithmetic on printed numbers, so share stays visible for the
# analyst -- only "comparable" and the coverage count flip. (This is today's real state for all 105
# saved debt_maturity extractions: show them as incomparable rather than confirm a basis to fill a chart.)
x = statement('debt_maturity')
x.pop('basis')
wall = workbench.maturity_wall([x])
row = wall['rows'][0]
assert row['share'] == 0.2 and not row['comparable'] and 'Basis not confirmed' in row['reason']
assert wall['coverage']['basis_unconfirmed'] == 1 and wall['coverage']['comparable'] == 0

# Missing total_debt: N/A share, counted, not comparable.
x = statement('debt_maturity')
next(f for f in x['fields'] if f['key'] == 'total_debt')['value'] = None
wall = workbench.maturity_wall([x])
row = wall['rows'][0]
assert row['share'] is None and not row['comparable'] and 'Missing total debt' in row['reason']
assert wall['coverage']['missing_total'] == 1

# Missing due_within_1_year: same shape, the other required field.
x = statement('debt_maturity')
next(f for f in x['fields'] if f['key'] == 'due_within_1_year')['value'] = None
wall = workbench.maturity_wall([x])
row = wall['rows'][0]
assert row['share'] is None and not row['comparable']
assert wall['coverage']['missing_w1y'] == 1

# Zero debt: a real, clean state -- share N/A (zero denominator) but not an error and not "missing".
x = statement('debt_maturity')
for key in ('total_debt', 'due_within_1_year', 'due_1_to_5_years', 'due_after_5_years'):
    next(f for f in x['fields'] if f['key'] == key)['value'] = 0
wall = workbench.maturity_wall([x])
row = wall['rows'][0]
assert row['share'] is None and row['comparable'] and row['reason'] == 'No debt outstanding.'
assert wall['coverage']['missing_total'] == 0 and wall['coverage']['comparable'] == 1

# Unverified evidence (human review removed, no evidence codes): share still shown, comparable flips.
x = statement('debt_maturity')
next(f for f in x['fields'] if f['key'] == 'total_debt').pop('human_review')
row = workbench.maturity_wall([x])['rows'][0]
assert row['share'] == 0.2 and not row['comparable'] and 'not yet verified' in row['reason']

# MSEK vs TSEK: same currency, different printed scale -- share correct, display harmonised to MSEK.
x = statement('debt_maturity')
next(f for f in x['fields'] if f['key'] == 'due_within_1_year').update(value=50000, unit='TSEK')
row = workbench.maturity_wall([x])['rows'][0]
assert row['share'] == 0.5
assert row['total'] == {'value': 100.0, 'unit': 'MSEK'} and row['due_within_1_year'] == {'value': 50.0, 'unit': 'MSEK'}

# EUR vs SEK: no FX conversion -- not comparable, share withheld, original units never relabelled.
x = statement('debt_maturity')
next(f for f in x['fields'] if f['key'] == 'due_within_1_year')['unit'] = 'MEUR'
row = workbench.maturity_wall([x])['rows'][0]
assert row['share'] is None and not row['comparable'] and 'currency' in row['reason']
assert row['due_within_1_year']['unit'] == 'MEUR'

# Unrecognised unit text: treated the same as a currency mismatch, never guessed.
x = statement('debt_maturity')
next(f for f in x['fields'] if f['key'] == 'total_debt')['unit'] = 'doubloons'
row = workbench.maturity_wall([x])['rows'][0]
assert row['share'] is None and not row['comparable']

# Parent company: a confirmed non-Group entity level is not itself a reason to block comparability.
x = statement('debt_maturity')
x['basis']['values']['consolidation'] = 'Parent'
row = workbench.maturity_wall([x])['rows'][0]
assert row['comparable'] and row['consolidation'] == 'Parent'

# Leases excluded: same -- a confirmed choice is exposed, not penalised.
x = statement('debt_maturity')
x['basis']['values']['leases'] = 'Excluded'
row = workbench.maturity_wall([x])['rows'][0]
assert row['comparable'] and row['leases'] == 'Excluded'

# Incomplete bucket breakdown (the long buckets absent) does not block total/<1y -- "partially
# computable": the full 3-bucket identity check may go unavailable elsewhere, but this metric needs
# only two fields.
x = statement('debt_maturity')
for key in ('due_1_to_5_years', 'due_after_5_years'):
    next(f for f in x['fields'] if f['key'] == key)['value'] = None
row = workbench.maturity_wall([x])['rows'][0]
assert row['comparable'] and row['share'] == 0.2

# Sort: comparable rows by share descending, then non-comparable rows last regardless of their own share.
a, b, c = statement('debt_maturity'), statement('debt_maturity'), statement('debt_maturity')
a['company'], b['company'], c['company'] = 'A', 'B', 'C'
next(f for f in a['fields'] if f['key'] == 'due_within_1_year')['value'] = 50  # share 0.5
next(f for f in c['fields'] if f['key'] == 'due_within_1_year')['value'] = 90  # share 0.9, but...
c.pop('basis')                                                                 # ...not comparable
wall = workbench.maturity_wall([a, b, c])
assert [r['company'] for r in wall['rows']] == ['A', 'B', 'C']

# Coverage aggregates across a mixed set without double-counting missing vs. unconfirmed.
ok = statement('debt_maturity')
no_basis = statement('debt_maturity')
no_basis.pop('basis')
no_total = statement('debt_maturity')
next(f for f in no_total['fields'] if f['key'] == 'total_debt')['value'] = None
no_w1y = statement('debt_maturity')
next(f for f in no_w1y['fields'] if f['key'] == 'due_within_1_year')['value'] = None
zero = statement('debt_maturity')
for key in ('total_debt', 'due_within_1_year'):
    next(f for f in zero['fields'] if f['key'] == key)['value'] = 0
wall = workbench.maturity_wall([ok, no_basis, no_total, no_w1y, zero])
assert wall['coverage'] == {'total': 5, 'comparable': 2, 'missing_total': 1, 'missing_w1y': 1, 'basis_unconfirmed': 1}

# Non-debt_maturity rows are ignored, never crashed on (income_statement fields have no total_debt key).
wall = workbench.maturity_wall([statement('income_statement')])
assert wall['rows'] == [] and wall['coverage']['total'] == 0

print('maturity_wall: baseline share, basis/evidence gates, zero debt, unit scale and currency, entity/lease transparency, partial buckets, sort order, coverage')

for unit, expected in (('MSEK', ('SEK', 'Millions')), ('SEK million', ('SEK', 'Millions')), ('SEKm', ('SEK', 'Millions')), ('Mkr', ('SEK', 'Millions')), ('€m', ('EUR', 'Millions')), ('MUSD', ('USD', 'Millions')),
                       ('KSEK', ('SEK', 'Thousands')), ('SEK thousand', ('SEK', 'Thousands')), ('USD IN THOUSANDS', ('USD', 'Thousands')), ("EUR’000", ('EUR', 'Thousands')), ('SEK 000', ('SEK', 'Thousands')),
                       ('Mdkr', ('SEK', 'Billions')), ('kr', ('SEK', 'Units')), ('SEK', ('SEK', '')), ('%', ('', '')), (None, ('', ''))):
    assert workbench._unit_basis(unit) == expected, (unit, workbench._unit_basis(unit))

with tempfile.TemporaryDirectory() as tmp, patch('pipeline.fetch.fetch_report', side_effect=AssertionError('PDF download forbidden')):
    os.environ['KB_DIR'] = tmp
    for year in (2025, 2024, 2023):
        stem = f'test_{year}'
        folder = Path(tmp) / stem
        folder.mkdir()
        (folder / 'meta.json').write_text(json.dumps(dict(company='Atlas Copco AB', fiscal_year=year, pages=1, filename=stem+'.pdf')))
        (folder / 'pages.jsonl').write_text(json.dumps({'page': 1, 'text': 'Saved statement text'}) + '\n')
        for section in ('income_statement', 'debt_maturity'):
            kb.save_extraction(stem, section, workbench.decorate(statement(section, year), app.load_schema(section)))
    x = app.kb_extraction('test_2025', 'debt_maturity')
    assert x['ready'] and not x['pdf_available'] and x['basis_suggested']['entity'] == 'Atlas Copco AB'
    assert app.review_queue() == []  # a confirmed statement; and an unconfirmed basis never queues either (basis_issues, not issues)
    assert app.comparison('test_2025', 'debt_maturity')['previous_year'] == 2024
    assert app.comparison('test_2025', 'debt_maturity', 'test_2023')['previous_year'] == 2023
    field = copy.deepcopy(x['fields'][0])
    changed = app.review_field(x['report_id'], app.ReviewBody(section=x['section'], key=field['key'], expected=field, decision='corrected', reviewer='Tester', note='Page 1 correction', value=200, unit='MSEK', period='2025'))
    assert changed['checks'][0]['status'] == 'failed' and changed['check_history'][-1]['previous'][0]['passed']
    assert {i['kind'] for i in app.review_queue()} == {'check'}
    corrected = app.review_field(x['report_id'], app.ReviewBody(section=x['section'], key=field['key'], expected=changed['fields'][0], decision='corrected', reviewer='Tester', note='Rechecked page 1', value=100, unit='MSEK', period='2025'))
    assert corrected['ready']
    body = app.BasisBody(section=x['section'], expected=x['basis'], values={**x['basis']['values'], 'leases': 'Excluded'}, reviewer='Tester', note='Debt excludes leases, note 12')
    updated = app.review_basis(x['report_id'], body)
    assert updated['basis_history'][-1]['previous']['values']['leases'] == 'Included' and updated['ready'] and updated['basis_issues'] == []
    try:
        app.review_basis(x['report_id'], body)
        raise AssertionError('stale basis accepted')
    except HTTPException as e:
        assert e.status_code == 409
    app.extractions.clear()
    assert app.kb_extraction('test_2025', 'debt_maturity')['basis'] == updated['basis']
    # Exports bind the selected section, regardless of last section opened.
    rows = list(csv.DictReader(io.StringIO(app.extraction_csv(x['report_id'], 'income_statement', 'test_2024').body.decode())))
    assert all(r['section'] == 'income_statement' for r in rows) and 'basis' in rows[0]
    deck = Presentation(io.BytesIO(app.extraction_pptx(x['report_id'], 'debt_maturity', 'test_2024').body))
    assert len(deck.slides) == 1 and 'previous_stem' in deck.slides[0].notes_slide.notes_text_frame.text
    # Duplicates require an explicit source, and gaps never silently become YoY.
    other = Path(tmp) / 'duplicate_2024'
    other.mkdir()
    for name in ('meta.json', 'pages.jsonl'):
        (other / name).write_bytes((Path(tmp) / 'test_2024' / name).read_bytes())
    kb.save_extraction('duplicate_2024', 'debt_maturity', statement('debt_maturity', 2024))
    assert not app.comparison('test_2025', 'debt_maturity')['rows']
    assert not app.comparison('test_2023', 'debt_maturity')['rows']

# Whole-KB exports deliberately read saved extracts, so this test has no PDF or model dependency.
with tempfile.TemporaryDirectory() as tmp:
    root = Path(tmp)
    library = root / 'reports'
    library.mkdir()
    (library / 'index.json').write_text('[]')
    with patch.dict(os.environ, {'KB_DIR': str(root / 'kb')}), patch.object(app, 'LIBRARY', library), patch.object(app, 'reports', {}), patch.object(app, 'extractions', {}), patch.object(app, 'library_paths', {}):
        for stem, company, reviewed in [('abb_2025', 'ABB Ltd', True), ('ericsson_2025', 'Ericsson', False), ('outside_2025', 'Volvo', False)]:
            x = statement('debt_maturity')
            x.update(stem=stem, report_id=f'lib-{stem}', company=company)
            if not reviewed:
                for field in x['fields']:
                    field.pop('human_review', None)
            if stem == 'ericsson_2025':  # buckets stop reconciling -> identity fails -> wall "incomplete" (share still 0.2)
                next(f for f in x['fields'] if f['key'] == 'due_1_to_5_years')['value'] = 60
            if stem == 'outside_2025':  # a second complete share (200 = 100 + 60 + 40 -> 0.5),
                for key, value in (('total_debt', 200), ('due_within_1_year', 100), ('due_1_to_5_years', 60), ('due_after_5_years', 40)):
                    next(f for f in x['fields'] if f['key'] == key)['value'] = value  # so a sector's median/min/max is non-trivial
            kb.save_report(stem, {'company': company, 'fiscal_year': 2025, 'pages': 1, 'sha256': 'test'}, ['Saved statement text'])
            kb.save_extraction(stem, 'debt_maturity', workbench.decorate(x, app.load_schema('debt_maturity')))
        client = TestClient(app.app)
        rows = list(csv.DictReader(io.StringIO(client.get('/api/kb/export.csv?section=debt_maturity&collection=wallenberg').text)))
        assert len(rows) == 2 and {row['stem'] for row in rows} == {'abb_2025', 'ericsson_2025'}
        assert {'company', 'stem', 'review_status', 'human_review', 'ready'} <= set(rows[0])
        abb = next(row for row in rows if row['stem'] == 'abb_2025')
        assert abb['review_status'] == 'confirmed' and abb['human_review'] == 'yes'
        filtered = client.get('/api/kb/export.csv?section=debt_maturity&collection=all&q=volvo')
        assert filtered.status_code == 200 and [row['stem'] for row in csv.DictReader(io.StringIO(filtered.text))] == ['outside_2025']
        # The deck carries a "Maturity wall by sector" page between the summary table and the
        # per-company slides. _sector_map is frozen so the sectors don't depend on data/ content.
        with patch.object(ppt, '_sector_map', return_value={'abb ltd': 'Industrials', 'ericsson': 'Telecommunications', 'volvo': 'Industrials'}):
            deck = Presentation(io.BytesIO(client.get('/api/kb/export.pptx?section=debt_maturity&collection=wallenberg').content))
        assert len(deck.slides) == 4
        assert any(shape.has_table and shape.table.cell(0, 0).text == 'Company' for shape in deck.slides[0].shapes)
        wall_slide = deck.slides[1]
        texts = ' | '.join(shape.text_frame.text for shape in wall_slide.shapes if shape.has_text_frame)
        assert 'Maturity wall by sector' in texts
        assert 'Industrials' in texts and 'Telecommunications · 0 of 1 with complete buckets' in texts
        assert '20.0% · buckets incomplete' in texts  # ericsson: share still arithmetic, drawn grey
        # "not read" appears only in the page's own legend -- both fixture totals exist, so no
        # company row may show the missing-total label.
        assert sum(1 for shape in wall_slide.shapes if shape.has_text_frame and 'not read' in shape.text_frame.text) == 1
        bars = [shape for shape in wall_slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE]
        assert len(bars) == 2  # one horizontal bar per saved company on the page
        # The visible filter accepts non-ASCII company names; never put its raw bytes in a Latin-1 response header.
        with patch.object(app, 'kb_export_extractions', return_value=[statement('debt_maturity')]):
            unicode_filter = client.get('/api/kb/export.csv', params={'section': 'debt_maturity', 'collection': 'all', 'q': 'Å'})
        assert unicode_filter.status_code == 200
        assert unicode_filter.headers['content-disposition'] == 'attachment; filename="kb_debt_maturity_all.csv"'
        # GET /api/kb/maturity-wall -- same decorated extracts as the exports above, so the
        # collection filter and the reviewed/unreviewed split already set up here double as its test.
        wall = client.get('/api/kb/maturity-wall?collection=wallenberg').json()
        assert wall['coverage'] == {'total': 2, 'comparable': 1, 'missing_total': 0, 'missing_w1y': 0, 'basis_unconfirmed': 0}
        assert {r['stem'] for r in wall['rows']} == {'abb_2025', 'ericsson_2025'}
        assert next(r for r in wall['rows'] if r['stem'] == 'abb_2025')['comparable']
        assert not next(r for r in wall['rows'] if r['stem'] == 'ericsson_2025')['comparable']
        wall_all = client.get('/api/kb/maturity-wall?collection=all').json()
        assert wall_all['coverage']['total'] == 3 and {r['stem'] for r in wall_all['rows']} == {'abb_2025', 'ericsson_2025', 'outside_2025'}
        # Every row names its data/companies.json sector and whether its buckets are complete
        # (stored identity check passed AND total AND <1y present); COMPANIES is patched inline --
        # fixtures never trust the universe file's live content. 3 companies, 2 sectors, one incomplete.
        with patch.object(app, 'COMPANIES', [{'name': 'ABB Ltd', 'sector': 'Industrials'},
                                             {'name': 'Ericsson', 'sector': 'Telecommunications'},
                                             {'name': 'Volvo', 'sector': 'Industrials'}]):
            wall = client.get('/api/kb/maturity-wall?collection=all').json()
        rows = {r['stem']: r for r in wall['rows']}
        assert rows['abb_2025']['sector'] == 'Industrials' and rows['abb_2025']['complete'] and rows['abb_2025']['share'] == 0.2
        assert rows['ericsson_2025']['sector'] == 'Telecommunications' and not rows['ericsson_2025']['complete'] and rows['ericsson_2025']['share'] == 0.2
        assert rows['outside_2025']['sector'] == 'Industrials' and rows['outside_2025']['complete'] and rows['outside_2025']['share'] == 0.5
        sectors = {s['sector']: s for s in wall['sectors']}
        assert sectors['Industrials'] == {'sector': 'Industrials', 'companies': 2, 'complete': 2, 'median_share': 0.35, 'min': 0.2, 'max': 0.5}
        assert sectors['Telecommunications'] == {'sector': 'Telecommunications', 'companies': 1, 'complete': 0, 'median_share': None, 'min': None, 'max': None}
        # A company the universe file doesn't know has no sector: it groups under null, sorted last,
        # and its counts still count -- only incomplete buckets are excluded from median/min/max.
        with patch.object(app, 'COMPANIES', [{'name': 'ABB Ltd', 'sector': 'Industrials'}]):
            wall = client.get('/api/kb/maturity-wall?collection=wallenberg').json()
        assert {r['stem']: r['sector'] for r in wall['rows']} == {'abb_2025': 'Industrials', 'ericsson_2025': None}
        assert [s['sector'] for s in wall['sectors']] == ['Industrials', None]
        assert wall['sectors'][0] == {'sector': 'Industrials', 'companies': 1, 'complete': 1, 'median_share': 0.2, 'min': 0.2, 'max': 0.2}
        assert wall['sectors'][1] == {'sector': None, 'companies': 1, 'complete': 0, 'median_share': None, 'min': None, 'max': None}

# The sector wall page paginates at 30 row-units (companies + sector headers). 35 one-sector
# companies spill onto two "Maturity wall by sector" pages between the summary and the per-company
# slides; build_deck is pure here -- no disk, no KB, no model.
many = []
for i in range(35):
    x = statement('debt_maturity')
    x.update(stem=f'bulk_{i}', report_id=f'lib-bulk_{i}', company=f'Bulk {i} AB')
    many.append(workbench.decorate(x, app.load_schema('debt_maturity')))
bulk = Presentation(io.BytesIO(ppt.build_deck(many)))
sector_slides = [s for s in bulk.slides
                 if any(sh.has_text_frame and 'Maturity wall by sector' in sh.text_frame.text for sh in s.shapes)]
assert len(sector_slides) == 2
assert len(bulk.slides) == 1 + len(sector_slides) + len(many)
assert sum(1 for sh in sector_slides[0].shapes if sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE) \
    + sum(1 for sh in sector_slides[1].shapes if sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE) == 35
assert '(cont.)' in ' | '.join(sh.text_frame.text for sh in sector_slides[1].shapes if sh.has_text_frame)

# Candidate pages -- the deterministic locator behind GET /candidates, served before the
# model runs. Isolated KB folder; the model entrypoint is rigged to fail so the endpoint's
# zero-model claim is asserted, not assumed.
with tempfile.TemporaryDirectory() as tmp:
    os.environ['KB_DIR'] = tmp
    stem = 'candidates_2025'
    folder = Path(tmp) / stem
    folder.mkdir()
    filler = 'Annual report 2025\nKarnell Group'  # on every page: boilerplate the locator strips
    note = 'Note 20 Borrowings Maturity profile of the loans total 1 234 due within 1 year 20 1 to 5 years 50 after 5 years 30'
    pages = [filler] * 10 + [note] + [filler] * 9
    (folder / 'meta.json').write_text(json.dumps(dict(company='Karnell Group', fiscal_year=2025, pages=len(pages), filename=stem + '.pdf')), encoding='utf-8')
    (folder / 'pages.jsonl').write_text(''.join(json.dumps({'page': i + 1, 'text': t}) + '\n' for i, t in enumerate(pages)), encoding='utf-8')
    with patch('app.extract_mod.extract', side_effect=AssertionError('candidates must not call the model')):
        out = app.report_candidates(app.saved_report_id(stem), 'debt_maturity')
    assert [c['page'] for c in out] == [11, 12]  # the note page, then the page after it (statements span two pages)
    heading = out[0]['heading']
    assert heading.startswith('Note 20 Borrowings') and len(heading) <= 80 and '  ' not in heading
    assert 'Karnell Group' not in heading  # the running header was stripped before the heading was cut
    assert out[1]['heading'] == ''  # a boilerplate-only page says nothing
    for call in (lambda: app.report_candidates('lib-nowhere_2025', 'debt_maturity'),
                 lambda: app.report_candidates(app.saved_report_id(stem), 'no_such_section')):
        try:
            call()
            raise AssertionError('candidates accepted an unknown report or section')
        except HTTPException as e:
            assert e.status_code == 404

# GET .../pages/{n}/locate -- zero-model, degrades quote -> longest line -> longest digit run.
# One synthetic PDF with a real text layer: pages 1-5 cover the three match tiers, a duplicated row
# and a clean miss; pages 6-7 separate a wrapped citation's rectangles from its occurrence count.
# Uploaded like any user PDF so require_pdf's own cache is what is actually being read.
with tempfile.TemporaryDirectory() as tmp, patch.dict(os.environ, {'KB_DIR': tmp}), \
        patch.object(app, 'UPLOADS', Path(tmp)), patch.object(app, 'reports', {}), \
        patch.object(app, 'texts_cache', {}), patch.object(app, 'library_paths', {}):
    pdf_file = Path(tmp) / 'locate-src.pdf'
    with pymupdf.open() as doc:
        doc.new_page().insert_text((72, 720), 'Total debt 1 234 MSEK due within one year')  # p.1: verbatim
        twice = doc.new_page()  # p.2: the same line printed twice (group and parent columns)
        twice.insert_text((72, 700), 'Total debt 1 234 MSEK')
        twice.insert_text((72, 650), 'Total debt 1 234 MSEK')
        doc.new_page().insert_text((72, 720), 'Total debt 1 234 MSEK due within one year, per the borrowings note filed with the group')  # p.3: quote's 2nd line only
        doc.new_page().insert_text((72, 720), '1 234')  # p.4: the number alone, no surrounding row
        doc.new_page().insert_text((72, 720), 'Unrelated boilerplate text with no numbers at all')  # p.5: no match at any tier
        wrap_once = doc.new_page()  # p.6: a genuine two-line citation, printed once
        wrap_once.insert_text((72, 720), 'Note 20 Borrowings')
        wrap_once.insert_text((72, 733), 'Total debt is 1 234 MSEK due within one year, filed with the group')
        wrap_twice = doc.new_page()  # p.7: the same two-line citation, printed twice (two real occurrences)
        wrap_twice.insert_text((72, 720), 'Note 20 Borrowings')
        wrap_twice.insert_text((72, 733), 'Total debt is 1 234 MSEK due within one year, filed with the group')
        wrap_twice.insert_text((72, 760), 'Note 20 Borrowings')
        wrap_twice.insert_text((72, 773), 'Total debt is 1 234 MSEK due within one year, filed with the group')
        doc.save(pdf_file)
    client = TestClient(app.app)
    upload = client.post('/api/reports', files={'file': ('locate.pdf', pdf_file.read_bytes(), 'application/pdf')})
    assert upload.status_code == 200, upload.text
    report_id = upload.json()['report_id']

    def locate(page, quote):
        r = client.get(f'/api/reports/{report_id}/pages/{page}/locate', params={'quote': quote})
        assert r.status_code == 200, r.text
        return r.json()

    exact = locate(1, 'Total debt 1 234 MSEK due within one year')
    assert exact['matched'] == 'quote' and len(exact['rects']) == 1 and exact['occurrences'] == 1
    x0, y0, x1, y1 = exact['rects'][0]
    assert 0 <= x0 < x1 <= exact['width'] and 0 <= y0 < y1 <= exact['height'], exact

    dup = locate(2, 'Total debt 1 234 MSEK')
    assert dup['matched'] == 'quote' and len(dup['rects']) == 2 and dup['occurrences'] == 2, dup  # printed twice -> both frame, neither is guessed

    wrapped = locate(3, 'Note 20\nTotal debt 1 234 MSEK due within one year, per the borrowings note filed with the group')
    assert wrapped['matched'] == 'line' and len(wrapped['rects']) == 1 and wrapped['occurrences'] == 1, wrapped  # the fabricated 1st line never printed; the 2nd (longest) did

    cell = locate(4, 'Total debt 1 234 MSEK due within one year')
    assert cell['matched'] == 'value' and len(cell['rects']) == 1 and cell['occurrences'] == 1, cell  # only the bare number sits on this page

    nothing = locate(5, 'Something entirely different, printed nowhere in this report')
    assert nothing['matched'] == 'none' and nothing['rects'] == [] and nothing['occurrences'] == 0, nothing

    # A genuine two-line citation must report as ONE occurrence even though it draws two rects (one
    # per printed line) -- otherwise an ordinary wrapped citation would wrongly cry "2 matches".
    two_line_quote = 'Note 20 Borrowings\nTotal debt is 1 234 MSEK due within one year, filed with the group'
    once = locate(6, two_line_quote)
    assert once['matched'] == 'quote' and len(once['rects']) == 2 and once['occurrences'] == 1, once
    twice = locate(7, two_line_quote)
    assert twice['matched'] == 'quote' and len(twice['rects']) == 4 and twice['occurrences'] == 2, twice  # a real second occurrence must still show as 2

    assert client.get(f'/api/reports/{report_id}/pages/99/locate', params={'quote': 'x'}).status_code == 404
    app.reports['up-nopdf'] = {'report_id': 'up-nopdf', 'filename': 'x.pdf', 'pages': 1, 'company': None, 'fiscal_year': None, 'stem': 'up-nopdf'}
    missing = client.get('/api/reports/up-nopdf/pages/1/locate', params={'quote': 'x'})
    assert missing.status_code == 409, missing.text
print('locate endpoint checks passed: quote/line/value tiers, multi-match, none, out-of-range, no PDF cached')

# Smoke test: build_pptx runs for both sections, with and without a confirmed basis, and with a
# comparison attached. The decks land in the OS temp directory, never in the repository.
for section in ('income_statement', 'debt_maturity'):
    x = statement(section)
    workbench.decorate(x, app.load_schema(section))
    if section == 'income_statement':
        x.pop('basis')
        workbench.decorate(x, app.load_schema(section))
    x['comparison'] = workbench.compare(x, statement(section, 2024))
    Path(tempfile.gettempdir(), f'arp-{section}.pptx').write_bytes(ppt.build_pptx(x))
print('Workbench checks passed: both statements, strict arithmetic, audit history, conflicts, queue, single and whole-KB exports, and no downloads')
